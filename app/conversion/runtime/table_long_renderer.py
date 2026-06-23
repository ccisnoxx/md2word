from __future__ import annotations

# app/conversion/runtime/table_long_renderer.py: 将跨页长表渲染为带重复表头的连续 DOCX 表格。

from typing import Any

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm


DEFAULT_LONG_TABLE_FIRST_PAGE_RESERVE_PT = 42.0
DEFAULT_LONG_TABLE_CONT_PAGE_RESERVE_PT = 12.0
A4_HEIGHT = Mm(297)


def is_long_simple_table(table_model: Any, max_complex_table_rows: int) -> bool:
    if table_model.max_columns != 3:
        return False
    if max_complex_table_rows > 0 and len(table_model.rows) < max_complex_table_rows:
        return False
    for row in table_model.rows[: min(len(table_model.rows), 120)]:
        for cell in row:
            if cell.rowspan != 1 or cell.colspan != 1:
                return False
    return True


def render_long_simple_table(
    container: Any,
    block: Any,
    table_model: Any,
    render_config: Any,
    usable_width: int,
    *,
    apply_named_style_config_fn: Any,
    fill_rich_text_fn: Any,
    prepare_text_fn: Any,
    fill_table_cell_fn: Any,
    remove_table_borders_fn: Any,
    get_style_config_fn: Any,
) -> None:
    if not table_model.rows:
        return

    header_row = table_model.rows[0]
    body_rows = table_model.rows[1:]
    caption_base = block.captions[0] if block.captions else "长表"

    caption_paragraph = container.add_paragraph()
    caption_paragraph.paragraph_format.keep_with_next = True
    apply_named_style_config_fn(caption_paragraph, render_config, "LongTableCaption")
    fill_rich_text_fn(
        caption_paragraph,
        prepare_text_fn(caption_base, render_config),
        render_config=render_config,
        style_name="LongTableCaption",
    )

    render_simple_table_chunk(
        container,
        [header_row, *body_rows],
        render_config,
        usable_width,
        fill_table_cell_fn=fill_table_cell_fn,
        remove_table_borders_fn=remove_table_borders_fn,
        repeat_header=True,
    )

    for footnote in block.footnotes:
        note_paragraph = container.add_paragraph()
        apply_named_style_config_fn(note_paragraph, render_config, "TableNote")
        fill_rich_text_fn(
            note_paragraph,
            prepare_text_fn(footnote, render_config),
            render_config=render_config,
            style_name="TableNote",
        )


def split_long_simple_table_rows(
    *,
    header_row: list[Any],
    body_rows: list[list[Any]],
    render_config: Any,
    usable_width: int,
    first_page_height: float,
    continuation_heights: list[float] | None = None,
    get_style_config_fn: Any,
    caption_text: str | None = None,
) -> list[list[list[Any]]]:
    if not body_rows:
        return []

    page_capacities = build_long_table_page_capacities(
        first_page_height=first_page_height,
        continuation_heights=continuation_heights or [],
        header_row=header_row,
        render_config=render_config,
        usable_width=usable_width,
        get_style_config_fn=get_style_config_fn,
        caption_text=caption_text,
    )
    row_weights = [
        estimate_row_height(
            row,
            render_config,
            usable_width,
            get_style_config_fn=get_style_config_fn,
        )
        for row in body_rows
    ]

    return split_rows_by_page_capacities(
        body_rows,
        row_weights,
        page_capacities,
    )


def build_long_table_page_capacities(
    *,
    first_page_height: float,
    continuation_heights: list[float],
    header_row: list[Any],
    render_config: Any,
    usable_width: int,
    get_style_config_fn: Any,
    caption_text: str | None = None,
) -> list[float]:
    capacities = [
        estimate_long_table_page_capacity(
            first_page_height,
            header_row,
            render_config,
            usable_width,
            is_first_page=True,
            get_style_config_fn=get_style_config_fn,
            caption_text=caption_text,
        )
    ]

    continuation_source_heights = continuation_heights or [first_page_height]
    capacities.extend(
        estimate_long_table_page_capacity(
            continuation_height,
            header_row,
            render_config,
            usable_width,
            is_first_page=False,
            get_style_config_fn=get_style_config_fn,
            caption_text=caption_text,
        )
        for continuation_height in continuation_source_heights
    )
    return capacities


def calibrate_row_weights_to_source_pages(
    row_weights: list[float],
    page_capacities: list[float],
    has_continuation_pages: bool,
) -> list[float]:
    if not has_continuation_pages or not row_weights or not page_capacities:
        return row_weights

    total_weight = sum(row_weights)
    total_capacity = sum(page_capacities)
    if total_weight <= 0 or total_capacity <= 0:
        return row_weights

    scaling_factor = total_capacity / total_weight
    if 0.92 <= scaling_factor <= 1.08:
        return row_weights

    return [weight * scaling_factor for weight in row_weights]


def split_rows_by_page_capacities(
    body_rows: list[list[Any]],
    row_weights: list[float],
    page_capacities: list[float],
) -> list[list[list[Any]]]:
    if not body_rows:
        return []

    chunks: list[list[list[Any]]] = []
    cursor = 0
    page_idx = 0
    fallback_capacity = page_capacities[-1] if page_capacities else 0.0

    while cursor < len(body_rows):
        capacity = page_capacities[page_idx] if page_idx < len(page_capacities) else fallback_capacity
        current_chunk: list[list[Any]] = []
        used_weight = 0.0

        while cursor < len(body_rows):
            weight = row_weights[cursor]
            if current_chunk and used_weight + weight > capacity:
                break
            current_chunk.append(body_rows[cursor])
            used_weight += weight
            cursor += 1

        if not current_chunk and cursor < len(body_rows):
            current_chunk.append(body_rows[cursor])
            cursor += 1

        chunks.append(current_chunk)
        page_idx += 1

    return chunks


def estimate_long_table_page_capacity(
    first_page_height: float,
    header_row: list[Any],
    render_config: Any,
    usable_width: int,
    *,
    is_first_page: bool,
    get_style_config_fn: Any,
    caption_text: str | None = None,
) -> float:
    base_page_height = emu_to_pt(int(A4_HEIGHT - render_config.section_layout.top_margin - render_config.section_layout.bottom_margin))
    header_height = estimate_row_height(
        header_row,
        render_config,
        usable_width,
        get_style_config_fn=get_style_config_fn,
    )
    reserve = DEFAULT_LONG_TABLE_FIRST_PAGE_RESERVE_PT if is_first_page else DEFAULT_LONG_TABLE_CONT_PAGE_RESERVE_PT
    caption_height = estimate_paragraph_height(
        caption_text or "",
        render_config,
        usable_width,
        style_name="LongTableCaption",
        get_style_config_fn=get_style_config_fn,
    )

    if is_first_page:
        ratio = max(min(first_page_height / 820.0, 1.0), 0.45)
        available = base_page_height * ratio - reserve - header_height - caption_height
    else:
        available = base_page_height - reserve - header_height - caption_height

    return max(available, header_height + 24)


def estimate_paragraph_height(
    text: str,
    render_config: Any,
    usable_width: int,
    *,
    style_name: str,
    get_style_config_fn: Any,
) -> float:
    if not text.strip():
        return 0.0

    style = get_style_config_fn(render_config, style_name)
    font_size = style.font_size if style and style.font_size else 10.5
    line_spacing = style.line_spacing if style and style.line_spacing else 1.0
    line_height = font_size * line_spacing + 0.2
    lines = estimate_cell_lines(text, emu_to_pt(usable_width), font_size)
    space_before = style.space_before if style and style.space_before is not None else 0.0
    space_after = style.space_after if style and style.space_after is not None else 0.0
    return lines * line_height + space_before + space_after


def estimate_row_height(
    row: list[Any],
    render_config: Any,
    usable_width: int,
    *,
    get_style_config_fn: Any,
) -> float:
    max_cols = max(sum(cell.colspan for cell in row), 1)
    col_width_pt = emu_to_pt(max(int(usable_width / max_cols), int(Emu(Mm(16)))))
    style = get_style_config_fn(render_config, "TableContent")
    font_size = style.font_size if style and style.font_size else 9.0
    line_spacing = style.line_spacing if style and style.line_spacing else 1.0
    # 这里按“单行技术表格”的真实观感收紧估算，否则续表分页会过于保守，
    # 导致页底长期保留大块空白。
    line_height = font_size * line_spacing + 0.2

    max_lines = 1
    for cell in row:
        text = cell.text or ""
        lines = estimate_cell_lines(text, col_width_pt * max(cell.colspan, 1), font_size)
        max_lines = max(max_lines, lines)

    return max_lines * line_height + 1


def estimate_cell_lines(text: str, width_pt: float, font_size_pt: float) -> int:
    if not text.strip():
        return 1

    wrapped_units = 0.0
    for char in text:
        if char == "\n":
            wrapped_units += max(width_pt / max(font_size_pt, 1), 1)
            continue
        wrapped_units += 0.90 if ord(char) > 127 else 0.45

    chars_per_line = max(width_pt / max(font_size_pt, 1), 4)
    return max(int((wrapped_units / chars_per_line) + 0.999), 1)


def emu_to_pt(value: int) -> float:
    return value / 12700.0


def render_simple_table_chunk(
    container: Any,
    rows: list[list[Any]],
    render_config: Any,
    usable_width: int,
    *,
    fill_table_cell_fn: Any,
    remove_table_borders_fn: Any,
    repeat_header: bool = False,
) -> None:
    max_cols = max((sum(cell.colspan for cell in row) for row in rows), default=0)
    if max_cols <= 0:
        return

    table = container.add_table(rows=len(rows), cols=max_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders_fn(table, render_config)
    configure_table_rows(table, repeat_header=repeat_header)

    column_width = max(int(usable_width // max_cols), int(Emu(Mm(16))))
    for col in table.columns:
        col.width = column_width

    for row_idx, row in enumerate(rows):
        table_row = table.rows[row_idx]
        for col_idx, cell_model in enumerate(row[:max_cols]):
            cell = table_row.cells[col_idx]
            fill_table_cell_fn(cell, cell_model, row_idx == 0, render_config)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def configure_table_rows(table: Any, *, repeat_header: bool) -> None:
    for row_idx, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        ensure_row_flag(row_properties, "w:cantSplit")
        if repeat_header and row_idx == 0:
            ensure_row_flag(row_properties, "w:tblHeader", value="true")


def ensure_row_flag(row_properties: Any, tag: str, *, value: str | None = None) -> None:
    element = row_properties.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        row_properties.append(element)
    if value is not None:
        element.set(qn("w:val"), value)

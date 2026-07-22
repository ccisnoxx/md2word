from __future__ import annotations

# app/conversion/runtime/text_rendering.py: 渲染正文、标题、列表和内联公式文本到 DOCX。

import re
from typing import Any

from bs4 import BeautifulSoup, NavigableString, Tag
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.conversion.docx_rules import text_semantics as _shared_text_semantics
from app.conversion.docx_rules.run_fonts import apply_run_fonts


BODY_FONT = Pt(10.5)
SECTION_TITLE_FONT = Pt(12.5)
TOP_TITLE_FONT = Pt(16)

HEADING_NUMBER_RE = re.compile(
    r"^(?:\d+[.、]|\d+\.\d+|附录[A-Z\d.\-]+|appendix\s+[A-Z\d.\-]+|(?:表|table|图|figure)\s*[A-Z\d.\-]+)",
    re.IGNORECASE,
)
TOC_TITLE_RE = re.compile(r"^(目录|table of contents|contents)\b", re.IGNORECASE)
TRAILING_PAGE_NOISE_RE = re.compile(r"(?:[.\s·。…]{2,}|\s+)\d+\s*$")
TOC_ENTRY_PARTS_RE = re.compile(r"^(?P<title>.*?)(?:\s*(?:\.{2,}|·{2,}|…{2,})\s*)(?P<page>\d+)\s*$")
INLINE_MARKUP_SPLIT_RE = re.compile(r"(<[^>]+>)")
MD2WORD_BULLET_NUM_ID_ATTR = "_md2word_bullet_num_id"

_RUNTIME_MODULE: Any | None = None


def bind_runtime(module: Any) -> None:
    global _RUNTIME_MODULE
    _RUNTIME_MODULE = module


def _runtime_value(name: str, default: Any) -> Any:
    if _RUNTIME_MODULE is None:
        return default
    return getattr(_RUNTIME_MODULE, name, default)


def _clean_text(text: str) -> str:
    return " ".join(text.replace("\n", " ").split())


def normalize_toc_text(text: str) -> str:
    cleaned = _clean_text(text)
    cleaned = TRAILING_PAGE_NOISE_RE.sub("", cleaned)
    cleaned = cleaned.replace("..", " ").replace(" .", " ")
    cleaned = re.sub(r"\s{2,}", " ", cleaned)
    return cleaned.strip()


def fill_toc_entry_paragraph(paragraph: Any, text: str, render_config: Any, usable_width: Any | None = None) -> None:
    paragraph.text = ""
    normalize_toc_text_fn = _runtime_value("normalize_toc_text", normalize_toc_text)
    split_toc_entry_parts_fn = _runtime_value("split_toc_entry_parts", split_toc_entry_parts)
    infer_heading_level_fn = _runtime_value("infer_heading_level", infer_heading_level)
    get_toc_style_name_fn = _runtime_value("get_toc_style_name", get_toc_style_name)
    apply_named_style_config_fn = _runtime_value("apply_named_style_config", apply_named_style_config)
    fill_rich_text_fn = _runtime_value("fill_rich_text", fill_rich_text)
    prepare_text_fn = _runtime_value("prepare_text", prepare_text)
    get_style_config_fn = _runtime_value("get_style_config", get_style_config)
    apply_run_defaults_fn = _runtime_value("apply_run_defaults", apply_run_defaults)

    title, page_number = split_toc_entry_parts_fn(text)
    normalized_title = normalize_toc_text_fn(title)
    level = infer_heading_level_fn(normalized_title)
    style_name = get_toc_style_name_fn(level)
    apply_named_style_config_fn(paragraph, render_config, style_name)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(compute_toc_indent_pt(normalized_title, render_config))
    if page_number:
        add_toc_dot_leader_tab_stop(paragraph, usable_width)
        fill_rich_text_fn(
            paragraph,
            prepare_text_fn(normalized_title, render_config),
            render_config=render_config,
            style_name=style_name,
        )
        page_run = paragraph.add_run(f"\t{page_number}")
        apply_run_defaults_fn(page_run, get_style_config_fn(render_config, style_name))
        return

    fill_rich_text_fn(
        paragraph,
        prepare_text_fn(normalized_title, render_config),
        render_config=render_config,
        style_name=style_name,
    )


def fill_toc_title_paragraph(paragraph: Any, text: str, render_config: Any) -> None:
    paragraph.text = ""
    apply_named_style_config_fn = _runtime_value("apply_named_style_config", apply_named_style_config)
    fill_rich_text_fn = _runtime_value("fill_rich_text", fill_rich_text)
    prepare_text_fn = _runtime_value("prepare_text", prepare_text)

    apply_named_style_config_fn(paragraph, render_config, "TocTitle")
    fill_rich_text_fn(paragraph, prepare_text_fn(text, render_config), render_config=render_config, style_name="TocTitle", bold=True)


def get_toc_style_name(level: int) -> str:
    if level <= 1:
        return "TocLevel1"
    if level == 2:
        return "TocLevel2"
    return "TocLevel3"


def compute_toc_indent_pt(text: str, render_config: Any) -> float:
    level = _runtime_value("infer_heading_level", infer_heading_level)(_runtime_value("normalize_toc_text", normalize_toc_text)(text))
    style = _runtime_value("get_style_config", get_style_config)(render_config, "Normal")
    font_size = style.font_size if style and style.font_size else 10.5
    return max(level - 1, 0) * font_size * 2


def split_toc_entry_parts(text: str) -> tuple[str, str | None]:
    cleaned = _clean_text(text)
    match = TOC_ENTRY_PARTS_RE.match(cleaned)
    if not match:
        return cleaned, None
    return _clean_text(match.group("title")), match.group("page")


def add_toc_dot_leader_tab_stop(paragraph: Any, position: Any | None = None) -> None:
    position = position or Pt(453.6)
    paragraph_element = getattr(paragraph, "_p", None)
    if paragraph_element is not None:
        paragraph_properties = paragraph_element.get_or_add_pPr()
        tabs = paragraph_properties.find(qn("w:tabs"))
        if tabs is None:
            tabs = OxmlElement("w:tabs")
            paragraph_properties.append(tabs)
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:leader"), "dot")
        tab.set(qn("w:pos"), str(int(position) // 635))
        tabs.append(tab)
        return

    tab_stops = getattr(getattr(paragraph, "paragraph_format", None), "tab_stops", None)
    if tab_stops is not None and hasattr(tab_stops, "add_tab_stop"):
        tab_stops.add_tab_stop(position)


def prepare_text(text: str, render_config: Any | None) -> str:
    prepared = normalize_formula_text(text)
    if render_config and render_config.terminology:
        prepared = render_config.terminology.replace_text(prepared)
    return prepared


def apply_style_if_exists(paragraph: Any, style_name: str) -> None:
    try:
        paragraph.style = style_name
    except (KeyError, ValueError):
        return


def apply_named_style_config(paragraph: Any, render_config: Any, style_name: str) -> None:
    apply_style_if_exists(paragraph, style_name)
    config = render_config.paragraph_styles.get(style_name)
    if not config:
        return

    fmt = paragraph.paragraph_format
    if config.alignment:
        paragraph.alignment = parse_alignment(config.alignment)
    if config.space_before is not None:
        fmt.space_before = Pt(config.space_before)
    if config.space_after is not None:
        fmt.space_after = Pt(config.space_after)
    if config.line_spacing is not None:
        fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        fmt.line_spacing = config.line_spacing
    if config.first_line_indent_pt is not None:
        fmt.first_line_indent = Pt(config.first_line_indent_pt)
    if config.left_indent_pt is not None:
        fmt.left_indent = Pt(config.left_indent_pt)
    if config.right_indent_pt is not None:
        fmt.right_indent = Pt(config.right_indent_pt)


def parse_alignment(value: str) -> WD_ALIGN_PARAGRAPH:
    normalized = value.strip().lower()
    mapping = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    return mapping.get(normalized, WD_ALIGN_PARAGRAPH.LEFT)


def get_style_config(render_config: Any | None, style_name: str | None) -> Any | None:
    if not render_config or not style_name:
        return None
    return render_config.paragraph_styles.get(style_name)


def apply_run_defaults(run: Any, style_config: Any | None = None) -> None:
    if style_config:
        apply_run_fonts(
            run,
            font_name=style_config.font_name,
            western_font_name=style_config.western_font_name,
            complex_script_font_name=style_config.complex_script_font_name,
        )
    if style_config and style_config.font_size is not None:
        run.font.size = Pt(style_config.font_size)
    elif run.font.size is None:
        run.font.size = BODY_FONT
    if style_config and style_config.bold is not None:
        run.bold = style_config.bold if run.bold is None else run.bold
    if style_config and style_config.italic is not None:
        run.italic = style_config.italic if run.italic is None else run.italic
    if style_config and style_config.color:
        run.font.color.rgb = parse_rgb_color(style_config.color)


def parse_rgb_color(value: str) -> RGBColor:
    raw = value.strip().lstrip("#")
    if len(raw) != 6:
        return RGBColor(0, 0, 0)
    return RGBColor.from_string(raw.upper())


def fill_rich_text(
    paragraph: Any,
    text: str,
    *,
    render_config: Any | None = None,
    style_name: str | None = None,
    bold: bool = False,
    italic: bool = False,
) -> None:
    paragraph.text = ""
    _runtime_value("render_inline_html_into_paragraph", render_inline_html_into_paragraph)(
        paragraph,
        text,
        render_config=render_config,
        style_name=style_name,
        bold=bold,
        italic=italic,
    )


def render_inline_html_into_paragraph(
    paragraph: Any,
    text_or_html: str,
    *,
    render_config: Any | None = None,
    style_name: str | None = None,
    bold: bool = False,
    italic: bool = False,
    subscript: bool = False,
    superscript: bool = False,
    table_header_bold: bool = True,
) -> None:
    style_config = _runtime_value("get_style_config", get_style_config)(render_config, style_name)
    soup = _runtime_value("BeautifulSoup", BeautifulSoup)(f"<root>{text_or_html}</root>", "html.parser")
    root = soup.find("root")
    if root is None:
        run = paragraph.add_run(_runtime_value("normalize_formula_text", normalize_formula_text)(text_or_html))
        _runtime_value("apply_run_defaults", apply_run_defaults)(run, style_config)
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        if subscript:
            run.font.subscript = True
        if superscript:
            run.font.superscript = True
        return

    _runtime_value("render_html_children", render_html_children)(
        root,
        paragraph,
        style_config=style_config,
        bold=bold,
        italic=italic,
        subscript=subscript,
        superscript=superscript,
        table_header_bold=table_header_bold,
    )


def render_html_children(
    parent: Tag,
    paragraph: Any,
    *,
    style_config: Any | None,
    bold: bool,
    italic: bool,
    subscript: bool,
    superscript: bool,
    table_header_bold: bool = True,
) -> None:
    for child in parent.children:
        if isinstance(child, NavigableString):
            _runtime_value("render_normalized_text_fragment", render_normalized_text_fragment)(
                paragraph,
                str(child),
                style_config=style_config,
                bold=bold,
                italic=italic,
                subscript=subscript,
                superscript=superscript,
            )
            continue

        if not isinstance(child, Tag):
            continue

        if child.name == "br":
            paragraph.add_run("\n")
            continue

        render_html_children(
            child,
            paragraph,
            style_config=style_config,
            bold=bold or child.name in {"b", "strong"} or (table_header_bold and child.name == "th"),
            italic=italic or child.name in {"i", "em"},
            subscript=subscript or child.name == "sub",
            superscript=superscript or child.name == "sup",
            table_header_bold=table_header_bold,
        )


def render_normalized_text_fragment(
    paragraph: Any,
    text: str,
    *,
    style_config: Any | None,
    bold: bool,
    italic: bool,
    subscript: bool,
    superscript: bool,
) -> None:
    had_leading_space = bool(text[:1].isspace())
    had_trailing_space = bool(text[-1:].isspace())
    normalized = _runtime_value("normalize_formula_text", normalize_formula_text)(text)
    if not normalized:
        return
    if had_leading_space and not normalized.startswith(" "):
        normalized = f" {normalized}"
    if had_trailing_space and not normalized.endswith(" "):
        normalized = f"{normalized} "

    if "<" not in normalized or ">" not in normalized:
        _runtime_value("append_text_run", append_text_run)(
            paragraph,
            normalized,
            style_config=style_config,
            bold=bold,
            italic=italic,
            subscript=subscript,
            superscript=superscript,
        )
        return

    if _runtime_value("render_inline_markup_fragment", render_inline_markup_fragment)(
        paragraph,
        normalized,
        style_config=style_config,
        bold=bold,
        italic=italic,
        subscript=subscript,
        superscript=superscript,
    ):
        return

    fragment = _runtime_value("BeautifulSoup", BeautifulSoup)(f"<root>{normalized}</root>", "html.parser").find("root")
    if fragment is None:
        _runtime_value("append_text_run", append_text_run)(
            paragraph,
            normalized,
            style_config=style_config,
            bold=bold,
            italic=italic,
            subscript=subscript,
            superscript=superscript,
        )
        return

    _runtime_value("render_markup_children", render_markup_children)(
        fragment,
        paragraph,
        style_config=style_config,
        bold=bold,
        italic=italic,
        subscript=subscript,
        superscript=superscript,
    )


def render_inline_markup_fragment(
    paragraph: Any,
    normalized: str,
    *,
    style_config: Any | None,
    bold: bool,
    italic: bool,
    subscript: bool,
    superscript: bool,
) -> bool:
    active_subscript = subscript
    active_superscript = superscript
    handled_markup = False

    for token in INLINE_MARKUP_SPLIT_RE.split(normalized):
        if not token:
            continue

        lowered = token.lower()
        if lowered == "<sub>":
            active_subscript = True
            handled_markup = True
            continue
        if lowered == "</sub>":
            active_subscript = subscript
            handled_markup = True
            continue
        if lowered == "<sup>":
            active_superscript = True
            handled_markup = True
            continue
        if lowered == "</sup>":
            active_superscript = superscript
            handled_markup = True
            continue
        if lowered in {"<br>", "<br/>", "<br />"}:
            paragraph.add_run("\n")
            handled_markup = True
            continue

        if token.startswith("<") and token.endswith(">"):
            return False

        _runtime_value("append_text_run", append_text_run)(
            paragraph,
            token,
            style_config=style_config,
            bold=bold,
            italic=italic,
            subscript=active_subscript,
            superscript=active_superscript,
        )

    return handled_markup


def render_markup_children(
    parent: Tag,
    paragraph: Any,
    *,
    style_config: Any | None,
    bold: bool,
    italic: bool,
    subscript: bool,
    superscript: bool,
) -> None:
    for child in parent.children:
        if isinstance(child, NavigableString):
            text = str(child)
            if not text:
                continue
            _runtime_value("append_text_run", append_text_run)(
                paragraph,
                text,
                style_config=style_config,
                bold=bold,
                italic=italic,
                subscript=subscript,
                superscript=superscript,
            )
            continue

        if not isinstance(child, Tag):
            continue

        if child.name == "br":
            paragraph.add_run("\n")
            continue

        render_markup_children(
            child,
            paragraph,
            style_config=style_config,
            bold=bold or child.name in {"b", "strong", "th"},
            italic=italic or child.name in {"i", "em"},
            subscript=subscript or child.name == "sub",
            superscript=superscript or child.name == "sup",
        )


def append_text_run(
    paragraph: Any,
    text: str,
    *,
    style_config: Any | None,
    bold: bool,
    italic: bool,
    subscript: bool,
    superscript: bool,
) -> None:
    run = paragraph.add_run(text)
    _runtime_value("apply_run_defaults", apply_run_defaults)(run, style_config)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True


def fill_title_paragraph(
    paragraph: Any,
    block: Any,
    page_width: float,
    render_config: Any,
    *,
    keep_with_next: bool = True,
) -> None:
    paragraph.text = ""
    paragraph.paragraph_format.keep_with_next = keep_with_next
    prepare_text_fn = _runtime_value("prepare_text", prepare_text)
    infer_heading_level_fn = _runtime_value("infer_heading_level", infer_heading_level)
    is_cover_title_fn = _runtime_value("is_cover_title", is_cover_title)
    apply_named_style_config_fn = _runtime_value("apply_named_style_config", apply_named_style_config)
    fill_rich_text_fn = _runtime_value("fill_rich_text", fill_rich_text)
    get_style_config_fn = _runtime_value("get_style_config", get_style_config)
    apply_run_defaults_fn = _runtime_value("apply_run_defaults", apply_run_defaults)

    prepared = prepare_text_fn(block.text, render_config)
    level = infer_heading_level_fn(prepared)

    if is_cover_title_fn(block, prepared, page_width):
        apply_named_style_config_fn(paragraph, render_config, "PartNumber")
        style_name = "PartNumber"
    else:
        style_name = f"Heading {level}"
        apply_named_style_config_fn(paragraph, render_config, style_name)

    fill_rich_text_fn(paragraph, prepared, render_config=render_config, style_name=style_name, bold=True)
    for run in paragraph.runs:
        apply_run_defaults_fn(run, get_style_config_fn(render_config, style_name))
        if is_cover_title_fn(block, prepared, page_width):
            if run.font.size is None:
                run.font.size = TOP_TITLE_FONT
        elif run.font.size is None:
            run.font.size = SECTION_TITLE_FONT


def is_cover_title(block: Any, text: str, page_width: float) -> bool:
    if block.page_idx not in {0, 1}:
        return False
    if block.y1 >= 220 or block.width < page_width * 0.22:
        return False
    if re.match(r"^\d+\s+\S+", text):
        return False
    if HEADING_NUMBER_RE.match(text):
        return False
    if TOC_TITLE_RE.match(text.strip()):
        return False
    return True


def fill_body_paragraph(paragraph: Any, text: str, render_config: Any) -> None:
    paragraph.text = ""
    apply_named_style_config_fn = _runtime_value("apply_named_style_config", apply_named_style_config)
    fill_rich_text_fn = _runtime_value("fill_rich_text", fill_rich_text)
    prepare_text_fn = _runtime_value("prepare_text", prepare_text)
    get_style_config_fn = _runtime_value("get_style_config", get_style_config)
    apply_run_defaults_fn = _runtime_value("apply_run_defaults", apply_run_defaults)

    apply_named_style_config_fn(paragraph, render_config, "Normal")
    fill_rich_text_fn(paragraph, prepare_text_fn(text, render_config), render_config=render_config, style_name="Normal")
    for run in paragraph.runs:
        apply_run_defaults_fn(run, get_style_config_fn(render_config, "Normal"))


def fill_list_paragraph(paragraph: Any, text: str, render_config: Any, *, level: int = 0) -> None:
    """Render a MinerU list item as real Word numbering, not manual bullet text."""
    paragraph.text = ""
    apply_named_style_config_fn = _runtime_value("apply_named_style_config", apply_named_style_config)
    apply_bullet_numbering_fn = _runtime_value("apply_bullet_numbering", apply_bullet_numbering)
    fill_rich_text_fn = _runtime_value("fill_rich_text", fill_rich_text)
    prepare_text_fn = _runtime_value("prepare_text", prepare_text)
    get_style_config_fn = _runtime_value("get_style_config", get_style_config)
    apply_run_defaults_fn = _runtime_value("apply_run_defaults", apply_run_defaults)

    apply_named_style_config_fn(paragraph, render_config, "Normal")
    apply_bullet_numbering_fn(paragraph, level=level)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.left_indent = Pt(12 + min(max(level, 0), 1) * 12)
    paragraph.paragraph_format.first_line_indent = Pt(-9)
    fill_rich_text_fn(paragraph, prepare_text_fn(text, render_config), render_config=render_config, style_name="Normal")
    for run in paragraph.runs:
        apply_run_defaults_fn(run, get_style_config_fn(render_config, "Normal"))


def apply_bullet_numbering(paragraph: Any, *, level: int = 0) -> None:
    """Attach a shared two-level bullet numbering definition to a paragraph."""
    normalized_level = min(max(level, 0), 1)
    num_id = _runtime_value("ensure_bullet_numbering", ensure_bullet_numbering)(paragraph)
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)

    _runtime_value("set_numbering_child_value", set_numbering_child_value)(num_pr, "w:ilvl", str(normalized_level))
    _runtime_value("set_numbering_child_value", set_numbering_child_value)(num_pr, "w:numId", str(num_id))


def ensure_bullet_numbering(paragraph: Any) -> str:
    """Create the local bullet numbering definition once per DOCX part."""
    part = paragraph.part
    cached_num_id = getattr(part, MD2WORD_BULLET_NUM_ID_ATTR, None)
    if cached_num_id is not None:
        return str(cached_num_id)

    numbering = part.numbering_part.element
    abstract_num_id = str(_runtime_value("next_numbering_id", next_numbering_id)(numbering, "w:abstractNum", "w:abstractNumId"))
    num_id = str(_runtime_value("next_numbering_id", next_numbering_id)(numbering, "w:num", "w:numId"))

    abstract_num = OxmlElement("w:abstractNum")
    abstract_num.set(qn("w:abstractNumId"), abstract_num_id)
    multi_level_type = OxmlElement("w:multiLevelType")
    multi_level_type.set(qn("w:val"), "hybridMultilevel")
    abstract_num.append(multi_level_type)

    for ilvl, marker, left in ((0, "•", "240"), (1, "–", "480")):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(ilvl))
        for tag, value in (
            ("w:start", "1"),
            ("w:numFmt", "bullet"),
            ("w:lvlText", marker),
            ("w:lvlJc", "left"),
        ):
            child = OxmlElement(tag)
            child.set(qn("w:val"), value)
            lvl.append(child)
        p_pr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), left)
        ind.set(qn("w:hanging"), "180")
        p_pr.append(ind)
        lvl.append(p_pr)
        abstract_num.append(lvl)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), num_id)
    abstract_num_id_element = OxmlElement("w:abstractNumId")
    abstract_num_id_element.set(qn("w:val"), abstract_num_id)
    num.append(abstract_num_id_element)

    numbering.append(abstract_num)
    numbering.append(num)
    setattr(part, MD2WORD_BULLET_NUM_ID_ATTR, num_id)
    return num_id


def next_numbering_id(numbering: Any, element_tag: str, attr_name: str) -> int:
    values: list[int] = []
    for element in numbering.findall(qn(element_tag)):
        raw_value = element.get(qn(attr_name))
        if raw_value is not None and raw_value.isdigit():
            values.append(int(raw_value))
    return max(values, default=0) + 1


def set_numbering_child_value(parent: Any, tag: str, value: str) -> None:
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    child.set(qn("w:val"), value)


normalize_formula_text = _shared_text_semantics.normalize_formula_text
normalize_tex_text = _shared_text_semantics.normalize_tex_text
convert_inline_tex_to_html = _shared_text_semantics.convert_inline_tex_to_html
convert_simple_tex_expr = _shared_text_semantics.convert_simple_tex_expr
clean_tex_fragment = _shared_text_semantics.clean_tex_fragment
compact_math_token_runs = _shared_text_semantics.compact_math_token_runs
collapse_double_escaped_tex_commands = _shared_text_semantics.collapse_double_escaped_tex_commands
unwrap_tex_text_wrappers = _shared_text_semantics.unwrap_tex_text_wrappers
replace_control_word = _shared_text_semantics.replace_control_word
replace_overline_expr = _shared_text_semantics.replace_overline_expr
infer_list_item_level = _shared_text_semantics.infer_list_item_level
is_list_item_text = _shared_text_semantics.is_list_item_text
normalize_list_item = _shared_text_semantics.normalize_list_item
infer_heading_level = _shared_text_semantics.infer_heading_level
looks_like_heading_text = _shared_text_semantics.looks_like_heading_text
looks_like_english_heading_text = _shared_text_semantics.looks_like_english_heading_text
parse_toc_entry_text = _shared_text_semantics.parse_toc_entry_text
is_toc_title_text = _shared_text_semantics.is_toc_title_text

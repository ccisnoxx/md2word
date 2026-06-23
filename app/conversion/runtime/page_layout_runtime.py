from __future__ import annotations

# app/conversion/runtime/page_layout_runtime.py: 按 MinerU 页面和 block 位置组织 DOCX 页面布局。

import re
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENTATION, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt


_RUNTIME_MODULE: Any | None = None
TOP_LEVEL_SECTION_TITLE_RE = re.compile(r"^\s*\d{1,2}\s+\S")
DENSE_PAGE_BOTTOM_RATIO = 0.96
DENSE_PAGE_OVERFLOW_RATIO = 1.08


def bind_runtime(module: Any) -> None:
    global _RUNTIME_MODULE
    _RUNTIME_MODULE = module


def _runtime_value(name: str, default: Any) -> Any:
    if _RUNTIME_MODULE is None:
        return default
    return getattr(_RUNTIME_MODULE, name, default)


class _NullLogger:
    def debug(self, *_args: Any, **_kwargs: Any) -> None:
        return

    def warning(self, *_args: Any, **_kwargs: Any) -> None:
        return


def _logger() -> Any:
    return _runtime_value("logger", _NullLogger())


def _region_cls() -> Any:
    cls = _runtime_value("Region", None)
    if cls is None:
        raise RuntimeError("Region is not bound")
    return cls


def _toc_title_re() -> Any:
    return _runtime_value("TOC_TITLE_RE", re.compile(r"^$"))


def _toc_entry_re() -> Any:
    return _runtime_value("TOC_ENTRY_RE", re.compile(r"^$"))


def _visual_types() -> set[str]:
    return _runtime_value("VISUAL_TYPES", set())


def _product_title_re() -> Any:
    return _runtime_value("PRODUCT_TITLE_RE", re.compile(r"^(?P<part>[A-Z0-9xX,+/\- ]{3,})\s+(?P<category>.+)$"))


def _appendix_title_re() -> Any:
    return re.compile(r"^(?:附录|appendix\b)", re.IGNORECASE)


def build_document(
    pages: list[Any],
    render_config: Any,
    template_document: Any | None,
) -> Document:
    create_document_fn = _runtime_value(
        "create_document",
        lambda template: template if template is not None else Document(),
    )
    clear_document_body_fn = _runtime_value("clear_document_body", lambda _document: None)
    configure_section_fn = _runtime_value("configure_section", configure_section)
    extract_document_identity_fn = _runtime_value("extract_document_identity", extract_document_identity)
    fill_template_intro_fn = _runtime_value("fill_template_intro", fill_template_intro)
    render_document_page_fn = _runtime_value("render_document_page", render_document_page)
    clone_first_section_layout_fn = _runtime_value("clone_first_section_layout", lambda _document: None)
    split_page_blocks_fn = _runtime_value("split_page_blocks", split_page_blocks)

    document = create_document_fn(template_document)
    clear_document_body_fn(document)
    first_renderable_page = next(
        (
            page
            for page in pages
            if split_page_blocks_fn(page.blocks)[3]
        ),
        pages[0] if pages else None,
    )
    if first_renderable_page is not None:
        configure_section_fn(document.sections[0], first_renderable_page.page_size, render_config.section_layout)

    part_number, category = extract_document_identity_fn(pages)
    fill_template_intro_fn(document, part_number, category, render_config)

    rendered_page_idx = 0
    toc_mode_active = False
    previous_rendered_page: Any | None = None
    for page_idx, page in enumerate(pages):
        rendered_page = render_document_page_fn(
            document,
            page,
            render_config,
            rendered_page_idx,
            toc_mode_active,
            previous_rendered_page=previous_rendered_page,
        )
        if rendered_page is None:
            _logger().debug("skip empty rendered page page={}", page_idx)
            continue
        rendered_page_idx += 1
        toc_mode_active = rendered_page
        previous_rendered_page = page

    clone_first_section_layout_fn(document)
    return document


def render_document_page(
    document: Document,
    page: Any,
    render_config: Any,
    rendered_page_idx: int,
    toc_mode_active: bool,
    previous_rendered_page: Any | None = None,
) -> bool | None:
    split_page_blocks_fn = _runtime_value("split_page_blocks", split_page_blocks)
    prepare_document_section_fn = _runtime_value("prepare_document_section", prepare_document_section)
    resolve_toc_page_state_fn = _runtime_value("resolve_toc_page_state", resolve_toc_page_state)
    render_document_page_body_fn = _runtime_value("render_document_page_body", render_document_page_body)

    header_blocks, footer_blocks, page_number_blocks, body_blocks = split_page_blocks_fn(page.blocks)
    if not body_blocks:
        return None

    prepare_document_section_fn(
        document,
        page,
        render_config,
        rendered_page_idx,
        header_blocks,
        footer_blocks,
        page_number_blocks,
        previous_rendered_page=previous_rendered_page,
    )
    current_is_toc = resolve_toc_page_state_fn(body_blocks, toc_mode_active)
    render_document_page_body_fn(
        document,
        body_blocks,
        render_config,
        current_is_toc,
        rendered_page_idx == 0,
    )
    return current_is_toc


def prepare_document_section(
    document: Document,
    page: Any,
    render_config: Any,
    rendered_page_idx: int,
    header_blocks: list[Any],
    footer_blocks: list[Any],
    page_number_blocks: list[Any],
    previous_rendered_page: Any | None = None,
) -> Any:
    configure_section_fn = _runtime_value("configure_section", configure_section)
    render_header_fn = _runtime_value("render_header", render_header)
    render_footer_fn = _runtime_value("render_footer", render_footer)
    should_start_new_section_fn = _runtime_value("should_start_new_section", should_start_new_section)

    start_new_section = should_start_new_section_fn(
        page,
        previous_rendered_page,
        render_config,
        header_blocks,
        footer_blocks,
        page_number_blocks,
    )
    section = (
        document.sections[-1]
        if rendered_page_idx == 0 or not start_new_section
        else document.add_section(WD_SECTION_START.NEW_PAGE)
    )

    if rendered_page_idx == 0 or start_new_section:
        configure_section_fn(section, page.page_size, render_config.section_layout)

    if rendered_page_idx > 0 and start_new_section:
        if render_config.use_mineru_header_footer:
            section.header.is_linked_to_previous = False
            section.footer.is_linked_to_previous = False
        else:
            section.header.is_linked_to_previous = True
            section.footer.is_linked_to_previous = True

    if render_config.use_mineru_header_footer and (rendered_page_idx == 0 or start_new_section):
        render_header_fn(section, header_blocks)
        render_footer_fn(section, footer_blocks, page_number_blocks)
    return section


def is_landscape_page(page_size: tuple[float, float] | None) -> bool:
    return bool(page_size and page_size[0] > page_size[1])


def should_start_new_section(
    page: Any,
    previous_rendered_page: Any | None,
    render_config: Any,
    header_blocks: list[Any],
    footer_blocks: list[Any],
    page_number_blocks: list[Any],
) -> bool:
    split_page_blocks_fn = _runtime_value("split_page_blocks", split_page_blocks)
    is_toc_page_fn = _runtime_value("is_toc_page", is_toc_page)
    is_toc_continuation_page_fn = _runtime_value("is_toc_continuation_page", is_toc_continuation_page)
    if previous_rendered_page is None:
        return False

    current_body_blocks = split_page_blocks_fn(page.blocks)[3]
    previous_body_blocks = split_page_blocks_fn(previous_rendered_page.blocks)[3]
    previous_is_toc = is_toc_page_fn(previous_body_blocks)
    current_is_toc = is_toc_page_fn(current_body_blocks) or (
        previous_is_toc and is_toc_continuation_page_fn(current_body_blocks)
    )

    if current_is_toc and not previous_is_toc:
        return True

    if not current_is_toc and any(
        block.type == "title" and _appendix_title_re().match((block.text or "").strip())
        for block in current_body_blocks
    ):
        return True

    should_start_new_page_before_dense_section_fn = _runtime_value(
        "should_start_new_page_before_dense_section",
        should_start_new_page_before_dense_section,
    )
    if should_start_new_page_before_dense_section_fn(
        current_body_blocks,
        previous_body_blocks,
        page,
        previous_rendered_page,
    ):
        return True

    if previous_is_toc and not current_is_toc:
        return True

    if is_landscape_page(page.page_size) != is_landscape_page(previous_rendered_page.page_size):
        return True

    if render_config.use_mineru_header_footer and (header_blocks or footer_blocks or page_number_blocks):
        return True

    return False


def should_start_new_page_before_dense_section(
    current_body_blocks: list[Any],
    previous_body_blocks: list[Any],
    current_page: Any,
    previous_page: Any,
) -> bool:
    """Keep a new top-level section from being squeezed after a dense visual page."""
    if not current_body_blocks or not previous_body_blocks:
        return False

    current_first_block = next(
        (block for block in sorted(current_body_blocks, key=lambda item: (item.y1, item.x1))),
        None,
    )
    if current_first_block is None or current_first_block.type != "title":
        return False

    title_text = (current_first_block.text or "").strip()
    if not TOP_LEVEL_SECTION_TITLE_RE.match(title_text):
        return False

    current_page_height = current_page.page_size[1] if current_page.page_size else None
    if current_page_height and current_first_block.y1 > current_page_height * 0.25:
        return False

    previous_page_height = previous_page.page_size[1] if previous_page.page_size else None
    if not previous_page_height:
        return False

    previous_bottom = max((block.y2 for block in previous_body_blocks), default=0.0)
    if previous_bottom < previous_page_height * DENSE_PAGE_BOTTOM_RATIO:
        return False
    if previous_bottom > previous_page_height * DENSE_PAGE_OVERFLOW_RATIO:
        return False

    return any(block.type in _visual_types() for block in previous_body_blocks)


def resolve_toc_page_state(body_blocks: list[Any], toc_mode_active: bool) -> bool:
    is_toc_page_fn = _runtime_value("is_toc_page", is_toc_page)
    is_toc_continuation_page_fn = _runtime_value("is_toc_continuation_page", is_toc_continuation_page)
    current_is_toc = is_toc_page_fn(body_blocks)
    if current_is_toc:
        return True
    if toc_mode_active and is_toc_continuation_page_fn(body_blocks):
        return True
    return False


def render_document_page_body(
    document: Document,
    body_blocks: list[Any],
    render_config: Any,
    current_is_toc: bool,
    skip_product_title: bool,
) -> None:
    render_toc_page_fn = _runtime_value("render_toc_page", render_toc_page)
    render_page_body_fn = _runtime_value("render_page_body", render_page_body)
    is_toc_page_fn = _runtime_value("is_toc_page", is_toc_page)
    split_toc_page_blocks_fn = _runtime_value("split_toc_page_blocks", split_toc_page_blocks)
    if current_is_toc:
        if is_toc_page_fn(body_blocks):
            toc_blocks, remaining_blocks = split_toc_page_blocks_fn(body_blocks)
            render_toc_page_fn(document, toc_blocks, render_config, skip_product_title)
            if remaining_blocks:
                render_page_body_fn(document, remaining_blocks, render_config, False)
        else:
            render_toc_page_fn(document, body_blocks, render_config, skip_product_title)
    else:
        render_page_body_fn(document, body_blocks, render_config, skip_product_title)


def configure_section(section: Any, page_size: tuple[float, float] | None, layout: Any) -> None:
    landscape = bool(page_size and page_size[0] > page_size[1])
    section.orientation = WD_ORIENTATION.LANDSCAPE if landscape else WD_ORIENTATION.PORTRAIT
    if page_size:
        section.page_width = Pt(page_size[0])
        section.page_height = Pt(page_size[1])
    else:
        section.page_width = _runtime_value("A4_HEIGHT", None) if landscape else _runtime_value("A4_WIDTH", None)
        section.page_height = _runtime_value("A4_WIDTH", None) if landscape else _runtime_value("A4_HEIGHT", None)
    section.left_margin = layout.left_margin
    section.right_margin = layout.right_margin
    section.top_margin = layout.top_margin
    section.bottom_margin = layout.bottom_margin
    section.header_distance = layout.header_distance
    section.footer_distance = layout.footer_distance


def split_page_blocks(page_blocks: list[Any]) -> tuple[list[Any], list[Any], list[Any], list[Any]]:
    header_blocks = [block for block in page_blocks if block.type == "page_header" and block.text]
    footer_blocks = [block for block in page_blocks if block.type == "page_footer" and block.text]
    page_number_blocks = [block for block in page_blocks if block.type == "page_number" and block.text]
    body_blocks = [
        block
        for block in page_blocks
        if block.type not in {"page_header", "page_footer", "page_number"}
    ]
    return header_blocks, footer_blocks, page_number_blocks, body_blocks


def extract_document_identity(pages: list[Any]) -> tuple[str, str]:
    if not pages:
        return "", ""

    top_titles = [
        block.text.strip()
        for block in sorted(pages[0].blocks, key=lambda block: (block.y1, block.x1))
        if block.type == "title" and block.text.strip()
    ]
    if not top_titles:
        return "", ""

    first_title = top_titles[0]
    match = _product_title_re().match(first_title)
    if match:
        return match.group("part").strip(), match.group("category").strip()

    return "", ""


def fill_template_intro(document: Document, part_number: str, category: str, render_config: Any) -> None:
    apply_named_style_config_fn = _runtime_value("apply_named_style_config", lambda *_args, **_kwargs: None)
    fill_rich_text_fn = _runtime_value("fill_rich_text", lambda *_args, **_kwargs: None)
    if part_number:
        paragraph = document.add_paragraph()
        apply_named_style_config_fn(paragraph, render_config, "PartNumber")
        fill_rich_text_fn(paragraph, part_number, render_config=render_config, style_name="PartNumber")

    if category:
        paragraph = document.add_paragraph()
        apply_named_style_config_fn(paragraph, render_config, "Category")
        fill_rich_text_fn(paragraph, category, render_config=render_config, style_name="Category")


def is_primary_product_title(block: Any, page_width: float) -> bool:
    """Only skip first-page titles that match the product-name intro pattern."""
    text = (block.text or "").strip()
    return bool(text and _product_title_re().match(text) and block.y1 < 220 and block.width >= page_width * 0.22)


def render_header(section: Any, header_blocks: list[Any]) -> None:
    render_positioned_lines_fn = _runtime_value("render_positioned_lines", render_positioned_lines)
    usable_width = section.page_width - section.left_margin - section.right_margin
    render_positioned_lines_fn(section.header, header_blocks, usable_width, _runtime_value("HEADER_FOOTER_FONT", Pt(8.5)))


def render_footer(section: Any, footer_blocks: list[Any], page_number_blocks: list[Any]) -> None:
    render_positioned_lines_fn = _runtime_value("render_positioned_lines", render_positioned_lines)
    usable_width = section.page_width - section.left_margin - section.right_margin
    render_positioned_lines_fn(
        section.footer,
        sorted(footer_blocks + page_number_blocks, key=lambda block: (block.y1, block.x1)),
        usable_width,
        _runtime_value("HEADER_FOOTER_FONT", Pt(8.5)),
    )


def render_positioned_lines(container: Any, blocks: list[Any], usable_width: int, font_size: Pt) -> None:
    group_rows_fn = _runtime_value("group_rows", group_rows)
    split_row_text_fn = _runtime_value("split_row_text", split_row_text)
    rows = group_rows_fn(blocks)
    if not rows:
        container.paragraphs[0].text = ""
        return

    for paragraph in list(container.paragraphs)[1:]:
        paragraph._element.getparent().remove(paragraph._element)

    base_paragraph = container.paragraphs[0]
    for row_idx, row in enumerate(rows):
        paragraph = base_paragraph if row_idx == 0 else container.add_paragraph()
        paragraph.text = ""
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.tab_stops.add_tab_stop(int(usable_width // 2), WD_TAB_ALIGNMENT.CENTER)
        paragraph.paragraph_format.tab_stops.add_tab_stop(int(usable_width), WD_TAB_ALIGNMENT.RIGHT)

        left, center, right = split_row_text_fn(row)
        line = left
        if center or right:
            line += f"\t{center}"
        if right:
            line += f"\t{right}"

        run = paragraph.add_run(line.strip())
        run.font.size = font_size


def group_rows(blocks: list[Any], tolerance: float = 14.0) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for block in sorted(blocks, key=lambda current: (current.y1, current.x1)):
        if not rows or abs(block.y1 - rows[-1][0].y1) > tolerance:
            rows.append([block])
        else:
            rows[-1].append(block)
    return rows


def split_row_text(row: list[Any]) -> tuple[str, str, str]:
    if not row:
        return "", "", ""

    page_width = max(block.x2 for block in row) or 1000.0
    left_parts: list[str] = []
    center_parts: list[str] = []
    right_parts: list[str] = []

    for block in sorted(row, key=lambda current: current.x1):
        ratio = block.center_x / page_width
        if ratio < 0.34:
            left_parts.append(block.text)
        elif ratio > 0.66:
            right_parts.append(block.text)
        else:
            center_parts.append(block.text)

    return " ".join(left_parts).strip(), " ".join(center_parts).strip(), " ".join(right_parts).strip()


def render_page_body(
    document: Document,
    body_blocks: list[Any],
    render_config: Any,
    skip_product_title: bool,
) -> None:
    is_toc_page_fn = _runtime_value("is_toc_page", is_toc_page)
    render_toc_page_fn = _runtime_value("render_toc_page", render_toc_page)
    prepare_body_render_plan_fn = _runtime_value("prepare_body_render_plan", prepare_body_render_plan)
    render_two_col_region_fn = _runtime_value("render_two_col_region", render_two_col_region)
    render_one_col_region_fn = _runtime_value("render_one_col_region", render_one_col_region)
    split_toc_page_blocks_fn = _runtime_value("split_toc_page_blocks", split_toc_page_blocks)

    if not body_blocks:
        return

    page_idx = body_blocks[0].page_idx if body_blocks else None
    _logger().debug("render page body page={} body_blocks={} skip_product_title={}", page_idx, len(body_blocks), skip_product_title)

    if is_toc_page_fn(body_blocks):
        toc_blocks, remaining_blocks = split_toc_page_blocks_fn(body_blocks)
        render_toc_page_fn(document, toc_blocks, render_config, skip_product_title)
        body_blocks = remaining_blocks
        skip_product_title = False
        if not body_blocks:
            return

    body_render_plan = prepare_body_render_plan_fn(document, body_blocks, skip_product_title)
    if body_render_plan is None:
        return

    page_width, usable_width, regions = body_render_plan

    for region in regions:
        if region.kind == "two_col":
            render_two_col_region_fn(document, region.blocks, page_width, usable_width, render_config)
        else:
            render_one_col_region_fn(document, region.blocks, page_width, usable_width, render_config)


def prepare_body_render_plan(
    document: Document,
    body_blocks: list[Any],
    skip_product_title: bool,
) -> tuple[float, int, list[Any]] | None:
    is_primary_product_title_fn = _runtime_value("is_primary_product_title", is_primary_product_title)
    build_regions_fn = _runtime_value("build_regions", build_regions)

    page_width = max(block.x2 for block in body_blocks) or 1000.0
    filtered_blocks = [
        block
        for block in body_blocks
        if not (skip_product_title and block.type == "title" and is_primary_product_title_fn(block, page_width))
    ]
    if not filtered_blocks:
        return None

    regions = build_regions_fn(filtered_blocks, page_width)
    section = document.sections[-1]
    usable_width = section.page_width - section.left_margin - section.right_margin
    return page_width, usable_width, regions


def is_toc_page(body_blocks: list[Any]) -> bool:
    count_toc_like_entries_fn = _runtime_value("count_toc_like_entries", count_toc_like_entries)
    normalize_toc_text_fn = _runtime_value("normalize_toc_text", lambda text: text)
    toc_title_re = _toc_title_re()

    if not body_blocks:
        return False

    for block in body_blocks:
        if block.type != "title":
            continue
        if toc_title_re.match(normalize_toc_text_fn(block.text)):
            return True

    text_like_blocks, toc_like_blocks = count_toc_like_entries_fn(body_blocks)
    return text_like_blocks >= 6 and toc_like_blocks / max(text_like_blocks, 1) >= 0.75


def looks_like_toc_entry_text(text: str) -> bool:
    toc_title_re = _toc_title_re()
    if toc_title_re.match(text):
        return True
    return bool(_toc_entry_re().match(text))


def is_toc_continuation_page(body_blocks: list[Any]) -> bool:
    count_toc_like_entries_fn = _runtime_value("count_toc_like_entries", count_toc_like_entries)
    if not body_blocks:
        return False
    if any(block.type in _visual_types() for block in body_blocks):
        return False

    text_like, toc_like = count_toc_like_entries_fn(body_blocks)
    return text_like >= 4 and toc_like / max(text_like, 1) >= 0.7


def count_toc_like_entries(body_blocks: list[Any]) -> tuple[int, int]:
    normalize_toc_text_fn = _runtime_value("normalize_toc_text", lambda text: text)
    looks_like_toc_entry_text_fn = _runtime_value("looks_like_toc_entry_text", looks_like_toc_entry_text)
    text_like = 0
    toc_like = 0
    for block in body_blocks:
        texts: list[str] = []
        if block.type == "list":
            texts = block.items
        elif block.type in {"title", "paragraph"}:
            texts = [block.text]
        for text in texts:
            normalized = normalize_toc_text_fn(text)
            if not normalized:
                continue
            text_like += 1
            if looks_like_toc_entry_text_fn(normalized):
                toc_like += 1
    return text_like, toc_like


def split_toc_page_blocks(body_blocks: list[Any]) -> tuple[list[Any], list[Any]]:
    """Split mixed TOC pages so later body content keeps normal semantics."""
    normalize_toc_text_fn = _runtime_value("normalize_toc_text", lambda text: text)
    sorted_blocks = sorted(body_blocks, key=lambda current: (current.y1, current.x1))
    seen_toc_title = False
    for idx, block in enumerate(sorted_blocks):
        if block.type == "title" and _toc_title_re().match(normalize_toc_text_fn(block.text)):
            seen_toc_title = True
            continue
        if seen_toc_title and block.type == "title" and block.y1 > 360:
            return sorted_blocks[:idx], sorted_blocks[idx:]
    return sorted_blocks, []


def render_toc_page(
    document: Document,
    body_blocks: list[Any],
    render_config: Any,
    skip_product_title: bool,
) -> None:
    normalize_toc_text_fn = _runtime_value("normalize_toc_text", lambda text: text)
    is_cover_title_fn = _runtime_value("is_cover_title", lambda *_args, **_kwargs: False)
    fill_toc_title_paragraph_fn = _runtime_value("fill_toc_title_paragraph", lambda *_args, **_kwargs: None)
    fill_toc_entry_paragraph_fn = _runtime_value("fill_toc_entry_paragraph", lambda *_args, **_kwargs: None)
    render_block_into_container_fn = _runtime_value("render_block_into_container", render_block_into_container)

    page_width = max(block.x2 for block in body_blocks) or 1000.0
    usable_width = document.sections[-1].page_width - document.sections[-1].left_margin - document.sections[-1].right_margin

    for block in sorted(body_blocks, key=lambda current: (current.y1, current.x1)):
        if block.type == "title" and is_cover_title_fn(block, normalize_toc_text_fn(block.text), page_width):
            continue

        if block.type == "title":
            toc_text = normalize_toc_text_fn(block.text)
            paragraph = document.add_paragraph()
            if _toc_title_re().match(toc_text):
                fill_toc_title_paragraph_fn(paragraph, toc_text, render_config)
            else:
                fill_toc_entry_paragraph_fn(paragraph, toc_text, render_config, usable_width)
            continue

        if block.type == "list":
            for item in block.items:
                paragraph = document.add_paragraph()
                fill_toc_entry_paragraph_fn(paragraph, item, render_config, usable_width)
            continue

        if block.type == "paragraph":
            paragraph = document.add_paragraph()
            fill_toc_entry_paragraph_fn(paragraph, block.text, render_config, usable_width)
            continue

        render_block_into_container_fn(document, block, page_width, usable_width, render_config)


def build_regions(body_blocks: list[Any], page_width: float) -> list[Any]:
    detect_two_col_bands_fn = _runtime_value("detect_two_col_bands", detect_two_col_bands)
    classify_body_block_fn = _runtime_value("classify_body_block", classify_body_block)
    region_cls = _region_cls()
    two_col_bands = detect_two_col_bands_fn(body_blocks, page_width)
    sorted_blocks = sorted(body_blocks, key=lambda block: (block.y1, block.x1))

    regions: list[Any] = []
    for block in sorted_blocks:
        role = classify_body_block_fn(block, page_width)
        kind = "two_col" if role != "wide" and any(start <= block.mid_y <= end for start, end in two_col_bands) else "one_col"
        if regions and regions[-1].kind == kind:
            regions[-1].blocks.append(block)
        else:
            regions.append(region_cls(kind=kind, blocks=[block]))
    return regions


def detect_two_col_bands(body_blocks: list[Any], page_width: float) -> list[tuple[float, float]]:
    classify_body_block_fn = _runtime_value("classify_body_block", classify_body_block)
    detect_stable_two_col_band_fn = _runtime_value("detect_stable_two_col_band", detect_stable_two_col_band)
    y_points = sorted({block.y1 for block in body_blocks} | {block.y2 for block in body_blocks})
    intervals: list[tuple[float, float, str]] = []

    for top, bottom in zip(y_points, y_points[1:]):
        if bottom - top < 4:
            continue
        active = [block for block in body_blocks if block.y1 < bottom and block.y2 > top]
        if not active:
            continue

        roles = {classify_body_block_fn(block, page_width) for block in active}
        label = "two_col" if "left" in roles and "right" in roles and "wide" not in roles else "one_col"
        intervals.append((top, bottom, label))

    merged: list[tuple[float, float, str]] = []
    for top, bottom, label in intervals:
        if merged and merged[-1][2] == label and top - merged[-1][1] <= 12:
            merged[-1] = (merged[-1][0], bottom, label)
        else:
            merged.append((top, bottom, label))

    bands = [(top, bottom) for top, bottom, label in merged if label == "two_col"]
    stable_band = detect_stable_two_col_band_fn(body_blocks, page_width)
    if stable_band is not None:
        bands.append(stable_band)
    return merge_two_col_bands(bands)


def detect_stable_two_col_band(body_blocks: list[Any], page_width: float) -> tuple[float, float] | None:
    """Detect separated source columns even when their y-ranges do not overlap continuously."""
    classify_body_block_fn = _runtime_value("classify_body_block", classify_body_block)
    left_blocks = [block for block in body_blocks if classify_body_block_fn(block, page_width) == "left"]
    right_blocks = [block for block in body_blocks if classify_body_block_fn(block, page_width) == "right"]
    if len(left_blocks) < 2 or len(right_blocks) < 2:
        return None

    top = min(block.y1 for block in [*left_blocks, *right_blocks])
    bottom = max(block.y2 for block in [*left_blocks, *right_blocks])
    if bottom - top < 120:
        return None

    return top, bottom


def merge_two_col_bands(bands: list[tuple[float, float]]) -> list[tuple[float, float]]:
    if not bands:
        return []
    merged: list[tuple[float, float]] = []
    for top, bottom in sorted(bands):
        if merged and top - merged[-1][1] <= 12:
            merged[-1] = (merged[-1][0], max(merged[-1][1], bottom))
        else:
            merged.append((top, bottom))
    return merged


def classify_body_block(block: Any, page_width: float) -> str:
    if block.width >= page_width * 0.6:
        return "wide"
    return "left" if block.center_x < page_width / 2 else "right"


def render_one_col_region(
    document: Document,
    blocks: list[Any],
    page_width: float,
    usable_width: int,
    render_config: Any,
) -> None:
    render_block_into_container_fn = _runtime_value("render_block_into_container", render_block_into_container)
    for idx, block in enumerate(blocks):
        next_block = blocks[idx + 1] if idx + 1 < len(blocks) else None
        render_block_into_container_fn(document, block, page_width, usable_width, render_config, next_block=next_block)


def render_two_col_region(
    document: Document,
    blocks: list[Any],
    page_width: float,
    usable_width: int,
    render_config: Any,
) -> None:
    classify_body_block_fn = _runtime_value("classify_body_block", classify_body_block)
    render_blocks_in_cell_fn = _runtime_value("render_blocks_in_cell", render_blocks_in_cell)
    clear_table_borders_fn = _runtime_value("clear_table_borders", clear_table_borders)
    left_blocks: list[Any] = []
    right_blocks: list[Any] = []

    for block in sorted(blocks, key=lambda current: (current.y1, current.x1)):
        role = classify_body_block_fn(block, page_width)
        if role == "right":
            right_blocks.append(block)
        else:
            left_blocks.append(block)

    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    clear_table_borders_fn(table)

    column_width = usable_width // 2
    content_width = column_width - Emu(int(Mm(4)))

    source_column_width = max(page_width / 2, 1.0)
    for idx, cell in enumerate(table.rows[0].cells):
        cell.width = column_width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cell.text = ""
        render_blocks_in_cell_fn(
            cell,
            left_blocks if idx == 0 else right_blocks,
            source_column_width,
            content_width,
            render_config,
        )


def clear_table_borders(table: Any) -> None:
    """Hide borders on layout-only tables used to hold source columns."""
    table_properties = table._tbl.tblPr
    borders = table_properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "none")
        element.set(qn("w:sz"), "0")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "auto")


def render_blocks_in_cell(
    cell: Any,
    blocks: list[Any],
    page_width: float,
    usable_width: int,
    render_config: Any,
) -> None:
    render_block_into_container_fn = _runtime_value("render_block_into_container", render_block_into_container)
    if not blocks:
        return

    first_paragraph = cell.paragraphs[0]
    first_used = False

    for idx, block in enumerate(blocks):
        target = first_paragraph if not first_used else None
        next_block = blocks[idx + 1] if idx + 1 < len(blocks) else None
        render_block_into_container_fn(
            cell,
            block,
            page_width,
            usable_width,
            render_config,
            first_paragraph=target,
            next_block=next_block,
        )
        first_used = True


def render_block_into_container(
    container: Any,
    block: Any,
    page_width: float,
    usable_width: int,
    render_config: Any,
    first_paragraph: Any | None = None,
    next_block: Any | None = None,
) -> None:
    render_visual_block_fn = _runtime_value("render_visual_block", render_visual_block)
    normalize_list_item_fn = _runtime_value("normalize_list_item", lambda text: text)
    infer_list_item_level_fn = _runtime_value("infer_list_item_level", infer_list_item_level)
    is_list_item_text_fn = _runtime_value("is_list_item_text", lambda _text: False)
    fill_list_paragraph_fn = _runtime_value("fill_list_paragraph", lambda *_args, **_kwargs: None)
    looks_like_heading_text_fn = _runtime_value("looks_like_heading_text", lambda _text: False)
    fill_title_paragraph_fn = _runtime_value("fill_title_paragraph", lambda *_args, **_kwargs: None)
    fill_body_paragraph_fn = _runtime_value("fill_body_paragraph", lambda *_args, **_kwargs: None)

    if block.type in _visual_types():
        render_visual_block_fn(container, block, page_width, usable_width, render_config, first_paragraph)
        return

    if block.type == "list":
        raw_items = [item for item in block.items if item]
        for idx, raw_item in enumerate(raw_items):
            paragraph = first_paragraph if idx == 0 and first_paragraph is not None else container.add_paragraph()
            fill_list_paragraph_fn(
                paragraph,
                normalize_list_item_fn(raw_item),
                render_config,
                level=infer_list_item_level_fn(raw_item),
            )
        return

    paragraph = first_paragraph if first_paragraph is not None else container.add_paragraph()
    if block.type == "paragraph" and is_list_item_text_fn(block.text):
        fill_list_paragraph_fn(
            paragraph,
            normalize_list_item_fn(block.text),
            render_config,
            level=infer_list_item_level_fn(block.text),
        )
    elif block.type == "title" or (block.type == "paragraph" and looks_like_heading_text_fn(block.text)):
        keep_with_next = not (next_block is not None and next_block.type in _visual_types())
        fill_title_paragraph_fn(paragraph, block, page_width, render_config, keep_with_next=keep_with_next)
    else:
        fill_body_paragraph_fn(paragraph, block.text, render_config)


def render_visual_block(
    container: Any,
    block: Any,
    page_width: float,
    usable_width: int,
    render_config: Any,
    first_paragraph: Any | None = None,
) -> None:
    render_table_from_html_fn = _runtime_value("render_table_from_html", lambda *_args, **_kwargs: False)
    apply_style_if_exists_fn = _runtime_value("apply_style_if_exists", lambda *_args, **_kwargs: None)
    compute_visual_width_fn = _runtime_value("compute_visual_width", lambda *_args, **_kwargs: usable_width)
    fill_body_paragraph_fn = _runtime_value("fill_body_paragraph", lambda *_args, **_kwargs: None)
    render_block_annotations_fn = _runtime_value("render_block_annotations", lambda *_args, **_kwargs: None)

    _logger().debug(
        "render visual block page={} block={} type={} bbox={} image_path={} has_html={} caption={}",
        block.page_idx,
        block.block_idx,
        block.type,
        block.bbox,
        block.image_path,
        bool(block.html),
        block.captions[0] if block.captions else "N/A",
    )
    if block.type == "table" and block.html:
        rendered = render_table_from_html_fn(container, block, usable_width, render_config)
        if rendered:
            return

    paragraph = first_paragraph if first_paragraph is not None else container.add_paragraph()
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = bool(block.captions or block.footnotes)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    apply_style_if_exists_fn(paragraph, "Normal")

    if block.image_path and block.image_path.exists():
        picture_width = compute_visual_width_fn(block, page_width, usable_width)
        paragraph.add_run().add_picture(str(block.image_path), width=Emu(picture_width))
    else:
        _logger().warning(
            "missing visual asset page={} block={} type={} bbox={} caption={} image_path={}",
            block.page_idx,
            block.block_idx,
            block.type,
            block.bbox,
            block.captions[0] if block.captions else "N/A",
            block.image_path,
        )
        fill_body_paragraph_fn(paragraph, block.text or f"[{block.type} missing image]", render_config)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    render_block_annotations_fn(container, block, render_config)


def infer_list_item_level(text: str) -> int:
    """Fallback list-level inference used when the CLI compatibility runtime is not bound."""
    return 1 if re.match(r"^\s*[-–—]\s+", text or "") else 0

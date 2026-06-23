from __future__ import annotations

# app/conversion/runtime/table_render_runtime.py: 渲染 MinerU HTML 表格、caption 和 footnote 到 DOCX。

import re
from typing import Any

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt


SECTION_CAPTION_RE = re.compile(r"^\s*(?P<num>\d+(?:\.\d+)+)\s+\S")


def render_block_captions_impl(
    container: Any,
    block: Any,
    render_config: Any,
    *,
    apply_named_style_config_fn: Any,
    fill_rich_text_fn: Any,
    prepare_text_fn: Any,
    keep_with_next: bool = False,
) -> None:
    for caption in block.captions:
        prepared = prepare_text_fn(caption, render_config)
        style_name = resolve_caption_style_name(prepared)
        caption_paragraph = container.add_paragraph()
        caption_paragraph.paragraph_format.keep_with_next = keep_with_next
        apply_named_style_config_fn(caption_paragraph, render_config, style_name)
        fill_rich_text_fn(
            caption_paragraph,
            prepared,
            render_config=render_config,
            style_name=style_name,
        )


def resolve_caption_style_name(caption: str) -> str:
    """Render numbered section-like table captions as headings, not centered captions."""
    match = SECTION_CAPTION_RE.match(caption or "")
    if match is None:
        return "Caption"
    level = min(match.group("num").count(".") + 1, 3)
    return f"Heading {level}"


def render_block_footnotes_impl(
    container: Any,
    block: Any,
    render_config: Any,
    *,
    apply_named_style_config_fn: Any,
    fill_rich_text_fn: Any,
    prepare_text_fn: Any,
) -> None:
    for footnote in block.footnotes:
        note_paragraph = container.add_paragraph()
        apply_named_style_config_fn(note_paragraph, render_config, "TableNote")
        fill_rich_text_fn(
            note_paragraph,
            prepare_text_fn(footnote, render_config),
            render_config=render_config,
            style_name="TableNote",
        )


def render_block_annotations_impl(
    container: Any,
    block: Any,
    render_config: Any,
    *,
    apply_named_style_config_fn: Any,
    fill_rich_text_fn: Any,
    prepare_text_fn: Any,
) -> None:
    render_block_captions_impl(
        container,
        block,
        render_config,
        apply_named_style_config_fn=apply_named_style_config_fn,
        fill_rich_text_fn=fill_rich_text_fn,
        prepare_text_fn=prepare_text_fn,
    )
    render_block_footnotes_impl(
        container,
        block,
        render_config,
        apply_named_style_config_fn=apply_named_style_config_fn,
        fill_rich_text_fn=fill_rich_text_fn,
        prepare_text_fn=prepare_text_fn,
    )


def render_large_table_if_needed_impl(
    container: Any,
    block: Any,
    table_model: Any,
    max_cols: int,
    html_len: int,
    render_config: Any,
    usable_width: int,
    *,
    max_complex_rows: int,
    max_complex_html_len: int,
    is_long_simple_table_fn: Any,
    render_long_simple_table_fn: Any,
    render_large_table_fallback_fn: Any,
    logger: Any,
    estimated_page_count: int | None = None,
    long_table_page_threshold: int = 2,
) -> bool:
    if estimated_page_count is None:
        if len(table_model.rows) <= max_complex_rows and html_len <= max_complex_html_len:
            return False
    elif estimated_page_count <= long_table_page_threshold:
        return False

    logger.warning(
        "large table fallback page={} block={} rows={} max_cols={} html_len={} estimated_pages={} caption={}",
        block.page_idx,
        block.block_idx,
        len(table_model.rows),
        max_cols,
        html_len,
        estimated_page_count if estimated_page_count is not None else "unknown",
        block.captions[0] if block.captions else "N/A",
    )
    if is_long_simple_table_fn(table_model, 0 if estimated_page_count is not None else max_complex_rows):
        logger.info(
            "long simple table chunk rendering page={} block={} rows={} max_cols={} estimated_pages={} caption={}",
            block.page_idx,
            block.block_idx,
            len(table_model.rows),
            max_cols,
            estimated_page_count if estimated_page_count is not None else "unknown",
            block.captions[0] if block.captions else "N/A",
        )
        render_long_simple_table_fn(container, block, table_model, render_config, usable_width)
        return True

    render_large_table_fallback_fn(container, block, table_model, render_config, usable_width)
    return True


def render_standard_html_table_impl(
    container: Any,
    table_model: Any,
    max_cols: int,
    usable_width: int,
    render_config: Any,
    *,
    remove_table_borders_fn: Any,
    fill_table_cell_fn: Any,
) -> int:
    table = container.add_table(rows=len(table_model.rows), cols=max_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders_fn(table, render_config)
    configure_table_rows(table, repeat_header=True)

    column_widths = compute_table_column_widths(max_cols, usable_width)
    for col_idx, col in enumerate(table.columns):
        col.width = column_widths[col_idx]

    occupied: dict[tuple[int, int], bool] = {}
    merge_failures = 0
    compact = max_cols >= 4 or len(table_model.rows) >= 10
    for row_idx, row in enumerate(table_model.rows):
        col_idx = 0
        for cell_model in row:
            while occupied.get((row_idx, col_idx)):
                col_idx += 1
            if col_idx >= max_cols:
                break

            base_cell = table.cell(row_idx, col_idx)
            end_row = min(row_idx + cell_model.rowspan - 1, len(table_model.rows) - 1)
            end_col = min(col_idx + cell_model.colspan - 1, max_cols - 1)
            merged_cell = base_cell
            if end_row != row_idx or end_col != col_idx:
                try:
                    target_cell = table.cell(end_row, end_col)
                    merged_cell = base_cell.merge(target_cell)
                except ValueError:
                    merge_failures += 1

            merged_cell.width = sum(column_widths[col_idx : end_col + 1])
            fill_table_cell_fn(
                merged_cell,
                cell_model,
                row_idx == 0 or bool(getattr(cell_model, "is_header", False)),
                render_config,
                compact=compact,
            )
            merged_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            for r in range(row_idx, end_row + 1):
                for c in range(col_idx, end_col + 1):
                    occupied[(r, c)] = True

            col_idx = end_col + 1

    return merge_failures


def compute_table_column_widths(max_cols: int, usable_width: int) -> list[int]:
    """Give 5-column datasheet tables a wider description column to avoid row inflation."""
    if max_cols == 5:
        weights = [0.11, 0.37, 0.21, 0.21, 0.10]
        widths = [int(usable_width * weight) for weight in weights]
        widths[1] += int(usable_width) - sum(widths)
        return widths

    column_width = max(int(usable_width // max_cols), int(Emu(Mm(16))))
    return [column_width for _ in range(max_cols)]


def expand_table_rows(table_model: Any) -> list[list[str]]:
    expanded: list[list[str]] = []
    spans: dict[int, tuple[int, str]] = {}
    max_cols = table_model.max_columns

    for row in table_model.rows:
        pending_cells = list(row)
        row_values: list[str] = []
        col_idx = 0

        while col_idx < max_cols:
            active = spans.get(col_idx)
            if active and active[0] > 0:
                row_values.append(active[1])
                spans[col_idx] = (active[0] - 1, active[1])
                col_idx += 1
                continue

            if active and active[0] <= 0:
                spans.pop(col_idx, None)

            if not pending_cells:
                row_values.append("")
                col_idx += 1
                continue

            cell = pending_cells.pop(0)
            cell_text = cell.text
            for offset in range(cell.colspan):
                row_values.append(cell_text if offset == 0 else "")
                if cell.rowspan > 1:
                    spans[col_idx + offset] = (cell.rowspan - 2, cell_text if offset == 0 else "")
            col_idx += cell.colspan

        expanded.append(row_values[:max_cols])

    return expanded


def render_large_table_fallback_impl(
    container: Any,
    block: Any,
    table_model: Any,
    render_config: Any,
    usable_width: int,
    *,
    compute_visual_width_fn: Any,
    remove_table_borders_fn: Any,
    fill_table_cell_fn: Any,
    render_block_captions_fn: Any,
    render_block_footnotes_fn: Any,
    html_cell_cls: Any,
) -> None:
    render_block_captions_fn(container, block, render_config, keep_with_next=True)
    if block.image_path and block.image_path.exists():
        picture_paragraph = container.add_paragraph()
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.paragraph_format.space_before = Pt(2)
        picture_paragraph.paragraph_format.space_after = Pt(2)
        picture_paragraph.add_run().add_picture(str(block.image_path), width=Emu(min(int(usable_width), compute_visual_width_fn(block, 1000.0, usable_width))))
    else:
        expanded_rows = expand_table_rows(table_model)
        if not expanded_rows:
            return

        max_cols = max((len(row) for row in expanded_rows), default=0)
        if max_cols <= 0:
            return

        table = container.add_table(rows=len(expanded_rows), cols=max_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        remove_table_borders_fn(table, render_config)
        configure_table_rows(table, repeat_header=True)

        for row_idx, row_values in enumerate(expanded_rows):
            for col_idx in range(max_cols):
                cell = table.cell(row_idx, col_idx)
                text = row_values[col_idx] if col_idx < len(row_values) else ""
                fill_table_cell_fn(
                    cell,
                    html_cell_cls(text=text, rowspan=1, colspan=1, html=text),
                    row_idx == 0,
                    render_config,
                )
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    render_block_footnotes_fn(container, block, render_config)


def configure_table_rows(table: Any, *, repeat_header: bool) -> None:
    for row_idx, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        ensure_row_flag(row_properties, "w:cantSplit")
        if repeat_header and row_idx == 0:
            ensure_row_flag(row_properties, "w:tblHeader", value="true")


def ensure_row_flag(row_properties: Any, tag: str, *, value: str | None = None) -> None:
    element = row_properties.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        row_properties.append(element)
    if value is not None:
        element.set(qn("w:val"), value)


def fill_table_cell_impl(
    cell: Any,
    cell_model: Any,
    is_header: bool,
    render_config: Any,
    *,
    apply_named_style_config_fn: Any,
    render_inline_html_into_paragraph_fn: Any,
    prepare_text_fn: Any,
    compact: bool = False,
) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.text = ""
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.LEFT
    apply_named_style_config_fn(paragraph, render_config, "TableContent")
    header_bold = is_header and should_bold_table_header(render_config)
    render_inline_html_into_paragraph_fn(
        paragraph,
        cell_model.html or prepare_text_fn(cell_model.text, None),
        render_config=render_config,
        style_name="TableContent",
        bold=header_bold,
        table_header_bold=header_bold,
    )
    if compact:
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        paragraph.paragraph_format.line_spacing = 0.92
        for run in paragraph.runs:
            run.font.size = Pt(8)


def should_bold_table_header(render_config: Any) -> bool:
    """Read table header bolding from template config, preserving bold by default."""
    table_style = getattr(render_config, "table_style", None)
    if table_style is None:
        return True
    return bool(getattr(table_style, "header_bold", True))

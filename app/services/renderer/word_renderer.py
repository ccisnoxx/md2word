from __future__ import annotations

# app/services/renderer/word_renderer.py: 将 DocumentIR v1/v2 按模板配置渲染为可编辑 DOCX。

import io
import hashlib
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from types import SimpleNamespace
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENTATION, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.image.exceptions import InvalidImageStreamError, UnexpectedEndOfFileError, UnrecognizedImageError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Twips

from app.conversion.docx_rules.terminology import TerminologyReplacer
from app.services.document_ir.models import DocumentIR, DocumentIRAsset, DocumentIRBlock
from app.services.renderer.image_renderer import WordImageRendererMixin
from app.services.renderer.image_evidence import caption_bookmark_inventory, new_image_report, finish_image_report
from app.services.renderer.layout_plan import build_layout_plan, build_image_plan
from app.services.renderer.render_plan import ResourceIdentity, build_render_plan, content_hash
from app.services.renderer.block_evidence import RenderCollector, WordRenderResult, ItemEvidence, paragraph_target
from app.services.renderer.block_anchors import bookmark_inventory, materialize_block_anchors
from app.services.renderer.table_renderer import WordTableRendererMixin, prepare_table_input
from app.services.renderer.template_config import TemplateConfigPackage, TemplateParagraphStyleConfig
from app.services.renderer.text_renderer import WordTextRendererMixin
from app.services.renderer.logical_table_plan import build_content_table_plan
from app.services.renderer.content_contract import CONTENT_VERSION, build_formula_source_coverage, formula_sources, content_failure
from app.services.document_ir.table_placeholder import decode_placeholders


ImageContentLoader = Callable[[DocumentIRAsset], bytes | None]
A4_PORTRAIT_WIDTH = Mm(210)
A4_PORTRAIT_HEIGHT = Mm(297)


@dataclass(frozen=True, slots=True)
class DocumentIRRenderOptions:
    """document_jobs 正式管线为单次渲染提供的选项。"""

    subscript_config_path: Path | None = None
    document_language: str | None = None
    template_version_id: str | None = None


@dataclass(frozen=True, slots=True)
class DocumentIRRenderResult:
    """同一次渲染生成的 DOCX 与私有公式、图片报告。"""

    content: bytes = field(repr=False)
    formula_report: dict[str, Any]
    image_layout_report: dict[str, Any] | None = None


@dataclass(slots=True)
class _RenderContext:
    """单次渲染的字段目标、图片计划和报告，不写入可并发复用的实例。"""

    caption_bookmarks: dict[tuple[str, int], tuple[str, int]]
    collector: RenderCollector | None = None
    formula_sources: list[dict[str, Any]] = field(default_factory=list)
    formula_entries: list[dict[str, Any]] = field(default_factory=list)
    used_formula_sources: set[int] = field(default_factory=set)
    formula_location: tuple[str, str] | None = None
    image_report: dict[str, Any] = field(default_factory=dict)
    image_spans: dict[str, Any] = field(default_factory=dict)
    image_blocks: dict[str, DocumentIRBlock] = field(default_factory=dict)
    image_records: dict[str, tuple[dict[str, Any], dict[str, Any]]] = field(default_factory=dict)
    consumed_images: set[str] = field(default_factory=set)
    image_bookmark_ids: dict[str, int] = field(default_factory=dict)
    source_blocks: list[DocumentIRBlock] = field(default_factory=list)
    fragment_markers: list[tuple[str, Any]] = field(default_factory=list)
    placeholders: list[Any] = field(default_factory=list)
    assets: dict[str, Any] = field(default_factory=dict)
    image_loader: Any = None


class DocumentIRWordRenderer(WordTextRendererMixin, WordTableRendererMixin, WordImageRendererMixin):
    """P0 Word 渲染器：覆盖标题、段落、基础表格和图片。"""

    def render_to_bytes(
        self,
        document_ir: DocumentIR,
        *,
        template_path: Path | None,
        config_path: Path | None,
        image_loader: ImageContentLoader,
        options: DocumentIRRenderOptions | None = None,
    ) -> bytes:
        """按模板 DOCX 与配置包生成 Word 二进制。"""

        return self.render_to_result(document_ir, template_path=template_path, config_path=config_path, image_loader=image_loader, options=options).content

    def render_to_result(
        self,
        document_ir: DocumentIR,
        *,
        template_path: Path | None,
        config_path: Path | None,
        image_loader: ImageContentLoader,
        options: DocumentIRRenderOptions | None = None,
    ) -> WordRenderResult:
        """以单次文档上下文生成字节和私有报告，兼容入口共用此核心。"""

        config_content = config_path.read_bytes() if config_path is not None else None
        template_content = template_path.read_bytes() if template_path is not None else None
        template_config = self._load_template_config(document_ir, config_path, options, content=config_content)
        from app.services.renderer.evidence import render_contract_failure
        logo_path = self._resolve_asset_path(template_config.header.logo_path, config_path)
        logo_content = logo_path.read_bytes() if template_config.header.enabled and logo_path is not None and logo_path.exists() else None
        from app.services.renderer.resource_identity import normalized_resource_identity
        resources = normalized_resource_identity(template_content=template_content, config=template_config,
            logo_content=logo_content, document_language=options.document_language if options and options.document_language else document_ir.metadata.document_language)
        layout_plan = build_layout_plan(document_ir)
        table_plan = build_content_table_plan(document_ir, layout_plan)
        grid_by_block = {fragment.source.block_id: fragment.render.grid for unit in table_plan.units
                        for fragment in unit.fragments if fragment.render is not None}
        tables = {block.id: (SimpleNamespace(rows=range(grid_by_block[block.id].rows),
                    max_columns=grid_by_block[block.id].columns), "standard") if block.id in grid_by_block
                    else prepare_table_input(block, template_config)
                  for block in document_ir.blocks if block.type == "table"}
        try:
            plan = build_render_plan(document_ir, layout_plan=layout_plan, effective_config=template_config,
                resource_identity=resources, table_models={key: value[0] for key, value in tables.items()},
                table_strategies={key: value[1] for key, value in tables.items()}, logical_table_plan=table_plan)
            collector = RenderCollector(document_ir, plan)
        except (ValueError, TypeError, UnicodeError) as error:
            raise render_contract_failure(error) from None
        # 资源快照和计划在创建 Word 对象前固定；后续 writer 只消费同一计划。
        document = Document(io.BytesIO(template_content)) if template_content is not None else Document()
        reserved_bookmarks = bookmark_inventory(document)
        self._clear_body(document)
        self._apply_page_setup(document, template_config)
        self._apply_header_footer(document, template_config, logo_bytes=logo_content)
        self._apply_document_settings(document)

        assets_by_id = {asset.id: asset for asset in document_ir.assets}
        render_context = self._build_render_context(document_ir, table_plan, plan.table_units)
        render_context.collector = collector
        render_context.assets, render_context.image_loader = assets_by_id, image_loader
        expected_formulas = len(render_context.formula_sources)
        if any(1_000_001 <= value <= 1_000_000 + expected_formulas for value in reserved_bookmarks[1]):
            raise content_failure("ir_render_evidence:target_mismatch")
        render_context.image_report = new_image_report(document_ir, template_config=template_config,
            template_content=template_content, config_content=config_content,
            template_version_id=options.template_version_id if options else None,
            language=(options.document_language if options and options.document_language else document_ir.metadata.document_language))
        render_context.source_blocks = sorted(document_ir.blocks, key=lambda block: block.order)
        render_context.image_blocks = {block.id: block for block in document_ir.blocks if block.type == "image"}
        render_context.image_spans = {block_id: span for span in build_image_plan(document_ir) for block_id in span.block_ids}
        document.part._md2word_formula_context = render_context
        blocks = {block.id: block for block in document_ir.blocks}
        decisions = {decision.block_id: decision for decision in plan.decisions}
        native_units = {unit.key: unit for unit in table_plan.units}
        unit_contracts = {unit.unit_id: unit for unit in plan.table_units}
        for operation in plan.operations:
            if operation.kind == "section":
                if operation.start_kind != "initial":
                    section = document.add_section(WD_SECTION_START.NEW_PAGE if operation.start_kind == "new_page" else WD_SECTION_START.CONTINUOUS)
                    if operation.start_kind == "continuous":
                        self._set_section_columns(section, operation.columns)
            elif operation.kind == "configure_page":
                self._configure_layout_page_section(document, template_config, landscape=operation.landscape)
            elif operation.kind == "decoration":
                self._render_page_decoration(document, operation, blocks, template_config, collector)
            elif operation.kind == "column_break":
                document.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)
            elif operation.kind == "logical_table":
                self._render_native_unit(document, native_units[operation.unit_id], unit_contracts[operation.unit_id],
                    blocks, template_config, render_context, page_break_before=decisions[operation.block_ids[0]].page_break_before)
            elif operation.kind == "block":
                block = blocks[operation.block_ids[0]]
                if block.type == "image":
                    self._render_image_span(document, render_context.image_spans[block.id], [blocks[key] for key in operation.block_ids],
                        assets_by_id, template_config, image_loader, render_context, page_width=operation.page_width)
                else:
                    self._render_block(document, block, assets_by_id, template_config, image_loader, render_context,
                        decision=decisions[block.id], table_model=tables.get(block.id, (None, None))[0], page_width=operation.page_width)
        self._materialize_fragment_markers(document, render_context, reserved_bookmarks)
        collector.created_anchors, collector.block_anchor_counts = materialize_block_anchors(
            document, collector.records, collector.targets, reserved_bookmarks)
        buffer = io.BytesIO()
        document.save(buffer)
        content = buffer.getvalue()
        try:
            evidence, ownership, anchors, anchor_ownership = collector.freeze(document, content, plan)
        except (ValueError, TypeError, KeyError, UnicodeError) as error:
            raise render_contract_failure(error) from None
        formula_report = {"schema_version": "formula-render/v1", "entries": render_context.formula_entries}
        coverage = build_formula_source_coverage(document_ir, content, table_plan, formula_report)
        return WordRenderResult(content, plan, evidence, CONTENT_VERSION, coverage, ownership,
            formula_report=formula_report,
            image_layout_report=finish_image_report(render_context.image_report),
            block_anchor_evidence=anchors, _anchor_ownership=anchor_ownership,
            block_anchor_counts=collector.block_anchor_counts)

    def _apply_document_settings(self, document: Any) -> None:
        settings = document.settings.element
        for tag in ("w:doNotAutoCompressPictures", "w:updateFields"):
            element = settings.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                settings.append(element)
            element.set(qn("w:val"), "true")

    def _content_width_and_columns(self, document: Any) -> tuple[int, int]:
        sections = getattr(document, "sections", None)
        if not sections:
            return int(Cm(14)), 1
        section = sections[-1]
        total_width = section.page_width - section.left_margin - section.right_margin
        if total_width <= 0:
            return int(Cm(14)), 1

        columns = section._sectPr.find(qn("w:cols"))
        raw_count = columns.get(qn("w:num")) if columns is not None else None
        column_count = max(int(raw_count), 1) if raw_count and raw_count.isdigit() else 1
        if column_count == 1:
            return int(total_width), 1
        raw_space = columns.get(qn("w:space")) if columns is not None else None
        column_space = int(raw_space) if raw_space and raw_space.isdigit() else 720
        available_width = total_width - Twips(column_space * (column_count - 1))
        return max(int(available_width / column_count), 1), column_count

    def _build_render_context(self, document_ir: DocumentIR, table_plan, contracts) -> _RenderContext:
        placeholders = decode_placeholders(document_ir)
        return _RenderContext(caption_bookmarks=caption_bookmark_inventory(document_ir, table_contracts=contracts,
                primary_table_block_ids={document_ir.blocks[unit.anchor_block_index].id
                                         for unit in table_plan.units if unit.is_group_primary}),
            formula_sources=[source.binding for source in formula_sources(document_ir, table_plan) if source.representative is None],
            placeholders=placeholders,
            image_bookmark_ids={f"reserved-placeholder:{index}": 2_000_000 + index for index in range(len(placeholders))})

    def _materialize_fragment_markers(self, document, context, reserved):
        """全部固定公式完成后分配 TAB 数字 ID，通用 block 锚点随后读取实际占用集。"""
        if len(context.used_formula_sources) != len(context.formula_sources):
            raise content_failure("ir_render_evidence:target_mismatch")
        names, ids = bookmark_inventory(document)
        if len(context.formula_entries) != len(context.formula_sources):
            raise content_failure("ir_render_evidence:target_mismatch")
        names.update(reserved[0])
        ids.update(reserved[1])
        next_id = max(max(ids, default=-1), 1_000_000 + len(context.formula_entries)) + 1
        if next_id + len(context.fragment_markers) > 2_147_483_647:
            raise content_failure("ir_render_evidence:target_mismatch")
        for name, paragraph in context.fragment_markers:
            if name in names:
                raise content_failure("ir_render_evidence:target_mismatch")
            start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
            start.set(qn("w:name"), name)
            start.set(qn("w:id"), str(next_id))
            end.set(qn("w:id"), str(next_id))
            paragraph.append(start)
            paragraph.append(end)
            names.add(name)
            next_id += 1

    def _load_template_config(
        self,
        document_ir: DocumentIR,
        config_path: Path | None,
        options: DocumentIRRenderOptions | None,
        *,
        content: bytes | None = None,
    ) -> TemplateConfigPackage:
        document_language = (
            options.document_language
            if options is not None and options.document_language
            else document_ir.metadata.document_language
        )
        template_config = (TemplateConfigPackage.load_content(content) if content is not None else TemplateConfigPackage.load(config_path)).for_document_language(document_language)
        subscript_config_path = options.subscript_config_path if options is not None else None
        if subscript_config_path is not None:
            terminology = TerminologyReplacer(subscript_config_path)
            if terminology.rules:
                template_config.terminology = terminology
        return template_config

    def _clear_body(self, document: Any) -> None:
        body = document._body._element
        sect_pr = body.sectPr
        for child in list(body):
            if child is sect_pr:
                continue
            body.remove(child)

    def _apply_page_setup(self, document: Any, template_config: TemplateConfigPackage) -> None:
        setup = template_config.page_setup
        for section in document.sections:
            self._apply_page_setup_to_section(section, setup)

    def _apply_page_setup_to_section(self, section: Any, setup: Any, *, apply_page_size: bool = True) -> None:
        if apply_page_size:
            if setup.page_width_cm is not None:
                section.page_width = Cm(setup.page_width_cm)
            if setup.page_height_cm is not None:
                section.page_height = Cm(setup.page_height_cm)
        if setup.top_margin is not None:
            section.top_margin = Cm(setup.top_margin)
        if setup.bottom_margin is not None:
            section.bottom_margin = Cm(setup.bottom_margin)
        if setup.left_margin is not None:
            section.left_margin = Cm(setup.left_margin)
        if setup.right_margin is not None:
            section.right_margin = Cm(setup.right_margin)
        if setup.header_distance is not None:
            section.header_distance = Cm(setup.header_distance)
        if setup.footer_distance is not None:
            section.footer_distance = Cm(setup.footer_distance)

    def _configure_layout_page_section(
        self,
        document: Any,
        template_config: TemplateConfigPackage,
        *,
        landscape: bool,
    ) -> None:
        if not document.sections:
            return
        section = document.sections[-1]
        section.orientation = WD_ORIENTATION.LANDSCAPE if landscape else WD_ORIENTATION.PORTRAIT
        portrait_width, portrait_height = self._portrait_page_dimensions(template_config)
        short_side = min(portrait_width, portrait_height)
        long_side = max(portrait_width, portrait_height)
        section.page_width = long_side if landscape else short_side
        section.page_height = short_side if landscape else long_side
        self._apply_page_setup_to_section(section, template_config.page_setup, apply_page_size=False)

    def _portrait_page_dimensions(self, template_config: TemplateConfigPackage) -> tuple[Any, Any]:
        setup = template_config.page_setup
        width = Cm(setup.page_width_cm) if setup.page_width_cm is not None else A4_PORTRAIT_WIDTH
        height = Cm(setup.page_height_cm) if setup.page_height_cm is not None else A4_PORTRAIT_HEIGHT
        return width, height

    def _apply_header_footer(
        self,
        document: Any,
        template_config: TemplateConfigPackage,
        *,
        logo_bytes: bytes | None,
    ) -> None:
        if not document.sections:
            return
        section = document.sections[0]
        if template_config.header.enabled:
            paragraph = section.header.paragraphs[0] if section.header.paragraphs else section.header.add_paragraph()
            paragraph.text = ""
            header_style = TemplateParagraphStyleConfig(
                font_name=template_config.header.title_font_name,
                western_font_name=template_config.header.title_western_font_name,
                complex_script_font_name=template_config.header.title_complex_script_font_name,
                font_size=template_config.header.title_font_size,
                color=template_config.header.title_color,
                alignment="left",
            )
            if logo_bytes is not None:
                run = paragraph.add_run()
                try:
                    run.add_picture(io.BytesIO(logo_bytes), height=Cm(template_config.header.logo_height_cm or 1.0))
                    paragraph.add_run("  ")
                except (InvalidImageStreamError, UnexpectedEndOfFileError, UnrecognizedImageError):
                    paragraph.text = ""
            if template_config.header.title_text:
                run = paragraph.add_run(template_config.header.title_text)
                self._apply_run_style(run, header_style)

        if template_config.footer.enabled:
            footer_style = TemplateParagraphStyleConfig(
                font_name=template_config.footer.font_name,
                western_font_name=template_config.footer.western_font_name,
                complex_script_font_name=template_config.footer.complex_script_font_name,
            )
            if template_config.footer.left_text or template_config.footer.right_text:
                footer = section.footer
                for child in list(footer._element):
                    footer._element.remove(child)
                paragraph = footer.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                self._add_text_run(paragraph, template_config.footer.left_text, footer_style)
                self._add_positional_tab(paragraph, "center")
                if template_config.footer.show_page_number:
                    self._add_complex_field(paragraph, " PAGE ", "1", footer_style)
                self._add_positional_tab(paragraph, "right")
                self._add_text_run(paragraph, template_config.footer.right_text, footer_style)
                return

            paragraph = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
            if template_config.footer.show_page_number and not paragraph.text.strip() and not paragraph.runs:
                paragraph.text = "第  页"
            for footer_paragraph in section.footer.paragraphs:
                footer_paragraph.alignment = self._parse_alignment(template_config.footer.alignment or "center")
                for run in footer_paragraph.runs:
                    self._apply_run_style(run, footer_style)

    def _add_positional_tab(self, paragraph: Any, alignment: str) -> None:
        run = paragraph.add_run()
        tab = OxmlElement("w:ptab")
        tab.set(qn("w:alignment"), alignment)
        tab.set(qn("w:relativeTo"), "margin")
        tab.set(qn("w:leader"), "none")
        run._r.append(tab)

    def _resolve_asset_path(self, raw_path: str | None, config_path: Path | None) -> Path | None:
        if not raw_path:
            return None
        candidate = Path(raw_path)
        if candidate.is_absolute():
            return candidate
        if config_path is not None:
            return (config_path.parent / candidate).resolve()
        return candidate.resolve()

    def _render_block(
        self, document: Any, block: DocumentIRBlock, assets_by_id: dict[str, DocumentIRAsset],
        template_config: TemplateConfigPackage, image_loader: ImageContentLoader, render_context: _RenderContext,
        *, decision: Any, table_model: Any = None, page_width: float | None = None,
    ) -> None:
        collector = render_context.collector
        if decision.mode == "native_text":
            texts = block.list_items or [] if block.type == "list" else [block.latex or block.text or "" if block.type == "formula" else block.text or ""]
            for index, (text, policy) in enumerate(zip(texts, decision.text_policies, strict=True)):
                if policy.mode == "coalesced":
                    collector.records[block.id].items.append(ItemEvidence(item_index=index, disposition="coalesced",
                        representative_block_id=policy.representative_block_id, representative_item_index=policy.representative_item_index))
                    collector.representatives[block.id, index] = (block.id, policy.representative_block_id, index, policy.representative_item_index)
                    continue
                render_context.formula_location = (block.id, f"list_items/{index}" if block.type == "list"
                    else "latex" if block.type == "formula" and block.latex else "text")
                paragraph = document.add_paragraph()
                self._fill_planned_text(paragraph, text, policy, template_config,
                    reference_bookmarks=render_context.caption_bookmarks)
                if decision.keep_with_next:
                    paragraph.paragraph_format.keep_with_next = True
                if decision.page_break_before and (block.type == "heading" or policy.mode == "table_context"):
                    paragraph.paragraph_format.page_break_before = True
                role = "list_item" if block.type == "list" else "formula" if block.type == "formula" else "body"
                collector.add(block.id, [paragraph_target(paragraph, role, index)])
                if block.type == "list":
                    collector.records[block.id].items.append(ItemEvidence(item_index=index, disposition="rendered"))
            if block.type == "list" and not texts:
                collector.records[block.id].reason_code = "empty_list"
        elif decision.table_strategy is not None:
            placeholders = [item for item in render_context.placeholders if item.owner_block_id == block.id]
            if placeholders and not (block.html or "").strip():
                collector.add(block.id, self._render_table_captions(document, block, template_config,
                    caption_bookmarks=render_context.caption_bookmarks, page_break_before=decision.page_break_before,
                    style_groups=decision.annotation_styles))
                if block.text and block.text not in block.captions:
                    paragraph = document.add_paragraph()
                    self._set_formula_source(document, block.id, "text")
                    self._fill_paragraph(paragraph, block.text, template_config.normal, terminology=template_config.terminology)
                    collector.add(block.id, [paragraph_target(paragraph)])
                self._render_table_placeholders(document, {item.fragment_id for item in placeholders},
                    assets_by_id, image_loader, render_context, template_config)
                collector.add(block.id, self._render_table_notes(document, block, template_config))
                return
            collector.add(block.id, self._render_table(document, block, template_config,
                caption_bookmarks=render_context.caption_bookmarks, reference_bookmarks=render_context.caption_bookmarks,
                page_break_before=decision.page_break_before, table_model=table_model, table_strategy=decision.table_strategy, annotation_styles=decision.annotation_styles))
    def _render_page_decoration(self, document, operation, blocks, template_config, collector):
        ids = operation.block_ids
        if operation.decoration_policy != "replace":
            for block_id in ids:
                record = collector.records[block_id]
                record.disposition = "suppressed"
                record.reason_code = "custom_footer" if operation.decoration_policy == "custom" else "decoration_disabled"
            return
        section = document.sections[-1]
        header = operation.decoration_kind == "header"
        decoration = section.header if header else section.footer
        paragraph = decoration.paragraphs[0] if decoration.paragraphs else decoration.add_paragraph()
        # 同一物理段落被复用时先撤销旧贡献，即使跨 section 链接也不能保留旧 owner。
        collector.replace_decoration(ids, paragraph, operation.decoration_kind)
        paragraph.text = ""
        settings = template_config.header if header else template_config.footer
        font_name = settings.title_font_name if header else settings.font_name
        western = settings.title_western_font_name if header else settings.western_font_name
        complex_font = settings.title_complex_script_font_name if header else settings.complex_script_font_name
        style = template_config.normal.model_copy(update={
            "font_name": template_config.normal.font_name if font_name is None else font_name,
            "western_font_name": self._resolve_decoration_font(western, font_name, template_config.normal.western_font_name),
            "complex_script_font_name": self._resolve_decoration_font(complex_font, font_name, template_config.normal.complex_script_font_name),
        })
        self._fill_paragraph(paragraph, " ".join(blocks[key].text or "" for key in ids).strip(), style,
            terminology=template_config.terminology)

    @staticmethod
    def _resolve_decoration_font(configured: str | None, legacy: str | None, inherited: str | None) -> str | None:
        """字体字段缺失时依次沿用旧单字体和正文；显式空字符串不回退。"""

        if configured is not None:
            return configured
        if legacy is not None:
            return legacy
        return inherited

    def _set_section_columns(self, section: Any, count: int) -> None:
        columns = section._sectPr.find(qn("w:cols"))
        if columns is None:
            columns = OxmlElement("w:cols")
            section._sectPr.append(columns)
        columns.set(qn("w:num"), str(count))

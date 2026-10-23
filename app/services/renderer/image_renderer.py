"""图片写入只消费一次计划；每图保留独立来源、比例和完整注释。"""

from __future__ import annotations

import hashlib
import io
from collections.abc import Callable
from typing import Any

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image
from docx.image.exceptions import InvalidImageStreamError, UnexpectedEndOfFileError, UnrecognizedImageError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips

from app.conversion.docx_rules.formula import sha256_text
from app.services.document_ir.models import DocumentIRAsset, DocumentIRBlock
from app.services.renderer.image_evidence import (IMAGE_TABLE_PREFIX, expected_annotation, image_anchor,
    image_dimensions, image_slots, section_geometry, single_image_width, source_image_position)
from app.services.renderer.layout_plan import ImageSpan, image_annotations
from app.services.renderer.template_config import TemplateConfigPackage

ImageContentLoader = Callable[[DocumentIRAsset], bytes | None]
_IMAGE_ERRORS = (InvalidImageStreamError, UnexpectedEndOfFileError, UnrecognizedImageError)


class WordImageRendererMixin:
    """无边框组和顺序单图共享逐图写入职责，不在写入阶段重新归组。"""

    def _render_image_span(self, document: Any, span: ImageSpan, blocks: list[DocumentIRBlock],
                           assets_by_id: dict, template_config: TemplateConfigPackage,
                           image_loader: ImageContentLoader, context: Any, *, page_width: float | None) -> None:
        geometry = section_geometry(document.sections[-1]._sectPr)
        inputs = context.image_report["render_inputs"]
        mode = "group" if span.mode == "group" and inputs["generation_policy"] == "grouped/v1" else "single_sequence"
        reasons = list(span.reasons)
        prepared = []
        for block in blocks:
            asset = assets_by_id.get(block.asset_ref or "")
            content = image_loader(asset) if asset is not None else None
            decoded = None
            if content:
                try:
                    decoded = Image.from_blob(content)
                except _IMAGE_ERRORS:
                    pass
                if decoded is not None and asset.checksum_sha256 and asset.checksum_sha256.removeprefix("sha256:") != hashlib.sha256(content).hexdigest():
                    raise ValueError("图片字节与源资产 checksum 不一致")
            prepared.append((content, decoded))
        if mode == "group" and any(image is None for _, image in prepared):
            mode, reasons = "single_sequence", ["asset_unavailable"]
        slots = None
        if mode == "group":
            width, slots = image_slots(blocks, page_width, geometry)
            if width <= 0 or any(slot["inner_width"] <= 0 or min(image_dimensions(decoded.px_width, decoded.px_height,
                    slot["inner_width"], geometry["effective_height"])) <= 0 for slot, (_, decoded) in zip(slots, prepared, strict=True)):
                mode, reasons, slots = "single_sequence", ["image_size_unfit"], None
        record = context.image_records.get(span.id)
        if record is None:
            frozen = {"span_id": span.id, "geometry": geometry}
            inputs["spans"].append(frozen)
            entry = {"span_id": span.id, "block_ids": list(span.block_ids), "page_index": span.page_index,
                     "planned_mode": span.mode, "metrics": dict(span.metrics), "mode": mode,
                     "reasons": reasons, "marker": IMAGE_TABLE_PREFIX + span.id if mode == "group" else None, "images": []}
            context.image_report["spans"].append(entry)
            context.image_records[span.id] = (entry, frozen)
        else:
            entry, frozen = record
            if geometry != frozen["geometry"]:
                # 不确定 run 可跨原生栏区域；仅为实际变化的成员冻结几何差异。
                frozen.setdefault("image_geometries", {}).update({block.id: geometry for block in blocks})
        cells = None
        if mode == "group":
            cells = self._image_layout_cells(document, width, slots, entry["marker"])
        for index, (block, (content, decoded)) in enumerate(zip(blocks, prepared, strict=True)):
            available = slots[index]["inner_width"] if slots else single_image_width(block, page_width, geometry)
            self._write_image(document, cells[index] if cells else document, block, content, decoded,
                              available, geometry["effective_height"], template_config, context, entry)

    def _image_layout_cells(self, document: Any, width: int, slots: list[dict], marker: str) -> list:
        table = document.add_table(rows=1, cols=len(slots))
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        props = table._tbl.tblPr
        indent = OxmlElement("w:tblInd")
        indent.set(qn("w:w"), "0")
        indent.set(qn("w:type"), "dxa")
        props.insert_element_before(indent, "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
        table_width = props.find(qn("w:tblW"))
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(width))
        caption = OxmlElement("w:tblCaption")
        caption.set(qn("w:val"), marker)
        props.append(caption)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement("w:" + side)
            border.set(qn("w:val"), "nil")
            borders.append(border)
        props.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
        split = OxmlElement("w:cantSplit")
        table.rows[0]._tr.get_or_add_trPr().append(split)
        for column, cell, slot in zip(table.columns, table.rows[0].cells, slots, strict=True):
            column.width = Twips(slot["width"])
            cell.width = Twips(slot["width"])
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                node = OxmlElement("w:" + side)
                node.set(qn("w:w"), str(slot.get(side, 0)))
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            cell._tc.get_or_add_tcPr().append(margins)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        return list(table.rows[0].cells)

    def _write_image(self, root: Any, container: Any, block: DocumentIRBlock, content: bytes | None,
                     decoded: Any, available_width: int, available_height: int,
                     config: TemplateConfigPackage, context: Any, entry: dict) -> None:
        annotations = image_annotations(block)
        paragraph = container.add_paragraph() if container is root else container.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.left_indent = Twips(0)
        paragraph.paragraph_format.right_indent = Twips(0)
        paragraph.paragraph_format.first_line_indent = Twips(0)
        anchor = image_anchor(block.id)
        self._image_bookmark(paragraph, anchor, context, start=True)
        width = height = 0
        if decoded is not None:
            width, height = image_dimensions(decoded.px_width, decoded.px_height, available_width, available_height)
            if min(width, height) <= 0:
                raise ValueError("图片无法放入有效内容区域")
            inline = paragraph.add_run().add_picture(io.BytesIO(content), width=width, height=height)
            # 客户端对省略距离可继承对象外距，必须显式清零以闭合 cell 内宽。
            for key in ("distT", "distB", "distL", "distR"):
                inline._inline.set(key, "0")
            inline._inline.docPr.set("name", anchor)
        else:
            paragraph.add_run("【图片缺失或无效，需复核】")
        self._image_bookmark(paragraph, anchor, context, start=False)
        paragraph.paragraph_format.keep_with_next = bool(annotations)
        image = {"block_id": block.id, "asset_ref": block.asset_ref, "asset_sha256": hashlib.sha256(content).hexdigest() if content else None,
                 "anchor": anchor, "width": width, "height": height,
                 "source": source_image_position(block), "annotations": []}
        entry["images"].append(image)
        for index, (field, text) in enumerate(annotations):
            annotation_anchor = image_anchor(block.id, field)
            expected_hash = expected_annotation(text, context.image_report["render_inputs"], context.caption_bookmarks)
            caption = container.add_paragraph()
            caption.paragraph_format.keep_with_next = index < len(annotations) - 1
            self._image_bookmark(caption, annotation_anchor, context, start=True)
            self._set_formula_source(root, block.id, field)
            if not self._fill_caption_paragraph(caption, text, config.caption, terminology=config.terminology, caption_bookmarks=context.caption_bookmarks):
                self._fill_paragraph(caption, text, config.caption, terminology=config.terminology, reference_bookmarks=context.caption_bookmarks)
            self._image_bookmark(caption, annotation_anchor, context, start=False)
            image["annotations"].append({"field": field, "source_sha256": sha256_text(text), "anchor": annotation_anchor,
                                         "expected_output_sha256": expected_hash})
        if not annotations and block.text and decoded is None:
            # 没有结构化注释的历史缺图仍保留其源文本和公式来源。
            self._set_formula_source(root, block.id, "text")
            self._fill_paragraph(paragraph, block.text, config.caption, terminology=config.terminology)

    @staticmethod
    def _image_bookmark(paragraph: Any, name: str, context: Any, *, start: bool) -> None:
        if start:
            context.image_bookmark_ids[name] = 2000000 + len(context.image_bookmark_ids)
        node = OxmlElement("w:bookmarkStart" if start else "w:bookmarkEnd")
        node.set(qn("w:id"), str(context.image_bookmark_ids[name]))
        if start:
            node.set(qn("w:name"), name)
        paragraph._p.append(node)

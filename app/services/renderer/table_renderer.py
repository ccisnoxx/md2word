from __future__ import annotations

# app/services/renderer/table_renderer.py: 提供 DocumentIR Word 渲染器的 HTML 表格与表注 helper。

from collections.abc import Iterator
from dataclasses import dataclass
from html import escape
import math
import re
from typing import Any, Literal

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt
from docx.text.run import Run
from docx.oxml.table import CT_Tc
from docx.table import _Cell
from docx.shared import Twips

from app.conversion.docx_rules.formula import MathSpan, content_math_spans, formula_block_span
from app.conversion.docx_rules.table_expansion import expanded_table_rows, requires_table_expansion
from app.conversion.runtime.table_html_parser import HtmlCell, HtmlTable, parse_html_table
from app.services.document_ir.models import DocumentIRBlock
from app.services.document_ir.table_annotations import is_table_context_annotation
from app.services.renderer.block_evidence import LocalTarget, paragraph_target
from app.services.renderer.template_config import TemplateConfigPackage, TemplateParagraphStyleConfig
from app.services.tables.html_normalizer import MAX_GRID_SLOTS
from app.services.tables.models import CellRun, TableCell, TableGrid


MIN_TABLE_COLUMN_WIDTH = Mm(16)
WIDE_TABLE_COLUMN_COUNT = 6
COMPACT_TABLE_FONT_SIZE = 8.0
EXTRA_COMPACT_TABLE_FONT_SIZE = 7.0
LONG_TABLE_CAPTION_RE = re.compile(r"^(?:续表|continued\s+table|[A-Z]\.\d+(?:\.\d+)?)\b", re.IGNORECASE)
TH_TAG_RE = re.compile(r"</?th\b[^>]*>", re.IGNORECASE)


@dataclass(frozen=True, slots=True)
class NativeCellPart:
    """来源账与 writer 共用的确定 native 视图；offset 按该视图连续计算。"""

    kind: Literal["inline", "protected_token", "line_break", "formula"]
    source: str
    offset: int
    bold: bool
    italic: bool
    spans: tuple[MathSpan, ...]
    literal: str | None = None


def native_cell_parts(runs: list[CellRun]) -> tuple[NativeCellPart, ...]:
    """只恢复同 marks 的有序 run 流；保护 token、换行和强公式维持原分界。"""
    parts: list[NativeCellPart] = []
    pending: list[CellRun] = []
    offset = 0

    def flush():
        nonlocal offset
        if not pending:
            return
        first = pending[0]
        chunks = []
        for run in pending:
            value = escape(run.source, quote=False)
            if run.type in {"subscript", "superscript"}:
                tag = "sub" if run.type == "subscript" else "sup"
                chunks.append(f"<{tag}>{value}</{tag}>")
            else:
                chunks.append(run.source if run.type == "formula" else value)
        source = "".join(chunks)
        for enabled, tag in ((first.italic, "i"), (first.bold, "b")):
            if enabled:
                source = f"<{tag}>{source}</{tag}>"
        # typed script 与显式数学均完整登记；共享扫描器负责脚本 base 的整体绑定。
        spans = content_math_spans(source)
        parts.append(NativeCellPart("inline", source, offset, first.bold, first.italic, spans))
        offset += len(source)
        pending.clear()

    for run in runs:
        strong_formula = run.type == "formula" and not content_math_spans(run.source)
        if run.type in {"protected_token", "line_break"} or strong_formula:
            flush()
            source = "<br/>" if run.type == "line_break" else run.source if strong_formula else escape(run.source, quote=False)
            spans = (formula_block_span(run.source),) if strong_formula else ()
            parts.append(NativeCellPart("formula" if strong_formula else run.type, source, offset,
                run.bold, run.italic, spans, run.source))
            offset += len(source)
            continue
        if pending and (pending[0].bold, pending[0].italic) != (run.bold, run.italic):
            flush()
        pending.append(run)
    flush()
    return tuple(parts)


def prepare_table_input(block: DocumentIRBlock, config: TemplateConfigPackage):
    """计划与 writer 共用实际 HTML 解析和既有大表阈值，输出模型仅驻内存。"""
    model = parse_html_table(block.html)
    strategy = "fallback" if model is None or not model.rows or model.max_columns <= 0 else (
        "expanded" if requires_table_expansion(model, block.html) else "standard")
    return model, strategy


def table_annotation_style(text: str) -> str:
    if is_table_context_annotation(text):
        return "normal"
    return "long_table_caption" if LONG_TABLE_CAPTION_RE.match((text or "").strip()) else "caption"


def table_config_for_columns(
    template_config: TemplateConfigPackage,
    max_cols: int,
) -> TemplateConfigPackage:
    """计划与写者共用半磅字号；仅显式下限允许宽表额外缩字，窄表间距保持。"""
    current_size = template_config.table_content.font_size
    minimum_size = template_config.table_content.min_font_size
    target_size = None if current_size is None else math.floor(current_size * 2) / 2
    if target_size is not None and minimum_size is not None and max_cols >= WIDE_TABLE_COLUMN_COUNT:
        compact_size = COMPACT_TABLE_FONT_SIZE if max_cols <= 8 else EXTRA_COMPACT_TABLE_FONT_SIZE
        target_size = max(math.ceil(minimum_size * 2) / 2, min(target_size, compact_size))
    style_updates = {"font_size": target_size}
    if max_cols >= WIDE_TABLE_COLUMN_COUNT:
        style_updates.update({
            "line_spacing": 1.0,
            "space_before": 0,
            "space_after": 0,
            "first_line_indent_pt": 0,
            "left_indent_pt": 0,
            "right_indent_pt": 0,
        })
    elif target_size == current_size:
        return template_config
    table_content = template_config.table_content.model_copy(update=style_updates)
    return template_config.model_copy(update={"table_content": table_content})


class TableRenderError(ValueError):
    """仅投影固定原因和当前 IR 数组索引，不回显 HTML、任意 ID 或内部异常。"""

    def __init__(self, code: Literal["table_render_input", "table_render_geometry", "table_render_failed",
                                     "table_build_limit", "table_html_encoding", "table_html_unsupported_structure"],
                 *, block_index: int) -> None:
        self.code = code
        self.block_index = block_index
        super().__init__(f"表格渲染失败：{code}，块索引={block_index}。")


def iter_grid_rows(grid: TableGrid, *, block_index: int) -> Iterator[tuple[int, list[TableCell]]]:
    """遍历源序锚点和活跃 rowspan；不分配覆盖全表的第二份占位矩阵。"""
    if grid.rows * grid.columns > MAX_GRID_SLOTS:
        raise TableRenderError("table_build_limit", block_index=block_index)
    active: list[TableCell] = []
    index = 0
    for row_index in range(grid.rows):
        active = [cell for cell in active if cell.row + cell.rowspan > row_index]
        row_cells = list(active)
        while index < len(grid.cells) and grid.cells[index].row == row_index:
            cell = grid.cells[index]
            if (cell.rowspan <= 0 or cell.colspan <= 0 or cell.row + cell.rowspan > grid.rows):
                raise TableRenderError("table_render_geometry", block_index=block_index)
            row_cells.append(cell)
            if cell.rowspan > 1:
                active.append(cell)
            index += 1
        row_cells.sort(key=lambda cell: cell.column)
        cursor = 0
        for cell in row_cells:
            if cell.column < cursor or cell.column + cell.colspan > grid.columns:
                raise TableRenderError("table_render_geometry", block_index=block_index)
            cursor = cell.column + cell.colspan
        yield row_index, row_cells
    if index != len(grid.cells):
        raise TableRenderError("table_render_geometry", block_index=block_index)



class WordTableRendererMixin:
    """封装可编辑 HTML 表格、合并单元格和表注渲染逻辑。"""

    def _render_table(
        self,
        document: Any,
        block: DocumentIRBlock,
        template_config: TemplateConfigPackage,
        *,
        caption_bookmarks: Any | None = None,
        reference_bookmarks: Any | None = None,
        page_break_before: bool = False,
        table_model: HtmlTable | None,
        table_strategy: str,
        annotation_styles: list[str],
    ) -> list[LocalTarget]:
        targets = self._render_table_captions(
            document,
            block,
            template_config,
            caption_bookmarks=caption_bookmarks,
            page_break_before=page_break_before,
            style_groups=annotation_styles,
        )
        if table_strategy == "fallback":
            self._set_formula_source(document, block.id, "text")
            paragraph = document.add_paragraph()
            self._fill_paragraph(
                paragraph,
                block.text or "",
                template_config.normal,
                terminology=template_config.terminology,
                reference_bookmarks=reference_bookmarks,
            )
            targets.append(paragraph_target(paragraph, "body"))
            targets.extend(self._render_table_notes(document, block, template_config))
            return targets

        max_cols = table_model.max_columns
        usable_width = self._available_table_width(document)
        table_template_config = table_config_for_columns(template_config, max_cols)
        if table_strategy == "expanded":
            table = self._render_expanded_table(
                document,
                table_model,
                max_cols,
                usable_width,
                table_template_config,
                block_id=block.id,
                reference_bookmarks=reference_bookmarks,
            )
        else:
            table = self._render_standard_html_table(
                document,
                table_model,
                max_cols,
                usable_width,
                table_template_config,
                block_id=block.id,
                reference_bookmarks=reference_bookmarks,
            )
        targets.append(LocalTarget(table._tbl, table.part, "table"))
        targets.extend(self._render_table_notes(document, block, template_config))
        return targets

    def _render_standard_html_table(
        self,
        document: Any,
        table_model: HtmlTable,
        max_cols: int,
        usable_width: int,
        template_config: TemplateConfigPackage,
        *,
        block_id: str,
        reference_bookmarks: Any | None = None,
    ) -> Any:
        """渲染常规 HTML table，保留 rowspan/colspan 合并结构。"""

        table = self._add_configured_table(document, len(table_model.rows), max_cols, usable_width, template_config)
        column_widths = self._table_column_widths(max_cols, usable_width)
        occupied: dict[tuple[int, int], bool] = {}
        for row_index, row in enumerate(table_model.rows):
            col_index = 0
            for source_column, cell_model in enumerate(row):
                while occupied.get((row_index, col_index)):
                    col_index += 1
                if col_index >= max_cols:
                    break

                base_cell = table.cell(row_index, col_index)
                end_row = min(row_index + cell_model.rowspan - 1, len(table_model.rows) - 1)
                end_col = min(col_index + cell_model.colspan - 1, max_cols - 1)
                merged_cell = base_cell
                if end_row != row_index or end_col != col_index:
                    merged_cell = base_cell.merge(table.cell(end_row, end_col))
                merged_cell.width = sum(column_widths[col_index : end_col + 1])
                self._set_formula_source(document, block_id, f"html/{row_index}/{source_column}")
                self._fill_table_cell(
                    merged_cell,
                    cell_model,
                    row_index == 0 or cell_model.is_header,
                    template_config,
                    reference_bookmarks=reference_bookmarks,
                )

                for occupied_row in range(row_index, end_row + 1):
                    for occupied_col in range(col_index, end_col + 1):
                        occupied[(occupied_row, occupied_col)] = True
                col_index = end_col + 1

        return table

    def _render_expanded_table(
        self,
        document: Any,
        table_model: HtmlTable,
        max_cols: int,
        usable_width: int,
        template_config: TemplateConfigPackage,
        *,
        block_id: str,
        reference_bookmarks: Any | None = None,
    ) -> Any:
        """大表回退策略：展开跨行跨列，避免巨大复杂表合并导致渲染失败。"""

        rows = expanded_table_rows(table_model)
        table = self._add_configured_table(document, len(rows), max_cols, usable_width, template_config)
        column_widths = self._table_column_widths(max_cols, usable_width)
        for row_index, row in enumerate(rows):
            for col_index in range(max_cols):
                occurrence = row[col_index]
                if occurrence is not None:
                    self._set_formula_source(document, block_id, f"html/{occurrence.source_row}/{occurrence.source_column}")
                table.cell(row_index, col_index).width = column_widths[col_index]
                self._fill_table_cell(
                    table.cell(row_index, col_index),
                    occurrence.cell if occurrence is not None else HtmlCell(text=""),
                    row_index == 0,
                    template_config,
                    reference_bookmarks=reference_bookmarks,
                )

        return table

    def _add_configured_table(
        self,
        document: Any,
        rows: int,
        cols: int,
        usable_width: int,
        template_config: TemplateConfigPackage,
    ) -> Any:
        table = document.add_table(rows=rows, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        self._set_table_width(table, usable_width)
        try:
            table.style = template_config.table_style.style_name
        except (KeyError, ValueError):
            table.style = "Table Grid"

        for column, column_width in zip(table.columns, self._table_column_widths(cols, usable_width), strict=True):
            column.width = column_width
        if template_config.table_style.border_color:
            self._set_table_borders(table, template_config.table_style.border_color)
        self._configure_data_table_rows(table)
        return table

    def _table_column_widths(self, cols: int, usable_width: int) -> list[int]:
        if cols <= 0:
            return []
        if cols == 5:
            weights = [0.11, 0.37, 0.21, 0.21, 0.10]
            widths = [int(usable_width * weight) for weight in weights]
            widths[1] += int(usable_width) - sum(widths)
            return widths

        if cols >= WIDE_TABLE_COLUMN_COUNT:
            return self._equal_widths(cols, usable_width)

        min_width = int(Emu(MIN_TABLE_COLUMN_WIDTH))
        if min_width * cols <= usable_width:
            column_width = max(int(usable_width // cols), min_width)
            widths = [column_width for _ in range(cols)]
            widths[-1] += int(usable_width) - sum(widths)
            return widths
        return self._equal_widths(cols, usable_width)

    def _equal_widths(self, cols: int, usable_width: int) -> list[int]:
        column_width = max(int(usable_width // cols), 1)
        widths = [column_width for _ in range(cols)]
        widths[-1] += int(usable_width) - sum(widths)
        return widths

    def _fill_table_cell(
        self,
        cell: Any,
        cell_model: HtmlCell,
        is_header: bool,
        template_config: TemplateConfigPackage,
        *,
        reference_bookmarks: Any | None = None,
    ) -> None:
        cell.text = ""
        paragraph = cell.paragraphs[0]
        cell_content = self._table_cell_content(cell_model, is_header, template_config)
        self._fill_paragraph(
            paragraph,
            cell_content,
            template_config.table_content,
            terminology=template_config.terminology,
            reference_bookmarks=reference_bookmarks,
        )
        self._apply_table_header_run_style(paragraph, is_header, template_config)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        if is_header and template_config.table_style.header_fill:
            self._set_cell_fill(cell, template_config.table_style.header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    def _render_table_captions(
        self,
        document: Any,
        block: DocumentIRBlock,
        template_config: TemplateConfigPackage,
        *,
        caption_bookmarks: Any | None = None,
        page_break_before: bool = False,
        style_groups: list[str],
    ) -> list[LocalTarget]:
        targets = []
        for index, text in enumerate(self._unique_annotation_texts(block.captions)):
            source_index = next(index for index, value in enumerate(block.captions) if value.strip() == text)
            self._set_formula_source(document, block.id, f"captions/{source_index}")
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            if page_break_before and index == 0:
                paragraph.paragraph_format.page_break_before = True
            group = style_groups[index]
            if group == "normal":
                self._fill_table_context_paragraph(paragraph, text, template_config)
                targets.append(paragraph_target(paragraph, "context", index))
                continue
            targets.append(paragraph_target(paragraph, "caption", index))
            style = getattr(template_config, group)
            if not self._fill_caption_paragraph(
                paragraph,
                text,
                style,
                terminology=template_config.terminology,
                caption_bookmarks=caption_bookmarks,
            ):
                self._fill_paragraph(
                    paragraph,
                    text,
                    style,
                    terminology=template_config.terminology,
                )

        return targets

    def _render_table_notes(
        self,
        document: Any,
        block: DocumentIRBlock,
        template_config: TemplateConfigPackage,
    ) -> list[LocalTarget]:
        targets = []
        for index, text in enumerate(self._unique_annotation_texts(block.footnotes)):
            source_index = next(i for i, value in enumerate(block.footnotes) if value.strip() == text)
            self._set_formula_source(document, block.id, f"footnotes/{source_index}")
            paragraph = document.add_paragraph()
            targets.append(paragraph_target(paragraph, "note", index))
            self._fill_paragraph(
                paragraph,
                text,
                template_config.table_note,
                terminology=template_config.terminology,
            )

        return targets

    def _unique_annotation_texts(self, values: list[str]) -> list[str]:
        seen: set[str] = set()
        texts: list[str] = []
        for value in values:
            text = value.strip()
            if text and text not in seen:
                texts.append(text)
                seen.add(text)
        return texts

    def _apply_table_header_run_style(
        self,
        paragraph: Any,
        is_header: bool,
        template_config: TemplateConfigPackage,
    ) -> None:
        if not is_header or not template_config.table_style.header_bold:
            return
        inside_formula = False
        for child in paragraph._p:
            if child.tag == qn("w:bookmarkStart") and (child.get(qn("w:name")) or "").startswith("md2w_formula_"):
                inside_formula = True
            elif child.tag == qn("w:bookmarkEnd") and inside_formula:
                inside_formula = False
            elif child.tag == qn("w:r") and not inside_formula:
                Run(child, paragraph).bold = True

    def _table_cell_content(
        self,
        cell_model: HtmlCell,
        is_header: bool,
        template_config: TemplateConfigPackage,
    ) -> str:
        content = cell_model.html or cell_model.text
        if is_header and not template_config.table_style.header_bold:
            return TH_TAG_RE.sub("", content)
        return content

    def _available_table_width(self, document: Any) -> int:
        return self._content_width_and_columns(document)[0]

    def _fill_table_context_paragraph(
        self,
        paragraph: Any,
        text: str,
        template_config: TemplateConfigPackage,
    ) -> None:
        self._fill_paragraph(
            paragraph,
            text,
            self._table_context_style(template_config),
            terminology=template_config.terminology,
        )

    def _table_context_style(self, template_config: TemplateConfigPackage) -> TemplateParagraphStyleConfig:
        alignment = template_config.normal.alignment
        if alignment not in {"left", "justify"}:
            alignment = "left"
        return template_config.normal.model_copy(
            update={
                "alignment": alignment,
                "first_line_indent_pt": 0,
                "space_before": template_config.normal.space_before if template_config.normal.space_before is not None else 0,
                "space_after": template_config.normal.space_after if template_config.normal.space_after is not None else 2,
            }
        )

    def _set_table_width(self, table: Any, width: int) -> None:
        table_properties = table._tbl.tblPr
        table_width = table_properties.find(qn("w:tblW"))
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            table_properties.insert(0, table_width)
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(Emu(width).twips))

    def _set_cell_fill(self, cell: Any, fill: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = tc_pr.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            tc_pr.append(shading)
        shading.set(qn("w:fill"), fill.strip().lstrip("#").upper())

    def _set_table_borders(self, table: Any, color: str) -> None:
        table_properties = table._tbl.tblPr
        borders = table_properties.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            table_properties.append(borders)
        normalized_color = color.strip().lstrip("#").upper()
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = borders.find(qn(f"w:{edge}"))
            if element is None:
                element = OxmlElement(f"w:{edge}")
                borders.append(element)
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "4")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), normalized_color)

    def _configure_data_table_rows(self, table: Any) -> None:
        for row_index, row in enumerate(table.rows):
            row_properties = row._tr.get_or_add_trPr()
            self._ensure_row_flag(row_properties, "w:cantSplit")
            if row_index == 0:
                self._ensure_row_flag(row_properties, "w:tblHeader", value="true")

    def _ensure_row_flag(self, row_properties: Any, tag: str, *, value: str | None = None) -> None:
        element = row_properties.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            row_properties.append(element)
        if value is not None:
            element.set(qn("w:val"), value)


    def _append_grid_cell(self, tr: Any, table: Any, width: int, span: int,
                          *, merge: str | None = None) -> Any:
        """局部 _Cell 不触发 Table._cells 的全表代理重建；CT_Tc.new 自带必需 p。"""
        tc = CT_Tc.new()
        tr.append(tc)
        cell = _Cell(tc, table)
        cell.width = width
        if span > 1:
            tc.grid_span = span
        if merge is not None:
            tc.vMerge = merge
        return cell


    def _format_table_cell(self, cell: Any, is_header: bool, template_config: TemplateConfigPackage) -> None:
        """大小表共用外观配置；合并延续不写内容但继承 owner 的表头外观。"""
        paragraph = cell.paragraphs[0]
        self._apply_table_header_run_style(paragraph, is_header, template_config)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        if is_header and template_config.table_style.header_fill:
            self._set_cell_fill(cell, template_config.table_style.header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    def _fill_native_grid_runs(self, paragraph, runs, config, *, reference_bookmarks=None):
        """保留 native run 的脚本/保护边界，普通文字继续原 text owner 叶链。"""
        style = config.table_content
        self._apply_named_style(paragraph, style.word_style_name)
        self._apply_paragraph_style(paragraph, style)
        for part in native_cell_parts(runs):
            run_style = style.model_copy(update={"bold": part.bold or style.bold, "italic": part.italic or style.italic})
            if part.kind == "line_break":
                paragraph.add_run().add_break()
            elif part.kind == "protected_token":
                self._add_text_run(paragraph, part.literal, run_style, bold=bool(run_style.bold), italic=bool(run_style.italic),
                    reference_bookmarks=reference_bookmarks)
            elif part.kind == "formula":
                self._append_formula(paragraph, part.spans[0], run_style,
                    bold=bool(run_style.bold), italic=bool(run_style.italic), _native_marks=True,
                    _native_references=reference_bookmarks)
            else:
                self._render_mixed_math(paragraph, part.source, run_style, config.terminology,
                    reference_bookmarks, _native_literal=True, _spans=part.spans)

    def _append_native_rows(self, table, fragment, width, config, context):
        """每个来源 cell 写一次；垂直合并的延续只继承几何，不复制文字或公式。"""
        widths = [Twips(Emu(value).twips) for value in self._table_column_widths(fragment.render.grid.columns, width)]
        created = []
        for source_row, row_cells in iter_grid_rows(fragment.render.grid, block_index=fragment.block_index):
            if source_row < fragment.removed_header_rows:
                continue
            tr = OxmlElement("w:tr")
            table._tbl.append(tr)
            created.append(tr)
            properties = tr.get_or_add_trPr()
            self._ensure_row_flag(properties, "w:cantSplit")
            cursor = 0
            for owner in row_cells:
                while cursor < owner.column:
                    self._append_grid_cell(tr, table, widths[cursor], 1)
                    cursor += 1
                continuation = owner.row < source_row
                cell = self._append_grid_cell(tr, table, sum(widths[owner.column:owner.column + owner.colspan]),
                    owner.colspan, merge="continue" if continuation else "restart" if owner.rowspan > 1 else None)
                header = owner.row == 0 or owner.is_header
                if not continuation:
                    self._set_formula_source(table, fragment.source.block_id,
                        f"html/{owner.source_row_index}/{owner.source_cell_index}")
                    try:
                        self._fill_native_grid_runs(cell.paragraphs[0], owner.runs, config,
                            reference_bookmarks=context.caption_bookmarks)
                    except ValueError as error:
                        # OOXML 拒绝非法字符时仍失败关闭，公开错误不携带单元格原文。
                        raise TableRenderError("table_render_failed", block_index=fragment.block_index) from error
                self._format_table_cell(cell, header, config)
                cursor = owner.column + owner.colspan
            while cursor < fragment.render.grid.columns:
                self._append_grid_cell(tr, table, widths[cursor], 1)
                cursor += 1
            for paragraph in tr.iter(qn("w:p")):
                paragraph.get_or_add_pPr().get_or_add_keepNext().val = False
        return tuple(created)

    def _render_native_unit(self, document, unit, contract, blocks, config, context, *, page_break_before=False):
        """表只由 unit 创建；行、题注和末注分别归实际来源 owner。"""
        collector = context.collector
        annotations, scopes = [], []

        def annotation(item):
            block = blocks[item.block_id]
            text = getattr(block, item.field)[item.item_index].strip()
            paragraph = document.add_paragraph()
            self._set_formula_source(document, block.id, f"{item.field}/{item.item_index}")
            style = config.table_note if item.role == "note" else getattr(config, table_annotation_style(text))
            if item.role != "note":
                paragraph.paragraph_format.keep_with_next = True
                if page_break_before and not annotations:
                    paragraph.paragraph_format.page_break_before = True
            if (item.role != "caption" or not unit.is_group_primary or not self._fill_caption_paragraph(paragraph, text, style,
                    terminology=config.terminology, caption_bookmarks=context.caption_bookmarks)):
                self._fill_paragraph(paragraph, text, style, terminology=config.terminology,
                    reference_bookmarks=context.caption_bookmarks if item.role != "caption" else None)
            target = paragraph_target(paragraph, item.role, item.item_index)
            collector.add(block.id, [target])
            annotations.append((block.id, item.field, item.item_index, target))

        for item in contract.annotations:
            if item.disposition == "rendered" and item.role != "note":
                annotation(item)
        if unit.blocking_quality_flags:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            reasons = []
            if "table_fragment_missing" in unit.blocking_quality_flags:
                reasons.append("来源表格存在空白或缺失片段，未补造内容")
            if "table_render_layout_incompatible" in unit.blocking_quality_flags:
                reasons.append("续表布局或中间说明不兼容，已保留各片原生表格")
            self._fill_paragraph(paragraph, "表格诊断：" + "；".join(reasons) + "。", config.table_note,
                terminology=config.terminology)
            collector.add(contract.anchor_block_id, [paragraph_target(paragraph, "body")])
        width = self._available_table_width(document)
        table_config = table_config_for_columns(config, unit.columns)
        table = self._add_configured_table(document, 0, unit.columns, width, table_config)
        machine = OxmlElement("w:tblCaption")
        machine.set(qn("w:val"), unit.key)
        table._tbl.tblPr.append(machine)
        fragments = {fragment.fragment_id: fragment for fragment in contract.fragments}
        for fragment in unit.fragments:
            if fragment.render is None or fragment.status == "header_only":
                continue
            rows = self._append_native_rows(table, fragment, width, table_config, context)
            for output_index, row in enumerate(rows, fragment.output_start):
                if output_index < unit.header_rows:
                    self._ensure_row_flag(row.get_or_add_trPr(), "w:tblHeader", value="true")
            first = next(rows[0].iter(qn("w:p")))
            scope = fragments[fragment.source.fragment_id]
            target = LocalTarget(first, document.part, "table", content_scope_id=scope.content_scope_id,
                scope_table=table._tbl, scope_rows=rows)
            collector.add(fragment.source.block_id, [target])
            collector.records[fragment.source.block_id].reason_code = "table_fragment_rows"
            scopes.append((fragment.source.block_id, target))
            context.fragment_markers.append((fragment.bookmark, first))
        if context.placeholders:
            self._render_table_placeholders(document, {fragment.source.fragment_id for fragment in unit.fragments},
                context.assets, context.image_loader, context, config)
        for item in contract.annotations:
            if item.disposition == "rendered" and item.role == "note":
                annotation(item)
        for fragment in unit.fragments:
            block_id = fragment.source.block_id
            if block_id is not None and any(target.content_scope_id is not None for target in collector.targets[block_id]):
                collector.records[block_id].reason_code = "table_fragment_rows"
            if block_id is not None and fragment.status == "header_only" and not collector.targets[block_id]:
                record = collector.records[block_id]
                record.disposition, record.reason_code = "coalesced", "table_header_coalesced"
                record.representative_block_id = contract.anchor_block_id
        collector.tables.append((unit.key, table._tbl, document.part, tuple(scopes), tuple(annotations)))

    def _render_table_placeholders(self, document, fragment_ids, assets, image_loader, context, config):
        """占位证明的固定 marker 保持，来源 owner 仍捕获真实诊断和图片段。"""
        import io
        from app.services.document_ir.table_placeholder import LABEL, TablePlaceholderError, checksum, source_label, validate_image
        from app.services.renderer.table_placeholder_evidence import marker
        for index, binding in enumerate(context.placeholders):
            if binding.fragment_id not in fragment_ids:
                continue
            content = image_loader(assets[binding.asset_ref])
            if content is None or checksum(content) != binding.asset_sha256:
                raise TablePlaceholderError("table_placeholder_integrity_failed")
            validate_image(content)
            paragraph = document.add_paragraph()
            self._apply_paragraph_style(paragraph, config.table_note)
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(2_000_000 + index))
            start.set(qn("w:name"), marker(binding))
            paragraph._p.append(start)
            paragraph.add_run(LABEL + "\n" + source_label(binding) + "\n")
            try:
                shape = paragraph.add_run().add_picture(io.BytesIO(content), width=self._available_table_width(document))
                shape._inline.docPr.set("descr", binding.placeholder_id)
            except Exception:
                raise TablePlaceholderError("table_placeholder_integrity_failed") from None
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(2_000_000 + index))
            paragraph._p.append(end)
            context.collector.add(binding.owner_block_id, [paragraph_target(paragraph, "image_fallback")])

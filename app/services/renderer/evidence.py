"""正式任务与本地转换共用的渲染完整性证据。"""

from __future__ import annotations

import hashlib
import re
from collections import Counter
from typing import TYPE_CHECKING, Any

from app.conversion.docx_rules.formula import content_math_spans, formula_block_span, sha256_text
from app.conversion.docx_rules.table_expansion import expanded_table_rows, requires_table_expansion
from app.conversion.runtime.table_html_parser import parse_html_table
from app.services.document_ir.integrity import DocumentIntegrityError, build_source_integrity, ensure_integrity, is_visible_body_block
from app.services.document_ir.models import DocumentIR
from app.services.document_ir.caption_reconciler import RULE_VERSION as CAPTION_RULE_VERSION
from app.services.renderer.docx_quality import (
    analyze_docx_quality_bytes,
    build_docx_quality_flags,
    validate_formula_report,
    verify_docx_formula_report,
)
from app.services.renderer.layout_plan import LayoutPlan, build_layout_plan, image_annotations
from app.services.renderer.image_evidence import verify_image_report
from app.services.renderer.logical_table_plan import (
    TABLE_RENDERING_VERSION, QUALITY_BLOCKERS, QUALITY_WARNINGS, SAFE_COUNT_KEYS,
)


if TYPE_CHECKING:
    from app.services.renderer.block_anchors import BlockAnchorEvidence
    from app.services.renderer.block_evidence import BlockRenderEvidence, WordRenderResult, _RenderOwnership
    from app.services.renderer.render_plan import RenderPlan

# 历史主件须逐项核对原aggregate合同，不能把新增trace或来源字段强加给旧件。
LEGACY_RENDERER_EVIDENCE_KEYS = frozenset({
    'schema_version',
    'renderer_mode',
    'layout_page_count',
    'layout_region_count',
    'layout_block_count',
    'expected_layout_block_count',
    'layout_block_count_delta',
    'layout_missing_block_ids',
    'layout_duplicate_block_ids',
    'layout_quality_flags',
    'source_integrity',
    'docx_quality',
    'quality_flags',
    'expected_inline_shape_count',
    'inline_shape_count',
    'inline_shape_count_delta',
    'expected_table_count',
    'table_count',
    'table_count_delta',
    'expected_heading_count',
    'heading_count',
    'heading_count_delta',
    'blocking_quality_flags',
})
_UNSET = object()


def build_renderer_metadata(document_ir: DocumentIR, layout_plan: LayoutPlan | None = None) -> dict[str, Any]:
    """构造产物和事件共用的布局证据，输入质量与执行完整性分开。"""

    resolved_layout_plan = layout_plan or build_layout_plan(document_ir)
    body_blocks = [block for block in document_ir.blocks if is_visible_body_block(block)]
    layout_block_ids = [
        block.id for page in resolved_layout_plan.pages for region in page.regions for block in region.blocks
    ]
    expected_block_ids = [block.id for block in body_blocks]
    metadata = {
        "schema_version": document_ir.schema_version,
        "renderer_mode": resolved_layout_plan.mode,
        "layout_page_count": len(resolved_layout_plan.pages),
        "layout_region_count": sum(len(page.regions) for page in resolved_layout_plan.pages),
        "layout_block_count": sum(
            len(region.blocks) for page in resolved_layout_plan.pages for region in page.regions
        ),
        "expected_layout_block_count": len(expected_block_ids),
        "layout_block_count_delta": len(layout_block_ids) - len(expected_block_ids),
        "layout_missing_block_ids": sorted(set(expected_block_ids) - set(layout_block_ids)),
        "layout_duplicate_block_ids": sorted(
            block_id for block_id, count in Counter(layout_block_ids).items() if count > 1
        ),
        "layout_quality_flags": document_ir.quality_flags,
        "source_integrity": build_source_integrity(document_ir),
        "mineru_validation": (document_ir.metadata.mineru_validation.model_dump(mode="json")
                              if document_ir.metadata.mineru_validation is not None else None),
        "mineru_validation_input": document_ir.metadata.mineru_validation_input,
    }
    caption_blocks = [block for block in document_ir.blocks if not block.metadata.get("source_block_id")
                      and isinstance(block.metadata.get("caption_reconciliation"), dict)]
    if caption_blocks:
        statuses: Counter[str] = Counter()
        reasons: Counter[str] = Counter()
        binding_count = candidate_count = 0
        examples: list[dict[str, Any]] = []
        relation_examples: list[dict[str, Any]] = []
        for block in caption_blocks:
            caption = block.metadata["caption_reconciliation"]
            statuses[caption["status"]] += 1
            reasons.update(caption.get("reasons", []))
            binding_count += len(caption.get("bindings", []))
            candidate_count += len(caption.get("candidates", []))
            for kind in ("bindings", "candidates"):
                for receipt in caption.get(kind, []):
                    if len(relation_examples) < 40:
                        relation_examples.append({
                            "kind": kind, "donor_block_id": receipt["donor_block_id"],
                            "target_block_id": receipt["target_block_id"], "source_locator": receipt["source_locator"],
                            "rule_id": receipt["rule_id"], "evidence_only": receipt["evidence_only"],
                            "reasons": receipt.get("reasons", [receipt["reason"]] if "reason" in receipt else []),
                        })
            if len(examples) < 40:
                examples.append({
                    "block_id": block.id, "source_locator": block.metadata.get("mineru_source"),
                    "status": caption["status"], "reasons": caption.get("reasons", []),
                })
        # 摘要不复制 originals、rich runs 或可见文本；完整关系仍由 IR 主源块持有。
        metadata["caption_reconciliation"] = {
            "rule_version": CAPTION_RULE_VERSION, "block_count": len(caption_blocks),
            "status_counts": dict(sorted(statuses.items())), "reason_counts": dict(sorted(reasons.items())),
            "binding_count": binding_count, "candidate_count": candidate_count,
            "examples": examples, "omitted_example_count": len(caption_blocks) - len(examples),
            "relation_examples": relation_examples,
            "omitted_relation_count": binding_count + candidate_count - len(relation_examples),
        }
    return metadata



def build_renderer_aggregate_evidence(
    document_ir: DocumentIR,
    word_content: bytes,
    layout_plan: LayoutPlan | None = None,
    *,
    formula_report: dict[str, Any] | None = None,
    image_layout_report: dict[str, Any] | None = None,
    expected_template_version_id: str | None = None,
    expected_generation_policy: str | None = None,
    _table_plan=None,
    _formula_sources: list[dict[str, Any]] | None = None,
    _placeholder_count: int = 0,
    _v2_content: bool = False,
    _legacy_ir_identity: bool = False,
    _placeholder_ownership=None,
) -> dict[str, Any]:
    """生成 DOCX 证据；新报告严格核验，历史无报告保持原字段。"""

    evidence = build_renderer_metadata(document_ir, layout_plan)
    if image_layout_report is not None:
        from app.services.renderer.logical_table_plan import table_unit_contracts
        evidence["image_layout_summary"] = verify_image_report(document_ir, word_content, image_layout_report,
            expected_template_version_id=expected_template_version_id, expected_generation_policy=expected_generation_policy,
            _legacy_ir_identity=_legacy_ir_identity,
            _table_contracts=table_unit_contracts(document_ir, _table_plan) if _v2_content else None,
            _table_primary_block_ids={document_ir.blocks[unit.anchor_block_index].id
                for unit in _table_plan.units if unit.is_group_primary} if _v2_content else None,
            _placeholder_ownership=_placeholder_ownership)
        evidence["image_layout_report"] = image_layout_report
    if formula_report is not None:
        summary = validate_formula_report(formula_report)
        expected_sources = formula_source_inventory(document_ir) if _formula_sources is None else _formula_sources
        recorded_sources = [entry["source"] for entry in formula_report["entries"]]
        if Counter(_binding_key(value) for value in expected_sources) != Counter(_binding_key(value) for value in recorded_sources):
            raise ValueError("公式报告与源 DocumentIR 不匹配")
        verify_docx_formula_report(word_content, formula_report)
        evidence["formula_report"] = formula_report
        evidence["formula_summary"] = summary
    docx_quality = analyze_docx_quality_bytes(word_content, formula_report=formula_report)
    evidence["docx_quality"] = docx_quality
    evidence["quality_flags"] = build_docx_quality_flags(docx_quality)
    if formula_report is not None and evidence["formula_summary"]["fallback_count"]:
        evidence["quality_flags"].append("formula_text_fallback")
    if image_layout_report is not None and evidence["image_layout_summary"]["uncertain_span_count"]:
        evidence["quality_flags"].append("image_group_uncertain")
    body_blocks = [block for block in document_ir.blocks if is_visible_body_block(block)]
    expected_inline_shapes = sum(1 for block in body_blocks if block.type == "image") + _placeholder_count
    expected_tables = (sum(1 for block in body_blocks if block.type == "table")
                       if _table_plan is None else len(_table_plan.units))
    expected_headings = sum(1 for block in body_blocks if block.type == "heading")
    actual_inline_shapes = int(docx_quality.get("inline_shape_count") or 0)
    actual_tables = int(docx_quality.get("table_count") or 0)
    actual_headings = int(docx_quality.get("heading_paragraphs") or 0)
    evidence.update(
        {
            "expected_inline_shape_count": expected_inline_shapes,
            "inline_shape_count": actual_inline_shapes,
            "inline_shape_count_delta": actual_inline_shapes - expected_inline_shapes,
            "expected_table_count": expected_tables,
            "table_count": actual_tables,
            "table_count_delta": actual_tables - expected_tables,
            "expected_heading_count": expected_headings,
            "heading_count": actual_headings,
            "heading_count_delta": actual_headings - expected_headings,
        }
    )
    flags = list(evidence["source_integrity"]["blocking_quality_flags"])
    if evidence["layout_missing_block_ids"]:
        flags.append("ir_layout_integrity:missing_blocks")
    if evidence["layout_duplicate_block_ids"]:
        flags.append("ir_layout_integrity:duplicate_blocks")
    asset_ids = {asset.id for asset in document_ir.assets}
    if any(block.type == "image" and (not block.asset_ref or block.asset_ref not in asset_ids) for block in body_blocks):
        flags.append("ir_docx_integrity:image_asset_missing")
    if actual_inline_shapes != expected_inline_shapes:
        flags.append("ir_docx_integrity:inline_shape_count_mismatch")
    if image_layout_report is not None:
        layout_count = evidence["image_layout_summary"]["accepted_group_count"]
        content_count = actual_tables - layout_count
        evidence.update({"expected_layout_table_count": layout_count, "layout_table_count": layout_count,
                         "layout_table_count_delta": 0, "expected_content_table_count": expected_tables,
                         "content_table_count": content_count, "content_table_count_delta": content_count - expected_tables,
                         "expected_total_table_count": expected_tables + layout_count,
                         "total_table_count": actual_tables, "total_table_count_delta": actual_tables - expected_tables - layout_count})
        if content_count != expected_tables:
            flags.append("ir_docx_integrity:table_count_mismatch")
    elif actual_tables != expected_tables:
        flags.append("ir_docx_integrity:table_count_mismatch")
    if document_ir.schema_version == "document-ir/v2" and actual_headings != expected_headings:
        flags.append("ir_docx_integrity:heading_count_mismatch")
    evidence["blocking_quality_flags"] = list(dict.fromkeys(flags))
    return evidence


def _binding_key(binding: dict[str, Any]) -> tuple[Any, ...]:
    return tuple(binding[key] for key in ("block_id", "field", "start", "end", "boundary", "source_sha256"))


def formula_source_inventory(document_ir: DocumentIR) -> list[dict[str, Any]]:
    """从权威 IR 重建每次公式出现的源绑定，不依赖渲染报告自证。"""
    sources: list[dict[str, Any]] = []

    def add(block_id: str, field: str, text: str, *, formula: bool = False) -> None:
        spans = (formula_block_span(text),) if formula else content_math_spans(text)
        sources.extend({"block_id": block_id, "field": field, "start": span.start, "end": span.end, "boundary": span.boundary, "source_sha256": sha256_text(span.source)} for span in spans)

    for block in document_ir.blocks:
        if not is_visible_body_block(block):
            continue
        if block.type == "formula":
            add(block.id, "latex" if block.latex else "text", block.latex or block.text or "", formula=True)
        elif block.type == "table":
            table = parse_html_table(block.html)
            if table:
                if requires_table_expansion(table, block.html):
                    for row in expanded_table_rows(table):
                        for occurrence in row:
                            if occurrence is not None:
                                add(block.id, f"html/{occurrence.source_row}/{occurrence.source_column}", occurrence.cell.html or occurrence.cell.text)
                else:
                    for row_index, row in enumerate(table.rows):
                        for col_index, cell in enumerate(row):
                            add(block.id, f"html/{row_index}/{col_index}", cell.html or cell.text)
            else:
                add(block.id, "text", block.text or "")
        elif block.type == "list":
            for index, text in enumerate(block.list_items or []):
                add(block.id, f"list_items/{index}", text)
        elif block.type != "image":
            add(block.id, "text", block.text or "")
        if block.type == "image":
            annotations = image_annotations(block)
            for field, text in annotations:
                add(block.id, field, text)
            if not annotations:
                add(block.id, "text", block.text or "")
        elif block.type == "table":
            seen: set[tuple[str, str]] = set()
            values = [(f"captions/{index}", text) for index, text in enumerate(block.captions)]
            values.extend((f"footnotes/{index}", text) for index, text in enumerate(block.footnotes))
            for field, text in values:
                # 表题与表注分别去重；图片库存与 writer 共用上面的唯一注释遍历。
                key = (field.split("/")[0], text.strip())
                if text.strip() and key not in seen:
                    seen.add(key)
                    add(block.id, field, text.strip())
    return sources


def ensure_renderer_integrity(evidence: dict[str, Any]) -> None:
    """阻止包含确定性 IR→DOCX 内容丢失的产物进入发布边界。"""

    if "table_rendering" in evidence and not is_valid_table_rendering_evidence(evidence["table_rendering"]):
        evidence = {**evidence, "blocking_quality_flags": [*evidence.get("blocking_quality_flags", []),
                    "ir_docx_integrity:table_fragment_map_mismatch"]}
    if "table_placeholder_count" in evidence:
        from app.services.renderer.table_placeholder_evidence import valid_placeholder_evidence
        count = evidence["table_placeholder_count"]
        proof = evidence.get("table_placeholder_rendering")
        if (type(count) is not int or count < 0 or count and proof is None
                or proof is not None and (not valid_placeholder_evidence(proof) or proof["count"] != count)):
            evidence = {**evidence, "blocking_quality_flags": [*evidence.get("blocking_quality_flags", []),
                        "table_placeholder_integrity_failed"]}
    if "table_evidence" in evidence:
        from app.services.tables.evidence import read_table_evidence, TableEvidenceError
        report = read_table_evidence(evidence["table_evidence"])
        observations = evidence.get("docx_quality", {}).get("table_observations")
        if not isinstance(observations, list) or any(output.actual.table_index >= len(observations)
                or output.actual.model_dump(mode="json") != observations[output.actual.table_index]
                for table in report.tables for output in table.outputs):
            raise TableEvidenceError("table_evidence_integrity")
    if evidence.get("blocking_quality_flags"):
        raise DocumentIntegrityError("DocumentIR→DOCX", build_renderer_event_summary(evidence))


RENDER_EVIDENCE_CODES = frozenset({
    "ir_render_evidence:binding_mismatch", "ir_render_evidence:block_coverage_mismatch",
    "ir_render_evidence:target_mismatch", "ir_render_evidence:invalid_contract",
})


def render_contract_failure(error: Exception) -> DocumentIntegrityError:
    """只公开固定机器码，不把 ValidationError/源身份写入异常载体。"""
    code = str(error) if str(error) in RENDER_EVIDENCE_CODES else "ir_render_evidence:invalid_contract"
    return DocumentIntegrityError("DocumentIR→DOCX", {"blocking_quality_flags": [code]})


def validate_render_result(document_ir: DocumentIR, word_content: bytes, render_result: WordRenderResult) -> None:
    """同次结果必须保有真实运行期归属，缺失时不能降级到历史读者。"""
    if word_content != render_result.content:
        raise render_contract_failure(ValueError("ir_render_evidence:binding_mismatch"))
    _validate_render_evidence(document_ir, word_content, render_result.plan, render_result.block_evidence,
        ownership=render_result._ownership, require_ownership=True)


def _validate_persisted_image_targets(index, records, blocks, image_ids) -> None:
    """原图片报告通过后，用其稳定 marker 容器核对逐图 target。"""
    from docx.oxml.ns import qn
    from app.services.renderer.image_evidence import image_anchor
    root = index.roots["word/document.xml"]
    markers = {}
    for node in root.iter(qn("w:bookmarkStart")):
        markers.setdefault(node.get(qn("w:name")), []).append(node)
    for block_id in image_ids:
        record = records[block_id]
        expected = [(image_anchor(block_id), None)] + [(image_anchor(block_id, field), i)
            for i, (field, _) in enumerate(image_annotations(blocks[block_id]))]
        if len(record.targets) != len(expected):
            raise ValueError("ir_render_evidence:target_mismatch")
        for target, (name, item_index) in zip(record.targets, expected, strict=True):
            starts = markers.get(name, [])
            if len(starts) != 1 or starts[0].getparent().tag != qn("w:p"):
                raise ValueError("ir_render_evidence:target_mismatch")
            paragraph = starts[0].getparent()
            if target.role == "image":
                drawings = list(paragraph.iter(qn("w:drawing")))
                if len(drawings) != 1:
                    raise ValueError("ir_render_evidence:target_mismatch")
                node = drawings[0]
            else:
                node = paragraph
            if (target.part, tuple(target.child_path), target.item_index) != (*index.paths[node], item_index):
                raise ValueError("ir_render_evidence:target_mismatch")


def validate_persisted_render_evidence(document_ir: DocumentIR, word_content: bytes, *,
                                       render_plan: RenderPlan | dict[str, Any] | None,
                                       block_render_evidence: BlockRenderEvidence | dict[str, Any] | None,
                                       ) -> tuple[RenderPlan, BlockRenderEvidence] | None:
    """消费历史可选 trace；坏 trace 不改变原主件恢复/下载合同。"""
    try:
        return _validate_render_evidence(document_ir, word_content, render_plan, block_render_evidence)
    except DocumentIntegrityError:
        return None


def _validate_render_evidence(document_ir: DocumentIR, word_content: bytes,
                              plan_value: RenderPlan | dict[str, Any] | None,
                              trace_value: BlockRenderEvidence | dict[str, Any] | None,
                              *, ownership: _RenderOwnership | None = None,
                              require_ownership: bool = False) -> tuple[RenderPlan, BlockRenderEvidence]:
    """共用最终包认证；持久化证明不重建同次运行期归属。"""
    from docx.oxml.ns import qn
    from app.services.renderer.block_evidence import (
        BlockRenderEvidence, PackageIndex, _RenderOwnership, provenance_for,
    )
    from app.services.renderer.render_plan import RenderPlan, content_hash, document_ir_hash, read_render_plan
    from app.services.renderer.block_evidence import read_block_render_evidence, resolve_table_scope
    try:
        plan = read_render_plan(plan_value)
        trace = read_block_render_evidence(trace_value)
        v2 = plan.schema_version == "render-plan/v2"
        if (trace.schema_version == "block-render-evidence/v2") != v2:
            raise ValueError("ir_render_evidence:invalid_contract")
        if (trace.docx_sha256 != hashlib.sha256(word_content).hexdigest()
                or plan.input_ir_sha256 != document_ir_hash(document_ir, legacy=not v2)
                or trace.input_ir_sha256 != plan.input_ir_sha256 or trace.resource_identity != plan.resource_identity
                or trace.plan_sha256 != plan.plan_sha256 or not trace.complete):
            raise ValueError("ir_render_evidence:binding_mismatch")
        if require_ownership and (not isinstance(ownership, _RenderOwnership)
                or (ownership.input_ir_sha256, ownership.resource_identity_sha256, ownership.plan_sha256, ownership.docx_sha256)
                != (plan.input_ir_sha256, content_hash(plan.resource_identity.model_dump(mode="json")), plan.plan_sha256, trace.docx_sha256)):
            raise ValueError("ir_render_evidence:target_mismatch")
        ids = [block.id for block in document_ir.blocks]
        record_ids = [record.block_id for record in trace.block_render_map]
        decision_ids = [decision.block_id for decision in plan.decisions]
        if len(ids) != len(set(ids)) or ids != record_ids or ids != decision_ids:
            raise ValueError("ir_render_evidence:block_coverage_mismatch")
        index = PackageIndex(word_content)
        table_members = set()
        if v2:
            from app.services.renderer.content_contract import validate_v2_table_trace
            table_members = validate_v2_table_trace(document_ir, plan, trace, index, ownership,
                require_ownership=require_ownership)
        if len(index.sections) != len(plan.sections):
            raise ValueError("ir_render_evidence:target_mismatch")
        for planned, actual in zip(plan.sections, index.sections, strict=True):
            if planned.orientation != "inherit" and planned.orientation != actual["orientation"] or planned.columns != "inherit" and planned.columns != actual["columns"]:
                raise ValueError("ir_render_evidence:target_mismatch")
        provenance = provenance_for(document_ir)
        records = {record.block_id: record for record in trace.block_render_map}
        decisions = {decision.block_id: decision for decision in plan.decisions}
        blocks = {block.id: block for block in document_ir.blocks}
        owners = {}
        decoration_ops = {block_id: operation for operation in plan.operations if operation.kind == "decoration" for block_id in operation.block_ids}
        for record in trace.block_render_map:
            decision, block = decisions[record.block_id], blocks[record.block_id]
            if record.content_visible != is_visible_body_block(block):
                raise ValueError("ir_render_evidence:binding_mismatch")
            # collector 按 item 顺序记录最后一个归并代表；混合列表仍拥有独立正文目标。
            representative = next((policy.representative_block_id for policy in reversed(decision.text_policies)
                if policy.mode == "coalesced"), None)
            if record.block_id not in table_members and record.representative_block_id != representative:
                raise ValueError("ir_render_evidence:target_mismatch")
            own_targets = [target for target in record.targets if target.role != "toc_representative"]
            if decision.mode == "native_text":
                expected = [(i, "list_item" if block.type == "list" else "formula" if block.type == "formula" else "body")
                    for i, policy in enumerate(decision.text_policies) if policy.mode != "coalesced"]
                if [(target.item_index, target.role) for target in own_targets] != expected or any(target.element_kind != "paragraph" for target in own_targets):
                    raise ValueError("ir_render_evidence:target_mismatch")
                coalesced = [i for i, policy in enumerate(decision.text_policies) if policy.mode == "coalesced"]
                if [target.item_index for target in record.targets if target.role == "toc_representative"] != coalesced:
                    raise ValueError("ir_render_evidence:target_mismatch")
                if decision.text_policies and record.disposition != ("rendered" if expected else "coalesced"):
                    raise ValueError("ir_render_evidence:target_mismatch")
                if block.type == "list":
                    expected_items = [(i, "coalesced" if policy.mode == "coalesced" else "rendered",
                        policy.representative_block_id, policy.representative_item_index) for i, policy in enumerate(decision.text_policies)]
                    if [(item.item_index, item.disposition, item.representative_block_id, item.representative_item_index) for item in record.items] != expected_items:
                        raise ValueError("ir_render_evidence:target_mismatch")
            elif decision.table_strategy is not None and record.block_id not in table_members:
                caption_count = len(dict.fromkeys(text.strip() for text in block.captions if text.strip()))
                note_count = len(dict.fromkeys(text.strip() for text in block.footnotes if text.strip()))
                roles = ["context" if style == "normal" else "caption" for style in decision.annotation_styles[:caption_count]]
                roles += ["body" if decision.table_strategy == "fallback" else "table"] + ["note"] * note_count
                if v2:
                    from app.services.document_ir.table_placeholder import decode_placeholders
                    placeholders = [item for item in decode_placeholders(document_ir) if item.owner_block_id == block.id]
                    if placeholders:
                        roles = roles[:caption_count] + (["body"] if block.text and block.text not in block.captions else [])
                        roles += ["image_fallback"] * len(placeholders) + ["note"] * note_count
                if [target.role for target in own_targets] != roles:
                    raise ValueError("ir_render_evidence:target_mismatch")
                if any(target.element_kind != ("table" if target.role == "table" else "paragraph") for target in own_targets):
                    raise ValueError("ir_render_evidence:target_mismatch")
            elif decision.image_strategy is not None:
                if not own_targets or own_targets[0].role not in {"image", "image_fallback"}:
                    raise ValueError("ir_render_evidence:target_mismatch")
                if own_targets[0].role == "image" and own_targets[0].element_kind != "drawing":
                    raise ValueError("ir_render_evidence:target_mismatch")
                if any(target.role != "caption" for target in own_targets[1:]):
                    raise ValueError("ir_render_evidence:target_mismatch")
                if decision.image_strategy == "bbox_relative":
                    expected_count = len(decision.annotation_styles)
                    if len(own_targets) != expected_count + 1 or any(target.element_kind != "paragraph" for target in own_targets[1:]):
                        raise ValueError("ir_render_evidence:target_mismatch")
                elif decision.annotation_styles and len(own_targets) <= 1:
                    raise ValueError("ir_render_evidence:target_mismatch")
            if (record.decision_id, record.anchor_ref, record.provenance) != (decision.decision_id, decision.anchor_ref, provenance[block.id]):
                raise ValueError("ir_render_evidence:binding_mismatch")
            if record.disposition in {"rendered", "coalesced"} and not record.targets and record.block_id not in table_members:
                raise ValueError("ir_render_evidence:target_mismatch")
            if record.disposition in {"suppressed", "superseded", "not_emitted"} and record.targets:
                raise ValueError("ir_render_evidence:target_mismatch")
            if record.disposition == "not_emitted" and not (decision.mode == "unassigned" or block.type == "list" and not block.list_items):
                raise ValueError("ir_render_evidence:target_mismatch")
            if record.disposition == "not_emitted" and is_visible_body_block(block):
                raise ValueError("ir_render_evidence:target_mismatch")
            if record.disposition == "suppressed":
                operation = decoration_ops.get(block.id)
                expected = "custom_footer" if operation and operation.decoration_policy == "custom" else "decoration_disabled"
                if operation is None or operation.decoration_policy == "replace" or record.reason_code != expected:
                    raise ValueError("ir_render_evidence:target_mismatch")
            if record.disposition == "superseded":
                operation = decoration_ops.get(block.id)
                current, seen = record, set()
                while current.disposition == "superseded":
                    if current.block_id in seen or current.replacement_block_id not in records:
                        raise ValueError("ir_render_evidence:target_mismatch")
                    seen.add(current.block_id)
                    current = records[current.replacement_block_id]
                if operation is None or operation.decoration_policy != "replace" or current.disposition != "rendered":
                    raise ValueError("ir_render_evidence:target_mismatch")
                section_index = int(operation.section_id.split("_")[1])
                if not any(t.role == operation.decoration_kind and section_index in t.actual["referenced_section_indices"] for t in current.targets):
                    raise ValueError("ir_render_evidence:target_mismatch")
            if block.type == "list" and len(record.items) != len(block.list_items or []):
                raise ValueError("ir_render_evidence:target_mismatch")
            for target in record.targets:
                if any(type(value) is not int or value < 0 for value in target.child_path):
                    raise ValueError("ir_render_evidence:target_mismatch")
                if v2 and target.element_kind == "table_rows":
                    from app.services.renderer.content_contract import project_table_scope
                    scope = resolve_table_scope(trace, target.content_scope_id)
                    if scope.block_id != block.id or (scope.part, scope.table_child_path) != (target.part, target.child_path):
                        raise ValueError("ir_render_evidence:target_mismatch")
                    actual, bookmarks = project_table_scope(index, scope)
                    kind = "table_rows"
                else:
                    if v2 and target.content_scope_id is not None:
                        raise ValueError("ir_render_evidence:target_mismatch")
                    kind, actual, bookmarks = index.project(target.part, tuple(target.child_path))
                if (kind, actual, bookmarks) != (target.element_kind, target.actual, target.bookmark_names):
                    raise ValueError("ir_render_evidence:target_mismatch")
                if target.role == "toc_representative":
                    item = target.item_index
                    if item is None or item >= len(decision.text_policies):
                        raise ValueError("ir_render_evidence:target_mismatch")
                    policy = decision.text_policies[item]
                    representative = records.get(policy.representative_block_id)
                    if policy.mode != "coalesced" or representative is None or not any(
                        (t.part, t.child_path, t.item_index) == (target.part, target.child_path, policy.representative_item_index)
                        and t.role != "toc_representative" for t in representative.targets):
                        raise ValueError("ir_render_evidence:target_mismatch")
                    continue
                section = int(decision.section_id.split("_")[1]) if decision.section_id is not None else None
                if section not in actual["referenced_section_indices"]:
                    raise ValueError("ir_render_evidence:target_mismatch")
                if decision.mode == "decoration":
                    op = decoration_ops[block.id]
                    group = "decoration:" + content_hash([op.decoration_kind, op.block_ids])
                    if target.role != op.decoration_kind or target.sharing_group != group or target.participant_order != op.block_ids.index(block.id):
                        raise ValueError("ir_render_evidence:target_mismatch")
                elif target.sharing_group is not None or target.part != "word/document.xml":
                    raise ValueError("ir_render_evidence:target_mismatch")
                locator = (target.part, tuple(target.child_path), target.content_scope_id if v2 else None)
                other = owners.get(locator)
                if other is not None and not (target.sharing_group is not None and other == target.sharing_group):
                    raise ValueError("ir_render_evidence:target_mismatch")
                owners[locator] = target.sharing_group or block.id
            names = list(dict.fromkeys(name for t in record.targets for name in t.bookmark_names))
            if record.bookmark_name != (names[0] if names else None):
                raise ValueError("ir_render_evidence:target_mismatch")
        # 每图/每条注释只拥有其真实 cell 或正文对象，不能以表计数或全局 OMML 集合替代归属。
        image_ids = [block_id for operation in plan.operations if operation.kind == "block"
            for block_id in operation.block_ids if blocks[block_id].type == "image"]
        if require_ownership:
            if image_ids != [image.block_id for image in ownership.images]:
                raise ValueError("ir_render_evidence:target_mismatch")
            for image in ownership.images:
                actual = [(tuple(target.child_path), target.element_kind, target.role, target.item_index)
                    for target in records[image.block_id].targets]
                if image.part != "word/document.xml" or any(target.part != image.part for target in records[image.block_id].targets) or tuple(actual) != image.targets:
                    raise ValueError("ir_render_evidence:target_mismatch")
        else:
            # 恢复沿 RND-503 已认证的原图/注释 marker 核对位置，不声称重新获得创建事实。
            _validate_persisted_image_targets(index, records, blocks, image_ids)
        # 实际输出顺序与计划一致，交换两个合法段落 locator 也不能通过。
        previous = None
        for operation in plan.operations:
            if v2 and operation.kind == "logical_table":
                own = [tuple(target.child_path) for block_id in operation.block_ids
                       for target in records[block_id].targets]
                if own:
                    if previous is not None and min(own) <= previous:
                        raise ValueError("ir_render_evidence:target_mismatch")
                    previous = max(own)
                continue
            if operation.kind not in {"block", "image_pair"}:
                continue
            for block_id in operation.block_ids:
                own = [tuple(t.child_path) for t in records[block_id].targets if t.role != "toc_representative"]
                if own:
                    position = min(own)
                    if previous is not None and position <= previous:
                        raise ValueError("ir_render_evidence:target_mismatch")
                    previous = max(own) if v2 else position
        counts = {state: sum(r.disposition == state for r in records.values()) for state in ("rendered", "coalesced", "suppressed", "superseded", "not_emitted")}
        counts.update(block_count=len(records), target_count=sum(len(r.targets) for r in records.values()))
        if trace.counts != counts:
            raise ValueError("ir_render_evidence:block_coverage_mismatch")
        return plan, trace
    except DocumentIntegrityError:
        raise
    except (ValueError, TypeError, KeyError, AttributeError, IndexError, UnicodeError) as exc:
        raise render_contract_failure(exc) from None


def build_renderer_evidence(document_ir: DocumentIR, word_content: bytes, layout_plan: LayoutPlan | None = None,
                            *, render_result: WordRenderResult | None = None, formula_report: dict[str, Any] | None = None,
                            image_layout_report: dict[str, Any] | None = None,
                            expected_template_version_id: str | None = None,
                            expected_generation_policy: str | None = None,
                            render_plan: RenderPlan | dict[str, Any] | None = None,
                            block_render_evidence: BlockRenderEvidence | dict[str, Any] | None = None,
                            block_anchor_evidence: BlockAnchorEvidence | dict[str, Any] | None = None,
                            renderer_content_contract_version: Any = _UNSET,
                            formula_source_coverage: Any = _UNSET) -> dict[str, Any]:
    """原内容认证保持严格；恢复时只消费已存 trace，可选锚点失败局部降级。"""
    if render_result is not None:
        if formula_report is not None and formula_report != render_result.formula_report or image_layout_report is not None and image_layout_report != render_result.image_layout_report:
            raise render_contract_failure(ValueError("ir_render_evidence:binding_mismatch"))
        formula_report, image_layout_report = render_result.formula_report, render_result.image_layout_report
        if (renderer_content_contract_version is not _UNSET
                and renderer_content_contract_version != render_result.renderer_content_contract_version
                or formula_source_coverage is not _UNSET and formula_source_coverage != render_result.formula_source_coverage):
            raise render_contract_failure(ValueError("ir_render_evidence:binding_mismatch"))
        renderer_content_contract_version = render_result.renderer_content_contract_version
        formula_source_coverage = render_result.formula_source_coverage
    from app.services.renderer.content_contract import CONTENT_VERSION, build_v2_content_evidence, content_failure
    v2 = renderer_content_contract_version is not _UNSET
    if v2:
        if renderer_content_contract_version != CONTENT_VERSION or formula_source_coverage is _UNSET:
            raise content_failure()
        evidence = build_v2_content_evidence(document_ir, word_content, layout_plan, coverage=formula_source_coverage,
            formula_report=formula_report, image_report=image_layout_report,
            expected_template_version_id=expected_template_version_id, expected_generation_policy=expected_generation_policy)
    else:
        if formula_source_coverage is not _UNSET:
            raise content_failure()
        evidence = build_renderer_aggregate_evidence(document_ir, word_content, layout_plan,
            formula_report=formula_report, image_layout_report=image_layout_report,
            expected_template_version_id=expected_template_version_id, expected_generation_policy=expected_generation_policy)
    if render_result is not None:
        if v2 and (render_result.plan.schema_version != "render-plan/v2"
                   or render_result.block_evidence.schema_version != "block-render-evidence/v2"):
            raise content_failure()
        validate_render_result(document_ir, word_content, render_result)
        plan, trace = render_result.plan, render_result.block_evidence
        block_anchor_evidence = render_result.block_anchor_evidence
    else:
        restored = validate_persisted_render_evidence(document_ir, word_content,
            render_plan=render_plan, block_render_evidence=block_render_evidence) if render_plan is not None and block_render_evidence is not None else None
        plan, trace = restored if restored is not None else (None, None)
        if plan is not None and (plan.schema_version == "render-plan/v2") != v2:
            plan, trace = None, None
    if plan is None:
        evidence["render_trace_status"] = "unavailable"
    else:
        evidence.update(render_trace_status="complete", render_plan=plan.model_dump(mode="json"),
            block_render_evidence=trace.model_dump(mode="json"))
        from app.services.renderer.block_anchors import validate_block_anchor_evidence
        # 新件缺运行期创建事实时不可借名字或持久化 reader 追认；历史件只读原证明。
        ownership = render_result._anchor_ownership if render_result is not None else None
        anchors = validate_block_anchor_evidence(block_anchor_evidence, plan=plan, block_evidence=trace,
            word_content=word_content, ownership=ownership) if render_result is None or ownership is not None else None
        if anchors is not None:
            evidence["block_anchor_evidence"] = anchors.model_dump(mode="json")
        elif v2 and render_result is not None:
            raise content_failure("ir_render_evidence:target_mismatch")
        if render_result is not None:
            evidence["block_anchor_summary"] = {"schema_version": "block-anchor-evidence/v2" if v2 else "block-anchor-evidence/v1",
                **{key: count for key, count in render_result.block_anchor_counts.items()
                   if key in {"created", "name_conflict", "target_unavailable", "limit_exceeded", "generation_disabled"}
                   and type(count) is int and count >= 0}}
    return evidence


def validate_persisted_renderer_content(document_ir: DocumentIR, word_content: bytes, metadata: dict[str, Any], *,
                                        expected_template_version_id: str | None = None) -> dict[str, Any]:
    """共同文件门之后按显式 mandatory 字段分派；可选导航不参与主件内容比较。"""
    from copy import deepcopy
    from app.services.renderer.content_contract import CONTENT_VERSION, content_failure
    from app.services.renderer import image_evidence
    v2 = "renderer_content_contract_version" in metadata
    if not v2 and "formula_source_coverage" in metadata:
        raise content_failure()
    if v2:
        if metadata["renderer_content_contract_version"] != CONTENT_VERSION:
            raise content_failure()
        actual = build_renderer_evidence(document_ir, word_content,
            renderer_content_contract_version=metadata["renderer_content_contract_version"],
            formula_source_coverage=metadata.get("formula_source_coverage"),
            formula_report=metadata.get("formula_report"), image_layout_report=metadata.get("image_layout_report"),
            expected_template_version_id=expected_template_version_id,
            render_plan=metadata.get("render_plan"), block_render_evidence=metadata.get("block_render_evidence"),
            block_anchor_evidence=metadata.get("block_anchor_evidence"))
        ensure_renderer_integrity(actual)
        optional = {"render_trace_status", "render_plan", "block_render_evidence", "block_anchor_evidence", "block_anchor_summary"}
        if any(key not in metadata or metadata[key] != value for key, value in actual.items() if key not in optional):
            raise content_failure("ir_render_evidence:binding_mismatch")
        return actual
    if image_evidence.has_image_markers(word_content) and not isinstance(metadata.get("image_layout_report"), dict):
        raise ValueError("主 DOCX 含图片布局标记但缺少报告。")
    for key in ("formula_report", "image_layout_report"):
        if key in metadata and not isinstance(metadata[key], dict):
            raise ValueError("主 DOCX 的既有内容报告无效。")
    table_plan = None
    placeholders = []
    table_contract = any(key in metadata for key in ("table_rendering", "table_placeholder_rendering", "table_evidence", "table_placeholder_count"))
    if table_contract:
        from app.services.renderer.logical_table_plan import build_logical_table_render_plan
        from app.services.document_ir.table_placeholder import decode_placeholders
        table_plan = build_logical_table_render_plan(document_ir, build_layout_plan(document_ir))
        placeholders = decode_placeholders(document_ir)
    actual = build_renderer_aggregate_evidence(document_ir, word_content,
        formula_report=metadata.get("formula_report"), image_layout_report=metadata.get("image_layout_report"),
        expected_template_version_id=expected_template_version_id, _table_plan=table_plan,
        _placeholder_count=len(placeholders), _legacy_ir_identity=True)
    if table_contract:
        from app.services.renderer.table_content_evidence import _table_rendering_evidence, _build_table_evidence
        from app.services.renderer.table_placeholder_evidence import build_placeholder_evidence
        actual["table_placeholder_count"] = len(placeholders)
        if table_plan is not None:
            actual["table_rendering"], flags = _table_rendering_evidence(document_ir, word_content, table_plan,
                actual["docx_quality"]["table_observations"])
            actual["blocking_quality_flags"] = list(dict.fromkeys([*actual["blocking_quality_flags"], *flags]))
        proof = build_placeholder_evidence(document_ir, word_content)
        if proof is not None:
            actual["table_placeholder_rendering"] = proof
        if "table_evidence" in metadata and not actual["blocking_quality_flags"]:
            actual["table_evidence"] = _build_table_evidence(document_ir, word_content, actual).model_dump(mode="json", warnings=False)
    ensure_renderer_integrity(actual)
    comparable = deepcopy(actual)
    if isinstance(metadata.get("docx_quality"), dict) and "table_observations" not in metadata["docx_quality"]:
        comparable["docx_quality"].pop("table_observations", None)
    keys = LEGACY_RENDERER_EVIDENCE_KEYS | {key for key in comparable if key in metadata}
    if any(key not in metadata or metadata[key] != comparable[key] for key in keys):
        raise ValueError("主 DOCX 的既有质量证据与实际字节不一致。")
    return actual


CONTROLLED_SOURCE_TYPES = frozenset({"text", "paragraph", "para", "title", "heading", "header", "footer", "page_header", "page_footer", "page_number", "table", "table_placeholder", "image", "chart", "figure", "list", "index", "equation", "formula", "interline_equation", "inline_equation", "equation_interline", "unsupported", "invalid", "unknown"})
FIXED_FLAGS = frozenset({"formula_text_fallback", "image_group_uncertain", "caption_binding_uncertain", "docx_quality_unavailable", "docx_analysis_error", "docx_no_editable_content", "docx_image_only_suspected", "middle_json_missing", "empty_content_list",
    "source_ir_integrity:block_count_mismatch", "source_ir_integrity:type_count_mismatch", "ir_layout_integrity:missing_blocks", "ir_layout_integrity:duplicate_blocks", "ir_docx_integrity:image_asset_missing", "ir_docx_integrity:inline_shape_count_mismatch", "ir_docx_integrity:table_count_mismatch", "ir_docx_integrity:heading_count_mismatch",
    "ir_docx_integrity:table_anchor_mismatch", "ir_docx_integrity:table_row_count_mismatch", "ir_docx_integrity:table_header_mismatch",
    "ir_docx_integrity:table_geometry_mismatch", "ir_docx_integrity:table_fragment_anchor_mismatch",
    "ir_docx_integrity:table_fragment_map_mismatch", "table_placeholder_integrity_failed", "table_noneditable_placeholder"
}) | RENDER_EVIDENCE_CODES | QUALITY_BLOCKERS | QUALITY_WARNINGS
COUNT_FLAG_RE = re.compile(r"(?:docx_latex_residue|docx_manual_list_markers|docx_plain_toc_entries|docx_unstyled_heading_candidates|final_promotional_page_filtered):[0-9]+\Z")


def controlled_quality_flags(value: Any) -> list[str]:
    """只返回已定义的诊断码；自由 block/fragment 身份留在私有证据。"""
    if not isinstance(value, list):
        return []
    result = []
    for flag in value:
        if not isinstance(flag, str):
            continue
        parts = flag.split(":")
        if (flag in FIXED_FLAGS or COUNT_FLAG_RE.fullmatch(flag)
                or len(parts) == 3 and parts[0] == "source_ir_integrity" and parts[1] in {"type_mismatch", "payload_missing"} and parts[2] in CONTROLLED_SOURCE_TYPES):
            result.append(flag)
        elif parts[0] in {"table_continuation_uncertain", "layout_bbox_missing", "table_continuation", "layout_two_column_candidate", "layout_wide_table", "layout_table_context_binding", "layout_caption_binding", "layout_footnote_binding", "unsupported_block_type"}:
            result.append(parts[0])
    return list(dict.fromkeys(result))


NUMERIC_KEYS = frozenset("expected_layout_table_count layout_table_count layout_table_count_delta expected_content_table_count content_table_count content_table_count_delta expected_total_table_count total_table_count total_table_count_delta layout_missing_block_count layout_duplicate_block_count page_count block_count layout_page_count layout_region_count layout_block_count expected_layout_block_count layout_block_count_delta expected_inline_shape_count inline_shape_count inline_shape_count_delta expected_table_count table_count table_count_delta expected_heading_count heading_count heading_count_delta expected_source_block_count ir_main_block_count source_block_count_delta paragraph_count wide_table_count media_count heading_paragraphs heading_like_paragraphs heading_like_unstyled_paragraphs numbering_paragraphs manual_list_marker_paragraphs toc_tab_leader_paragraphs plain_toc_entry_paragraphs keep_next_paragraphs table_header_rows table_no_split_rows shaded_table_cells table_cell_zero_spacing_paragraphs landscape_sections latex_residue_count text_chars".split())
BOOLEAN_KEYS = frozenset({"editable_text_available", "editable_content_available", "image_only_suspected"})


def _safe_counts(value):
    if not isinstance(value, dict):
        return {}
    return {key: item for key, item in value.items() if key in NUMERIC_KEYS and (type(item) is int or item is None)
        or key in BOOLEAN_KEYS and type(item) is bool}


def _type_counts(value):
    result = {}
    for key, count in value.items() if isinstance(value, dict) else []:
        if type(count) is int and count >= 0:
            label = key if key in CONTROLLED_SOURCE_TYPES else "unknown"
            result[label] = result.get(label, 0) + count
    return result


def build_renderer_event_summary(evidence: dict[str, Any]) -> dict[str, Any]:
    """完整字段白名单；新增私有 metadata 默认不会进入事件/异常。"""
    summary = _safe_counts(evidence)
    if evidence.get("renderer_content_contract_version") == "renderer-content/v2":
        summary["renderer_content_contract_version"] = "renderer-content/v2"
    if evidence.get("schema_version") in {"document-ir/v1", "document-ir/v2"}:
        summary["schema_version"] = evidence["schema_version"]
    if evidence.get("renderer_mode") in {"layout_v2", "legacy_v1_fallback"}:
        summary["renderer_mode"] = evidence["renderer_mode"]
    for key in ("quality_flags", "layout_quality_flags", "blocking_quality_flags"):
        if key in evidence:
            value = evidence[key]
            flags = controlled_quality_flags(value)
            summary[key] = flags
            previous = evidence.get(key + "_unknown_count")
            summary[key + "_unknown_count"] = previous if type(previous) is int and previous >= 0 else (sum(isinstance(item, str) and not controlled_quality_flags([item]) for item in value) if isinstance(value, list) else 0)
    for key in ("layout_missing_block_ids", "layout_duplicate_block_ids"):
        if isinstance(evidence.get(key), list):
            summary[key.removesuffix("_ids") + "_count"] = len(evidence[key])
    for key in ("source_integrity", "docx_quality"):
        value = evidence.get(key)
        if isinstance(value, dict):
            projected = _safe_counts(value)
            if key == "source_integrity":
                projected["blocking_quality_flags"] = controlled_quality_flags(value.get("blocking_quality_flags"))
                for field in ("expected_source_type_counts", "ir_source_type_counts"):
                    projected[field] = _type_counts(value.get(field))
            elif value.get("analysis_error"):
                projected["analysis_error"] = "document_xml_parse_error"
            summary[key] = projected
    if "mineru_validation" in evidence:
        summary["mineru_validation"] = _safe_mineru_report(evidence["mineru_validation"])
    if "mineru_validation_input" in evidence:
        value = evidence["mineru_validation_input"]
        summary["mineru_validation_input"] = _safe_validation_input(value)
    if evidence.get("render_trace_status") in {"complete", "unavailable"}:
        summary["render_trace_status"] = evidence["render_trace_status"]
    for key in ("input_ir_sha256", "plan_sha256", "docx_sha256"):
        value = evidence.get(key)
        if isinstance(value, str) and re.fullmatch(r"[0-9a-f]{64}", value):
            summary[key] = value
    for key, versions in (("render_plan_schema_version", {"render-plan/v1", "render-plan/v2"}),
                          ("block_render_evidence_schema_version", {"block-render-evidence/v1", "block-render-evidence/v2"})):
        if evidence.get(key) in versions | {"unavailable"}:
            summary[key] = evidence[key]
    counts = evidence.get("render_trace_counts")
    if isinstance(counts, dict):
        summary["render_trace_counts"] = {k: v for k, v in counts.items() if k in {"rendered", "coalesced", "suppressed", "superseded", "not_emitted", "block_count", "target_count"} and type(v) is int and v >= 0}
    plan, trace = evidence.get("render_plan"), evidence.get("block_render_evidence")
    if isinstance(plan, dict) and isinstance(trace, dict):
        for key in ("input_ir_sha256", "plan_sha256", "docx_sha256"):
            value = trace.get(key)
            if isinstance(value, str) and re.fullmatch(r"[0-9a-f]{64}", value):
                summary[key] = value
        summary["render_plan_schema_version"] = plan["schema_version"] if plan.get("schema_version") in {"render-plan/v1", "render-plan/v2"} else "unavailable"
        summary["block_render_evidence_schema_version"] = trace["schema_version"] if trace.get("schema_version") in {"block-render-evidence/v1", "block-render-evidence/v2"} else "unavailable"
        counts = trace.get("counts")
        if isinstance(counts, dict):
            summary["render_trace_counts"] = {k: v for k, v in counts.items() if k in {"rendered", "coalesced", "suppressed", "superseded", "not_emitted", "block_count", "target_count"} and type(v) is int and v >= 0}
    anchor_summary = evidence.get("block_anchor_summary")
    if isinstance(anchor_summary, dict) and anchor_summary.get("schema_version") in {"block-anchor-evidence/v1", "block-anchor-evidence/v2"}:
        summary["block_anchor_summary"] = {"schema_version": anchor_summary["schema_version"],
            **{key: count for key, count in anchor_summary.items()
               if key in {"created", "name_conflict", "target_unavailable", "limit_exceeded", "generation_disabled"}
               and type(count) is int and count >= 0}}
    from app.conversion.docx_rules.formula import FORMULA_REASONS
    report_fields = (
        ("formula_summary", "formula-summary/v1", {"formula_count", "omml_count", "semantic_count", "fallback_count"}, FORMULA_REASONS),
        ("image_layout_summary", "image-layout/v1", {"accepted_group_count", "single_image_count", "uncertain_span_count", "omitted_example_count"},
         {"geometry_unavailable", "geometry_invalid", "page_or_section_boundary", "source_order_conflict", "row_alignment_insufficient",
          "horizontal_overlap", "gap_too_large", "competing_row_membership", "caption_binding_uncertain", "external_caption_binding", "asset_unavailable", "image_size_unfit"}),
    )
    for key, version, numeric_fields, reasons in report_fields:
        value = evidence.get(key)
        if not isinstance(value, dict) or value.get("schema_version") != version:
            continue
        safe = {"schema_version": version, **{name: item for name, item in value.items()
            if name in numeric_fields and type(item) is int and item >= 0}}
        digest = value.get("report_sha256")
        if isinstance(digest, str) and re.fullmatch(r"[0-9a-f]{64}", digest):
            safe["report_sha256"] = digest
        counts = value.get("reason_counts")
        if isinstance(counts, dict):
            safe["reason_counts"] = {reason: count for reason, count in counts.items()
                if reason in reasons and type(count) is int and count >= 0}
        if key == "formula_summary":
            safe["warning_codes"] = controlled_quality_flags(value.get("warning_codes"))
        summary[key] = safe
    native = evidence.get("table_rendering")
    if isinstance(native, dict) and native.get("version") == TABLE_RENDERING_VERSION:
        status = native.get("quality_status")
        summary["table_rendering"] = {"version": TABLE_RENDERING_VERSION,
            "quality_status": status if isinstance(status, str) and status in {"passed", "warning", "blocked"} else None,
            "blocking_quality_flags": controlled_quality_flags(native.get("blocking_quality_flags")),
            "warning_flags": controlled_quality_flags(native.get("warning_flags")),
            **{key: native[key] for key in SAFE_COUNT_KEYS if type(native.get(key)) is int and native[key] >= 0}}
    proof = evidence.get("table_placeholder_rendering")
    if isinstance(proof, dict) and proof.get("version") == "table-placeholder-rendering/v1":
        summary["table_placeholder_rendering"] = {"version": "table-placeholder-rendering/v1",
            "count": proof.get("count") if type(proof.get("count")) is int and proof["count"] >= 0 else None,
            "issue_codes": controlled_quality_flags(proof.get("issue_codes"))}
    table_report = evidence.get("table_evidence")
    if isinstance(table_report, dict) and table_report.get("schema_version") == "table-evidence/v1":
        from app.services.tables.evidence import ISSUE_CODES
        counts = table_report.get("counts")
        count_source = counts if isinstance(counts, dict) else table_report
        issues = table_report.get("issue_codes")
        if not isinstance(issues, list):
            tables = table_report.get("tables")
            issues = []
            for table in tables if isinstance(tables, list) else []:
                flags = table.get("flags") if isinstance(table, dict) else None
                if isinstance(flags, list):
                    issues.extend(flag.get("code") for flag in flags if isinstance(flag, dict))
        summary["table_evidence"] = {"schema_version": "table-evidence/v1",
            **{key: table_report[key] for key, allowed in (("coverage", {"complete", "unavailable"}),
               ("status", {"pending", "passed", "warning", "blocked"}))
               if isinstance(table_report.get(key), str) and table_report[key] in allowed},
            **{key: value for key, value in count_source.items() if key in {"table_count", "fragment_count", "native_count",
                "placeholder_count", "blocker_count", "warning_count"} and type(value) is int and value >= 0},
            "issue_codes": list(dict.fromkeys(code for code in issues if isinstance(code, str) and code in ISSUE_CODES))}
    return summary


# INPUT-006 的安全输出合同；不接受 schema 内自由字典中的任意机器字符串。
REPORT_CODES = frozenset({
    "mineru_package_invalid", "mineru_package_limit_exceeded", "mineru_path_invalid", "mineru_content_list_missing",
    "mineru_json_invalid", "mineru_content_empty", "mineru_page_invalid", "mineru_page_missing", "mineru_page_empty",
    "mineru_page_gap", "mineru_block_invalid", "mineru_block_payload_invalid", "unsupported_block_type",
    "legacy_page_index_coerced", "mineru_page_order_nonmonotonic", "mineru_bbox_missing", "mineru_bbox_invalid",
    "report_limit_exceeded", "mineru_margin_payload_empty", "missing_image_asset", "invalid_image_asset",
    "missing_table_reference", "invalid_table_reference", "empty_table_html", "invalid_table_html", "suspicious_table_html",
    "image_asset_normalized", "image_asset_pdf_fallback",
})
REPORT_REASONS = frozenset({
    "absolute_or_traversal", "actual_expanded_bytes", "ambiguous_asset_reference", "conflicting_index",
    "contentful_table_placeholder", "duplicate_normalized_path", "encrypted_entry", "entry_count",
    "expected_asset_path_string", "expected_list", "expected_runs", "expected_string", "expected_type_string",
    "file_directory_conflict", "image_bytes", "invalid", "invalid_block", "invalid_content_node", "invalid_content_type",
    "invalid_index", "invalid_page", "invalid_path_syntax", "invalid_span", "invalid_text_node", "invalid_type",
    "isolated_unicode_surrogate", "issue_details", "json_bytes", "missing", "missing_content_list", "missing_index",
    "missing_or_invalid_type", "missing_readable_formula", "missing_readable_payload", "missing_readable_table_payload",
    "mixed_or_invalid_page", "nested_table", "no_blocks", "no_consumable_rows", "non_finite", "non_numeric",
    "non_positive_area", "not_allowed_raster_format", "not_object", "outside_package_root", "outside_selected_list_payload",
    "outside_selected_rich_payload", "report_bytes", "single_file_bytes", "special_zip_entry", "symlink", "total_bytes",
    "type_too_long", "unaccounted_normalization", "unreadable_archive", "unreadable_json", "unreadable_or_unsupported_image",
    "unsupported_tag", "unsupported_top_level", "wrong_length_or_type", "final_promotional_page",
    "explicit_table_placeholder", "strict_empty_table_placeholder", "word_incompatible_raster",
})
REPORT_DETAIL_VALUES = {
    "code": REPORT_CODES,
    "severity": {"warning", "blocker"},
    "classification": {"html_present_and_parseable", "html_present_but_suspicious", "html_missing_reference_present", "html_missing_reference_missing"},
    "reference_status": {"valid", "missing", "invalid", "unsafe"},
    "reference_resolution": {"package_root", "content_list_directory"},
    "rule_id": {"final_promotional_page/v1", "mineru_empty_table_continuation/v1"},
    "reason": REPORT_REASONS,
}
REPORT_VERSIONS = {
    "schema_version": {"mineru-artifact-validation/v1"},
    "validator_version": {"mineru-artifact-validator/v1", "mineru-artifact-validator/v2", "mineru-artifact-validator/v3"},
    "policy_version": {"mineru-input-policy/v1", "mineru-input-policy/v2"},
}
REPORT_NUMERIC_KEYS = frozenset("page_index source_index_in_page source_index row_count cell_count warning_count blocker_count fatal_count omitted_issue_count table_check_count omitted_table_check_count raw_page_count raw_block_count valid_block_count invalid_block_count invalid_page_count image_check_count table_reference_check_count omitted_type_count block_count filtered_promotional_page_index filtered_block_count merged_table_placeholder_count omitted_merged_table_placeholder_count omitted_filtered_block_count previous_page_index index_gap width height".split())
REPORT_BOOLEAN_KEYS = frozenset({"complete", "can_convert", "details_truncated", "stops_conversion", "parseable", "readable_payload"})
REPORT_LOCATOR_KEYS = frozenset({"json_pointer", "page_index", "source_index_in_page", "source_index", "source_type"})


def _safe_validation_input(value):
    if not isinstance(value, dict):
        return None
    result = {}
    for key in ("mineru_raw_result_file_id", "mineru_raw_result_sha256", "retry_count"):
        item = value.get(key)
        if key == "retry_count" and type(item) is int and item >= 0:
            result[key] = item
        elif key.endswith("file_id") and isinstance(item, str) and re.fullmatch(r"file_[a-zA-Z0-9_-]{1,80}", item):
            result[key] = item
        elif key.endswith("sha256") and isinstance(item, str) and re.fullmatch(r"(?:sha256:)?[0-9a-f]{64}", item):
            result[key] = item
    return result


def _report_locator(value):
    from app.services.document_ir.artifact_validator import REPORT_FIELD_NAMES
    if not isinstance(value, dict):
        return {}
    result = {}
    for key in REPORT_LOCATOR_KEYS:
        item = value.get(key)
        if key in REPORT_NUMERIC_KEYS and (item is None or type(item) is int and item >= 0):
            result[key] = item
        elif key == "source_type" and (item is None or isinstance(item, str) and item in CONTROLLED_SOURCE_TYPES):
            result[key] = item
        elif key == "json_pointer" and isinstance(item, str) and re.fullmatch(r"(?:/(?:content_list|[0-9]+))*", item):
            result[key] = item
    field = value.get("field")
    if field is None:
        if "field" in value:
            result["field"] = None
    elif isinstance(field, str) and all(part in REPORT_FIELD_NAMES or part.isdecimal() or re.fullmatch(r"unknown_[0-9a-f]{12}", part) for part in field.split("/")):
        result["field"] = field
    return result


def _report_counts(value):
    return {key: count for key, count in value.items() if key in REPORT_CODES and type(count) is int and count >= 0} if isinstance(value, dict) else {}


def _report_detail(value):
    if not isinstance(value, dict):
        return {}
    result = _report_locator(value)
    for key in ("code", "severity", "classification", "reference_status", "reference_resolution", "rule_id", "reason"):
        item = value.get(key)
        if item is None and key in value or isinstance(item, str) and item in REPORT_DETAIL_VALUES[key]:
            result[key] = item
    for key in REPORT_NUMERIC_KEYS | REPORT_BOOLEAN_KEYS:
        item = value.get(key)
        if key in value and (key in REPORT_NUMERIC_KEYS and (type(item) is int or item is None) or key in REPORT_BOOLEAN_KEYS and type(item) is bool):
            result[key] = item
    facts = value.get("facts")
    if isinstance(facts, dict):
        safe = {}
        for key, item in facts.items():
            if key in REPORT_NUMERIC_KEYS and type(item) is int or key in REPORT_BOOLEAN_KEYS and type(item) is bool:
                safe[key] = item
            elif key == "reason" and isinstance(item, str) and item in REPORT_REASONS:
                safe[key] = item
            elif key == "actual_type" and isinstance(item, str) and item in {"str", "dict", "list", "int", "float", "bool", "NoneType"}:
                safe[key] = item
            elif key == "reasons" and isinstance(item, list):
                safe[key] = [reason for reason in item if isinstance(reason, str) and reason in REPORT_REASONS]
        result["facts"] = safe
    if isinstance(value.get("target"), dict):
        result["target"] = _report_locator(value["target"])
    return result


def _safe_mineru_report(value):
    """报告 schema 内自由字典仍按 INPUT-006 实际字段/有限值逐层投影。"""
    from app.services.document_ir.artifact_validator import MinerUValidationReport
    from app.services.document_ir.mineru_package import validate_relative_path
    from pydantic import ValidationError
    if value is None:
        return None
    try:
        report = MinerUValidationReport.model_validate(value).model_dump(mode="json")
    except (ValidationError, ValueError, TypeError):
        return None
    result = {}
    for key in ("schema_version", "validator_version", "policy_version"):
        if isinstance(report[key], str) and report[key] in REPORT_VERSIONS[key]:
            result[key] = report[key]
    for key in ("package_sha256", "content_list_sha256"):
        item = report[key]
        if item is None or re.fullmatch(r"[0-9a-f]{64}", item):
            result[key] = item
    # 这是 INPUT-006 明确允许的包内路径，不接受服务器绝对路径或 object key。
    item = report["content_list_path"]
    if item is None:
        result["content_list_path"] = None
    else:
        try:
            validate_relative_path(item)
            result["content_list_path"] = item
        except ValueError:
            pass
    for key in REPORT_NUMERIC_KEYS | REPORT_BOOLEAN_KEYS:
        if key in report:
            result[key] = report[key]
    for key in ("warnings", "blockers", "table_checks"):
        result[key] = [_report_detail(item) for item in report[key]]
    result["first_fatal"] = _report_detail(report["first_fatal"]) if report["first_fatal"] is not None else None
    result["issue_counts"] = {severity: _report_counts(report["issue_counts"].get(severity)) for severity in ("warning", "blocker")}
    for key in ("inventory", "effective_inventory"):
        value = report[key]
        safe = {k: v for k, v in value.items() if k in REPORT_NUMERIC_KEYS and type(v) is int}
        if isinstance(value.get("input_shape"), str) and value["input_shape"] in {"flat", "nested", "wrapper_flat", "wrapper_nested"}:
            safe["input_shape"] = value["input_shape"]
        if "type_counts" in value:
            safe["type_counts"] = _type_counts(value["type_counts"])
        result[key] = safe
    value = report["normalization"]
    safe = {k: v for k, v in value.items() if k in REPORT_NUMERIC_KEYS and (type(v) is int or v is None)}
    for key in ("filtered_blocks", "merged_table_placeholders"):
        if isinstance(value.get(key), list):
            safe[key] = [_report_detail(item) for item in value[key]]
    if isinstance(value.get("resolved_issue_counts"), dict):
        safe["resolved_issue_counts"] = {"blocker": _report_counts(value["resolved_issue_counts"].get("blocker"))}
    result["normalization"] = safe
    return result


def is_valid_table_rendering_evidence(value: Any) -> bool:
    """完整证据在登记前再次核对覆盖，防止中途丢 map 被根 flag 空列表掩盖。"""
    if not isinstance(value, dict) or value.get("version") != TABLE_RENDERING_VERSION:
        return False
    if not isinstance(value.get("quality_status"), str) or value["quality_status"] not in {"passed", "warning", "blocked"}:
        return False
    blockers, warnings = value.get("blocking_quality_flags"), value.get("warning_flags")
    if (not isinstance(blockers, list) or any(not isinstance(flag, str) or flag not in QUALITY_BLOCKERS for flag in blockers) or
            not isinstance(warnings, list) or any(not isinstance(flag, str) or flag not in QUALITY_WARNINGS for flag in warnings)):
        return False
    if value["quality_status"] != ("blocked" if blockers else "warning" if warnings else "passed"):
        return False
    if any(type(value.get(key)) is not int or value[key] < 0 for key in SAFE_COUNT_KEYS):
        return False
    tables = value.get("tables")
    if not isinstance(tables, list) or len(tables) != value["output_table_count"]:
        return False
    fragments = []
    local_blockers: set[str] = set()
    local_warnings: set[str] = set()
    for table in tables:
        if not isinstance(table, dict) or not isinstance(table.get("fragments"), list):
            return False
        table_blockers, table_warnings = table.get("blocking_quality_flags"), table.get("warning_flags")
        if (
            not isinstance(table_blockers, list)
            or any(not isinstance(flag, str) or flag not in QUALITY_BLOCKERS for flag in table_blockers)
            or not isinstance(table_warnings, list)
            or any(not isinstance(flag, str) or flag not in QUALITY_WARNINGS for flag in table_warnings)
            or table.get("group_status") != ("blocked" if table_blockers else "warning" if table_warnings else "passed")
        ):
            return False
        local_blockers.update(table_blockers)
        local_warnings.update(table_warnings)
        if (type(table.get("expected_rows")) is not int or table["expected_rows"] <= 0 or
                table.get("actual_rows") != table["expected_rows"] or
                table.get("actual_header_rows") != table.get("expected_header_rows")):
            return False
        cursor = 0
        has_missing_fragment = False
        for fragment in table["fragments"]:
            if not isinstance(fragment, dict) or not isinstance(fragment.get("fragment_id"), str):
                return False
            status = fragment.get("status")
            if not isinstance(status, str) or status not in {"rendered", "header_only", "empty", "missing"}:
                return False
            interval = fragment.get("output_rows")
            if (not isinstance(interval, list) or len(interval) != 2 or
                    any(type(number) is not int or number < 0 for number in interval) or
                    fragment.get("actual_output_rows") != interval):
                return False
            if status in {"empty", "missing"}:
                if interval != [0, 0] or fragment.get("block_id") is not None or fragment.get("bookmark") is not None:
                    return False
                if fragment.get("html_state") != status or fragment.get("reason") != "table_fragment_missing":
                    return False
                has_missing_fragment = True
            else:
                if interval[0] != cursor or interval[1] < interval[0] or (status == "rendered" and interval[1] == cursor):
                    return False
                cursor = interval[1]
            fragments.append(fragment)
        if cursor != table["expected_rows"]:
            return False
        if has_missing_fragment != ("table_fragment_missing" in table_blockers):
            return False
    if len(fragments) != value["fragment_count"]:
        return False
    if set(blockers) != local_blockers or set(warnings) != local_warnings:
        return False
    ids = [fragment["fragment_id"] for fragment in fragments]
    return (len(set(ids)) == len(ids) and
            type(value.get("placeholder_table_block_count", 0)) is int and value.get("placeholder_table_block_count", 0) >= 0 and
            sum(fragment.get("block_id") is not None for fragment in fragments) + value.get("placeholder_table_block_count", 0) == value["source_table_block_count"] and
            value["source_table_block_count"] - value["consumed_block_count"] - value.get("placeholder_table_block_count", 0) == value["output_table_count"] and
            sum(fragment.get("status") == "empty" for fragment in fragments) == value["empty_fragment_count"] and
            sum(fragment.get("status") == "missing" for fragment in fragments) == value["missing_fragment_count"])

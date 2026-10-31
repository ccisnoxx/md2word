from __future__ import annotations

# app/services/renderer/text_renderer.py: 提供 DocumentIR Word 渲染器的段落与 run 样式 helper。

import re
from dataclasses import replace
from html import escape
from collections.abc import Mapping
from typing import Any

from bs4 import BeautifulSoup, NavigableString, Tag

from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.conversion.docx_rules.run_fonts import apply_run_fonts
from docx.text.run import Run
from docx.text.paragraph import Paragraph

from app.conversion.docx_rules.formula import (
    FORMULA_OUTPUT_VERSION,
    MathSpan,
    _OmmlParser,
    formula_block_span,
    render_formula,
    scan_math_spans,
    sha256_text,
)
from app.conversion.docx_rules.terminology import TerminologyReplacer
from app.services.renderer.semantic_text import (
    infer_list_item_level,
    is_toc_title_text,
    is_list_item_text,
    looks_like_heading_text,
    normalize_list_item,
)
from app.services.renderer.template_config import TemplateConfigPackage, TemplateParagraphStyleConfig


MD2WORD_RENDERER_BULLET_NUM_ID_ATTR = "_md2word_renderer_bullet_num_id"
CAPTION_NUMBER_PATTERN = r"(?P<number>\d+)(?![A-Za-z0-9]|[.．]\d|[-–—][A-Za-z0-9])"
CAPTION_RE = re.compile(
    r"^(?P<prefix>\s*)(?P<label>图|表|Figure|Table)(?P<spacing>\s*)"
    + CAPTION_NUMBER_PATTERN
    + r"(?P<tail>.*)$",
    re.IGNORECASE,
)
REFERENCE_RE = re.compile(
    r"(?P<label>图|表|Figure|Table)(?P<spacing>\s*)" + CAPTION_NUMBER_PATTERN,
    re.IGNORECASE,
)
CaptionBookmarks = Mapping[tuple[str, int], tuple[str, int]]


class WordTextRendererMixin:
    """封装段落、run 和基础样式写入逻辑，供 Word renderer 门面复用。"""

    def _fill_paragraph(
        self,
        paragraph: Any,
        text: str,
        style: TemplateParagraphStyleConfig,
        *,
        prefix: str = "",
        word_style_name: str | None = None,
        terminology: TerminologyReplacer | None = None,
        reference_bookmarks: CaptionBookmarks | None = None,
    ) -> None:
        """写入段落文本；支持 MinerU 常见的内联 HTML 标记。"""

        self._apply_named_style(paragraph, word_style_name or style.word_style_name)
        self._apply_paragraph_style(paragraph, style)
        if prefix:
            self._add_text_run(paragraph, prefix, style)
        self._render_mixed_math(paragraph, text, style, terminology, reference_bookmarks)

    def _fill_heading_paragraph(
        self,
        paragraph: Any,
        text: str,
        level: int | None,
        template_config: TemplateConfigPackage,
    ) -> None:
        style = template_config.heading_style_for_level(level)
        self._fill_paragraph(
            paragraph,
            text,
            style,
            word_style_name=template_config.heading_word_style_name_for_level(level),
            terminology=template_config.terminology,
        )

    def _fill_list_paragraph(
        self,
        paragraph: Any,
        text: str,
        template_config: TemplateConfigPackage,
        *,
        reference_bookmarks: CaptionBookmarks | None = None,
    ) -> None:
        level = infer_list_item_level(text)
        self._fill_paragraph(
            paragraph,
            normalize_list_item(text),
            template_config.normal,
            terminology=template_config.terminology,
            reference_bookmarks=reference_bookmarks,
        )
        self._apply_bullet_numbering(paragraph, level=level)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph_format = paragraph.paragraph_format
        paragraph_format.space_before = Pt(0)
        paragraph_format.space_after = Pt(0)
        paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        paragraph_format.line_spacing = 1.0
        paragraph_format.left_indent = Pt(12 + min(max(level, 0), 1) * 12)
        paragraph_format.first_line_indent = Pt(-9)

    def _fill_planned_text(self, paragraph, text, policy, template_config, *, reference_bookmarks=None):
        """只执行已选策略，自动推断的优先级由 select_text_policy 唯一管理。"""
        if policy.mode == "toc":
            self._fill_toc_entry_paragraph(paragraph, policy, template_config)
        elif policy.mode == "bullet":
            self._fill_list_paragraph(paragraph, text, template_config, reference_bookmarks=reference_bookmarks)
        elif policy.mode == "heading":
            self._fill_heading_paragraph(paragraph, text, policy.heading_level, template_config)
        elif policy.mode == "formula":
            self._fill_formula_paragraph(paragraph, text, template_config)
        elif policy.mode == "table_context":
            self._fill_table_context_paragraph(paragraph, text, template_config)
        else:
            self._fill_paragraph(paragraph, text, getattr(template_config, policy.style_group),
                terminology=template_config.terminology, reference_bookmarks=reference_bookmarks)

    def _fill_toc_entry_paragraph(self, paragraph, policy, template_config) -> None:
        toc_style = getattr(template_config, policy.style_group)
        self._configure_toc_styles(paragraph.part.styles, template_config)
        self._apply_named_style(paragraph, toc_style.word_style_name)
        self._apply_paragraph_style(paragraph, toc_style)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Pt(0)
        self._add_complex_field(paragraph, r' TOC \o "1-3" \h \z \u ', "更新目录", toc_style)

    def _fill_formula_paragraph(
        self,
        paragraph: Any,
        text: str,
        template_config: TemplateConfigPackage,
    ) -> None:
        self._apply_named_style(paragraph, template_config.normal.word_style_name)
        self._apply_paragraph_style(paragraph, template_config.normal)
        self._append_formula(paragraph, formula_block_span(text), template_config.normal)

    def _parse_caption_text(self, text: str) -> tuple[str, str, str, int, str] | None:
        match = CAPTION_RE.match(text or "")
        if match is None:
            return None
        return (
            match.group("prefix"),
            match.group("label"),
            match.group("spacing"),
            int(match.group("number")),
            match.group("tail"),
        )

    def _fill_caption_paragraph(
        self,
        paragraph: Any,
        text: str,
        style: TemplateParagraphStyleConfig,
        *,
        terminology: TerminologyReplacer | None = None,
        caption_bookmarks: CaptionBookmarks | None = None,
    ) -> bool:
        parsed = self._parse_caption_text(text)
        if parsed is None:
            return False

        prefix, label, spacing, source_number, tail = parsed
        kind = "figure" if label.lower() == "figure" or label == "图" else "table"
        self._apply_named_style(paragraph, style.word_style_name)
        self._apply_paragraph_style(paragraph, style)
        self._add_text_run(paragraph, f"{prefix}{label}{spacing}", style)
        bookmark = caption_bookmarks.get((kind, source_number)) if caption_bookmarks else None
        if bookmark is not None:
            bookmark_start = OxmlElement("w:bookmarkStart")
            bookmark_start.set(qn("w:id"), str(bookmark[1]))
            bookmark_start.set(qn("w:name"), bookmark[0])
            paragraph._p.append(bookmark_start)
        sequence_name = label if label in {"图", "表"} else ("Figure" if kind == "figure" else "Table")
        self._add_complex_field(paragraph, f" SEQ {sequence_name} \\* ARABIC ", str(source_number), style)
        if bookmark is not None:
            bookmark_end = OxmlElement("w:bookmarkEnd")
            bookmark_end.set(qn("w:id"), str(bookmark[1]))
            paragraph._p.append(bookmark_end)
        self._render_mixed_math(paragraph, tail, style, terminology, None)
        return True

    def _append_formula(
        self,
        paragraph: Any,
        span: MathSpan,
        style: TemplateParagraphStyleConfig,
        *,
        bold: bool = False,
        italic: bool = False,
        _native_marks: bool = False,
        _native_references: CaptionBookmarks | None = None,
    ) -> None:
        result = render_formula(span)
        for element in result.elements:
            if element.tag == qn("w:r"):
                run = Run(element, paragraph)
                semantic = (run.bold, run.italic, run.font.subscript, run.font.superscript)
                self._apply_run_style(run, style)
                run.bold = bold or semantic[0]
                run.italic = italic or semantic[1]
                run.font.subscript = semantic[2]
                run.font.superscript = semantic[3]
            elif _native_marks and (bold or italic):
                for run in element.iter(qn("m:r")):
                    existing = run.find("./" + qn("m:rPr") + "/" + qn("m:sty"))
                    value = existing.get(qn("m:val"), "p") if existing is not None else "p"
                    combined_bold, combined_italic = bold or "b" in value, italic or "i" in value
                    # 复用原数学样式 owner，在预期投影冻结前合并外层真实 marks。
                    _OmmlParser._apply_math_style([run], "bi" if combined_bold and combined_italic
                        else "b" if combined_bold else "i")
        if _native_marks and _native_references and result.mode == "semantic_runs":
            # 同 part 的临时段落不进入最终 DOM；原文字 owner 先完成脚本与 REF 的实际输出。
            prepared = Paragraph(OxmlElement("w:p"), paragraph._parent)
            for element in result.elements:
                run = Run(element, paragraph)
                self._add_text_run(prepared, run.text, style, bold=bool(run.bold), italic=bool(run.italic),
                    subscript=bool(run.font.subscript), superscript=bool(run.font.superscript),
                    reference_bookmarks=_native_references, _native_reference_script=True)
            result = replace(result, elements=tuple(prepared._p))
        # 预期投影在元素进入最终 document DOM 以前确定。
        expected_output_sha256 = result.expected_output_sha256
        context = getattr(paragraph.part, "_md2word_formula_context", None)
        anchor = None
        if context is not None:
            source_hash = sha256_text(span.source)
            source_index = next((index for index, source in enumerate(context.formula_sources)
                if index not in context.used_formula_sources
                and (source["block_id"], source["field"]) == context.formula_location
                and source["source_sha256"] == source_hash and source["boundary"] == span.boundary), None)
            if source_index is None:
                raise ValueError("公式写入位置与源 DocumentIR 不匹配")
            context.used_formula_sources.add(source_index)
            occurrence = len(context.formula_entries) + 1
            anchor = f"md2w_formula_{occurrence:06d}"
            context.formula_entries.append({
                "occurrence": occurrence,
                "source": dict(context.formula_sources[source_index]),
                "anchor": anchor,
                "mode": result.mode,
                "reason": result.reason,
                "output_projection_version": FORMULA_OUTPUT_VERSION,
                "expected_output_sha256": expected_output_sha256,
            })
            bookmark = OxmlElement("w:bookmarkStart")
            bookmark.set(qn("w:id"), str(1000000 + occurrence))
            bookmark.set(qn("w:name"), anchor)
            paragraph._p.append(bookmark)
        paragraph._p.extend(result.elements)
        if anchor is not None:
            bookmark = OxmlElement("w:bookmarkEnd")
            bookmark.set(qn("w:id"), str(1000000 + occurrence))
            paragraph._p.append(bookmark)

    def _set_formula_source(self, document: Any, block_id: str, field: str) -> None:
        """在实际内容写入点登记来源；同原式不能靠库存顺序猜测所属位置。"""
        context = getattr(document.part, "_md2word_formula_context", None)
        if context is not None:
            context.formula_location = (block_id, field)

    def _render_mixed_math(
        self,
        paragraph: Any,
        text: str,
        style: TemplateParagraphStyleConfig,
        terminology: TerminologyReplacer | None,
        reference_bookmarks: CaptionBookmarks | None,
        *,
        _native_literal: bool = False,
        _spans: tuple[MathSpan, ...] | None = None,
    ) -> bool:
        spans = scan_math_spans(text) if _spans is None else _spans
        placeholder_tag = "md2w-formula"
        while placeholder_tag in text.lower():
            placeholder_tag += "-internal"
        chunks: list[str] = []
        position = 0
        for index, span in enumerate(spans):
            chunks.extend((text[position:span.start], f'<{placeholder_tag} data-index="{index}"></{placeholder_tag}>'))
            position = span.end
        chunks.append(text[position:])
        soup = BeautifulSoup("<root>" + "".join(chunks) + "</root>", "html.parser")
        root = soup.find("root")
        if root is None:
            raise ValueError("内联文本容器无法解析")

        def visit(parent: Tag, *, bold: bool = False, italic: bool = False, subscript: bool = False, superscript: bool = False) -> None:
            for child in parent.children:
                if isinstance(child, NavigableString):
                    raw = str(child)
                    prepared = (escape(raw, quote=False) if _native_literal and (subscript or superscript)
                                else self._prepare_text_for_rendering(raw, terminology, _native_literal=_native_literal))
                    if _native_literal or self._looks_like_inline_markup(prepared):
                        prepared_soup = BeautifulSoup(prepared, "html.parser")
                        self._render_inline_children(prepared_soup, paragraph, style, bold=bold, italic=italic, subscript=subscript, superscript=superscript, reference_bookmarks=reference_bookmarks)
                    else:
                        self._add_text_run(paragraph, prepared, style, bold=bold, italic=italic, subscript=subscript, superscript=superscript, reference_bookmarks=reference_bookmarks)
                elif isinstance(child, Tag):
                    name = child.name.lower()
                    if name == placeholder_tag:
                        self._append_formula(paragraph, spans[int(child["data-index"])], style, bold=bold,
                            italic=italic, _native_marks=_native_literal,
                            _native_references=reference_bookmarks if _native_literal else None)
                    elif name == "br":
                        paragraph.add_run().add_break()
                    else:
                        visit(child, bold=bold or name in {"b", "strong", "th"}, italic=italic or name in {"i", "em"}, subscript=subscript or name == "sub", superscript=superscript or name == "sup")

        visit(root)
        return True

    def _add_complex_field(
        self,
        paragraph: Any,
        instruction: str,
        result: str,
        style: TemplateParagraphStyleConfig,
    ) -> Any:
        begin_run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        begin_run._r.append(begin)

        instruction_run = paragraph.add_run()
        instruction_text = OxmlElement("w:instrText")
        instruction_text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        instruction_text.text = instruction
        instruction_run._r.append(instruction_text)

        separate_run = paragraph.add_run()
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        separate_run._r.append(separate)

        result_run = paragraph.add_run(result)
        self._apply_run_style(result_run, style)

        end_run = paragraph.add_run()
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        end_run._r.append(end)
        return result_run

    def _prepare_text_for_rendering(self, text: str, terminology: TerminologyReplacer | None,
                                  *, _native_literal: bool = False) -> str:
        """仅处理普通正文的单位间距与术语，数学片段由共享入口先行隔离。"""
        normalized = text.replace("\xa0", " ")
        normalized = re.sub(r"([±+\-])\s+(\d)", r"\1\2", normalized)
        normalized = re.sub(r"(\d)\s*°\s*([CFK])\b", r"\1°\2", normalized)
        normalized = re.sub(r"([µμ])\s+([A-Za-z])\b", r"\1\2", normalized)
        normalized = re.sub(r"([A-Za-zµμ])\s+Ω\b", r"\1Ω", normalized)
        if _native_literal:
            return (terminology.replace_text(normalized, _native_literal=True) if terminology is not None
                    else escape(normalized, quote=False))
        return terminology.replace_text(normalized) if terminology is not None else normalized

    def _looks_like_inline_markup(self, text: str) -> bool:
        lowered = text.lower()
        return any(
            token in lowered
            for token in ("<sub", "<sup", "<br", "<b", "<strong", "<i", "<em", "<td", "<th", "<span", "<p")
        )

    def _render_inline_children(
        self,
        parent: Tag,
        paragraph: Any,
        style: TemplateParagraphStyleConfig,
        *,
        bold: bool = False,
        italic: bool = False,
        subscript: bool = False,
        superscript: bool = False,
        reference_bookmarks: CaptionBookmarks | None = None,
    ) -> None:
        for child in parent.children:
            if isinstance(child, NavigableString):
                self._add_text_run(
                    paragraph,
                    str(child),
                    style,
                    bold=bold,
                    italic=italic,
                    subscript=subscript,
                    superscript=superscript,
                    reference_bookmarks=reference_bookmarks,
                )
                continue
            if not isinstance(child, Tag):
                continue

            tag_name = child.name.lower()
            if tag_name == "br":
                paragraph.add_run().add_break()
                continue
            self._render_inline_children(
                child,
                paragraph,
                style,
                bold=bold or tag_name in {"b", "strong", "th"},
                italic=italic or tag_name in {"i", "em"},
                subscript=subscript or tag_name == "sub",
                superscript=superscript or tag_name == "sup",
                reference_bookmarks=reference_bookmarks,
            )

    def _add_text_run(
        self,
        paragraph: Any,
        text: str,
        style: TemplateParagraphStyleConfig,
        *,
        bold: bool = False,
        italic: bool = False,
        subscript: bool = False,
        superscript: bool = False,
        reference_bookmarks: CaptionBookmarks | None = None,
        _native_reference_script: bool = False,
    ) -> None:
        if not text:
            return
        if reference_bookmarks:
            position = 0
            for match in REFERENCE_RE.finditer(text):
                label = match.group("label")
                kind = "figure" if label.lower() == "figure" or label == "图" else "table"
                bookmark = reference_bookmarks.get((kind, int(match.group("number"))))
                if bookmark is None:
                    continue
                self._add_text_run(
                    paragraph,
                    text[position : match.start()] + label + match.group("spacing"),
                    style,
                    bold=bold,
                    italic=italic,
                    subscript=subscript,
                    superscript=superscript,
                )
                field_run = self._add_complex_field(paragraph, f" REF {bookmark[0]} \\h ", match.group("number"), style)
                field_run.bold = bold or None
                field_run.italic = italic or None
                if _native_reference_script:
                    # 两个属性共用 vertAlign；native 结果只能设置实际生效的互斥脚本。
                    if subscript:
                        field_run.font.subscript = True
                    elif superscript:
                        field_run.font.superscript = True
                else:
                    field_run.font.subscript = subscript or None
                    field_run.font.superscript = superscript or None
                position = match.end()
            if position:
                self._add_text_run(
                    paragraph,
                    text[position:],
                    style,
                    bold=bold,
                    italic=italic,
                    subscript=subscript,
                    superscript=superscript,
                )
                return
        run = paragraph.add_run(text)
        self._apply_run_style(run, style)
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        if subscript:
            run.font.subscript = True
        if superscript:
            run.font.superscript = True

    def _apply_paragraph_style(self, paragraph: Any, style: TemplateParagraphStyleConfig) -> None:
        if style.alignment:
            paragraph.alignment = self._parse_alignment(style.alignment)
        paragraph_format = paragraph.paragraph_format
        if style.space_before is not None:
            paragraph_format.space_before = Pt(style.space_before)
        if style.space_after is not None:
            paragraph_format.space_after = Pt(style.space_after)
        if style.line_spacing is not None:
            paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            paragraph_format.line_spacing = style.line_spacing
        if style.first_line_indent_pt is not None:
            paragraph_format.first_line_indent = Pt(style.first_line_indent_pt)
        if style.left_indent_pt is not None:
            paragraph_format.left_indent = Pt(style.left_indent_pt)
        if style.right_indent_pt is not None:
            paragraph_format.right_indent = Pt(style.right_indent_pt)

    def _apply_named_style(self, paragraph: Any, style_name: str | None) -> None:
        if not style_name:
            return
        try:
            paragraph.style = style_name
        except (KeyError, ValueError):
            return

    def _configure_toc_styles(self, styles: Any, template_config: TemplateConfigPackage) -> None:
        for level in range(1, 4):
            style_id = f"TOC{level}"
            word_style = next((style for style in styles if style.style_id == style_id), None)
            if word_style is None:
                word_style = styles.add_style(f"TOC {level}", WD_STYLE_TYPE.PARAGRAPH, builtin=True)

            config = template_config.toc_style_for_level(level)
            if config.word_style_name and config.word_style_name != word_style.name:
                try:
                    word_style.base_style = styles[config.word_style_name]
                except KeyError:
                    pass
            apply_run_fonts(
                word_style,
                font_name=config.font_name,
                western_font_name=config.western_font_name,
                complex_script_font_name=config.complex_script_font_name,
            )
            if config.font_size is not None:
                word_style.font.size = Pt(config.font_size)
            if config.bold is not None:
                word_style.font.bold = config.bold
            if config.italic is not None:
                word_style.font.italic = config.italic
            if config.color:
                word_style.font.color.rgb = self._parse_color(config.color)

            paragraph_format = word_style.paragraph_format
            if config.alignment:
                paragraph_format.alignment = self._parse_alignment(config.alignment)
            if config.space_before is not None:
                paragraph_format.space_before = Pt(config.space_before)
            if config.space_after is not None:
                paragraph_format.space_after = Pt(config.space_after)
            if config.line_spacing is not None:
                paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                paragraph_format.line_spacing = config.line_spacing
            if config.first_line_indent_pt is not None:
                paragraph_format.first_line_indent = Pt(config.first_line_indent_pt)
            if config.left_indent_pt is not None:
                paragraph_format.left_indent = Pt(config.left_indent_pt)
            if config.right_indent_pt is not None:
                paragraph_format.right_indent = Pt(config.right_indent_pt)

    def _apply_run_style(self, run: Any, style: TemplateParagraphStyleConfig) -> None:
        apply_run_fonts(
            run,
            font_name=style.font_name,
            western_font_name=style.western_font_name,
            complex_script_font_name=style.complex_script_font_name,
        )
        if style.font_size is not None:
            run.font.size = Pt(style.font_size)
        if style.bold is not None:
            run.bold = style.bold
        if style.italic is not None:
            run.italic = style.italic
        if style.color:
            run.font.color.rgb = self._parse_color(style.color)

    def _parse_alignment(self, value: str) -> WD_ALIGN_PARAGRAPH:
        normalized = value.strip().lower()
        if normalized == "center":
            return WD_ALIGN_PARAGRAPH.CENTER
        if normalized == "right":
            return WD_ALIGN_PARAGRAPH.RIGHT
        if normalized == "justify":
            return WD_ALIGN_PARAGRAPH.JUSTIFY
        return WD_ALIGN_PARAGRAPH.LEFT

    def _parse_color(self, value: str) -> RGBColor:
        raw = value.strip().lstrip("#")
        if len(raw) != 6:
            return RGBColor(0, 0, 0)
        try:
            return RGBColor.from_string(raw.upper())
        except ValueError:
            return RGBColor(0, 0, 0)

    def _apply_bullet_numbering(self, paragraph: Any, *, level: int = 0) -> None:
        normalized_level = min(max(level, 0), 1)
        num_id = self._ensure_bullet_numbering(paragraph)
        paragraph_properties = paragraph._p.get_or_add_pPr()
        num_properties = paragraph_properties.find(qn("w:numPr"))
        if num_properties is None:
            num_properties = OxmlElement("w:numPr")
            paragraph_properties.append(num_properties)

        self._set_numbering_child_value(num_properties, "w:ilvl", str(normalized_level))
        self._set_numbering_child_value(num_properties, "w:numId", str(num_id))

    def _ensure_bullet_numbering(self, paragraph: Any) -> str:
        part = paragraph.part
        cached_num_id = getattr(part, MD2WORD_RENDERER_BULLET_NUM_ID_ATTR, None)
        if cached_num_id is not None:
            return str(cached_num_id)

        numbering = part.numbering_part.element
        abstract_num_id = str(self._next_numbering_id(numbering, "w:abstractNum", "w:abstractNumId"))
        num_id = str(self._next_numbering_id(numbering, "w:num", "w:numId"))

        abstract_num = OxmlElement("w:abstractNum")
        abstract_num.set(qn("w:abstractNumId"), abstract_num_id)
        multi_level_type = OxmlElement("w:multiLevelType")
        multi_level_type.set(qn("w:val"), "hybridMultilevel")
        abstract_num.append(multi_level_type)

        for ilvl, marker, left in ((0, "•", "240"), (1, "–", "480")):
            level = OxmlElement("w:lvl")
            level.set(qn("w:ilvl"), str(ilvl))
            for tag, value in (
                ("w:start", "1"),
                ("w:numFmt", "bullet"),
                ("w:lvlText", marker),
                ("w:lvlJc", "left"),
            ):
                child = OxmlElement(tag)
                child.set(qn("w:val"), value)
                level.append(child)
            paragraph_properties = OxmlElement("w:pPr")
            indent = OxmlElement("w:ind")
            indent.set(qn("w:left"), left)
            indent.set(qn("w:hanging"), "180")
            paragraph_properties.append(indent)
            level.append(paragraph_properties)
            abstract_num.append(level)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), num_id)
        abstract_num_id_element = OxmlElement("w:abstractNumId")
        abstract_num_id_element.set(qn("w:val"), abstract_num_id)
        num.append(abstract_num_id_element)

        numbering.append(abstract_num)
        numbering.append(num)
        setattr(part, MD2WORD_RENDERER_BULLET_NUM_ID_ATTR, num_id)
        return num_id

    def _next_numbering_id(self, numbering: Any, element_tag: str, attr_name: str) -> int:
        values: list[int] = []
        for element in numbering.findall(qn(element_tag)):
            raw_value = element.get(qn(attr_name))
            if raw_value is not None and raw_value.isdigit():
                values.append(int(raw_value))
        return max(values, default=0) + 1

    def _set_numbering_child_value(self, parent: Any, tag: str, value: str) -> None:
        child = parent.find(qn(tag))
        if child is None:
            child = OxmlElement(tag)
            parent.append(child)
        child.set(qn("w:val"), value)

"""v2 必需内容合同：固定来源、真实表行 scope 与全源公式覆盖。"""
from __future__ import annotations

from collections import defaultdict, deque
from dataclasses import dataclass
import hashlib
from typing import Any

from docx.oxml.ns import qn
from lxml import etree

from app.conversion.docx_rules.formula import content_math_spans, formula_block_span, sha256_text
from app.services.document_ir.integrity import DocumentIntegrityError, is_visible_body_block
from app.services.renderer.render_plan import canonical_bytes, content_hash, document_ir_hash

CONTENT_VERSION = "renderer-content/v2"
FORMULA_SOURCE_VERSION = "formula-source-coverage/v1"
MAX_FORMULA_SOURCE_ENTRIES = 100_000
MAX_FORMULA_SOURCE_BYTES = 8 * 1024 * 1024
MAX_CONTENT_SCOPES = 100_000
MAX_CONTENT_SCOPE_BYTES = 8 * 1024 * 1024
_SOURCE_KEYS = ("block_id", "field", "start", "end", "boundary", "source_sha256")


def content_failure(code: str = "ir_render_evidence:invalid_contract") -> DocumentIntegrityError:
    """新必需内容错误有固定类型/版本，消费者不按异常正文猜分类。"""
    return DocumentIntegrityError("DocumentIR→DOCX", {
        "renderer_content_contract_version": CONTENT_VERSION, "blocking_quality_flags": [code]})


def validate_scope_budget(value: Any) -> None:
    if (not isinstance(value, list) or len(value) > MAX_CONTENT_SCOPES
            or len(canonical_bytes(value)) > MAX_CONTENT_SCOPE_BYTES):
        raise content_failure()


def validate_formula_source_budget(value: Any) -> None:
    if (not isinstance(value, dict) or not isinstance(value.get("entries"), list)
            or len(value["entries"]) > MAX_FORMULA_SOURCE_ENTRIES
            or len(canonical_bytes(value)) > MAX_FORMULA_SOURCE_BYTES):
        raise content_failure()


def project_table_scope(index, scope) -> tuple[dict, list[str]]:
    """仅投影该唯一表的直接行；首段和行集合不能借用别的 fragment。"""
    table = index.nodes.get((scope.part, tuple(scope.table_child_path)))
    if table is None or table.tag != qn("w:tbl"):
        raise content_failure("ir_render_evidence:target_mismatch")
    rows = table.findall(qn("w:tr"))
    start, end = scope.output_rows
    selected = rows[start:end]
    if (not 0 <= start < end <= len(rows) or len(selected) != len(scope.row_child_paths)
            or [list(index.paths[row][1]) for row in selected] != scope.row_child_paths):
        raise content_failure("ir_render_evidence:target_mismatch")
    paragraph = next(selected[0].iter(qn("w:p")), None)
    if paragraph is None or list(index.paths[paragraph][1]) != scope.first_paragraph_child_path:
        raise content_failure("ir_render_evidence:target_mismatch")
    names = list(dict.fromkeys(node.get(qn("w:name")) for row in selected
                              for node in row.iter(qn("w:bookmarkStart"))))
    actual = {"output_rows": [start, end], "row_xml_sha256": [hashlib.sha256(etree.tostring(row)).hexdigest()
              for row in selected], "cell_counts": [len(row.findall(qn("w:tc"))) for row in selected]}
    actual["referenced_section_indices"] = index.project(scope.part, tuple(scope.table_child_path))[1]["referenced_section_indices"]
    return actual, names


@dataclass(frozen=True, slots=True)
class FormulaSource:
    binding: dict[str, Any]
    key: tuple[str, str, int]
    span: Any
    representative: tuple[str, str, int] | None = None
    alias_reason: str | None = None


def cell_math_spans(cell):
    """从 writer 同一确定 native 视图枚举，逐 run 切分不能截断原有数学 base。"""
    from app.services.renderer.table_renderer import native_cell_parts
    for part in native_cell_parts(cell.runs):
        for span in part.spans:
            yield part.offset + span.start, part.offset + span.end, span


def formula_sources(document, table_plan) -> list[FormulaSource]:
    """从去重前 IR/native grid 独立枚举，每个合法 alias 指向真实源 occurrence。"""
    from app.services.renderer.layout_plan import image_annotations
    from app.services.renderer.logical_table_plan import table_unit_contracts
    contracts = table_unit_contracts(document, table_plan)
    tables = {}
    aliases = {}
    for unit, contract in zip(table_plan.units, contracts, strict=True):
        primary = next(fragment for fragment in unit.fragments if fragment.render is not None)
        for fragment in unit.fragments:
            if fragment.render is None:
                continue
            tables[fragment.source.block_id] = fragment.render.grid
            if fragment.removed_header_rows:
                for cell in fragment.render.grid.cells:
                    if cell.row >= fragment.removed_header_rows:
                        continue
                    matches = [other for other in primary.render.grid.cells
                               if (other.row, other.column) == (cell.row, cell.column)]
                    if len(matches) != 1:
                        raise content_failure("ir_render_evidence:target_mismatch")
                    other = matches[0]
                    if [(r.type, r.source, r.bold, r.italic) for r in cell.runs] != [(r.type, r.source, r.bold, r.italic) for r in other.runs]:
                        raise content_failure("ir_render_evidence:target_mismatch")
                    aliases[fragment.source.block_id, f"html/{cell.source_row_index}/{cell.source_cell_index}"] = (
                        primary.source.block_id, f"html/{other.source_row_index}/{other.source_cell_index}", "table_header_coalesced")
        for item in contract.annotations:
            if item.disposition == "coalesced":
                aliases[item.block_id, f"{item.field}/{item.item_index}"] = (
                    item.representative_block_id, f"{item.representative_field}/{item.representative_item_index}",
                    "table_annotation_coalesced")
    result = []

    def add(block_id, field, text, *, formula=False, spans=None):
        values = spans if spans is not None else [(span.start, span.end, span) for span in
                 ((formula_block_span(text),) if formula else content_math_spans(text))]
        for ordinal, (start, end, span) in enumerate(values):
            alias = aliases.get((block_id, field))
            result.append(FormulaSource({"block_id": block_id, "field": field, "start": start, "end": end,
                "boundary": span.boundary, "source_sha256": sha256_text(span.source)}, (block_id, field, ordinal), span,
                (alias[0], alias[1], ordinal) if alias else None, alias[2] if alias else None))
            if len(result) > MAX_FORMULA_SOURCE_ENTRIES:
                raise content_failure()

    for block in document.blocks:
        if not is_visible_body_block(block):
            continue
        if block.type == "table" and block.id in tables:
            for cell in tables[block.id].cells:
                add(block.id, f"html/{cell.source_row_index}/{cell.source_cell_index}", "", spans=list(cell_math_spans(cell)))
        elif block.type == "formula":
            add(block.id, "latex" if block.latex else "text", block.latex or block.text or "", formula=True)
        elif block.type == "list":
            for index, text in enumerate(block.list_items or []):
                add(block.id, f"list_items/{index}", text)
        elif block.type == "image":
            selected = {text: field for field, text in image_annotations(block)}
            values = [(f"captions/{i}", text) for i, text in enumerate(block.captions)]
            if isinstance(block.metadata.get("caption"), str):
                values.append(("metadata/caption", block.metadata["caption"]))
            values.extend((f"footnotes/{i}", text) for i, text in enumerate(block.footnotes))
            if not selected:
                add(block.id, "text", block.text or "")
            for field, text in values:
                if text.strip():
                    chosen = selected[text.strip()]
                    if chosen != field:
                        aliases[block.id, field] = (block.id, chosen, "annotation_duplicate")
                    add(block.id, field, text.strip())
        else:
            add(block.id, "text", block.text or "")
        if block.type == "table":
            for field in ("captions", "footnotes"):
                seen = {}
                for index, text in enumerate(getattr(block, field)):
                    text = text.strip()
                    if not text:
                        continue
                    name = f"{field}/{index}"
                    if (block.id, name) not in aliases and text in seen:
                        aliases[block.id, name] = (block.id, seen[text], "annotation_duplicate")
                    seen.setdefault(text, name)
                    add(block.id, name, text)
    lookup = {source.key: source for source in result}
    if len(lookup) != len(result):
        raise content_failure()
    for source in result:
        if source.representative is None:
            continue
        other = lookup.get(source.representative)
        if (other is None or other.representative is not None
                or (source.binding["source_sha256"], source.binding["boundary"])
                != (other.binding["source_sha256"], other.binding["boundary"])):
            raise content_failure("ir_render_evidence:target_mismatch")
    return result


def build_formula_source_coverage(document, content: bytes, table_plan, report: Any) -> dict:
    from app.services.renderer.docx_quality import validate_formula_report
    validate_formula_report(report)
    sources = formula_sources(document, table_plan)
    available = defaultdict(deque)
    for entry in report["entries"]:
        available[tuple(entry["source"][key] for key in _SOURCE_KEYS)].append(entry["occurrence"])
    lookup = {source.key: index for index, source in enumerate(sources)}
    entries = []
    for index, source in enumerate(sources):
        occurrence = None
        representative = lookup[source.representative] if source.representative is not None else None
        if representative is None:
            occurrences = available[tuple(source.binding[key] for key in _SOURCE_KEYS)]
            if not occurrences:
                raise content_failure("ir_render_evidence:target_mismatch")
            occurrence = occurrences.popleft()
        entries.append({"source_occurrence_index": index, "source": source.binding,
            "disposition": "rendered" if representative is None else "coalesced", "rendered_occurrence": occurrence,
            "representative_source": representative, "alias_reason": source.alias_reason})
    if any(occurrences for occurrences in available.values()):
        raise content_failure("ir_render_evidence:target_mismatch")
    value = {"schema_version": FORMULA_SOURCE_VERSION, "input_ir_sha256": document_ir_hash(document),
        "docx_sha256": hashlib.sha256(content).hexdigest(), "table_render_input_sha256": table_plan.input_sha256,
        "entries": entries}
    result = {**value, "coverage_sha256": content_hash(value)}
    validate_formula_source_budget(result)
    return result


def validate_formula_source_coverage(document, content: bytes, table_plan, report: Any, coverage: Any) -> dict:
    validate_formula_source_budget(coverage)
    expected = build_formula_source_coverage(document, content, table_plan, report)
    if coverage != expected:
        raise content_failure("ir_render_evidence:binding_mismatch")
    return expected


def physical_table_order(document, layout, table_plan, image_report) -> list[str]:
    """沿原纯调度重建内容/图片布局表交错顺序，不从实际 ordinal 反推期待值。"""
    from types import SimpleNamespace
    from app.services.renderer.render_plan import ResourceIdentity, build_render_plan
    from app.services.renderer.template_config import TemplateConfigPackage
    from app.services.renderer.layout_plan import build_image_plan
    from app.services.renderer.image_evidence import IMAGE_TABLE_PREFIX
    grids = {fragment.source.block_id: fragment.render.grid for unit in table_plan.units
             for fragment in unit.fragments if fragment.render is not None}
    models = {key: SimpleNamespace(rows=range(grid.rows), max_columns=grid.columns) for key, grid in grids.items()}
    strategies = {block.id: "standard" if block.id in grids else "fallback"
                  for block in document.blocks if block.type == "table"}
    identity = ResourceIdentity(template_sha256=None, effective_config_sha256="0" * 64,
        symbol_rules_sha256=None, logo_sha256=None, document_language=document.metadata.document_language)
    plan = build_render_plan(document, layout_plan=layout, effective_config=TemplateConfigPackage(),
        resource_identity=identity, table_strategies=strategies, table_models=models, logical_table_plan=table_plan)
    spans = {block_id: span for span in build_image_plan(document) for block_id in span.block_ids}
    grouped = {entry["span_id"] for entry in image_report["spans"] if entry["mode"] == "group"}
    result = []
    for operation in plan.operations:
        if operation.kind == "logical_table":
            result.append(operation.unit_id)
        elif operation.kind == "block" and operation.block_ids[0] in spans:
            span = spans[operation.block_ids[0]]
            if span.id in grouped:
                result.append(IMAGE_TABLE_PREFIX + span.id)
    return result


def semantic_projection(paragraph, *, include_fields=False) -> list:
    """比较文字、脚本与数学结构；native cell 还须绑定同源 REF 指令。"""
    result = []

    def add_text(text, script=None):
        if not text:
            return
        if result and result[-1][0] == "text" and result[-1][2] == script:
            result[-1][1] += text
        else:
            result.append(["text", text, script])

    def math_node(node):
        if node.tag in {qn("w:rPr"), qn("m:rPr"), qn("m:ctrlPr")}:
            return None
        return [etree.QName(node).localname,
            [[etree.QName(key).localname, value] for key, value in sorted(node.attrib.items())],
            node.text or None, [value for child in node if (value := math_node(child)) is not None]]

    for child in paragraph:
        if child.tag in {qn("w:pPr"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")}:
            continue
        if child.tag == qn("m:oMath"):
            result.append(["math", math_node(child)])
            continue
        if child.tag != qn("w:r"):
            raise content_failure("ir_render_evidence:target_mismatch")
        vertical = child.find("./" + qn("w:rPr") + "/" + qn("w:vertAlign"))
        script = vertical.get(qn("w:val")) if vertical is not None else None
        for node in child:
            if node.tag in {qn("w:instrText"), qn("w:fldChar")}:
                if include_fields:
                    result.append([etree.QName(node).localname, node.text or "", sorted(node.attrib.items())])
                continue
            if node.tag == qn("w:rPr"):
                continue
            if node.tag == qn("w:t"):
                add_text(node.text or "", script)
            elif node.tag == qn("w:tab"):
                add_text("\t", script)
            elif node.tag in {qn("w:br"), qn("w:cr")}:
                add_text("\n", script)
            else:
                raise content_failure("ir_render_evidence:target_mismatch")
    return result


def verify_native_content(document, content: bytes, table_plan, formula_report, image_report, coverage):
    """固定 grid 与原冻结规则重放叶级投影，不渲染主件、不读取当前配置或写 Artifact。"""
    from docx import Document
    from app.conversion.docx_rules.terminology import TerminologyReplacer
    from app.services.renderer.block_evidence import PackageIndex
    from app.services.renderer.word_renderer import DocumentIRWordRenderer
    from app.services.renderer.logical_table_plan import table_unit_contracts
    from app.services.renderer.table_renderer import iter_grid_rows
    from app.services.renderer.template_config import TemplateConfigPackage
    rules = TerminologyReplacer.from_frozen_rules(image_report["render_inputs"]["format_rules"]["rules"])
    config = TemplateConfigPackage()
    config.terminology = rules
    writer = DocumentIRWordRenderer()
    scratch = Document()
    index = PackageIndex(content)
    body = index.roots["word/document.xml"].find(qn("w:body"))
    children = list(body)
    contracts = table_unit_contracts(document, table_plan)
    from app.services.renderer.image_evidence import caption_bookmark_inventory
    references = caption_bookmark_inventory(document, table_contracts=contracts,
        primary_table_block_ids={document.blocks[unit.anchor_block_index].id
                                 for unit in table_plan.units if unit.is_group_primary})
    by_id = {block.id: block for block in document.blocks}
    occurrences = defaultdict(list)
    for entry in coverage["entries"]:
        if entry["disposition"] == "rendered":
            source = entry["source"]
            occurrences[source["block_id"], source["field"]].append(entry["rendered_occurrence"])
    native = {}
    annotations = {}

    def expected_annotation(item, *, primary):
        paragraph = scratch.add_paragraph()
        text = getattr(by_id[item.block_id], item.field)[item.item_index].strip()
        style = config.table_note if item.role == "note" else config.normal if item.role == "context" else config.caption
        if item.role != "caption" or not primary or not writer._fill_caption_paragraph(paragraph, text, style, terminology=rules):
            writer._fill_paragraph(paragraph, text, style, terminology=rules)
        value = semantic_projection(paragraph._p)
        paragraph._p.getparent().remove(paragraph._p)
        return value

    for unit, contract in zip(table_plan.units, contracts, strict=True):
        matches = [node for node in body.findall(qn("w:tbl"))
            if (caption := node.find("./" + qn("w:tblPr") + "/" + qn("w:tblCaption"))) is not None
            and caption.get(qn("w:val")) == unit.key]
        if len(matches) != 1:
            raise content_failure("ir_render_evidence:target_mismatch")
        table = matches[0]
        native[unit.key] = table
        rows = table.findall(qn("w:tr"))
        for fragment in unit.fragments:
            if fragment.render is None:
                continue
            for source_row, owners in iter_grid_rows(fragment.render.grid, block_index=fragment.block_index):
                if source_row < fragment.removed_header_rows:
                    continue
                row = rows[fragment.output_start + source_row - fragment.removed_header_rows]
                cells = row.findall(qn("w:tc"))
                logical = {}
                column = 0
                for cell in cells:
                    logical[column] = cell
                    span = cell.find("./" + qn("w:tcPr") + "/" + qn("w:gridSpan"))
                    column += int(span.get(qn("w:val"))) if span is not None else 1
                used = set()
                for owner in owners:
                    cell = logical[owner.column]
                    used.add(owner.column)
                    paragraphs = cell.findall(qn("w:p"))
                    if len(paragraphs) != 1 or any(node.tag not in {qn("w:tcPr"), qn("w:p")} for node in cell):
                        raise content_failure("ir_render_evidence:target_mismatch")
                    paragraph = scratch.add_paragraph()
                    if owner.row == source_row:
                        writer._fill_native_grid_runs(paragraph, owner.runs, config, reference_bookmarks=references)
                    expected = semantic_projection(paragraph._p, include_fields=True)
                    paragraph._p.getparent().remove(paragraph._p)
                    if semantic_projection(paragraphs[0], include_fields=True) != expected:
                        raise content_failure("ir_render_evidence:target_mismatch")
                    if owner.row == source_row:
                        field = f"html/{owner.source_row_index}/{owner.source_cell_index}"
                        expected_occurrences = occurrences[fragment.source.block_id, field]
                    else:
                        expected_occurrences = []
                    actual_names = [node.get(qn("w:name")) for node in cell.iter(qn("w:bookmarkStart"))
                                    if node.get(qn("w:name"), "").startswith("md2w_formula_")]
                    if actual_names != [f"md2w_formula_{value:06d}" for value in expected_occurrences]:
                        raise content_failure("ir_render_evidence:target_mismatch")
                for column, cell in logical.items():
                    if column not in used and any(semantic_projection(paragraph, include_fields=True)
                                                 for paragraph in cell.findall(qn("w:p"))):
                        raise content_failure("ir_render_evidence:target_mismatch")
        prefix = [item for item in contract.annotations if item.disposition == "rendered" and item.role != "note"]
        suffix = [item for item in contract.annotations if item.disposition == "rendered" and item.role == "note"]
        position = children.index(table)
        before = position - len(prefix) - bool(unit.blocking_quality_flags)
        if before < 0:
            raise content_failure("ir_render_evidence:target_mismatch")
        bindings = []
        for offset, item in enumerate(prefix):
            paragraph = children[before + offset]
            if paragraph.tag != qn("w:p") or semantic_projection(paragraph) != expected_annotation(item, primary=unit.is_group_primary):
                raise content_failure("ir_render_evidence:target_mismatch")
            bindings.append((item.block_id, item.field, item.item_index, index.paths[paragraph][1]))
        from app.services.document_ir.table_placeholder import decode_placeholders
        placeholders = [item for item in decode_placeholders(document)
                        if item.fragment_id in {fragment.source.fragment_id for fragment in unit.fragments}]
        for offset, item in enumerate(suffix):
            paragraph = children[position + 1 + len(placeholders) + offset]
            if paragraph.tag != qn("w:p") or semantic_projection(paragraph) != expected_annotation(item, primary=unit.is_group_primary):
                raise content_failure("ir_render_evidence:target_mismatch")
            bindings.append((item.block_id, item.field, item.item_index, index.paths[paragraph][1]))
        annotations[unit.key] = tuple(bindings)
    return index, native, annotations


def build_v2_content_evidence(document, content, layout, *, coverage, formula_report, image_report,
                              expected_template_version_id=None, expected_generation_policy=None):
    """先认证全源覆盖，再进入原物理公式/图片门；不预先执行 legacy 表数或库存。"""
    from app.services.renderer.evidence import build_renderer_aggregate_evidence
    from app.services.renderer.logical_table_plan import build_content_table_plan
    from app.services.renderer.layout_plan import build_layout_plan
    from app.services.renderer.table_content_evidence import _table_rendering_evidence, _build_table_evidence
    from app.services.renderer.table_placeholder_evidence import build_placeholder_evidence
    from app.services.document_ir.table_placeholder import decode_placeholders
    try:
        if not isinstance(image_report, dict) or not isinstance(formula_report, dict):
            raise content_failure()
        resolved = layout or build_layout_plan(document)
        tables = build_content_table_plan(document, resolved)
        validated = validate_formula_source_coverage(document, content, tables, formula_report, coverage)
        rendered = [entry["source"] for entry in validated["entries"] if entry["disposition"] == "rendered"]
        placeholders = decode_placeholders(document)
        proof, placeholder_ownership = build_placeholder_evidence(document, content, _with_ownership=True)
        evidence = build_renderer_aggregate_evidence(document, content, resolved, formula_report=formula_report,
            image_layout_report=image_report, expected_template_version_id=expected_template_version_id,
            expected_generation_policy=expected_generation_policy, _table_plan=tables,
            _formula_sources=rendered, _placeholder_count=len(placeholders), _v2_content=True,
            _placeholder_ownership=placeholder_ownership)
        evidence.update(renderer_content_contract_version=CONTENT_VERSION, formula_source_coverage=validated,
                        table_placeholder_count=len(placeholders))
        table_report, flags = _table_rendering_evidence(document, content, tables,
            evidence["docx_quality"]["table_observations"], v2=True,
            expected_table_order=physical_table_order(document, resolved, tables, image_report))
        evidence["table_rendering"] = table_report
        if proof is not None:
            evidence["table_placeholder_rendering"] = proof
        evidence["blocking_quality_flags"] = list(dict.fromkeys([*evidence["blocking_quality_flags"], *flags]))
        if not evidence["blocking_quality_flags"]:
            verify_native_content(document, content, tables, formula_report, image_report, validated)
            evidence["table_evidence"] = _build_table_evidence(document, content, evidence).model_dump(mode="json", warnings=False)
        return evidence
    except DocumentIntegrityError:
        raise
    except (ValueError, TypeError, KeyError, IndexError, AttributeError, UnicodeError):
        raise content_failure() from None


def validate_v2_table_trace(document, plan, trace, index, ownership, *, require_ownership):
    """纯计划确定期待的唯一表和行，真实 runtime 归属作为独立第二输入逐值对账。"""
    from app.services.renderer.block_evidence import TableRowsScope
    from app.services.renderer.layout_plan import build_layout_plan
    from app.services.renderer.logical_table_plan import build_content_table_plan, table_unit_contracts
    from app.services.document_ir.table_placeholder import decode_placeholders
    from app.services.renderer.table_placeholder_evidence import marker
    tables = build_content_table_plan(document, build_layout_plan(document))
    contracts = table_unit_contracts(document, tables)
    if plan.table_units != contracts or [unit.unit_id for unit in trace.table_units] != [unit.unit_id for unit in contracts]:
        raise content_failure("ir_render_evidence:target_mismatch")
    records = {record.block_id: record for record in trace.block_render_map}
    all_members = {block_id for contract in contracts for block_id in contract.member_block_ids}
    body = index.roots["word/document.xml"].find(qn("w:body"))
    children = list(body)
    expected_scopes = []
    expected_runtime = []
    for unit, contract, proof in zip(tables.units, contracts, trace.table_units, strict=True):
        target = proof.table_target
        table = index.nodes.get((target.part, tuple(target.child_path)))
        if table is None or table.tag != qn("w:tbl") or table.getparent() is not body:
            raise content_failure("ir_render_evidence:target_mismatch")
        caption = table.find("./" + qn("w:tblPr") + "/" + qn("w:tblCaption"))
        if (caption is None or caption.get(qn("w:val")) != unit.key
                or index.project(target.part, tuple(target.child_path)) != (target.element_kind, target.actual, target.bookmark_names)):
            raise content_failure("ir_render_evidence:target_mismatch")
        expected = {block_id: [] for block_id in contract.member_block_ids}
        annotations = []
        prefix = [item for item in contract.annotations if item.disposition == "rendered" and item.role != "note"]
        suffix = [item for item in contract.annotations if item.disposition == "rendered" and item.role == "note"]
        position = children.index(table)
        before = position - len(prefix) - bool(unit.blocking_quality_flags)
        if before < 0:
            raise content_failure("ir_render_evidence:target_mismatch")

        def add_annotation(item, paragraph):
            if paragraph.tag != qn("w:p"):
                raise content_failure("ir_render_evidence:target_mismatch")
            part, path = index.paths[paragraph]
            expected[item.block_id].append((part, path, "paragraph", item.role, item.item_index, None))
            annotations.append((item.block_id, item.field, item.item_index, path))

        for offset, item in enumerate(prefix):
            add_annotation(item, children[before + offset])
        if unit.blocking_quality_flags:
            node = children[position - 1]
            if node.tag != qn("w:p"):
                raise content_failure("ir_render_evidence:target_mismatch")
            expected[contract.anchor_block_id].append((*index.paths[node], "paragraph", "body", None, None))
        rows = table.findall(qn("w:tr"))
        frozen = []
        for fragment in contract.fragments:
            if fragment.content_scope_id is None:
                continue
            start, end = fragment.output_rows
            selected = rows[start:end]
            if not selected or len(selected) != end - start:
                raise content_failure("ir_render_evidence:target_mismatch")
            first = next(selected[0].iter(qn("w:p")))
            scope = TableRowsScope(scope_id=fragment.content_scope_id, unit_id=unit.key, block_id=fragment.block_id,
                part=target.part, table_child_path=target.child_path, output_rows=[start, end],
                row_child_paths=[list(index.paths[row][1]) for row in selected],
                first_paragraph_child_path=list(index.paths[first][1]))
            expected_scopes.append(scope)
            expected[fragment.block_id].append((target.part, tuple(target.child_path), "table_rows", "table", None, scope.scope_id))
            frozen.append((scope.scope_id, scope.block_id, tuple(tuple(path) for path in scope.row_child_paths),
                           tuple(scope.first_paragraph_child_path)))
        placeholders = [item for item in decode_placeholders(document)
                        if item.fragment_id in {fragment.source.fragment_id for fragment in unit.fragments}]
        for offset, binding in enumerate(placeholders):
            paragraph = children[position + 1 + offset]
            if paragraph.tag != qn("w:p") or not any(node.get(qn("w:name")) == marker(binding)
                                                       for node in paragraph.iter(qn("w:bookmarkStart"))):
                raise content_failure("ir_render_evidence:target_mismatch")
            expected[binding.owner_block_id].append((*index.paths[paragraph], "paragraph", "image_fallback", None, None))
        for offset, item in enumerate(suffix):
            add_annotation(item, children[position + 1 + len(placeholders) + offset])
        for block_id in contract.member_block_ids:
            record = records[block_id]
            actual = [(item.part, tuple(item.child_path), item.element_kind, item.role, item.item_index, item.content_scope_id)
                      for item in record.targets]
            if actual != expected[block_id]:
                raise content_failure("ir_render_evidence:target_mismatch")
            if actual:
                if record.disposition != "rendered" or record.representative_block_id is not None:
                    raise content_failure("ir_render_evidence:target_mismatch")
                if any(item[-1] is not None for item in actual) and record.reason_code != "table_fragment_rows":
                    raise content_failure("ir_render_evidence:target_mismatch")
            elif (record.disposition, record.reason_code, record.representative_block_id) != (
                    "coalesced", "table_header_coalesced", contract.anchor_block_id):
                raise content_failure("ir_render_evidence:target_mismatch")
        expected_runtime.append((unit.key, target.part, tuple(target.child_path), tuple(frozen), tuple(annotations)))
    if trace.content_scopes != expected_scopes:
        raise content_failure("ir_render_evidence:target_mismatch")
    if require_ownership:
        if ownership is None or [(unit.unit_id, unit.part, unit.table_child_path, unit.scopes, unit.annotations)
                                  for unit in ownership.tables] != expected_runtime:
            raise content_failure("ir_render_evidence:target_mismatch")
    return all_members

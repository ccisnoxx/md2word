"""公式有限字面转义、正体函数与认证范围残留检查的直接回归。"""

from __future__ import annotations

import copy
import io
import unittest
import zipfile
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.conversion.docx_rules.formula import (
    FORMULA_FALLBACK_NOTICE,
    formula_block_span,
    output_sha256,
    render_formula,
    scan_math_spans,
    sha256_text,
)
from app.services.document_ir.models import DocumentIR, DocumentIRBlock, DocumentIRMetadata
from app.services.renderer.docx_quality import (
    analyze_docx_quality_bytes,
    validate_formula_report,
    verify_docx_formula_report,
)
from app.services.renderer.evidence import build_renderer_evidence, ensure_renderer_integrity
from app.services.renderer.word_renderer import DocumentIRWordRenderer


NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
}
ENOB = r"ENOB=\frac{\ln(FSR/Noise\_RMS)}{\ln(2)}+\sin(x)\,V"


def _text(element) -> str:
    return "".join(node.text or "" for node in element.iter() if node.tag in {qn("w:t"), qn("m:t")})


def _document(*blocks: DocumentIRBlock) -> DocumentIR:
    return DocumentIR(metadata=DocumentIRMetadata(
        job_id="literal_job", source_file_id="literal_source", document_language="zh",
    ), blocks=list(blocks))


def _render(document_ir: DocumentIR):
    return DocumentIRWordRenderer().render_to_result(
        document_ir, template_path=None, config_path=None, image_loader=lambda _: None,
    )


def _root(content: bytes) -> ET.Element:
    with zipfile.ZipFile(io.BytesIO(content)) as package:
        return ET.fromstring(package.read("word/document.xml"))


def _rewrite(content: bytes, change) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as original, zipfile.ZipFile(output, "w") as target:
        for entry in original.infolist():
            value = original.read(entry.filename)
            if entry.filename == "word/document.xml":
                root = ET.fromstring(value)
                change(root)
                value = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(entry, value)
    return output.getvalue()


def _bound_fixture(source: str, *, raw_math_runs: tuple[str, ...] | None = None):
    """构造旧降级件或故意含残留的数学件；新语法正向路径使用正式 renderer。"""
    document = Document()
    paragraph = document.add_paragraph()
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "1000001")
    start.set(qn("w:name"), "md2w_formula_000001")
    paragraph._p.append(start)
    if raw_math_runs is None:
        nodes = [paragraph.add_run(FORMULA_FALLBACK_NOTICE)._r, paragraph.add_run(source)._r]
        mode, reason = "text_fallback", "invalid_syntax"
    else:
        math = OxmlElement("m:oMath")
        for value in raw_math_runs:
            run = OxmlElement("m:r")
            properties = OxmlElement("m:rPr")
            style = OxmlElement("m:sty")
            style.set(qn("m:val"), "p")
            properties.append(style)
            run.append(properties)
            text = OxmlElement("m:t")
            text.text = value
            run.append(text)
            math.append(run)
        paragraph._p.append(math)
        nodes = [math]
        mode, reason = "omml", None
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "1000001")
    paragraph._p.append(end)
    report = {"schema_version": "formula-render/v1", "entries": [{
        "occurrence": 1,
        "source": {"block_id": "fixture", "field": "latex", "start": 0, "end": len(source),
                   "boundary": "block", "source_sha256": sha256_text(source)},
        "anchor": "md2w_formula_000001", "mode": mode, "reason": reason,
        "output_projection_version": "formula-output/v1", "expected_output_sha256": output_sha256(nodes),
    }]}
    output = io.BytesIO()
    document.save(output)
    return output.getvalue(), report


class FormulaLiteralParserTests(unittest.TestCase):
    def math(self, source: str):
        result = render_formula(formula_block_span(source))
        self.assertEqual((result.mode, result.reason), ("omml", None))
        return result.elements[0]

    def test_punctuation_has_exact_characters_and_plain_style(self):
        math = self.math(r"A\,B\_RMS")
        self.assertEqual(_text(math), "A\u2009B_RMS")
        self.assertEqual([ord(char) for char in _text(math) if char in "\u2009_"], [0x2009, 0x005F])
        self.assertEqual(math.findall(".//m:sSub", NS), [])
        self.assertEqual({style.get(qn("m:val")) for style in math.findall(".//m:sty", NS)}, {"p"})

    def test_six_function_commands_have_plain_math_runs(self):
        for function in ("ln", "log", "exp", "sin", "cos", "tan"):
            with self.subTest(function=function):
                math = self.math("\\" + function + "(x)")
                self.assertEqual(_text(math), function + "(x)")
                run = math.find("./m:r", NS)
                self.assertEqual(_text(run), function)
                self.assertEqual(run.find("m:rPr/m:sty", NS).get(qn("m:val")), "p")

    def test_explicit_styles_text_spaces_and_script_arguments_survive(self):
        math = self.math(r"\mathbf{Noise\_RMS}_{\mathit{n}}+\mathit{\ln(x)}^{\mathbf{2}}+\text{ A\, B\_RMS }")
        self.assertEqual(_text(math), "Noise_RMSn+ln(x)2+ A\u2009 B_RMS ")
        for path, text, style in (("m:sSub/m:e", "Noise_RMS", "b"),
                                  ("m:sSub/m:sub", "n", "i"),
                                  ("m:sSup/m:e", "ln(x)", "i"),
                                  ("m:sSup/m:sup", "2", "b")):
            with self.subTest(path=path):
                element = math.find("./" + path, NS)
                self.assertEqual(_text(element), text)
                self.assertEqual({node.get(qn("m:val")) for node in element.findall(".//m:sty", NS)}, {style})
        self.assertTrue(all(node.get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"
                            for node in math.findall(".//m:t", NS)))

    def test_function_and_group_are_complete_script_bases(self):
        for source, tag, base in ((r"\ln_{2}(x)", "sSub", "ln"),
                                  (r"{\ln(x)+y}^{2}", "sSup", "ln(x)+y")):
            with self.subTest(source=source):
                math = self.math(source)
                element = math.find("./m:" + tag + "/m:e/m:box", NS)
                self.assertIsNotNone(element)
                self.assertEqual(_text(element), base)
                self.assertEqual(_text(math.find("./m:" + tag + "/m:" + ("sub" if tag == "sSub" else "sup"), NS)), "2")

    def test_unsupported_commands_and_structures_keep_all_failure_reasons(self):
        cases = ((r"\ln(x)+\unknown{a<2 & b}", "unsupported_command"),
                 (r"\sinh(x)", "unsupported_command"),
                 (r"\operatorname{ln}(x)", "unsupported_command"),
                 (r"A\;B", "invalid_syntax"),
                 (r"Noise\_RMS+\frac{1}{2", "invalid_syntax"),
                 (r"\begin{matrix}\ln(x)&y\end{matrix}", "unsupported_structure"),
                 (r"\left.\ln(x)\right|", "unsupported_structure"),
                 (r"\ln(x)\tag{\frac{1}{2}}", "unsupported_structure"),
                 (r"\text{<b>A B</b>}\_RMS", "unsafe_normalization"))
        for source, reason in cases:
            with self.subTest(source=source):
                result = render_formula(formula_block_span(source))
                self.assertEqual((result.mode, result.reason), ("text_fallback", reason))
                self.assertEqual("".join(_text(node) for node in result.elements), FORMULA_FALLBACK_NOTICE + source)

    def test_new_commands_do_not_expand_ordinary_math_scanning(self):
        ordinary = (r"C:\work\file_name", r"C:\theta\file", r"\\server\share\file", "REG_VDD",
                    "ordinary_name", r"\_RMS", r"\,", r"\$5", "$5 and $10")
        for source in ordinary + tuple("\\" + name + "(x)" for name in ("ln", "log", "exp", "sin", "cos", "tan")):
            with self.subTest(source=source):
                self.assertEqual(scan_math_spans(source), ())
        source = r"REG_VDD C:\work\file_name ln log exp sin cos tan"
        result = _render(_document(DocumentIRBlock(id="ordinary", type="paragraph", order=0, text=source)))
        self.assertEqual(result.formula_report["entries"], [])
        self.assertEqual(_text(_root(result.content)), source)
        self.assertEqual(_root(result.content).findall(".//m:oMath", NS), [])


class FormulaLiteralRenderingTests(unittest.TestCase):
    def test_enob_in_block_inline_and_html_table_has_bound_native_output(self):
        delimited = "$" + ENOB + "$"
        ir = _document(
            DocumentIRBlock(id="formula", type="formula", order=0, latex=ENOB),
            DocumentIRBlock(id="inline", type="paragraph", order=1, text="前 " + delimited + " 后"),
            DocumentIRBlock(id="table", type="table", order=2,
                            html="<table><tr><td>" + delimited + "</td></tr></table>"),
        )
        result = _render(ir)
        evidence = build_renderer_evidence(ir, result.content, render_result=result)
        ensure_renderer_integrity(evidence)
        self.assertEqual(evidence["formula_summary"]["omml_count"], 3)
        self.assertEqual(evidence["formula_summary"]["fallback_count"], 0)
        self.assertEqual(evidence["docx_quality"]["latex_residue_count"], 0)
        self.assertEqual(result.formula_report["schema_version"], "formula-render/v1")
        maths = _root(result.content).findall(".//m:oMath", NS)
        self.assertEqual(len(maths), 3)
        sources = (("formula", "latex", ENOB, 0, "block"),
                   ("inline", "text", delimited, 2, "dollar"),
                   ("table", "html/0/0", delimited, 0, "dollar"))
        for occurrence, (math, entry, source) in enumerate(zip(maths, result.formula_report["entries"], sources, strict=True), 1):
            with self.subTest(occurrence=occurrence):
                block, field, raw, start, boundary = source
                self.assertEqual(_text(math), "ENOB=ln(FSR/Noise_RMS)ln(2)+sin(x)\u2009V")
                self.assertEqual(_text(math.find(".//m:num", NS)), "ln(FSR/Noise_RMS)")
                self.assertEqual(_text(math.find(".//m:den", NS)), "ln(2)")
                self.assertEqual(math.findall(".//m:sSub", NS), [])
                self.assertEqual(entry["source"], {"block_id": block, "field": field, "start": start,
                    "end": start + len(raw), "boundary": boundary, "source_sha256": sha256_text(raw)})
                self.assertEqual((entry["occurrence"], entry["anchor"], entry["mode"], entry["reason"]),
                                 (occurrence, f"md2w_formula_{occurrence:06d}", "omml", None))
                self.assertEqual(entry["output_projection_version"], "formula-output/v1")
                self.assertEqual(entry["expected_output_sha256"], output_sha256([math]))
                self.assertEqual({node.get(qn("m:val")) for node in math.findall(".//m:sty", NS)}, {"p"})

    def test_without_report_literal_underscore_keeps_original_residue(self):
        result = _render(_document(DocumentIRBlock(id="formula", type="formula", order=0, latex=r"Noise\_RMS")))
        original = analyze_docx_quality_bytes(result.content)
        certified = analyze_docx_quality_bytes(result.content, formula_report=result.formula_report)
        self.assertEqual(original["latex_residue_examples"], ["_RMS"])
        self.assertEqual(original["latex_residue_count"], 1)
        self.assertEqual(certified["latex_residue_count"], 0)
        self.assertEqual(original["text_chars"], certified["text_chars"])
        self.assertEqual(certified["text_chars"], len("Noise_RMS"))
        self.assertEqual(_text(_root(result.content)), "Noise_RMS")

    def test_equal_ordinary_text_in_same_paragraph_remains_a_residue(self):
        ir = _document(DocumentIRBlock(id="inline", type="paragraph", order=0,
                       text=r"Noise_RMS $Noise\_RMS$ Noise_RMS"))
        result = _render(ir)
        quality = analyze_docx_quality_bytes(result.content, formula_report=result.formula_report)
        self.assertEqual(quality["latex_residue_examples"], ["_RMS"])
        self.assertEqual(quality["latex_residue_count"], 1)
        root = _root(result.content)
        self.assertEqual(_text(root), "Noise_RMS Noise_RMS Noise_RMS")
        self.assertEqual(_text(root.find(".//m:oMath", NS)), "Noise_RMS")
        self.assertIn("Noise_RMS", "".join(node.text or "" for node in root.findall(".//w:t", NS)))

    def test_literal_underscore_does_not_hide_adjacent_ordinary_command(self):
        ir = _document(DocumentIRBlock(id="inline", type="paragraph", order=0,
                       text=r"$Noise\_$\unknown"))
        result = _render(ir)
        quality = analyze_docx_quality_bytes(result.content, formula_report=result.formula_report)
        root = _root(result.content)
        self.assertEqual(len(result.formula_report["entries"]), 1)
        self.assertEqual(result.formula_report["entries"][0]["mode"], "omml")
        self.assertEqual(_text(root.find(".//m:oMath", NS)), "Noise_")
        self.assertEqual("".join(node.text or "" for node in root.findall(".//w:t", NS)), r"\unknown")
        self.assertEqual(_text(root), r"Noise_\unknown")
        self.assertEqual(quality["latex_residue_examples"], [r"_\unknown"])
        self.assertEqual(quality["latex_residue_count"], 1)

    def test_offsets_follow_paragraphs_fallback_exclusion_and_table_cells(self):
        ir = _document(
            DocumentIRBlock(id="prefix", type="paragraph", order=0, text="   前置段落   "),
            DocumentIRBlock(id="mixed", type="paragraph", order=1,
                            text=r"   前 $\unknown{_OLD}$ 中 $Noise\_RMS$ 后   "),
            DocumentIRBlock(id="table", type="table", order=2,
                            html=r"<table><tr><td>$Noise\_RMS$</td><td></td><td>$Noise\_RMS$</td></tr></table>"),
        )
        result = _render(ir)
        quality = analyze_docx_quality_bytes(result.content, formula_report=result.formula_report)
        self.assertEqual(quality["latex_residue_examples"], [])
        self.assertEqual(quality["latex_residue_count"], 0)
        self.assertEqual(validate_formula_report(result.formula_report)["fallback_count"], 1)
        def add_ordinary(root):
            body = root.find("w:body", NS)
            paragraph = ET.Element(qn("w:p"))
            run = ET.SubElement(paragraph, qn("w:r"))
            ET.SubElement(run, qn("w:t")).text = "Noise_RMS"
            body.insert(0, paragraph)
        quality = analyze_docx_quality_bytes(_rewrite(result.content, add_ordinary), formula_report=result.formula_report)
        self.assertEqual(quality["latex_residue_examples"], ["_RMS"])
        self.assertEqual(quality["latex_residue_count"], 1)

    def test_bad_or_missing_report_bindings_are_rejected(self):
        result = _render(_document(DocumentIRBlock(id="formula", type="formula", order=0, latex=r"Noise\_RMS")))
        variants = []
        for field, value in (("expected_output_sha256", "0" * 64), ("anchor", "wrong"),
                             ("mode", "semantic_runs"), ("reason", "unsupported_command")):
            report = copy.deepcopy(result.formula_report)
            report["entries"][0][field] = value
            variants.append(report)
        variants.extend(({"schema_version": "bad", "entries": []},
                         {"schema_version": "formula-render/v1", "entries": []}))
        for report in variants:
            with self.subTest(report=report), self.assertRaises(ValueError):
                analyze_docx_quality_bytes(result.content, formula_report=report)

    def test_unregistered_math_cannot_gain_an_exception(self):
        result = _render(_document(DocumentIRBlock(id="formula", type="formula", order=0, latex=r"Noise\_RMS")))
        def append_unregistered(root):
            paragraph = root.find(".//w:p", NS)
            paragraph.append(copy.deepcopy(paragraph.find("m:oMath", NS)))
        with self.assertRaisesRegex(ValueError, "未登记"):
            analyze_docx_quality_bytes(_rewrite(result.content, append_unregistered), formula_report=result.formula_report)

    def test_character_space_and_math_style_tampering_break_binding(self):
        ir = _document(DocumentIRBlock(id="formula", type="formula", order=0, latex=r"\ln(A\,Noise\_RMS)"))
        result = _render(ir)
        def change_underscore(root):
            next(node for node in root.findall(".//m:t", NS) if node.text == "_").text = "-"
        def change_thin_space(root):
            next(node for node in root.findall(".//m:t", NS) if node.text == "\u2009").text = " "
        def change_style(root):
            root.find(".//m:sty", NS).set(qn("m:val"), "i")
        for change in (change_underscore, change_thin_space, change_style):
            with self.subTest(change=change.__name__), self.assertRaisesRegex(ValueError, "预期绑定"):
                build_renderer_evidence(ir, _rewrite(result.content, change), formula_report=result.formula_report)

    def test_new_syntax_does_not_hide_unknown_subexpression_in_three_entries(self):
        expression = r"\ln(x)+Noise\_RMS+\unknown{z}"
        ir = _document(
            DocumentIRBlock(id="formula", type="formula", order=0, latex=expression),
            DocumentIRBlock(id="inline", type="paragraph", order=1, text="$" + expression + "$"),
            DocumentIRBlock(id="table", type="table", order=2,
                            html="<table><tr><td>$" + expression + "$</td></tr></table>"),
        )
        result = _render(ir)
        evidence = build_renderer_evidence(ir, result.content, render_result=result)
        ensure_renderer_integrity(evidence)
        self.assertEqual(evidence["formula_summary"]["reason_counts"], {"unsupported_command": 3})
        self.assertEqual(evidence["docx_quality"]["latex_residue_count"], 0)
        self.assertIn("formula_text_fallback", evidence["quality_flags"])
        self.assertEqual(_root(result.content).findall(".//m:oMath", NS), [])
        for entry, raw in zip(result.formula_report["entries"], (expression, "$" + expression + "$", "$" + expression + "$"), strict=True):
            self.assertEqual((entry["mode"], entry["reason"]), ("text_fallback", "unsupported_command"))
            self.assertEqual(entry["source"]["source_sha256"], sha256_text(raw))
            self.assertIn(FORMULA_FALLBACK_NOTICE + raw, _text(_root(result.content)))

    def test_old_fallback_remains_verified_without_reparsing_or_upgrade(self):
        source = r"Noise\_RMS+\ln(x)\,V"
        content, report = _bound_fixture(source)
        original_report = copy.deepcopy(report)
        ir = _document(DocumentIRBlock(id="fixture", type="formula", order=0, latex=source))
        evidence = build_renderer_evidence(ir, content, formula_report=report)
        ensure_renderer_integrity(evidence)
        self.assertEqual(evidence["formula_summary"]["fallback_count"], 1)
        self.assertEqual(evidence["formula_summary"]["omml_count"], 0)
        self.assertEqual(evidence["formula_summary"]["reason_counts"], {"invalid_syntax": 1})
        self.assertEqual(evidence["docx_quality"]["latex_residue_count"], 0)
        self.assertEqual(_text(_root(content)), FORMULA_FALLBACK_NOTICE + source)
        self.assertEqual(_root(content).findall(".//m:oMath", NS), [])
        self.assertEqual(report, original_report)
        verify_docx_formula_report(content, report)

    def test_old_fallback_still_binds_notice_and_exact_source(self):
        content, report = _bound_fixture(r"Noise\_RMS+\ln(x)\,V")
        for index in (0, 1):
            def truncate(root):
                node = root.findall(".//w:t", NS)[index]
                node.text = (node.text or "")[:-1]
            with self.subTest(index=index), self.assertRaisesRegex(ValueError, "预期绑定"):
                analyze_docx_quality_bytes(_rewrite(content, truncate), formula_report=report)


class CertifiedMathResidueTests(unittest.TestCase):
    def test_other_residues_remain_visible_in_byte_bound_math(self):
        for value, expected in ((r"Noise\_RMS", "_RMS"), (r"\unknown", r"\unknown"),
                                ("$x$", "$x$"), ("x^2", "^2")):
            with self.subTest(value=value):
                content, report = _bound_fixture("x", raw_math_runs=(value,))
                quality = analyze_docx_quality_bytes(content, formula_report=report)
                self.assertEqual(quality["latex_residue_examples"], [expected])
                self.assertEqual(quality["latex_residue_count"], 1)
                self.assertEqual(_text(_root(content)), value)
                self.assertEqual(quality["text_chars"], len(value))

    def test_raw_backslash_across_math_runs_prevents_underscore_exception(self):
        content, report = _bound_fixture("x", raw_math_runs=("Noise\\", "_", "RMS"))
        quality = analyze_docx_quality_bytes(content, formula_report=report)
        self.assertEqual(quality["latex_residue_examples"], ["_RMS"])
        self.assertEqual(quality["latex_residue_count"], 1)
        self.assertEqual(_text(_root(content)), r"Noise\_RMS")


if __name__ == "__main__":
    unittest.main()

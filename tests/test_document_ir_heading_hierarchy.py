from __future__ import annotations

import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

from app.domains.templates.layout_config_schema import TemplateConfigPackage, TemplateParagraphStyleConfig
from app.domains.translation.document_ir_translator import DocumentIRTranslationService
from app.domains.translation.models import TranslationProviderRequest, TranslationProviderResult, TranslationTextItem
from app.services.document_ir import MinerUToDocumentIRService
from app.services.document_ir.models import (
    DocumentIR,
    DocumentIRBlock,
    DocumentIRBlockLayout,
    DocumentIRMetadata,
    DocumentIRPage,
)
from app.services.renderer.docx_quality import analyze_docx_evidence_bytes
from app.services.renderer.evidence import build_renderer_evidence
from app.services.renderer.word_renderer import DocumentIRWordRenderer


class _MappedProvider:
    """按 item id 返回固定译文，验证身份拆分读取翻译后的最终文本。"""

    def __init__(self, values: dict[str, str]) -> None:
        self.values = values

    def translate(self, request: TranslationProviderRequest) -> TranslationProviderResult:
        return TranslationProviderResult(
            items=tuple(TranslationTextItem(id=item.id, text=self.values[item.id]) for item in request.items)
        )


class DocumentIRHeadingHierarchyTests(unittest.TestCase):
    def test_converter_restores_numbered_depth_and_preserves_unnumbered_nested_level(self) -> None:
        blocks = [
            self._v2_title("5 Specifications", level=2, top=80),
            self._v2_title("5.1 Absolute Maximum Ratings", level=2, top=130),
            self._v2_title("8.2.1 Layout Guidelines", level=2, top=180),
            self._v2_title("8.2.1.1 Layout Example", level=2, top=230),
            self._v2_title("产品概况", level=2, top=280),
            self._v2_title("Chapter 2", level=2, top=330),
            self._v2_title("Appendix A", level=2, top=380),
            self._v2_title("第十二章", level=2, top=430),
        ]

        document = self._convert_v2([blocks])

        headings = [block for block in document.blocks if block.type == "heading"]
        self.assertEqual([block.text for block in headings], [
            "5 Specifications",
            "5.1 Absolute Maximum Ratings",
            "8.2.1 Layout Guidelines",
            "8.2.1.1 Layout Example",
            "产品概况",
            "Chapter 2",
            "Appendix A",
            "第十二章",
        ])
        self.assertEqual([block.level for block in headings], [1, 2, 3, 3, 2, 1, 1, 1])

    def test_converter_marks_only_unique_top_wide_first_page_product_title(self) -> None:
        cases = {
            "unique_top_wide": ([self._v2_title("SC2167 Variable Resolution R/D Converter", level=2, top=70, left=80, right=920)], 1),
            "not_unique": (
                [
                    self._v2_title("SC2167 Variable Resolution R/D Converter", level=2, top=70, left=80, right=920),
                    self._v2_title("AD2S1210 Resolver-to-Digital Converter", level=2, top=110, left=80, right=920),
                ],
                0,
            ),
            "not_top": ([self._v2_title("SC2167 Variable Resolution R/D Converter", level=2, top=260, left=80, right=920)], 0),
            "not_wide": ([self._v2_title("SC2167 Variable Resolution R/D Converter", level=2, top=70, left=80, right=420)], 0),
        }
        for name, (blocks, expected_count) in cases.items():
            with self.subTest(name=name):
                document = self._convert_v2([blocks])
                identities = [
                    block for block in document.blocks if block.metadata.get("annotation_role") == "document_identity"
                ]
                self.assertEqual(len(identities), expected_count)
                if expected_count:
                    self.assertEqual(identities[0].type, "paragraph")
                    self.assertIsNone(identities[0].level)
                else:
                    self.assertTrue(all(block.type == "heading" for block in document.blocks))

    def test_translated_identity_uses_part_and_category_styles_without_entering_toc(self) -> None:
        source = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_identity", source_file_id="file_identity", document_language="en"),
            pages=[DocumentIRPage(index=0, width=1000.0, height=1200.0)],
            blocks=[
                DocumentIRBlock(
                    id="identity",
                    type="paragraph",
                    order=0,
                    text="SC2167 Variable Resolution R/D Converter",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(80, 70, 920, 105), role="body"),
                    metadata={"annotation_role": "document_identity", "semantic_decision_reason": "first_page_document_identity"},
                ),
                DocumentIRBlock(
                    id="toc_anchor",
                    type="paragraph",
                    order=1,
                    text="",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(80, 150, 920, 180), role="body"),
                    metadata={"annotation_role": "toc_anchor", "semantic_decision_reason": "synthetic_toc_anchor"},
                ),
                DocumentIRBlock(
                    id="heading",
                    type="heading",
                    order=2,
                    text="Product Overview",
                    level=2,
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(80, 220, 400, 250), role="body"),
                ),
            ],
        )
        translated = DocumentIRTranslationService(
            _MappedProvider(
                {
                    "identity.text": "SC2167 可变分辨率 R/D 转换器",
                    "heading.text": "产品概况",
                }
            )
        ).translate(source, term_version=None)
        config = TemplateConfigPackage(
            part_number=TemplateParagraphStyleConfig(font_size=18.0, bold=True),
            category=TemplateParagraphStyleConfig(font_size=9.0, italic=True),
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            translated,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
            template_config=config,
        )
        rendered = Document(io.BytesIO(content))
        evidence = build_renderer_evidence(translated, content, template_config=config)
        _quality, content_evidence = analyze_docx_evidence_bytes(content)

        self.assertEqual([paragraph.text for paragraph in rendered.paragraphs[:2]], ["SC2167", "可变分辨率 R/D 转换器"])
        self.assertEqual(rendered.paragraphs[0].runs[0].font.size.pt, 18.0)
        self.assertTrue(rendered.paragraphs[0].runs[0].bold)
        self.assertEqual(rendered.paragraphs[1].runs[0].font.size.pt, 9.0)
        self.assertTrue(rendered.paragraphs[1].runs[0].italic)
        self.assertTrue(rendered.paragraphs[0].paragraph_format.keep_with_next)
        self.assertFalse(rendered.paragraphs[1].paragraph_format.keep_with_next)
        self.assertEqual(evidence["expected_heading_count"], 1)
        self.assertEqual(evidence["heading_count"], 1)
        self.assertEqual(evidence["visible_payload_evidence"]["counts"]["paragraph"], {"expected": 2, "covered": 2})
        self.assertEqual(evidence["blocking_quality_flags"], [])
        self.assertTrue(all("SC2167" not in text for text in content_evidence["toc_cache_texts"]))
        self.assertTrue(any("产品概况" in text for text in content_evidence["toc_cache_texts"]))

    def test_renderer_fails_when_final_identity_text_can_no_longer_be_split(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_bad_identity", source_file_id="file_bad_identity", document_language="zh"),
            pages=[DocumentIRPage(index=0, width=1000.0, height=1200.0)],
            blocks=[
                DocumentIRBlock(
                    id="identity",
                    type="paragraph",
                    order=0,
                    text="翻译后的产品说明",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(80, 70, 920, 105), role="body"),
                    metadata={"annotation_role": "document_identity", "semantic_decision_reason": "first_page_document_identity"},
                )
            ],
        )

        with self.assertRaisesRegex(ValueError, "文档身份文本无法拆分"):
            DocumentIRWordRenderer().render_to_bytes(
                document,
                template_path=None,
                config_path=None,
                image_loader=lambda _asset: None,
            )

    def _convert_v2(self, pages: list[list[dict[str, object]]]) -> DocumentIR:
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            layout = {
                "pdf_info": [
                    {"page_idx": page_index, "page_size": [1000, 1200]}
                    for page_index in range(len(pages))
                ]
            }
            with zipfile.ZipFile(zip_path, "w") as archive:
                archive.writestr("result_content_list_v2.json", json.dumps(pages, ensure_ascii=False))
                archive.writestr("layout.json", json.dumps(layout, ensure_ascii=False))
            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_heading_hierarchy",
                source_file_id="file_heading_hierarchy",
                document_language="zh",
            )
        return result.document

    @staticmethod
    def _v2_title(
        text: str,
        *,
        level: int,
        top: int,
        left: int = 80,
        right: int = 500,
    ) -> dict[str, object]:
        return {
            "type": "title",
            "content": {
                "level": level,
                "title_content": [{"type": "text", "content": text}],
            },
            "bbox": [left, top, right, top + 30],
        }


if __name__ == "__main__":
    unittest.main()

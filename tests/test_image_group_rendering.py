"""RND-503 的源计划、原生布局、冻结输入和损坏拒绝回归。"""

import copy
from concurrent.futures import ThreadPoolExecutor
import hashlib
import io
import json
from pathlib import Path
import struct
import tempfile
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
import zipfile
import zlib

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from pydantic import ValidationError

from app.conversion.docx_rules.formula import canonical_sha256
from app.domains.templates.layout_config_schema import TemplateConfigPackage
from app.services.document_ir.models import DocumentIR, DocumentIRMetadata, DocumentIRPage, DocumentIRBlock, DocumentIRBlockLayout, DocumentIRAsset
from app.services.document_ir.mineru_converter import MinerUToDocumentIRService
from app.services.renderer import image_evidence
from app.services.renderer.docx_quality import analyze_docx_quality_bytes
from app.services.renderer.evidence import build_renderer_evidence, ensure_renderer_integrity
from app.services.renderer.layout_plan import build_image_plan, image_row_metrics, _aligned, IMAGE_GROUP_THRESHOLDS
from app.services.renderer.word_renderer import DocumentIRWordRenderer, DocumentIRRenderOptions

NS = image_evidence._NS


def neutral_png(width=120, height=60, color=(60, 100, 140)):
    """只含固定中性色像素的有效 PNG，不依赖额外图像库。"""
    def chunk(kind, data):
        return struct.pack('!I', len(data)) + kind + data + struct.pack('!I', zlib.crc32(kind + data) & 0xffffffff)
    pixels = (b'\0' + bytes(color) * width) * height
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('!2I5B', width, height, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(pixels)) + chunk(b'IEND', b'')


def image_ir(count=2, *, annotations=True):
    blocks = [DocumentIRBlock(id=f'image_{i}', type='image', order=i, page_index=0, asset_ref=f'asset_{i}',
        layout=DocumentIRBlockLayout(page_index=0, source_index=i, bbox=(10 + 120*i, 100, 110 + 120*i, 200)),
        captions=[f'Figure {i+1}: DemoToken <b>bold</b> $x^2$'] if annotations else [],
        footnotes=[f'Note {i+1}: extra', 'See Figure 1'] if annotations else []) for i in range(count)]
    return DocumentIR(schema_version='document-ir/v2', metadata=DocumentIRMetadata(job_id='job', source_file_id='source', document_language='en'),
        pages=[DocumentIRPage(index=0, width=400, height=600)], blocks=blocks,
        assets=[DocumentIRAsset(id=f'asset_{i}', kind='image', source_path=f'images/{i}.png',
            checksum_sha256='sha256:' + hashlib.sha256(neutral_png(120 if i%2 == 0 else 50, 60 if i%2 == 0 else 150)).hexdigest()) for i in range(count)])


def image_loader(asset):
    index = int(asset.id.split('_')[-1])
    return neutral_png(120 if index%2 == 0 else 50, 60 if index%2 == 0 else 150)


def render(ir, **kwargs):
    return DocumentIRWordRenderer().render_to_result(ir, template_path=kwargs.pop('template_path', None),
        config_path=kwargs.pop('config_path', None), image_loader=kwargs.pop('image_loader', image_loader), **kwargs)


def evidence(ir, result, **kwargs):
    return build_renderer_evidence(ir, result.content, formula_report=result.formula_report,
                                   image_layout_report=result.image_layout_report, **kwargs)


def change_docx(content, change, *, part='word/document.xml'):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(output, 'w') as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == part:
                root = ET.fromstring(data)
                change(root)
                data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            target.writestr(item, data)
    return output.getvalue()


def legacy_image_docx(content):
    """构造无新图片标记的旧单图 writer 文件，专用于历史恢复测试。"""
    def remove(root):
        if any(node.get(qn('w:val'), '').startswith(image_evidence.IMAGE_TABLE_PREFIX) for node in root.findall('.//w:tblCaption', NS)):
            raise ValueError('历史单图夹具不能包含新布局表')
        ids = {node.get(qn('w:id')) for node in root.findall('.//w:bookmarkStart', NS) if node.get(qn('w:name'), '').startswith('md2w_img_')}
        for parent in root.iter():
            for node in list(parent):
                if node.tag in {qn('w:bookmarkStart'), qn('w:bookmarkEnd')} and node.get(qn('w:id')) in ids:
                    parent.remove(node)
        for node in root.findall('.//wp:docPr', NS):
            if node.get('name', '').startswith('md2w_img_'):
                node.set('name', 'Picture')
    return change_docx(content, remove)


def write_numbered_image_zip(path, numbers=(5, 6)):
    """正式 ZIP 形态的可靠独立图题；重用时不预修真实不可靠源串。"""
    blocks = [{'type': 'image', 'img_path': f'images/{i}.png', 'image_caption': [f'图 {number} 合成中性示例'],
               'image_footnote': [f'注 {number} 保留'], 'bbox': [10+i*120, 100, 110+i*120, 200], 'page_idx': 0}
              for i, number in enumerate(numbers)]
    with zipfile.ZipFile(path, 'w') as package:
        package.writestr('sample_content_list.json', json.dumps(blocks, ensure_ascii=False))
        package.writestr('sample_middle.json', json.dumps({'pdf_info': [{'page_idx': 0, 'page_size': [400, 600]}]}))
        for i in range(len(numbers)):
            package.writestr(f'images/{i}.png', neutral_png(120 if i%2 == 0 else 50, 60 if i%2 == 0 else 150))


class ImageGroupPlanTests(unittest.TestCase):
    def test_two_three_and_independent_rows_are_deterministic_without_mutation(self):
        for count in (2, 3):
            ir = image_ir(count)
            before = ir.model_dump_json()
            self.assertEqual([span.block_ids for span in build_image_plan(ir)], [tuple(b.id for b in ir.blocks)])
            self.assertEqual(build_image_plan(ir), build_image_plan(ir))
            self.assertEqual(ir.model_dump_json(), before)
        ir = image_ir(2)
        extra = [b.model_copy(deep=True) for b in ir.blocks]
        for i, b in enumerate(extra):
            b.id += '_row2'; b.order += 2
            box = b.layout.bbox
            b.layout.bbox = (box[0], 300, box[2], 400)
        ir.blocks.extend(extra)
        self.assertEqual([span.mode for span in build_image_plan(ir)], ['group', 'group'])

    def test_all_vertical_threshold_equalities_and_both_sides(self):
        baseline = {'common_overlap': 1.0, 'height_ratio': 1.0, 'center_offset': 0.0, 'top_offset': 0.0}
        for key in baseline:
            for delta in (-0.00001, 0, 0.00001):
                metrics = dict(baseline, **{key: IMAGE_GROUP_THRESHOLDS[key] + delta})
                expected = delta >= 0 if key in {'common_overlap', 'height_ratio'} else delta <= 0
                with self.subTest(key=key, delta=delta):
                    self.assertEqual(_aligned(metrics), expected)
        for gap, accepted in ((31.999, True), (32, True), (32.001, False)):
            ir = image_ir()
            ir.blocks[1].layout.bbox = (110+gap, 100, 210+gap, 200)
            self.assertEqual(build_image_plan(ir)[0].mode == 'group', accepted)

    def test_whole_component_chain_and_column_major_reject_without_moving_source(self):
        ir = image_ir(3)
        for i, b in enumerate(ir.blocks):
            box = b.layout.bbox; b.layout.bbox = (box[0], 100+i*9, box[2], 200+i*9)
        plan = build_image_plan(ir)
        self.assertEqual(len(plan), 1)
        self.assertIn('competing_row_membership', plan[0].reasons)
        ir = image_ir(2)
        ir.blocks.extend([b.model_copy(deep=True) for b in ir.blocks])
        for i, b in enumerate(ir.blocks):
            b.id = str(i); b.order = i
        ir.blocks[1].layout.bbox = (10, 300, 110, 400)
        ir.blocks[2].layout.bbox = (130, 100, 230, 200)
        ir.blocks[3].layout.bbox = (130, 300, 230, 400)
        self.assertEqual([span.block_ids for span in build_image_plan(ir)], [('0',), ('1',), ('2',), ('3',)])
        result = render(ir)
        self.assertEqual([b['block_id'] for span in result.image_layout_report['spans'] for b in span['images']], ['0', '1', '2', '3'])

    def test_invalid_missing_source_conflict_overlap_and_uncertainty_veto(self):
        cases = [(None, 'geometry_unavailable'), ((float('nan'), 100, 220, 200), 'geometry_invalid'),
                 ((-1, 100, 220, 200), 'geometry_invalid'), ((90, 100, 190, 200), 'horizontal_overlap'),
                 ((0, 100, 100, 200), 'source_order_conflict')]
        for box, reason in cases:
            ir = image_ir(); ir.blocks[1].layout.bbox = box
            with self.subTest(reason=reason):
                self.assertTrue(any(reason in span.reasons for span in build_image_plan(ir)))
                result = render(ir)
                self.assertFalse(Document(io.BytesIO(result.content)).tables)
                ensure_renderer_integrity(evidence(ir, result))
        for receipt, reason in (({'status': 'uncertain', 'candidates': []}, 'caption_binding_uncertain'),
                                ({'status': 'evidence_only', 'bindings': [{'evidence_only': True}]}, 'external_caption_binding')):
            ir = image_ir(); ir.blocks[0].metadata['caption_reconciliation'] = receipt
            self.assertIn(reason, build_image_plan(ir)[0].reasons)
        ir = image_ir(); ir.quality_flags.append('caption_binding_uncertain')
        self.assertIn('caption_binding_uncertain', build_image_plan(ir)[0].reasons)
        ir = image_ir(); ir.blocks[1].order = 0
        self.assertIn('source_order_conflict', build_image_plan(ir)[0].reasons)

    def test_single_legacy_page_boundary_and_nonimage_interrupt_are_conservative(self):
        for ir in (image_ir(1), image_ir(1).model_copy(update={'pages': []}), image_ir(1).model_copy(update={'schema_version': 'document-ir/v1'})):
            self.assertEqual(build_image_plan(ir)[0].reasons, ())
            self.assertNotIn('image_group_uncertain', evidence(ir, render(ir))['quality_flags'])
        ir = image_ir(); ir.blocks[1].page_index = 1; ir.blocks[1].layout.page_index = 1
        ir.pages.append(DocumentIRPage(index=1, width=600, height=400))
        self.assertTrue(all(span.mode == 'single' for span in build_image_plan(ir)))
        result = render(ir); ensure_renderer_integrity(evidence(ir, result))
        self.assertNotEqual(result.image_layout_report['render_inputs']['spans'][0]['geometry'], result.image_layout_report['render_inputs']['spans'][1]['geometry'])
        ir = image_ir(); ir.blocks[1].order = 2
        ir.blocks.insert(1, DocumentIRBlock(id='text', type='paragraph', order=1, page_index=0, text='Interruption', layout=DocumentIRBlockLayout(page_index=0,bbox=(10,210,250,230))))
        self.assertTrue(all(span.mode == 'single' for span in build_image_plan(ir)))

    def test_nonfinite_geometry_report_is_json_safe_without_rewriting_ir(self):
        for coordinate in (float('nan'),float('inf'),float('-inf')):
            ir=image_ir(); ir.blocks[0].layout.bbox=(coordinate,100,110,200)
            result=render(ir)
            # 报告进入 file metadata 的标准 JSON 往返不能依赖 NaN 相等性。
            report=json.loads(json.dumps(result.image_layout_report,allow_nan=False))
            ev=build_renderer_evidence(ir,result.content,formula_report=result.formula_report,image_layout_report=report)
            ensure_renderer_integrity(ev)
            self.assertEqual(report['spans'][0]['images'][0]['source']['bbox'][0],None)
            # 可读取的历史原始 JSON 非有限值仍按同一 IR 验证；不重写原文件。
            raw=json.dumps(ir.model_dump(mode='json')).encode()
            historical=DocumentIR.model_validate_json(raw)
            ensure_renderer_integrity(build_renderer_evidence(historical,result.content,
                formula_report=result.formula_report,image_layout_report=report))
            # Pydantic 的新 IR 序列化会把非有限值变 null；该无效 typed bbox 本来就不能再读。
            with self.assertRaises(ValidationError): DocumentIR.model_validate_json(ir.model_dump_json())


class ImageGroupWriterTests(unittest.TestCase):
    def test_uncertain_span_records_only_changed_member_geometry_in_native_flow(self):
        ir=image_ir()
        known, unknown=ir.blocks
        known.order=1; known.layout.bbox=(10,110,180,210)
        unknown.order=2; unknown.layout.bbox=None
        def text(name,order,box):
            return DocumentIRBlock(id=name,type='paragraph',order=order,page_index=0,text=name,
                layout=DocumentIRBlockLayout(page_index=0,bbox=box))
        ir.blocks=[text('left_before',0,(10,10,180,100)),known,unknown,
            text('left_after',3,(10,220,180,310)),text('right_top',4,(220,10,390,150)),
            text('right_bottom',5,(220,160,390,310))]
        result=render(ir)
        frozen=result.image_layout_report['render_inputs']['spans'][0]
        self.assertEqual(frozen['geometry']['columns'],2)
        self.assertEqual(set(frozen['image_geometries']),{unknown.id})
        self.assertEqual(frozen['image_geometries'][unknown.id]['columns'],1)
        self.assertEqual([item['block_id'] for item in result.image_layout_report['spans'][0]['images']],[known.id,unknown.id])
        ensure_renderer_integrity(evidence(ir,result))
        incomplete=copy.deepcopy(result.image_layout_report)
        incomplete['render_inputs']['spans'][0].pop('image_geometries')
        incomplete['render_inputs_sha256']=canonical_sha256(incomplete['render_inputs'])
        with self.assertRaisesRegex(ValueError,'section 几何'):
            build_renderer_evidence(ir,result.content,formula_report=result.formula_report,image_layout_report=incomplete)

    def test_reverse_source_images_keep_source_sequence_instead_of_column_reordering(self):
        ir=image_ir()
        ir.blocks[0].layout.bbox=(220,100,320,200)
        ir.blocks[1].layout.bbox=(10,100,110,200)
        result=render(ir)
        ensure_renderer_integrity(evidence(ir,result))
        self.assertEqual(result.image_layout_report['spans'][0]['reasons'],['source_order_conflict'])
        self.assertEqual([item['block_id'] for item in result.image_layout_report['spans'][0]['images']],['image_0','image_1'])
        self.assertEqual(result.image_layout_report['render_inputs']['spans'][0]['geometry']['columns'],1)

    def test_layout_table_overrides_template_indent_and_alignment(self):
        ir=image_ir()
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'table-indent.docx'
            template=Document()
            style=template.styles['Normal Table']._element
            props=style.find(qn('w:tblPr'))
            if props is None:
                props=OxmlElement('w:tblPr'); style.append(props)
            indent=OxmlElement('w:tblInd'); indent.set(qn('w:w'),'20000'); indent.set(qn('w:type'),'dxa'); props.append(indent)
            alignment=OxmlElement('w:jc'); alignment.set(qn('w:val'),'right'); props.append(alignment)
            template.save(path)
            result=render(ir,template_path=path)
        table=Document(io.BytesIO(result.content)).tables[0]
        self.assertIsNotNone(table._tbl.tblPr.find(qn('w:tblInd')))
        self.assertEqual(table._tbl.tblPr.find(qn('w:tblInd')).attrib,{qn('w:w'):'0',qn('w:type'):'dxa'})
        self.assertEqual(table._tbl.tblPr.find(qn('w:jc')).get(qn('w:val')),'left')
        ensure_renderer_integrity(evidence(ir,result))

    def test_table_position_inheritance_and_exact_row_height_are_authenticated(self):
        ir=image_ir(); result=render(ir)
        def without_direct(root):
            props=root.find('.//w:tblPr',NS)
            for tag in ('tblInd','jc'):
                for node in props.findall('./w:'+tag,NS):props.remove(node)
        old_content=change_docx(result.content,without_direct)
        ensure_renderer_integrity(build_renderer_evidence(ir,old_content,formula_report=result.formula_report,
            image_layout_report=result.image_layout_report))
        def inherited_indent(root):
            style=next(node for node in root.findall('./w:style',NS) if node.get(qn('w:type'))=='table' and node.get(qn('w:default'))=='1')
            props=style.find('./w:tblPr',NS)
            if props is None:props=ET.SubElement(style,qn('w:tblPr'))
            indent=ET.SubElement(props,qn('w:tblInd')); indent.set(qn('w:w'),'20000'); indent.set(qn('w:type'),'dxa')
        def direct_indent(root):
            props=root.find('.//w:tblPr',NS)
            indent=ET.SubElement(props,qn('w:tblInd')); indent.set(qn('w:w'),'20000'); indent.set(qn('w:type'),'dxa')
        def exact_height(root):
            props=root.find('.//w:trPr',NS)
            height=ET.SubElement(props,qn('w:trHeight')); height.set(qn('w:hRule'),'exact'); height.set(qn('w:val'),'10')
        def direct_alignment(root):
            node=ET.SubElement(root.find('.//w:tblPr',NS),qn('w:jc')); node.set(qn('w:val'),'both')
        def inherited_alignment(root):
            style=next(node for node in root.findall('./w:style',NS) if node.get(qn('w:type'))=='table' and node.get(qn('w:default'))=='1')
            props=style.find('./w:tblPr',NS)
            if props is None:props=ET.SubElement(style,qn('w:tblPr'))
            node=ET.SubElement(props,qn('w:jc')); node.set(qn('w:val'),'both')
        for kind,content in (('inherited_indent',change_docx(old_content,inherited_indent,part='word/styles.xml')),
                        ('direct_indent',change_docx(old_content,direct_indent)),('exact_height',change_docx(old_content,exact_height)),
                        ('direct_alignment',change_docx(old_content,direct_alignment)),
                        ('inherited_alignment',change_docx(old_content,inherited_alignment,part='word/styles.xml'))):
            with self.subTest(kind=kind), self.assertRaises(ValueError):
                build_renderer_evidence(ir,content,formula_report=result.formula_report,image_layout_report=result.image_layout_report)

        def explicit_style(root):
            node=ET.SubElement(root.find('.//w:tblPr',NS),qn('w:tblStyle')); node.set(qn('w:val'),'ImageDerived')
        explicit_content=change_docx(old_content,explicit_style)
        def based_on(root,value):
            base=ET.SubElement(root,qn('w:style')); base.set(qn('w:type'),'table'); base.set(qn('w:styleId'),'ImageBase')
            props=ET.SubElement(base,qn('w:tblPr'))
            indent=ET.SubElement(props,qn('w:tblInd')); indent.set(qn('w:w'),str(value)); indent.set(qn('w:type'),'dxa')
            derived=ET.SubElement(root,qn('w:style')); derived.set(qn('w:type'),'table'); derived.set(qn('w:styleId'),'ImageDerived')
            parent=ET.SubElement(derived,qn('w:basedOn')); parent.set(qn('w:val'),'ImageBase')
        zero_based_on=change_docx(explicit_content,lambda root:based_on(root,0),part='word/styles.xml')
        ensure_renderer_integrity(build_renderer_evidence(ir,zero_based_on,formula_report=result.formula_report,
            image_layout_report=result.image_layout_report))
        shifted_based_on=change_docx(explicit_content,lambda root:based_on(root,20000),part='word/styles.xml')
        with self.assertRaisesRegex(ValueError,'有效缩进或对齐'):
            build_renderer_evidence(ir,shifted_based_on,formula_report=result.formula_report,image_layout_report=result.image_layout_report)

    def test_legacy_zero_indent_tables_accept_supported_effective_alignment(self):
        ir=image_ir(); result=render(ir)
        def without_direct(root):
            props=root.find('.//w:tblPr',NS)
            for tag in ('tblInd','jc'):
                for node in props.findall('./w:'+tag,NS):props.remove(node)
        old_content=change_docx(result.content,without_direct)
        for alignment in ('left','center','right'):
            for source in ('direct','default','basedOn'):
                with self.subTest(alignment=alignment,source=source):
                    def position(root):
                        if source=='direct':
                            props=root.find('.//w:tblPr',NS)
                        elif source=='default':
                            style=next(node for node in root.findall('./w:style',NS)
                                if node.get(qn('w:type'))=='table' and node.get(qn('w:default'))=='1')
                            props=style.find('./w:tblPr',NS)
                            if props is None:props=ET.SubElement(style,qn('w:tblPr'))
                        else:
                            base=ET.SubElement(root,qn('w:style')); base.set(qn('w:type'),'table'); base.set(qn('w:styleId'),'ImageBase')
                            props=ET.SubElement(base,qn('w:tblPr'))
                            derived=ET.SubElement(root,qn('w:style')); derived.set(qn('w:type'),'table'); derived.set(qn('w:styleId'),'ImageDerived')
                            parent=ET.SubElement(derived,qn('w:basedOn')); parent.set(qn('w:val'),'ImageBase')
                        node=ET.SubElement(props,qn('w:jc')); node.set(qn('w:val'),alignment)
                    content=old_content
                    if source=='basedOn':
                        def explicit_style(root):
                            node=ET.SubElement(root.find('.//w:tblPr',NS),qn('w:tblStyle')); node.set(qn('w:val'),'ImageDerived')
                        content=change_docx(content,explicit_style)
                    content=change_docx(content,position,part='word/document.xml' if source=='direct' else 'word/styles.xml')
                    ensure_renderer_integrity(build_renderer_evidence(ir,content,formula_report=result.formula_report,
                        image_layout_report=result.image_layout_report))

    def test_legacy_table_indent_units_and_defaults_are_authenticated(self):
        ir=image_ir(); result=render(ir)
        def without_direct_indent(root):
            props=root.find('.//w:tblPr',NS)
            for node in props.findall('./w:tblInd',NS):props.remove(node)
        old_content=change_docx(result.content,without_direct_indent)
        for unit,value,rejection in (
                ('pct','20000',None),('auto','20000',None),('nil','20000',None),
                ('dxa','20000','有效缩进或对齐'),('unknown','20000','缩进单位'),
                (None,'0',None),('dxa',None,None),(None,None,None),
                (None,'20000','有效缩进或对齐')):
            for source in ('direct','default','basedOn'):
                with self.subTest(unit=unit,value=value,source=source):
                    def indentation(root):
                        if source=='direct':
                            props=root.find('.//w:tblPr',NS)
                        elif source=='default':
                            style=next(node for node in root.findall('./w:style',NS)
                                if node.get(qn('w:type'))=='table' and node.get(qn('w:default'))=='1')
                            props=style.find('./w:tblPr',NS)
                            if props is None:props=ET.SubElement(style,qn('w:tblPr'))
                            for node in props.findall('./w:tblInd',NS):props.remove(node)
                        else:
                            base=ET.SubElement(root,qn('w:style')); base.set(qn('w:type'),'table'); base.set(qn('w:styleId'),'ImageBase')
                            props=ET.SubElement(base,qn('w:tblPr'))
                            derived=ET.SubElement(root,qn('w:style')); derived.set(qn('w:type'),'table'); derived.set(qn('w:styleId'),'ImageDerived')
                            parent=ET.SubElement(derived,qn('w:basedOn')); parent.set(qn('w:val'),'ImageBase')
                        node=ET.SubElement(props,qn('w:tblInd'))
                        if unit is not None:node.set(qn('w:type'),unit)
                        if value is not None:node.set(qn('w:w'),value)
                    content=old_content
                    if source=='basedOn':
                        def explicit_style(root):
                            node=ET.SubElement(root.find('.//w:tblPr',NS),qn('w:tblStyle')); node.set(qn('w:val'),'ImageDerived')
                        content=change_docx(content,explicit_style)
                    content=change_docx(content,indentation,part='word/document.xml' if source=='direct' else 'word/styles.xml')
                    if rejection is None:
                        ensure_renderer_integrity(build_renderer_evidence(ir,content,formula_report=result.formula_report,
                            image_layout_report=result.image_layout_report))
                    else:
                        with self.assertRaisesRegex(ValueError,rejection):
                            build_renderer_evidence(ir,content,formula_report=result.formula_report,
                                image_layout_report=result.image_layout_report)

    def test_image_paragraph_cancels_inherited_template_indentation(self):
        ir=image_ir()
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'indented.docx'
            template=Document()
            normal=template.styles['Normal'].paragraph_format
            normal.left_indent=Cm(1); normal.right_indent=Cm(0.6); normal.first_line_indent=Cm(0.8)
            template.save(path)
            inherited_indent=Document(path).styles['Normal'].paragraph_format.first_line_indent
            result=render(ir,template_path=path)
        ev=evidence(ir,result); ensure_renderer_integrity(ev)
        doc=Document(io.BytesIO(result.content))
        self.assertEqual(doc.styles['Normal'].paragraph_format.first_line_indent,inherited_indent)
        for cell in doc.tables[0].rows[0].cells:
            fmt=cell.paragraphs[0].paragraph_format
            self.assertEqual((fmt.left_indent,fmt.right_indent,fmt.first_line_indent),(0,0,0))
        def inherited_again(root): root.find('.//w:tc/w:p/w:pPr/w:ind',NS).set(qn('w:firstLine'),'440')
        with self.assertRaisesRegex(ValueError,'缩进'):
            build_renderer_evidence(ir,change_docx(result.content,inherited_again),image_layout_report=result.image_layout_report)

    def test_three_cells_inline_ratios_fields_notes_and_content_table_accounting(self):
        ir = image_ir(3)
        ir.blocks.append(DocumentIRBlock(id='table', type='table', order=4, page_index=0, html='<table><tr><td>Data</td></tr></table>', layout=DocumentIRBlockLayout(page_index=0,bbox=(0,300,400,400))))
        result = render(ir)
        ev = evidence(ir, result); ensure_renderer_integrity(ev)
        doc = Document(io.BytesIO(result.content))
        self.assertEqual(len(doc.tables), 2)
        self.assertEqual(len(doc.tables[0].rows), 1)
        props=doc.tables[0]._tbl.tblPr
        self.assertLess(list(props).index(props.find(qn('w:tblBorders'))),list(props).index(props.find(qn('w:tblLayout'))))
        for index, cell in enumerate(doc.tables[0].rows[0].cells):
            self.assertIn(f'Figure {index+1}', cell.text)
            self.assertIn(f'Note {index+1}', cell.text)
            self.assertEqual(len(cell.paragraphs), 4)
            self.assertFalse(cell.paragraphs[-1].paragraph_format.keep_with_next)
            tc_props=cell._tc.tcPr
            self.assertLess(list(tc_props).index(tc_props.find(qn('w:tcMar'))),list(tc_props).index(tc_props.find(qn('w:vAlign'))))
        self.assertEqual(len(doc.inline_shapes), 3)
        for index, inline in enumerate(doc.inline_shapes):
            self.assertTrue(all(inline._inline.get(key)=='0' for key in ('distT','distB','distL','distR')))
            px, py = (120, 60) if index%2 == 0 else (50, 150)
            self.assertLessEqual(abs(inline.width * py - inline.height * px), max(px, py))
        self.assertEqual(ev['table_count_delta'], 1)
        self.assertEqual((ev['layout_table_count'], ev['content_table_count'], ev['total_table_count']), (1,1,2))
        self.assertEqual(ev['total_table_count_delta'], 0)
        self.assertEqual(len(result.formula_report['entries']), 3)
        self.assertEqual([e['source']['field'] for e in result.formula_report['entries']], ['captions/0']*3)

    def test_pure_image_layout_cannot_supply_native_data_table_editability(self):
        ir = image_ir(3, annotations=False); result = render(ir)
        quality = analyze_docx_quality_bytes(result.content)
        self.assertEqual(quality['table_count'], 1)
        self.assertEqual(quality['candidate_layout_table_count'], 1)
        self.assertFalse(quality['editable_content_available'])
        self.assertTrue(quality['image_only_suspected'])
        self.assertEqual(quality['wide_table_count'], 0)
        ensure_renderer_integrity(evidence(ir,result))
        ir.blocks.append(DocumentIRBlock(id='missing_table',type='table',order=3,html='<table><tr><td>A</td></tr></table>'))
        # 纯图布局表不能替代本应存在的内容表。
        altered = copy.deepcopy(result.image_layout_report)
        altered['source_ir_sha256'] = canonical_sha256(ir.model_dump(mode='json'))
        ev = build_renderer_evidence(ir,result.content,formula_report=result.formula_report,image_layout_report=altered)
        self.assertIn('ir_docx_integrity:table_count_mismatch',ev['blocking_quality_flags'])

    def test_missing_corrupt_or_wrong_checksum_never_drops_annotations(self):
        for content in (None, b'invalid-image'):
            ir = image_ir()
            result = render(ir, image_loader=lambda a: image_loader(a) if a.id=='asset_0' else content)
            doc = Document(io.BytesIO(result.content))
            self.assertEqual(len(doc.tables),0)
            texts = '\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('Note 2: extra',texts)
            self.assertIn('Figure 2',texts)
            self.assertIn('图片缺失或无效',texts)
            with self.assertRaises(ValueError): evidence(ir,result)
        ir=image_ir(); ir.assets[0].checksum_sha256='sha256:'+'0'*64
        with self.assertRaisesRegex(ValueError,'checksum'): render(ir)

    def test_source_annotations_not_report_expected_hash_determine_validation(self):
        ir=image_ir(); result=render(ir)
        report=copy.deepcopy(result.image_layout_report)
        changed=change_docx(result.content,lambda root: setattr(root.find('.//w:tc/w:p[2]//w:t',NS),'text','Changed'))
        # 即使把报告 expected 改成篡改后的实际投影，源重建仍会拒绝。
        with zipfile.ZipFile(io.BytesIO(changed)) as package:
            root=ET.fromstring(package.read('word/document.xml'))
        report['spans'][0]['images'][0]['annotations'][0]['expected_output_sha256']=image_evidence.annotation_projection(root.find('.//w:tc/w:p[2]',NS))
        with self.assertRaises(ValueError): build_renderer_evidence(ir,changed,formula_report=result.formula_report,image_layout_report=report)

    def test_final_ooxml_tampering_is_rejected(self):
        ir=image_ir(); result=render(ir)
        def remove_note(root):
            cell=root.find('.//w:tc',NS); cell.remove(cell.findall('./w:p',NS)[2])
        def swap_images(root):
            nodes=root.findall('.//a:blip',NS); key=qn('r:embed'); a,b=[n.get(key) for n in nodes]; nodes[0].set(key,b); nodes[1].set(key,a)
        def swap_captions(root):
            cells=root.findall('.//w:tc',NS); a=cells[0].findall('./w:p',NS)[1]; b=cells[1].findall('./w:p',NS)[1]; cells[0].remove(a); cells[1].remove(b); cells[0].insert(2,b); cells[1].insert(2,a)
        def duplicate_table(root): root.find('./w:body',NS).insert(0,copy.deepcopy(root.find('.//w:tbl',NS)))
        def fake_marker(root): root.find('.//w:tblCaption',NS).set(qn('w:val'),'md2w_image_layout:fake')
        def anchor(root): root.find('.//wp:inline',NS).tag=qn('wp:anchor')
        def width(root): root.find('.//wp:extent',NS).set('cx','99999999')
        def picture_width(root): root.find('.//a:xfrm/a:ext',NS).set('cx','99999999')
        def image_distance(root): root.find('.//wp:inline',NS).attrib.pop('distR')
        def column(root): root.find('.//w:tcW',NS).set(qn('w:w'),'1')
        def section(root): root.find('.//w:pgMar',NS).set(qn('w:left'),'1')
        for mutation in (remove_note,swap_images,swap_captions,duplicate_table,fake_marker,anchor,width,picture_width,image_distance,column,section):
            with self.subTest(mutation=mutation.__name__),self.assertRaises(ValueError):
                build_renderer_evidence(ir,change_docx(result.content,mutation),formula_report=result.formula_report,image_layout_report=result.image_layout_report)

    def test_complete_group_or_independent_single_span_exchange_is_rejected(self):
        for grouped in (True, False):
            ir = image_ir(2)
            extra = [block.model_copy(deep=True) for block in ir.blocks]
            for index, block in enumerate(extra):
                block.id += '_row2'; block.order += 2
                box = block.layout.bbox; block.layout.bbox = (box[0], 300, box[2], 400)
            ir.blocks.extend(extra)
            if not grouped:
                ir.blocks = [ir.blocks[0], ir.blocks[2]]
            result = render(ir)
            def swap(root):
                body = root.find('./w:body', NS)
                if grouped:
                    a, b = body.findall('./w:tbl', NS)
                    body.remove(b); body.insert(list(body).index(a), b)
                else:
                    paragraphs = body.findall('./w:p', NS)
                    half = len(paragraphs) // 2
                    for paragraph in paragraphs[half:]:
                        body.remove(paragraph)
                    for index, paragraph in enumerate(paragraphs[half:]):
                        body.insert(index, paragraph)
            with self.subTest(grouped=grouped), self.assertRaisesRegex(ValueError, '顺序'):
                build_renderer_evidence(ir, change_docx(result.content, swap), formula_report=result.formula_report,
                                        image_layout_report=result.image_layout_report)

    def test_picture_crop_transform_and_cell_width_unit_tampering_are_rejected(self):
        ir = image_ir()
        result = render(ir)

        def source_crop(root):
            fill = root.find('.//pic:blipFill', NS)
            fill.insert(1, ET.Element(qn('a:srcRect'), {'r': '50000'}))

        def fill_offset(root): root.find('.//a:fillRect', NS).set('r', '50000')
        def horizontal_flip(root): root.find('.//a:xfrm', NS).set('flipH', '1')
        def vertical_flip(root): root.find('.//a:xfrm', NS).set('flipV', 'true')
        def rotation(root): root.find('.//a:xfrm', NS).set('rot', '5400000')
        def picture_offset(root): root.find('.//a:xfrm/a:off', NS).set('x', '100000')
        def percent_width(root): root.find('.//w:tcW', NS).set(qn('w:type'), 'pct')

        for mutation in (source_crop, fill_offset, horizontal_flip, vertical_flip, rotation, picture_offset, percent_width):
            with self.subTest(mutation=mutation.__name__), self.assertRaises(ValueError):
                build_renderer_evidence(ir, change_docx(result.content, mutation),
                    formula_report=result.formula_report, image_layout_report=result.image_layout_report)

        def explicit_defaults(root):
            transform = root.find('.//a:xfrm', NS)
            transform.attrib.update({'rot': '0', 'flipH': 'false', 'flipV': '0'})
            root.find('.//pic:blipFill', NS).insert(1, ET.Element(qn('a:srcRect'), {'r': '0'}))
            root.find('.//a:fillRect', NS).set('r', '0')

        # 显式零值与 writer 的省略默认值等价，不能把完整图误判为损坏。
        ensure_renderer_integrity(build_renderer_evidence(ir, change_docx(result.content, explicit_defaults),
            formula_report=result.formula_report, image_layout_report=result.image_layout_report))

    def test_duplicate_table_caption_whitespace_preserves_unique_ref(self):
        ir = image_ir(1, annotations=False)
        ir.blocks.extend([DocumentIRBlock(id='table', type='table', order=1, page_index=0,
            captions=['Table 1: Neutral', ' Table 1: Neutral '], html='<table><tr><td>Data</td></tr></table>',
            layout=DocumentIRBlockLayout(page_index=0, bbox=(10,300,390,400))),
            DocumentIRBlock(id='reference', type='paragraph', order=2, page_index=0, text='See Table 1',
            layout=DocumentIRBlockLayout(page_index=0, bbox=(10,410,390,450)))])
        result = render(ir)
        ensure_renderer_integrity(evidence(ir,result))
        with zipfile.ZipFile(io.BytesIO(result.content)) as package:
            xml = package.read('word/document.xml').decode()
        self.assertIn('w:name="_md2word_table_1"',xml)
        self.assertIn(' REF _md2word_table_1 \\h ',xml)

    def test_frozen_template_and_format_rules_survive_files_removed_and_versions_reject(self):
        ir=image_ir(); outputs=[]
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            for index in (1,2):
                config=root/f'{index}.yaml'; rules=root/f'r{index}.yaml'; template=root/f'{index}.docx'
                Document().save(template)
                config.write_text(f'schema_version: template-config/v1\npage_setup:\n  page_width_cm: {20-index}\n  left_margin: {index}\n  right_margin: {index}\ncaption:\n  bold: true\n',encoding='utf-8')
                rules.write_text(f'subscript_rules:\n  symbols:\n    DemoToken: Frozen{index}\n',encoding='utf-8')
                result=render(ir,template_path=template,config_path=config,options=DocumentIRRenderOptions(subscript_config_path=rules,template_version_id=f'template@{index}'))
                self.assertEqual(TemplateConfigPackage.load(config),TemplateConfigPackage.load_content(config.read_bytes()))
                self.assertEqual(result.image_layout_report['render_inputs']['template']['config']['sha256'],hashlib.sha256(config.read_bytes()).hexdigest())
                self.assertNotIn(str(root),str(result.image_layout_report))
                outputs.append(result)
        for index,result in enumerate(outputs,1):
            ensure_renderer_integrity(evidence(ir,result,expected_template_version_id=f'template@{index}'))
            self.assertIn(f'Frozen{index}',Document(io.BytesIO(result.content)).tables[0].cell(0,0).text)
            with self.assertRaises(ValueError): evidence(ir,result,expected_template_version_id='wrong@1')
        self.assertNotEqual(outputs[0].image_layout_report['render_inputs_sha256'],outputs[1].image_layout_report['render_inputs_sha256'])
        for key in ('schema_version','projection_version','writer_version'):
            report=copy.deepcopy(outputs[0].image_layout_report); report['render_inputs'][key]='unknown/v999'
            report['render_inputs_sha256']=canonical_sha256(report['render_inputs'])
            with self.assertRaises(ValueError): build_renderer_evidence(ir,outputs[0].content,image_layout_report=report)
        report=copy.deepcopy(outputs[0].image_layout_report); del report['render_inputs']
        with self.assertRaises(ValueError): build_renderer_evidence(ir,outputs[0].content,image_layout_report=report)

    def test_b1_writer_and_bidirectional_reader_and_parallel_contexts(self):
        ir=image_ir(); grouped=render(ir)
        with patch.object(image_evidence,'IMAGE_GENERATION_POLICY','single-only/v1'):
            ensure_renderer_integrity(evidence(ir,grouped))
            single=render(ir)
            self.assertFalse(Document(io.BytesIO(single.content)).tables)
            self.assertEqual(single.image_layout_report['render_inputs']['generation_policy'],'single-only/v1')
            self.assertNotIn('image_group_uncertain',evidence(ir,single)['quality_flags'])
        ensure_renderer_integrity(evidence(ir,single))
        renderer=DocumentIRWordRenderer()
        def task(index):
            local=image_ir(index%2+2)
            result=renderer.render_to_result(local,template_path=None,config_path=None,image_loader=image_loader)
            ensure_renderer_integrity(evidence(local,result))
            return len(result.image_layout_report['spans'][0]['images'])
        with ThreadPoolExecutor(max_workers=4) as pool:
            self.assertEqual(list(pool.map(task,range(8))),[2,3]*4)

    def test_group_downgrade_cannot_be_claimed_by_report_alone(self):
        ir=image_ir()
        with patch.object(image_evidence,'IMAGE_GENERATION_POLICY','single-only/v1'):
            result=render(ir)
        for reason in ('asset_unavailable','image_size_unfit'):
            report=copy.deepcopy(result.image_layout_report)
            report['render_inputs']['generation_policy']='grouped/v1'
            report['render_inputs_sha256']=canonical_sha256(report['render_inputs'])
            report['spans'][0]['reasons']=[reason]
            with self.subTest(reason=reason), self.assertRaises(ValueError):
                build_renderer_evidence(ir,result.content,formula_report=result.formula_report,image_layout_report=report,
                    expected_generation_policy='grouped/v1')

    def test_injected_body_and_private_snapshot_values_are_rejected_safely(self):
        ir=image_ir(); result=render(ir)
        def inject(root):
            paragraph=root.find('.//w:tc/w:p',NS)
            run=ET.Element(qn('w:r')); text=ET.SubElement(run,qn('w:t')); text.text='Unregistered body'
            paragraph.insert(len(paragraph)-1,run)
        with self.assertRaises(ValueError):
            build_renderer_evidence(ir,change_docx(result.content,inject),image_layout_report=result.image_layout_report)
        report=copy.deepcopy(result.image_layout_report)
        report['render_inputs']['caption_style']['font_size']='PRIVATE_FORMAT_SENTINEL'
        report['render_inputs_sha256']=canonical_sha256(report['render_inputs'])
        with self.assertRaises(ValueError) as caught:
            build_renderer_evidence(ir,result.content,image_layout_report=report)
        self.assertNotIn('PRIVATE_FORMAT_SENTINEL',str(caught.exception))

    def test_formal_zip_reliable_five_six_and_twenty_four_twenty_five(self):
        for numbers in ((5,6),(24,25)):
            with self.subTest(numbers=numbers),tempfile.TemporaryDirectory() as directory:
                path=Path(directory)/'images.zip'; write_numbered_image_zip(path,numbers)
                converted=MinerUToDocumentIRService().convert_zip(path,job_id='job',source_file_id='source',document_language='zh')
                blobs={asset.asset_id:asset.content for asset in converted.image_assets}
                result=render(converted.document,image_loader=lambda asset:blobs[asset.id])
                ensure_renderer_integrity(evidence(converted.document,result))
                table=Document(io.BytesIO(result.content)).tables[0]
                for index,number in enumerate(numbers):
                    self.assertIn(f'图 {number}',table.cell(0,index).text)
                    self.assertIn(f'注 {number}',table.cell(0,index).text)

    def test_accepted_deidentified_uncertain_fixture_keeps_source_once(self):
        fixture=Path(__file__).parent/'fixtures/component_pdf_samples/samples/source-redacted/datasheet/zk7044-caption-pair/mineru'
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'uncertain.zip'
            with zipfile.ZipFile(path,'w') as package:
                for item in sorted(fixture.rglob('*')):
                    if item.is_file(): package.write(item,str(item.relative_to(fixture)))
            converted=MinerUToDocumentIRService().convert_zip(path,job_id='job',source_file_id='source',document_language='zh')
            original=converted.document.model_dump_json()
            blobs={asset.asset_id:asset.content for asset in converted.image_assets}
            result=render(converted.document,image_loader=lambda asset:blobs[asset.id])
            ev=evidence(converted.document,result); ensure_renderer_integrity(ev)
            doc=Document(io.BytesIO(result.content))
            self.assertFalse(doc.tables)
            self.assertEqual(converted.document.model_dump_json(),original)
            self.assertIn('caption_binding_uncertain',converted.document.quality_flags)
            self.assertIn('image_group_uncertain',ev['quality_flags'])
            text='\n'.join(p.text for p in doc.paragraphs)
            self.assertEqual(text.count('图5'),1)
            self.assertEqual(text.count('图6'),1)

    def test_external_caption_evidence_stays_in_its_original_paragraph_once(self):
        blocks=[{'type':'image','img_path':f'images/{i}.png','image_caption':[],'bbox':[10+i*120,100,110+i*120,200],'page_idx':0} for i in range(2)]
        blocks.append({'type':'text','text':'图6 外部独立题注','bbox':[130,205,230,225],'page_idx':0})
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'external.zip'
            with zipfile.ZipFile(path,'w') as package:
                package.writestr('content_list.json',json.dumps(blocks,ensure_ascii=False))
                package.writestr('layout.json',json.dumps({'pdf_info':[{'page_idx':0,'page_size':[400,600]}]}))
                for i in range(2): package.writestr(f'images/{i}.png',neutral_png())
            converted=MinerUToDocumentIRService().convert_zip(path,job_id='job',source_file_id='source',document_language='zh')
            blobs={asset.asset_id:asset.content for asset in converted.image_assets}
            result=render(converted.document,image_loader=lambda asset:blobs[asset.id])
            ev=evidence(converted.document,result); ensure_renderer_integrity(ev)
            doc=Document(io.BytesIO(result.content))
            self.assertFalse(doc.tables)
            self.assertEqual('\n'.join(p.text for p in doc.paragraphs).count('外部独立题注'),1)
            self.assertIn('external_caption_binding',ev['image_layout_summary']['reason_counts'])

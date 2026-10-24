"""RND-501 的计划先行、资源快照和纯 JSON 合同。"""
from __future__ import annotations

import hashlib
import io
import json
import math
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch

import yaml
from docx import Document
from pydantic import ValidationError

from app.services.document_ir.integrity import DocumentIntegrityError
from app.services.document_ir.models import DocumentIR, DocumentIRMetadata, DocumentIRBlock
from app.services.renderer import DocumentIRWordRenderer, DocumentIRRenderOptions
from app.services.renderer.evidence import build_renderer_evidence, validate_render_result
from app.services.renderer.layout_plan import build_layout_plan
from app.services.renderer.render_plan import RenderPlan, ResourceIdentity, build_render_plan, content_hash, document_ir_hash
from app.services.renderer.template_config import TemplateConfigPackage


def make_ir(blocks=None):
    return DocumentIR(metadata=DocumentIRMetadata(job_id="job", source_file_id="source", document_language="en"),
        blocks=blocks if blocks is not None else [DocumentIRBlock(id="body", type="paragraph", order=0, text="Editable paragraph for deterministic rendering.")])


def render(ir, **kwargs):
    return DocumentIRWordRenderer().render_to_result(ir, template_path=kwargs.pop("template_path", None),
        config_path=kwargs.pop("config_path", None), image_loader=kwargs.pop("image_loader", lambda asset: None), **kwargs)


class RenderPlanTests(TestCase):
    def test_repeat_json_roundtrip_and_effective_config_change(self):
        ir = make_ir()
        before = ir.model_dump_json()
        first, second = render(ir), render(ir)
        self.assertEqual(first.plan, second.plan)
        self.assertEqual(RenderPlan.model_validate_json(first.plan.model_dump_json()), first.plan)
        self.assertEqual(ir.model_dump_json(), before)
        changed = first.plan.model_dump(mode="json")
        changed["extra"] = "forbidden"
        with self.assertRaises(ValidationError):
            RenderPlan.model_validate(changed)
        changed = first.plan.model_dump(mode="json")
        changed["styles"]["normal"]["font_size"] = float("nan")
        with self.assertRaises((ValidationError, ValueError)):
            RenderPlan.model_validate(changed)
        with TemporaryDirectory() as directory:
            path = Path(directory) / "config.yaml"
            config = TemplateConfigPackage(normal={"font_size": 10.0})
            path.write_text(yaml.safe_dump(config.model_dump(mode="json")), encoding="utf-8")
            sized = render(ir, config_path=path)
            self.assertNotEqual(first.plan.plan_sha256, sized.plan.plan_sha256)
            self.assertEqual(sized.plan.decisions[0].decision_id, first.plan.decisions[0].decision_id)
            self.assertEqual(sized.plan.decisions[0].anchor_ref, first.plan.decisions[0].anchor_ref)
            self.assertNotIn(directory, sized.plan.model_dump_json())

    def test_plan_precedes_document_and_template_snapshot_is_consumed_once(self):
        from app.services.renderer import word_renderer
        with TemporaryDirectory() as directory:
            path = Path(directory) / "template.docx"
            template = Document()
            template.core_properties.title = "original"
            template.save(path)
            original_bytes = path.read_bytes()
            original_builder = word_renderer.build_render_plan
            observed = []
            def planned(*args, **kwargs):
                plan = original_builder(*args, **kwargs)
                observed.append("plan")
                template.core_properties.title = "changed-after-plan"
                template.save(path)
                return plan
            original_document = word_renderer.Document
            def create(*args, **kwargs):
                self.assertEqual(observed, ["plan"])
                observed.append("document")
                return original_document(*args, **kwargs)
            with patch.object(word_renderer, "build_render_plan", side_effect=planned), patch.object(word_renderer, "Document", side_effect=create):
                result = render(make_ir(), template_path=path)
            self.assertEqual(result.plan.resource_identity.template_sha256, hashlib.sha256(original_bytes).hexdigest())
            self.assertEqual(Document(io.BytesIO(result.content)).core_properties.title, "original")
            self.assertEqual(observed, ["plan", "document"])

    def test_language_and_loaded_symbol_rules_are_path_independent(self):
        ir = make_ir([DocumentIRBlock(id="body", type="paragraph", order=0, text="VCC signal should use symbol rules.")])
        config = TemplateConfigPackage(normal={"font_size": 10.0}, language_overrides={"zh": {"normal": {"font_size": 12.0}}})
        with TemporaryDirectory() as directory:
            root = Path(directory)
            config_path = root / "layout.yaml"
            config_path.write_text(yaml.safe_dump(config.model_dump(mode="json")), encoding="utf-8")
            rules = "case_sensitive: true\nsubscript_rules:\n  symbols:\n    VCC: V<sub>CC</sub>\n"
            a, b = root / "a.yaml", root / "b.yaml"
            a.write_text(rules, encoding="utf-8")
            b.write_text(rules, encoding="utf-8")
            first = render(ir, config_path=config_path, options=DocumentIRRenderOptions(subscript_config_path=a, document_language="zh"))
            second = render(ir, config_path=config_path, options=DocumentIRRenderOptions(subscript_config_path=b, document_language="zh"))
            self.assertEqual(first.plan, second.plan)
            self.assertEqual(first.plan.styles["normal"]["font_size"], 12.0)
            self.assertIsNotNone(first.plan.resource_identity.symbol_rules_sha256)
            self.assertNotIn("V<sub>", first.plan.model_dump_json())
            a.write_text(rules.replace("true", "false"), encoding="utf-8")
            changed = render(ir, config_path=config_path, options=DocumentIRRenderOptions(subscript_config_path=a, document_language="zh"))
            self.assertNotEqual(changed.plan.resource_identity.symbol_rules_sha256, first.plan.resource_identity.symbol_rules_sha256)

    def test_standalone_plan_and_writer_share_table_strategy(self):
        ir = make_ir([DocumentIRBlock(id="table", type="table", order=0,
            html="<table>" + "<tr><td>Value</td></tr>" * 401 + "</table>")])
        result = render(ir)
        resource = result.plan.resource_identity
        standalone = build_render_plan(ir, layout_plan=build_layout_plan(ir), effective_config=TemplateConfigPackage(), resource_identity=resource)
        self.assertEqual(standalone, result.plan)
        self.assertEqual(result.plan.decisions[0].mode, "table_expanded")
        self.assertEqual(result.plan.decisions[0].table_policy["row_count"], 401)
        self.assertEqual(len(Document(io.BytesIO(result.content)).tables[0].rows), 401)

    def test_bytes_compatibility_and_same_instance_interleaving(self):
        renderer = DocumentIRWordRenderer()
        ir = make_ir()
        content = renderer.render_to_bytes(ir, template_path=None, config_path=None, image_loader=lambda asset: None)
        self.assertIsInstance(content, bytes)
        self.assertTrue(Document(io.BytesIO(content)).paragraphs)
        inputs = [make_ir([DocumentIRBlock(id=f"body_{index}", type="paragraph", order=0, text=f"Unique paragraph {index} preserved in its own render.")]) for index in range(4)]
        with ThreadPoolExecutor(max_workers=4) as pool:
            results = list(pool.map(lambda item: renderer.render_to_result(item, template_path=None, config_path=None, image_loader=lambda asset: None), inputs))
        for item, result in zip(inputs, results, strict=True):
            validate_render_result(item, result.content, result)
            self.assertEqual(result.block_evidence.block_render_map[0].block_id, item.blocks[0].id)
            self.assertEqual(Document(io.BytesIO(result.content)).paragraphs[0].text, item.blocks[0].text)
        self.assertEqual(len({result.plan.plan_sha256 for result in results}), 4)

    def test_duplicate_source_identity_fails_without_creating_word(self):
        ir = make_ir([DocumentIRBlock(id="same", type="paragraph", order=0, text="A"), DocumentIRBlock(id="same", type="paragraph", order=1, text="B")])
        with patch("app.services.renderer.word_renderer.Document") as document, self.assertRaises(DocumentIntegrityError) as raised:
            render(ir)
        document.assert_not_called()
        self.assertEqual(raised.exception.evidence["blocking_quality_flags"], ["ir_render_evidence:block_coverage_mismatch"])

    def test_nonfinite_legacy_identity_is_lossless_and_finite_identity_unchanged(self):
        from tests.test_image_group_rendering import image_ir, image_loader
        finite = image_ir()
        self.assertEqual(document_ir_hash(finite), content_hash(finite.model_dump(mode="json")))
        hashes = []
        for coordinate in (float("nan"), float("inf"), float("-inf")):
            ir = image_ir()
            ir.blocks[0].layout.bbox = (coordinate, 100, 110, 200)
            original_layout = ir.blocks[0].layout.bbox
            result = render(ir, image_loader=image_loader)
            validate_render_result(ir, result.content, result)
            hashes.append(result.plan.input_ir_sha256)
            self.assertIs(ir.blocks[0].layout.bbox, original_layout)
            self.assertTrue(math.isnan(original_layout[0]) if math.isnan(coordinate) else original_layout[0] == coordinate)
            self.assertEqual(result.block_evidence.block_render_map[0].provenance.bbox, [None, 100.0, 110.0, 200.0])
            self.assertEqual(document_ir_hash(ir), result.plan.input_ir_sha256)
            self.assertEqual(document_ir_hash(ir.model_copy(deep=True)), result.plan.input_ir_sha256)
            self.assertNotIn("NaN", result.plan.model_dump_json())
            self.assertNotIn("Infinity", result.plan.model_dump_json())
        self.assertEqual(len(set(hashes)), 3)
        # 自由 metadata 的相同文字/容器不能冒充非有限 float 的类型标记。
        distinct = []
        for value in (None, "nan", ["nonfinite", "nan"], {"nonfinite": "nan"}, float("nan"), float("inf"), float("-inf")):
            ir = make_ir()
            ir.blocks[0].metadata["legacy_value"] = value
            distinct.append(document_ir_hash(ir))
        self.assertEqual(len(set(distinct)), len(distinct))

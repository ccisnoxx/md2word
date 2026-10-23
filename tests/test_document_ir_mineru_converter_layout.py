from __future__ import annotations

# tests/test_document_ir_mineru_converter_layout.py: 覆盖 MinerU -> DocumentIR v2 layout signal 提取。

import io
import copy
import hashlib
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from typing import Any
from dataclasses import replace
from unittest.mock import patch
from xml.etree import ElementTree

from docx import Document

from app.conversion.docx_rules.mineru_content import find_final_promotional_page
from app.services.document_ir.integrity import DocumentIntegrityError
from app.services.document_ir import MinerUToDocumentIRService, caption_reconciler
from app.services.document_ir.artifact_validator import MinerUArtifactValidationError
from app.services.document_ir.models import DocumentIR
from app.services.renderer.evidence import build_renderer_evidence, ensure_renderer_integrity
from app.services.renderer import DocumentIRWordRenderer


# 使用可被真实图片解码器完整读取的中性PNG，保护输入资产校验。
MINIMAL_PNG_BYTES = b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x02\x00\x00\x00\x02\x08\x02\x00\x00\x00\xfd\xd4\x9as\x00\x00\x00\x13IDATx\x9cc\xfc\xff\xff?\x03\x03\x03\x13\x03\x18\x00\x00$\x06\x03\x01]\xa2N\x88\x00\x00\x00\x00IEND\xaeB`\x82'


class MinerUToDocumentIRLayoutTests(unittest.TestCase):

    def test_nested_caption_donor_is_consumed_once(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "nested.zip"
            self._write_zip(
                path,
                blocks=[[
                    {"type": "chart", "content": {"image_source": {"path": "images/chart.png"}, "chart_caption": []},
                     "bbox": [10, 80, 170, 220]},
                    {"type": "chart", "content": {"image_source": {"path": "images/chart.png"},
                                                       "chart_caption": ["图 5 响应", "图 6 增益"]},
                     "bbox": [210, 80, 370, 220]},
                ]],
                layout={"pdf_info": [{"page_idx": 0, "page_size": [400, 300]}]},
            )
            result = MinerUToDocumentIRService().convert_zip(
                path, job_id="nested", source_file_id="source", document_language="zh",
            )
        self.assertEqual([block.captions for block in result.document.blocks], [["图 5 响应"], ["图 6 增益"]])
        self.assertEqual([block.text for block in result.document.blocks], ["图 5 响应", "图 6 增益"])


    def test_caption_sequences_preserve_sources_and_render_unique_native_fields(self) -> None:
        cases = [
            (5, ["图 5 响应 10 MHz", "图 6 增益 2 kΩ"]),
            (24, "图 24 响应 10 MHz\n图 25 增益 2 kΩ"),
            (24, [[{"type": "text", "content": "图 24 响应 10 MHz"}],
                  [{"type": "text", "content": "图 25 增益 2 kΩ"}]]),
            (5, [{"type": "text", "content": "图 5 响应 ("},
                 {"type": "equation_inline", "content": r"\pm"},
                 {"type": "text", "content": " 18 V)\n图 6 增益 2 kΩ"}]),
        ]
        for number, donor in cases:
            with self.subTest(number=number, donor=donor), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "caption.zip"
                blocks = [[
                    {"type": "text", "text": f"见图 {number} 和图 {number + 1}。", "bbox": [10, 10, 390, 40]},
                    {"type": "image", "content": {"image_source": {"path": "images/chart.png"}, "image_caption": []},
                     "bbox": [10, 80, 170, 220]},
                    {"type": "figure", "content": {"image_source": {"path": "images/chart.png"}, "figure_caption": donor},
                     "bbox": [210, 80, 370, 220]},
                ]]
                original = copy.deepcopy(blocks)
                self._write_zip(path, blocks=blocks, layout={"pdf_info": [{"page_idx": 0, "page_size": [400, 300]}]})
                checksum = hashlib.sha256(path.read_bytes()).hexdigest()
                result = MinerUToDocumentIRService().convert_zip(path, job_id="captions", source_file_id="source", document_language="zh")
                repeated = MinerUToDocumentIRService().convert_zip(path, job_id="captions", source_file_id="source", document_language="zh")
                self.assertEqual(result.document, repeated.document)
                self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), checksum)
                self.assertEqual(blocks, original)
            document = DocumentIR.model_validate_json(result.document.model_dump_json())
            images = document.blocks[1:]
            self.assertTrue(all(len(block.captions) == 1 and block.text == block.captions[0] for block in images))
            self.assertNotIn("caption_binding_uncertain", document.quality_flags)
            evidence = images[0].metadata["caption_reconciliation"]
            self.assertEqual(evidence["originals"], [])
            self.assertEqual(evidence["bindings"][0]["donor_block_id"], "block_0003")
            self.assertEqual(evidence["bindings"][0]["target_block_id"], "block_0002")
            self.assertEqual(evidence["bindings"][0]["source_locator"]["json_pointer"], "/0/2")
            self.assertEqual(images[1].metadata["caption_reconciliation"]["originals"][0]["value"], donor)
            self.assertNotIn("caption", images[1].metadata)
            self.assertEqual(document.metadata.source_type_counts, {"figure": 1, "image": 1, "text": 1})
            content = DocumentIRWordRenderer().render_to_bytes(
                document, template_path=None, config_path=None, image_loader=lambda _asset: MINIMAL_PNG_BYTES,
            )
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                xml = archive.read("word/document.xml").decode()
            self.assertEqual(xml.count(" SEQ 图 \\* ARABIC "), 2)
            for target in (number, number + 1):
                self.assertEqual(xml.count(f'w:name="_md2word_figure_{target}"'), 1)
                self.assertIn(f" REF _md2word_figure_{target} \\h ", xml)
            self.assertNotIn("18 V)", [paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs])
            renderer_evidence = build_renderer_evidence(document, content)
            ensure_renderer_integrity(renderer_evidence)
            self.assertEqual(renderer_evidence["inline_shape_count"], 2)
            self.assertEqual(renderer_evidence["caption_reconciliation"]["binding_count"], 2)
            self.assertNotIn("originals", json.dumps(renderer_evidence["caption_reconciliation"]))
            self.assertNotIn("响应", json.dumps(renderer_evidence["caption_reconciliation"], ensure_ascii=False))
            self.assertEqual(renderer_evidence["mineru_validation"]["inventory"]["raw_block_count"], 3)


    def test_accepted_redacted_caption_pair_stays_uncertain_without_source_boundaries(self) -> None:
        fixture = Path(__file__).parent / "fixtures/component_pdf_samples/samples/source-redacted/datasheet/zk7044-caption-pair/mineru"
        original = json.loads((fixture / "content_list_v2.json").read_text())
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "redacted.zip"
            with zipfile.ZipFile(path, "w") as archive:
                for file in sorted(fixture.rglob("*")):
                    if file.is_file():
                        archive.write(file, file.relative_to(fixture).as_posix())
            result = MinerUToDocumentIRService().convert_zip(path, job_id="redacted", source_file_id="source", document_language="zh")
        blocks = result.document.blocks
        raw = original[0][0]["content"]["chart_caption"]
        self.assertEqual(blocks[0].captions, [raw[0]["content"]])
        self.assertEqual(blocks[1].captions, [])
        self.assertEqual(blocks[0].metadata["caption_reconciliation"]["originals"][0]["value"], raw)
        self.assertIn("split_boundary_uncertain", blocks[0].metadata["caption_reconciliation"]["reasons"])
        self.assertEqual(result.document.quality_flags.count("caption_binding_uncertain"), 1)
        image_bytes = {asset.asset_id: asset.content for asset in result.image_assets}
        content = DocumentIRWordRenderer().render_to_bytes(
            result.document, template_path=None, config_path=None, image_loader=lambda asset: image_bytes[asset.id],
        )
        self.assertEqual(sum(raw[0]["content"] in paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs), 1)
        evidence = build_renderer_evidence(result.document, content)
        ensure_renderer_integrity(evidence)
        self.assertEqual(evidence["inline_shape_count"], 2)
        self.assertEqual(evidence["caption_reconciliation"]["binding_count"], 0)


    def test_numbered_tables_and_contexts_keep_native_tables_and_notes(self) -> None:
        condition = "at TA = 25°C (unless otherwise noted)"
        cases = [
            ([[], ["Table 1 Limits", "Table 2 Ratings"]], 2, False),
            ([f"Table 1 Limits\n{condition}"], 1, False),
            (["Table 1 Limits unless otherwise noted"], 1, True),
        ]
        for captions, expected_tables, uncertain in cases:
            with self.subTest(captions=captions), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "tables.zip"
                tables = [{"type": "table", "content": {"table_caption": caption,
                           "table_footnote": [f"Note {index}: typical values."],
                           "html": f"<table><tr><th>Parameter</th><td>Value{index}</td></tr></table>"},
                           "bbox": [10, 80 + index * 100, 190, 150 + index * 100]}
                          for index, caption in enumerate(captions)]
                source_blocks = tables + ([{"type": "text", "text": "See Table 1 and Table 2.", "bbox": [10, 300, 190, 330]}]
                                          if expected_tables == 2 else [])
                self._write_zip(path, blocks=[source_blocks], layout={"pdf_info": [{"page_idx": 0, "page_size": [200, 400]}]})
                result = MinerUToDocumentIRService().convert_zip(path, job_id="tables", source_file_id="source", document_language="en")
            document = result.document
            main = [block for block in document.blocks if block.type == "table"]
            self.assertEqual(len(main), expected_tables)
            self.assertEqual([block.html for block in main], [table["content"]["html"] for table in tables])
            self.assertEqual("caption_binding_uncertain" in document.quality_flags, uncertain)
            if expected_tables == 2:
                self.assertEqual([block.captions for block in main], [["Table 1 Limits"], ["Table 2 Ratings"]])
            contexts = [block for block in document.blocks if block.metadata.get("annotation_role") == "table_context"]
            if contexts:
                context = contexts[0]
                self.assertEqual(context.text, condition)
                self.assertNotIn("caption_reconciliation", context.metadata)
                self.assertEqual(context.metadata["source_block_id"], main[0].id)
            if uncertain:
                self.assertEqual(len(document.blocks), 1)
                self.assertEqual(main[0].captions, captions)
                self.assertIn("mixed_context_boundary_uncertain", main[0].metadata["caption_reconciliation"]["reasons"])
            content = DocumentIRWordRenderer().render_to_bytes(
                document, template_path=None, config_path=None, image_loader=lambda _asset: None,
            )
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                root = ElementTree.fromstring(archive.read("word/document.xml"))
            word_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            if expected_tables == 2:
                instructions = [node.text for node in root.iter(f"{word_ns}instrText")]
                self.assertEqual(instructions.count(" SEQ Table \\* ARABIC "), 2)
                bookmarks = [node.get(f"{word_ns}name") for node in root.iter(f"{word_ns}bookmarkStart")]
                for number in (1, 2):
                    self.assertEqual(bookmarks.count(f"_md2word_table_{number}"), 1)
                    self.assertIn(f" REF _md2word_table_{number} \\h ", instructions)
            body = root.find(f"{word_ns}body")
            body_texts = ["".join(node.text or "" for node in child.iter(f"{word_ns}t")) for child in body]
            for index in range(expected_tables):
                title_index = next(position for position, text in enumerate(body_texts) if f"Table {index + 1}" in text)
                table_index = next(position for position, text in enumerate(body_texts) if f"Value{index}" in text)
                note_index = next(position for position, text in enumerate(body_texts) if f"Note {index}:" in text)
                self.assertLess(title_index, table_index)
                self.assertLess(table_index, note_index)
            if condition in body_texts:
                self.assertLess(body_texts.index(condition), next(index for index, text in enumerate(body_texts) if "Value0" in text))
                self.assertEqual(body_texts.count(condition), 1)
            evidence = build_renderer_evidence(document, content)
            ensure_renderer_integrity(evidence)
            self.assertEqual(evidence["table_count"], expected_tables)
            self.assertEqual(evidence["caption_reconciliation"]["block_count"], expected_tables)


    def test_split_table_conditions_keep_repeated_source_occurrences(self) -> None:
        condition = "at TA = 25°C"
        raw_caption = f"Table 1 Limits\n{condition}\n{condition}"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "repeated-conditions.zip"
            self._write_zip(path, blocks=[[{
                "type": "table", "content": {"table_caption": raw_caption,
                "html": "<table><tr><td>Value</td></tr></table>"}, "bbox": [10, 80, 190, 150],
            }]], layout={"pdf_info": [{"page_idx": 0, "page_size": [200, 300]}]})
            result = MinerUToDocumentIRService().convert_zip(
                path, job_id="repeated-conditions", source_file_id="source", document_language="en",
            )
        contexts = [block for block in result.document.blocks if block.metadata.get("annotation_role") == "table_context"]
        self.assertEqual([block.text for block in contexts], [condition, condition])
        table = next(block for block in result.document.blocks if block.type == "table")
        self.assertEqual(table.metadata["caption_reconciliation"]["originals"][0]["value"], raw_caption)
        self.assertEqual(len(table.metadata["caption_reconciliation"]["bindings"]), 3)
        content = DocumentIRWordRenderer().render_to_bytes(
            result.document, template_path=None, config_path=None, image_loader=lambda _asset: None,
        )
        self.assertEqual([paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs].count(condition), 2)
        ensure_renderer_integrity(build_renderer_evidence(result.document, content))


    def test_duplicate_table_titles_with_context_stay_uncertain_and_visible(self) -> None:
        title = "Table 1 Limits"
        raw_caption = f"{title}\n{title}\nat TA = 25°C"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "repeated-titles.zip"
            self._write_zip(path, blocks=[[{
                "type": "table", "content": {"table_caption": raw_caption,
                "html": "<table><tr><td>Value</td></tr></table>"}, "bbox": [10, 80, 190, 150],
            }]], layout={"pdf_info": [{"page_idx": 0, "page_size": [200, 300]}]})
            result = MinerUToDocumentIRService().convert_zip(
                path, job_id="repeated-titles", source_file_id="source", document_language="en",
            )
        self.assertEqual(len(result.document.blocks), 1)
        table = result.document.blocks[0]
        self.assertEqual(table.metadata["caption_reconciliation"]["status"], "uncertain")
        self.assertIn("caption_binding_uncertain", result.document.quality_flags)
        self.assertEqual(table.captions, [raw_caption.replace("\n", " ")])
        content = DocumentIRWordRenderer().render_to_bytes(
            result.document, template_path=None, config_path=None, image_loader=lambda _asset: None,
        )
        self.assertEqual("\n".join(paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs).count(title), 2)
        ensure_renderer_integrity(build_renderer_evidence(result.document, content))


    def test_adjacent_paragraph_is_not_copied_into_visible_captions(self) -> None:
        blocks = [[
            {"type": "chart", "img_path": "images/chart.png", "bbox": [10, 10, 170, 100]},
            {"type": "paragraph", "content": {"paragraph_content": [{"type": "text", "content": "图 5 响应"}]},
             "bbox": [10, 105, 170, 120]},
        ]]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "paragraph.zip"
            self._write_zip(path, blocks=blocks, layout={"pdf_info": [{"page_idx": 0, "page_size": [200, 300]}]})
            result = MinerUToDocumentIRService().convert_zip(path, job_id="paragraph", source_file_id="source", document_language="zh")
        image, paragraph = result.document.blocks
        self.assertEqual(image.captions, [])
        self.assertIsNone(image.text)
        self.assertEqual((paragraph.id, paragraph.order, paragraph.text, paragraph.layout.bbox),
                         ("block_0002", 1, "图 5 响应", (10.0, 105.0, 170.0, 120.0)))
        receipt = image.metadata["caption_reconciliation"]["bindings"][0]
        self.assertTrue(receipt["evidence_only"])
        self.assertEqual(receipt["donor_block_id"], paragraph.id)
        content = DocumentIRWordRenderer().render_to_bytes(
            result.document, template_path=None, config_path=None, image_loader=lambda _asset: MINIMAL_PNG_BYTES,
        )
        self.assertEqual([paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs].count("图 5 响应"), 1)
        ensure_renderer_integrity(build_renderer_evidence(result.document, content))


    def test_caption_candidate_damage_preserves_ir201_report_and_stops_conversion(self) -> None:
        splitter = caption_reconciler._split_group

        def damage(group, family):
            fragments = splitter(group, family)
            return [replace(fragment, text=fragment.text.replace("私有甲", "损坏")) for fragment in fragments]

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "damaged.zip"
            self._write_zip(path, blocks=[[
                {"type": "chart", "img_path": "images/chart.png", "bbox": [10, 10, 90, 90]},
                {"type": "chart", "img_path": "images/chart.png", "chart_caption": "图 5 私有甲\n图 6 私有乙",
                 "bbox": [110, 10, 190, 90]},
            ]], layout={"pdf_info": [{"page_idx": 0, "page_size": [200, 300]}]})
            with patch.object(caption_reconciler, "_split_group", side_effect=damage):
                with self.assertRaises(DocumentIntegrityError) as caught:
                    MinerUToDocumentIRService().convert_zip(path, job_id="damage", source_file_id="source", document_language="zh")
        evidence = caught.exception.evidence
        self.assertEqual(evidence["mineru_validation"]["inventory"]["raw_block_count"], 2)
        self.assertTrue(evidence["mineru_validation"]["complete"])
        self.assertTrue(evidence["mineru_validation"]["can_convert"])
        self.assertNotIn("私有", json.dumps(evidence, ensure_ascii=False))
        self.assertIn("caption_reconciliation_integrity:fragment_conservation", evidence["blocking_quality_flags"])


    def test_legacy_inferred_objects_preserve_generic_captions_and_evidence(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "legacy-inferred.zip"
            self._write_zip(path, blocks=[[
                {"type": "text", "img_path": "images/chart.png", "caption": ["图 5 响应"], "bbox": [10, 10, 190, 80]},
                {"type": "text", "table_body": "<table><tr><td>Value</td></tr></table>",
                 "caption": ["Table 1 Limits"], "bbox": [10, 100, 190, 180]},
            ]], layout={"pdf_info": [{"page_idx": 0, "page_size": [200, 300]}]})
            result = MinerUToDocumentIRService().convert_zip(path, job_id="legacy", source_file_id="source", document_language="en")
        self.assertEqual([block.type for block in result.document.blocks], ["image", "table"])
        self.assertEqual([block.captions for block in result.document.blocks], [["图 5 响应"], ["Table 1 Limits"]])
        for index, block in enumerate(result.document.blocks):
            self.assertEqual(block.metadata["caption_reconciliation"]["originals"][0]["json_pointer"], f"/0/{index}/caption")
        content = DocumentIRWordRenderer().render_to_bytes(
            result.document, template_path=None, config_path=None, image_loader=lambda _asset: MINIMAL_PNG_BYTES,
        )
        evidence = build_renderer_evidence(result.document, content)
        ensure_renderer_integrity(evidence)
        self.assertEqual((evidence["inline_shape_count"], evidence["table_count"]), (1, 1))
        self.assertNotIn("caption_binding_uncertain", result.document.quality_flags)

    def test_convert_zip_prefers_v2_content_and_keeps_legacy_only_fallback(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            both_path = temp_path / "both.zip"
            self._write_zip(
                both_path,
                blocks=[{"type": "text", "text": "legacy content", "page_idx": 0}],
                v2_blocks=[
                    [
                        {
                            "type": "paragraph",
                            "content": {"paragraph_content": [{"type": "text", "content": "v2 content"}]},
                            "bbox": [10, 10, 100, 30],
                        }
                    ]
                ],
                layout={"pdf_info": [{"page_idx": 0, "page_size": [200, 300]}]},
            )
            legacy_path = temp_path / "legacy.zip"
            self._write_zip(
                legacy_path,
                blocks=[{"type": "text", "text": "legacy content", "page_idx": 0}],
                layout={"pdf_info": [{"page_idx": 0, "page_size": [200, 300]}]},
            )

            both = MinerUToDocumentIRService().convert_zip(
                both_path,
                job_id="docjob_v2_priority",
                source_file_id="file_v2_priority",
                document_language="en",
            )
            legacy = MinerUToDocumentIRService().convert_zip(
                legacy_path,
                job_id="docjob_legacy_fallback",
                source_file_id="file_legacy_fallback",
                document_language="en",
            )

        self.assertEqual([block.text for block in both.document.blocks], ["v2 content"])
        self.assertEqual([block.text for block in legacy.document.blocks], ["legacy content"])

    def test_convert_zip_scales_physical_page_size_uniformly_to_cover_bbox(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(
                zip_path,
                blocks=[{"type": "text", "text": "Portrait", "bbox": [0, 0, 942, 879], "page_idx": 0}],
                layout={"pdf_info": [{"page_idx": 0, "page_size": [595, 841]}]},
            )

            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_scaled_portrait",
                source_file_id="file_scaled_portrait",
                document_language="en",
            )

        page = result.document.pages[0]
        self.assertEqual(page.width, 942)
        self.assertAlmostEqual(page.height, 841 * (942 / 595))
        self.assertAlmostEqual(page.width / page.height, 595 / 841)
        self.assertGreaterEqual(page.height, 879)
        self.assertEqual(page.metadata["orientation"], "portrait")

    def test_convert_zip_structures_v2_multiline_lists_and_preserves_math_sources(self) -> None:
        fraction = r"\frac {\left(V _ {I N} - V _ {S}\right)}{R _ {S} + 5 0 0 \Omega} \leq 5 \mathrm{mA}"
        total_delay = r"\mathrm{总} t _ {s} = \sqrt {\left(t _ {s , C M 7 5 0 1}\right) ^ {2}}"
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(
                zip_path,
                blocks=[],
                v2_blocks=[
                    [
                        {
                            "type": "paragraph",
                            "content": {
                                "paragraph_content": [
                                    {"type": "text", "content": "- Low offset "},
                                    {"type": "equation_inline", "content": r"12\mu \mathrm{V}"},
                                    {"type": "text", "content": "\n- Low noise\n- Stable output"},
                                ]
                            },
                            "bbox": [10, 20, 180, 90],
                        },
                        {
                            "type": "paragraph",
                            "content": {
                                "paragraph_content": [
                                    {"type": "text", "content": "■ Sensors\nThermocouple\nRTD"}
                                ]
                            },
                            "bbox": [10, 100, 180, 150],
                        },
                        {
                            "type": "paragraph",
                            "content": {"paragraph_content": [{"type": "text", "content": "ordinary\nbody"}]},
                            "bbox": [10, 160, 180, 190],
                        },
                        {
                            "type": "paragraph",
                            "content": {
                                "paragraph_content": [
                                    {"type": "text", "content": "Input "},
                                    {"type": "equation_inline", "content": r"V _ {I N}"},
                                    {"type": "text", "content": " stays visible"},
                                ]
                            },
                            "bbox": [10, 200, 180, 230],
                        },
                        {
                            "type": "list",
                            "content": {
                                "list_items": [
                                    {
                                        "item_content": [
                                            {"type": "text", "content": "Limit "},
                                            {"type": "equation_inline", "content": r"I _ {O U T}"},
                                        ]
                                    }
                                ]
                            },
                            "bbox": [10, 232, 180, 238],
                        },
                        {
                            "type": "equation_interline",
                            "content": {"math_content": fraction, "math_type": "latex", "image_source": {"path": "images/a.png"}},
                            "bbox": [10, 240, 180, 270],
                        },
                        {
                            "type": "equation_interline",
                            "content": {"math_content": total_delay, "math_type": "latex"},
                            "bbox": [10, 280, 180, 310],
                        },
                    ]
                ],
                layout={"pdf_info": [{"page_idx": 0, "page_size": [200, 400]}]},
            )

            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_v2_rich_text",
                source_file_id="file_v2_rich_text",
                document_language="en",
            )

        blocks = result.document.blocks
        self.assertEqual(blocks[0].type, "list")
        self.assertEqual(blocks[0].list_items, [r"- Low offset $12\mu \mathrm{V}$", "- Low noise", "- Stable output"])
        self.assertEqual(blocks[1].list_items, ["■ Sensors\nThermocouple\nRTD"])
        self.assertEqual(blocks[2].type, "paragraph")
        self.assertEqual(blocks[2].text, "ordinary body")
        self.assertEqual(blocks[3].text, r"Input $V _ {I N}$ stays visible")
        self.assertEqual(blocks[4].list_items, [r"Limit $I _ {O U T}$"])
        self.assertEqual(blocks[5].latex, fraction)
        self.assertEqual(blocks[6].latex, total_delay)

    def test_convert_zip_maps_nested_v2_chart_to_image_and_records_source_inventory(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(
                zip_path,
                blocks=[],
                v2_blocks=[
                    [
                        {
                            "type": "chart",
                            "content": {
                                "image_source": {"path": "images/chart.png"},
                                "content": "",
                                "chart_caption": [
                                    {"type": "text", "content": "图 1 输出 "},
                                    {"type": "equation_inline", "content": r"V _ {O U T}"},
                                ],
                                "chart_footnote": [{"type": "text", "content": "测试条件：25°C"}],
                            },
                            "bbox": [10, 20, 180, 160],
                        }
                    ]
                ],
                layout={"pdf_info": [{"page_idx": 0, "page_size": [200, 300]}]},
                images={"images/chart.png": MINIMAL_PNG_BYTES},
            )

            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_nested_chart",
                source_file_id="file_nested_chart",
                document_language="zh",
            )

        block = result.document.blocks[0]
        self.assertEqual(block.type, "image")
        self.assertEqual(block.metadata["source_type"], "chart")
        self.assertEqual(block.asset_ref, result.document.assets[0].id)
        self.assertEqual(result.image_assets[0].content, MINIMAL_PNG_BYTES)
        self.assertEqual(block.captions, [r"图 1 输出$V _ {O U T}$"])
        self.assertEqual(block.footnotes, ["测试条件：25°C"])
        self.assertEqual(result.document.metadata.source_block_count, 1)
        self.assertEqual(result.document.metadata.source_type_counts, {"chart": 1})
        self.assertNotIn("unsupported_block_type:chart", result.document.quality_flags)

    def test_convert_zip_preserves_multiple_logical_caption_groups(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(
                zip_path,
                blocks=[],
                v2_blocks=[
                    [
                        {
                            "type": "chart",
                            "content": {
                                "image_source": {"path": "images/strings.png"},
                                "chart_caption": ["图 1 第一条", "图 2 第二条"],
                            },
                            "bbox": [10, 20, 180, 160],
                        }
                    ],
                    [
                        {
                            "type": "chart",
                            "content": {
                                "image_source": {"path": "images/groups.png"},
                                "chart_caption": [
                                    [{"type": "text", "content": "图 3 第三条"}],
                                    [
                                        {"type": "text", "content": "图 4 输出 "},
                                        {"type": "equation_inline", "content": r"V _ {O U T}"},
                                    ],
                                ],
                            },
                            "bbox": [10, 20, 180, 160],
                        }
                    ],
                ],
                layout={
                    "pdf_info": [
                        {"page_idx": 0, "page_size": [200, 300]},
                        {"page_idx": 1, "page_size": [200, 300]},
                    ]
                },
                images={"images/strings.png": MINIMAL_PNG_BYTES, "images/groups.png": MINIMAL_PNG_BYTES},
            )

            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_caption_groups",
                source_file_id="file_caption_groups",
                document_language="zh",
            )

        self.assertEqual(result.document.blocks[0].captions, ["图 1 第一条", "图 2 第二条"])
        self.assertEqual(result.document.blocks[1].captions, ["图 3 第三条", r"图 4 输出$V _ {O U T}$"])

    def test_rich_chart_captions_render_without_isolated_fragments_and_keep_paired_layout(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(
                zip_path,
                blocks=[],
                v2_blocks=[
                    [
                        {
                            "type": "chart",
                            "content": {
                                "image_source": {"path": "images/left.png"},
                                "chart_caption": [
                                    {"type": "text", "content": "图 33 通道隔离度 (  "},
                                    {"type": "equation_inline", "content": r"\pm"},
                                    {"type": "text", "content": " 18V)"},
                                ],
                            },
                            "bbox": [10, 80, 170, 220],
                        },
                        {
                            "type": "chart",
                            "content": {
                                "image_source": {"path": "images/right.png"},
                                "chart_caption": [
                                    {"type": "text", "content": "图 34 输出摆幅 (  "},
                                    {"type": "equation_inline", "content": r"\pm"},
                                    {"type": "text", "content": " 15V)"},
                                ],
                            },
                            "bbox": [190, 82, 350, 220],
                        },
                    ]
                ],
                layout={"pdf_info": [{"page_idx": 0, "page_size": [400, 600]}]},
                images={"images/left.png": MINIMAL_PNG_BYTES, "images/right.png": MINIMAL_PNG_BYTES},
            )
            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_rich_caption_pair",
                source_file_id="file_rich_caption_pair",
                document_language="zh",
            )
            image_bytes = {asset.asset_id: asset.content for asset in result.image_assets}
            content = DocumentIRWordRenderer().render_to_bytes(
                result.document,
                template_path=None,
                config_path=None,
                image_loader=lambda asset: image_bytes.get(asset.id),
            )

        self.assertEqual(result.document.blocks[0].captions, [r"图 33 通道隔离度 ($\pm$ 18V)"])
        rendered = Document(io.BytesIO(content))
        self.assertNotIn("±", [paragraph.text.strip() for paragraph in rendered.paragraphs])
        self.assertNotIn("18V)", [paragraph.text.strip() for paragraph in rendered.paragraphs])
        with zipfile.ZipFile(io.BytesIO(content)) as docx_zip:
            document_xml = docx_zip.read("word/document.xml")
        root = ElementTree.fromstring(document_xml)
        word_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        math_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
        paragraph_texts = [
            "".join(node.text or "" for node in paragraph.iter() if node.tag in {f"{word_ns}t", f"{math_ns}t"})
            for paragraph in root.iter(f"{word_ns}p")
        ]
        figure_33_paragraph = next(text for text in paragraph_texts if "通道隔离度" in text)
        self.assertIn("±", figure_33_paragraph)
        self.assertIn("18V)", figure_33_paragraph)
        xml = document_xml.decode("utf-8")
        self.assertNotIn('<w:br w:type="column"', xml)
        self.assertNotIn("<w:tabs>", xml)
        self.assertEqual(len(rendered.tables), 1)
        self.assertEqual(xml.count(" SEQ 图 \\* ARABIC "), 2)

    def test_convert_zip_fails_closed_when_declared_formula_payload_is_missing(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(
                zip_path,
                blocks=[],
                v2_blocks=[
                    [
                        {
                            "type": "equation_interline",
                            "content": {"math_type": "latex", "image_source": {"path": "images/missing.png"}},
                            "bbox": [10, 20, 180, 60],
                        }
                    ]
                ],
                layout={"pdf_info": [{"page_idx": 0, "page_size": [200, 300]}]},
            )

            with self.assertRaisesRegex(MinerUArtifactValidationError, "mineru_block_payload_invalid"):
                MinerUToDocumentIRService().convert_zip(
                    zip_path,
                    job_id="docjob_missing_formula",
                    source_file_id="file_missing_formula",
                    document_language="zh",
                )


    def test_shared_final_promotional_page_rule_is_conservative(self) -> None:
        image = {"type": "image", "img_path": "images/qr.png", "bbox": [100, 100, 300, 300], "page_idx": 2}
        positive = [
            {"type": "text", "text": "可使用微信扫码进入芯片助手小程序", "page_idx": 2},
            image,
        ]
        self.assertEqual(find_final_promotional_page(positive), 2)

        cases = {
            "单个关键词": [{"type": "text", "text": "微信公众号", "page_idx": 2}, image],
            "无资产": [{"type": "text", "text": "微信扫码进入小程序", "page_idx": 2}],
            "方形技术图": [{"type": "text", "text": "产品封装图", "page_idx": 2}, image],
            "免责声明": [
                {"type": "text", "text": "微信扫码进入小程序" + "本资料仅供技术开发人员参考。" * 20, "page_idx": 2},
                image,
            ],
            "版本表": [
                {"type": "text", "text": "微信扫码进入小程序", "page_idx": 2},
                {"type": "table", "table_body": "<table><tr><td>V1</td></tr></table>", "page_idx": 2},
                image,
            ],
            "编号标题": [
                {"type": "title", "text": "7 Version update 微信扫码进入小程序", "page_idx": 2},
                image,
            ],
            "编号题注": [
                {"type": "text", "text": "微信扫码进入小程序", "page_idx": 2},
                {**image, "image_caption": ["图 8 技术框图"]},
            ],
            "公式": [
                {"type": "text", "text": "微信扫码进入小程序", "page_idx": 2},
                {"type": "equation", "latex": "V=IR", "page_idx": 2},
                image,
            ],
            "列表": [
                {"type": "text", "text": "微信扫码进入小程序", "page_idx": 2},
                {"type": "list", "list_items": ["技术参数"], "page_idx": 2},
                image,
            ],
            "非最后页": [
                {"type": "text", "text": "微信扫码进入小程序", "page_idx": 1},
                {**image, "page_idx": 1},
                {"type": "text", "text": "最后一页技术正文", "page_idx": 2},
            ],
        }
        for label, blocks in cases.items():
            with self.subTest(label=label):
                self.assertIsNone(find_final_promotional_page(blocks))

    def test_convert_zip_removes_final_promotional_page_blocks_pages_and_assets(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(
                zip_path,
                blocks=[
                    {"type": "text", "text": "正文", "bbox": [72, 80, 200, 100], "page_idx": 0},
                    {
                        "type": "image",
                        "img_path": "images/chart.png",
                        "image_caption": ["图 1 正文图片"],
                        "bbox": [72, 120, 300, 300],
                        "page_idx": 0,
                    },
                    {"type": "text", "text": "可使用微信扫码进入芯片助手小程序", "bbox": [72, 80, 300, 120], "page_idx": 1},
                    {
                        "type": "image",
                        "img_path": "images/qr.png",
                        "bbox": [150, 180, 350, 380],
                        "page_idx": 1,
                    },
                    {"type": "footer", "text": "www.example.com", "bbox": [72, 740, 300, 760], "page_idx": 1},
                ],
                layout={
                    "pdf_info": [
                        {"page_idx": 0, "page_size": [612, 792]},
                        {"page_idx": 1, "page_size": [612, 792]},
                    ]
                },
                images={"images/chart.png": MINIMAL_PNG_BYTES, "images/qr.png": MINIMAL_PNG_BYTES},
            )

            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_promo",
                source_file_id="file_promo",
                document_language="zh",
            )

        self.assertEqual([page.index for page in result.document.pages], [0])
        self.assertEqual({block.page_index for block in result.document.blocks}, {0})
        self.assertEqual([asset.source_path for asset in result.document.assets], ["images/chart.png"])
        self.assertEqual([asset.source_path for asset in result.image_assets], ["images/chart.png"])
        self.assertIn("final_promotional_page_filtered:1", result.document.quality_flags)

    def test_convert_zip_uses_physical_final_page_from_layout_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(
                zip_path,
                blocks=[
                    {"type": "text", "text": "微信扫码进入小程序", "page_idx": 1},
                    {"type": "image", "img_path": "images/qr.png", "page_idx": 1},
                ],
                layout={"pdf_info": [{"page_idx": index, "page_size": [612, 792]} for index in range(3)]},
                images={"images/qr.png": MINIMAL_PNG_BYTES},
            )

            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_non_final_promo",
                source_file_id="file_non_final_promo",
                document_language="zh",
            )

        self.assertEqual([page.index for page in result.document.pages], [0, 1, 2])
        self.assertNotIn("final_promotional_page_filtered:1", result.document.quality_flags)
        self.assertEqual(len(result.image_assets), 1)

    def test_convert_zip_filters_final_promotional_page_without_layout_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(
                zip_path,
                blocks=[
                    {"type": "text", "text": "微信扫码进入小程序", "page_idx": 3},
                    {"type": "image", "img_path": "images/qr.png", "page_idx": 3},
                ],
                layout=None,
                images={"images/qr.png": MINIMAL_PNG_BYTES},
            )

            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_promo_without_layout",
                source_file_id="file_promo_without_layout",
                document_language="zh",
            )

        self.assertEqual(result.document.pages, [])
        self.assertEqual(result.document.blocks, [])
        self.assertEqual(result.document.assets, [])
        self.assertEqual(result.image_assets, [])
        self.assertIn("final_promotional_page_filtered:3", result.document.quality_flags)

    def test_convert_zip_outputs_v2_pages_layout_roles_and_annotations(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(
                zip_path,
                blocks=[
                    {"type": "title", "text": "1 Features", "text_level": 1, "bbox": [72, 80, 180, 100], "page_idx": 0},
                    {"type": "list", "list_items": ["Low noise", "High slew rate"], "bbox": [72, 120, 260, 180], "page_idx": 0},
                    {
                        "type": "chart",
                        "img_path": "images/chart.png",
                        "chart_caption": ["Noise Density"],
                        "chart_footnote": ["Typical values."],
                        "bbox": [72, 200, 360, 320],
                        "page_idx": 0,
                    },
                    {
                        "type": "table",
                        "table_caption": ["Device Information"],
                        "table_footnote": ["All packages are listed in the addendum."],
                        "table_body": "<table><tr><td>Part</td></tr></table>",
                        "bbox": [72, 340, 540, 430],
                        "page_idx": 0,
                        "continuation_heights": [120],
                    },
                    {"type": "header", "text": "Header text", "bbox": [72, 20, 200, 40], "page_idx": 0},
                    {"type": "footer", "text": "Footer text", "bbox": [72, 740, 300, 760], "page_idx": 0},
                ],
                layout={"pdf_info": [{"page_idx": 0, "page_size": [612, 792]}]},
            )

            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_layout",
                source_file_id="file_layout",
                document_language="en",
            )

        document = result.document
        self.assertEqual(document.schema_version, "document-ir/v2")
        self.assertEqual(document.pages[0].width, 612)
        self.assertEqual(document.pages[0].height, 792)
        self.assertEqual(document.pages[0].metadata["page_size_source"], "layout.json")
        self.assertEqual(document.pages[0].metadata["orientation"], "portrait")
        self.assertEqual(document.blocks[0].layout.bbox, (72.0, 80.0, 180.0, 100.0))
        self.assertEqual(document.blocks[0].layout.column_hint, "left")
        self.assertEqual(document.blocks[1].list_items, ["Low noise", "High slew rate"])
        self.assertEqual(document.blocks[1].layout.column_hint, "left")
        self.assertEqual(document.blocks[2].type, "image")
        self.assertEqual(document.blocks[2].text, "Noise Density")
        self.assertEqual(document.blocks[2].layout.column_hint, "left")
        self.assertEqual(document.blocks[2].captions, ["Noise Density"])
        self.assertEqual(document.blocks[2].footnotes, ["Typical values."])
        self.assertEqual(document.blocks[2].metadata["layout_role_reason"], "column_hint=left;bbox_width=288.00;page_width=612.00")
        self.assertEqual(document.blocks[3].captions, ["Device Information"])
        self.assertEqual(document.blocks[3].footnotes, ["All packages are listed in the addendum."])
        self.assertEqual(document.blocks[3].layout.continuation_heights, [120.0])
        self.assertEqual(document.blocks[3].layout.column_hint, "wide")
        self.assertEqual(document.blocks[3].metadata["layout_role_reason"], "column_hint=wide;bbox_width=468.00;page_width=612.00")
        self.assertEqual(document.blocks[4].layout.role, "page_header")
        self.assertEqual(document.blocks[5].layout.role, "page_footer")
        self.assertNotIn("unsupported_block_type:header", document.quality_flags)
        self.assertIn("table_continuation:block_0004", document.quality_flags)
        self.assertIn("layout_wide_table:block_0004", document.quality_flags)
        self.assertIn("layout_caption_binding:block_0003", document.quality_flags)
        self.assertIn("layout_footnote_binding:block_0004", document.quality_flags)
        self.assertEqual(len(result.image_assets), 1)

    def test_convert_zip_merges_table_placeholder_continuation_and_landscape_orientation(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(
                zip_path,
                blocks=[
                    {
                        "type": "table",
                        "table_caption": ["Wide table"],
                        "table_body": "<table><tr><td>Part</td></tr></table>",
                        "bbox": [72, 120, 720, 240],
                        "page_idx": 0,
                    },
                    {"type": "table_placeholder", "bbox": [72, 50, 720, 250], "page_idx": 1},
                ],
                layout={
                    "pdf_info": [
                        {"page_idx": 0, "page_size": [792, 612]},
                        {"page_idx": 1, "page_size": [792, 612]},
                    ]
                },
            )

            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_continuation",
                source_file_id="file_continuation",
                document_language="en",
            )

        document = result.document
        self.assertEqual([page.metadata["orientation"] for page in document.pages], ["landscape", "landscape"])
        self.assertEqual(len(document.blocks), 1)
        self.assertEqual(document.blocks[0].type, "table")
        self.assertEqual(document.blocks[0].layout.column_hint, "wide")
        self.assertEqual(document.blocks[0].layout.continuation_heights, [200.0])
        self.assertIn("table_continuation:block_0001", document.quality_flags)
        self.assertIn("layout_wide_table:block_0001", document.quality_flags)

    def test_convert_zip_merges_real_mineru_empty_table_continuations(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(
                zip_path,
                blocks=[
                    {
                        "type": "table",
                        "table_body": "<table><tr><td>Part</td></tr></table>",
                        "bbox": [72, 120, 540, 300],
                        "page_idx": 0,
                    },
                    {
                        "type": "table",
                        "img_path": "",
                        "html": "",
                        "bbox": [72, 50, 540, 250],
                        "page_idx": 1,
                    },
                    {
                        "type": "table",
                        "content": {
                            "image_source": {"path": "images/"},
                            "table_caption": [],
                            "table_footnote": [],
                            "html": "",
                            "table_type": "simple_table",
                            "table_nest_level": 1,
                        },
                        "bbox": [72, 40, 540, 280],
                        "page_idx": 2,
                    },
                    {
                        "type": "table",
                        "content": {
                            "image_source": {"path": "images/"},
                            "table_caption": [],
                            "table_footnote": [],
                            "html": "<table><tr><td>保留正文</td></tr></table>",
                            "table_type": "simple_table",
                            "table_nest_level": 1,
                        },
                        "bbox": [72, 300, 540, 480],
                        "page_idx": 2,
                    },
                ],
                layout={
                    "pdf_info": [
                        {"page_idx": 0, "page_size": [612, 792]},
                        {"page_idx": 1, "page_size": [612, 792]},
                        {"page_idx": 2, "page_size": [612, 792]},
                    ]
                },
            )

            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_real_continuation",
                source_file_id="file_real_continuation",
                document_language="en",
            )

        self.assertEqual(len(result.document.blocks), 2)
        self.assertEqual(result.document.blocks[0].type, "table")
        self.assertEqual(result.document.blocks[0].layout.continuation_heights, [200.0, 240.0])
        self.assertEqual(result.document.blocks[1].html, "<table><tr><td>保留正文</td></tr></table>")
        self.assertEqual(result.document.metadata.source_block_count, 2)
        self.assertEqual(result.document.metadata.source_type_counts, {"table": 2})

    def test_convert_zip_keeps_contentful_or_orphan_empty_tables_fail_closed(self) -> None:
        real_table = {
            "type": "table",
            "table_body": "<table><tr><td>Part</td></tr></table>",
            "bbox": [72, 120, 540, 300],
            "page_idx": 0,
        }

        def nested_empty_table(**content: Any) -> dict[str, Any]:
            return {
                "type": "table",
                "content": {
                    "image_source": {"path": "images/"},
                    "html": "",
                    "table_type": "simple_table",
                    **content,
                },
                "bbox": [72, 40, 540, 280],
                "page_idx": 1,
            }

        cases = {
            "明确 text": [real_table, nested_empty_table(text="表格正文仍需保留")],
            "明确 description": [real_table, nested_empty_table(description="表格说明仍需保留")],
            "明确题注": [real_table, nested_empty_table(table_caption=["表 1 参数"])],
            "明确脚注": [real_table, nested_empty_table(table_footnote=["注：典型值"])],
            "孤立空表": [{**nested_empty_table(), "page_idx": 0}],
        }
        for label, blocks in cases.items():
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temp_dir:
                zip_path = Path(temp_dir) / "mineru.zip"
                self._write_zip(
                    zip_path,
                    blocks=blocks,
                    layout={
                        "pdf_info": [
                            {"page_idx": page_index, "page_size": [612, 792]}
                            for page_index in range(max(block["page_idx"] for block in blocks) + 1)
                        ]
                    },
                )

                with self.assertRaisesRegex(DocumentIntegrityError, "payload_missing:table"):
                    MinerUToDocumentIRService().convert_zip(
                        zip_path,
                        job_id="docjob_ambiguous_empty_table",
                        source_file_id="file_ambiguous_empty_table",
                        document_language="zh",
                    )

    def test_convert_zip_reassociates_only_strict_numbered_image_groups(self) -> None:
        blocks = [
            {"type": "chart", "img_path": "images/chart.png", "chart_caption": [], "bbox": [10, 10, 90, 90], "page_idx": 0},
            {
                "type": "chart",
                "img_path": "images/chart.png",
                "chart_caption": ["图 3 响应", "图 5 响应"],
                "bbox": [110, 10, 190, 90],
                "page_idx": 0,
            },
            {
                "type": "chart",
                "img_path": "images/chart.png",
                "chart_caption": ["图 4 响应"],
                "bbox": [10, 110, 90, 190],
                "page_idx": 0,
            },
            {
                "type": "chart",
                "img_path": "images/chart.png",
                "chart_caption": ["图 6 响应"],
                "bbox": [110, 110, 190, 190],
                "page_idx": 0,
            },
            {"type": "chart", "img_path": "images/chart.png", "chart_caption": ["图 7 正确"], "bbox": [10, 10, 90, 90], "page_idx": 1},
            {"type": "chart", "img_path": "images/chart.png", "chart_caption": ["图 8 正确"], "bbox": [110, 10, 190, 90], "page_idx": 1},
            {"type": "chart", "img_path": "images/chart.png", "chart_caption": [], "bbox": [10, 10, 90, 90], "page_idx": 2},
            {"type": "chart", "img_path": "images/chart.png", "chart_caption": ["图 9 数量不足"], "bbox": [110, 10, 190, 90], "page_idx": 2},
            {"type": "chart", "img_path": "images/chart.png", "chart_caption": [], "bbox": [10, 10, 90, 90], "page_idx": 3},
            {
                "type": "chart",
                "img_path": "images/chart.png",
                "chart_caption": ["图 11 不连续", "图 13 不连续"],
                "bbox": [110, 10, 190, 90],
                "page_idx": 3,
            },
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(
                zip_path,
                blocks=blocks,
                layout={"pdf_info": [{"page_idx": index, "page_size": [200, 300]} for index in range(4)]},
            )
            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_captions",
                source_file_id="file_captions",
                document_language="zh",
            )

        captions_by_page = {
            page_index: [block.captions for block in result.document.blocks if block.page_index == page_index]
            for page_index in range(4)
        }
        self.assertEqual(captions_by_page[0], [["图 3 响应"], ["图 4 响应"], ["图 5 响应"], ["图 6 响应"]])
        self.assertEqual(captions_by_page[1], [["图 7 正确"], ["图 8 正确"]])
        self.assertEqual(captions_by_page[2], [[], ["图 9 数量不足"]])
        self.assertEqual(captions_by_page[3], [[], ["图 11 不连续", "图 13 不连续"]])

    def test_convert_zip_demotes_table_condition_caption_to_context_paragraph(self) -> None:
        condition = (
            "at TA = +25°C, VS = ±5 V to ±15 V, VCM = VOUT = VS / 2, "
            "and RL = 10 kΩ connected to VS / 2 (unless otherwise noted)"
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(
                zip_path,
                blocks=[
                    {
                        "type": "table",
                        "table_caption": [
                            {"type": "text", "content": condition[:64]},
                            {"type": "text", "content": condition[64:]},
                        ],
                        "table_body": "<table><tr><th>PARAMETER</th><th>VALUE</th></tr></table>",
                        "bbox": [40, 120, 572, 260],
                        "page_idx": 0,
                    }
                ],
                layout={"pdf_info": [{"page_idx": 0, "page_size": [612, 792]}]},
            )

            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_context",
                source_file_id="file_context",
                document_language="en",
            )

        document = result.document
        self.assertEqual([block.type for block in document.blocks], ["paragraph", "table"])
        self.assertEqual(document.blocks[0].id, "block_0001_context_0001")
        self.assertEqual(document.blocks[0].text, condition)
        self.assertEqual(document.blocks[0].metadata["annotation_role"], "table_context")
        self.assertEqual(document.blocks[0].metadata["source_block_id"], "block_0001")
        self.assertEqual(document.blocks[1].id, "block_0001")
        self.assertEqual(document.blocks[1].captions, [])
        self.assertIsNone(document.blocks[1].text)
        self.assertIn("layout_table_context_binding:block_0001", document.quality_flags)
        self.assertNotIn("layout_caption_binding:block_0001", document.quality_flags)
        self.assertIn("layout_wide_table:block_0001", document.quality_flags)

    def test_convert_zip_accepts_nested_v2_pages_and_rich_content_fields(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(
                zip_path,
                blocks=[
                    [
                        {
                            "type": "text",
                            "text": "1 Features........................ 1\n2 Specifications.................. 2\n3 Pins............................ 3",
                            "bbox": [72, 80, 280, 160],
                        },
                        {
                            "type": "equation_interline",
                            "content": [{"type": "text", "content": r"V_{OS}"}],
                            "bbox": [72, 180, 180, 210],
                        },
                    ],
                    [
                        {
                            "type": "table",
                            "content": [{"html": "<table><tr><td>VOS</td></tr></table>"}],
                            "table_caption": [[{"type": "text", "content": "Electrical Characteristics"}]],
                            "table_footnote": [
                                {"type": "text", "content": "Note:"},
                                {"type": "equation_inline", "content": r"V_{OS}"},
                            ],
                            "bbox": [72, 80, 540, 260],
                        },
                        {
                            "type": "image",
                            "content": [{"img_path": "images/chart.png"}],
                            "caption": [{"type": "text", "content": "Figure 1. Noise"}],
                            "bbox": [72, 300, 360, 480],
                        },
                    ],
                ],
                layout={
                    "pdf_info": [
                        {"page_idx": 0, "page_size": [612, 792]},
                        {"page_idx": 1, "page_size": [612, 792]},
                    ]
                },
            )

            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_nested",
                source_file_id="file_nested",
                document_language="en",
            )

        document = result.document
        self.assertEqual([page.index for page in document.pages], [0, 1])
        self.assertEqual([block.page_index for block in document.blocks], [0, 0, 1, 1])
        self.assertEqual(document.blocks[0].type, "list")
        self.assertEqual(
            document.blocks[0].list_items,
            ["1 Features........................ 1", "2 Specifications.................. 2", "3 Pins............................ 3"],
        )
        self.assertEqual(document.blocks[1].type, "formula")
        self.assertEqual(document.blocks[1].latex, "V_{OS}")
        self.assertEqual(document.blocks[2].type, "table")
        self.assertEqual(document.blocks[2].html, "<table><tr><td>VOS</td></tr></table>")
        self.assertEqual(document.blocks[2].captions, ["Electrical Characteristics"])
        self.assertEqual(document.blocks[2].footnotes, ["Note: $V_{OS}$"])
        self.assertEqual(document.blocks[2].layout.column_hint, "wide")
        self.assertEqual(document.blocks[3].type, "image")
        self.assertEqual(document.blocks[3].captions, ["Figure 1. Noise"])
        self.assertEqual(result.image_assets[0].source_path, "images/chart.png")

    def test_convert_zip_records_missing_layout_source_and_bbox(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            zip_path = Path(temp_dir) / "mineru.zip"
            self._write_zip(zip_path, blocks=[{"type": "text", "text": "Body", "page_idx": 0}], layout=None)

            result = MinerUToDocumentIRService().convert_zip(
                zip_path,
                job_id="docjob_missing_layout",
                source_file_id="file_missing_layout",
                document_language="en",
            )

        self.assertEqual(result.document.schema_version, "document-ir/v2")
        self.assertEqual(result.document.pages[0].index, 0)
        self.assertIn("middle_json_missing", result.document.quality_flags)
        self.assertIn("layout_bbox_missing:block_0001", result.document.quality_flags)

    def _write_zip(
        self,
        zip_path: Path,
        *,
        blocks: list[Any],
        v2_blocks: list[Any] | None = None,
        layout: dict | None,
        images: dict[str, bytes] | None = None,
    ) -> None:
        with zipfile.ZipFile(zip_path, "w") as archive:
            archive.writestr("content_list.json", json.dumps(blocks, ensure_ascii=False))
            if v2_blocks is not None:
                archive.writestr("nested/result_content_list_v2.json", json.dumps(v2_blocks, ensure_ascii=False))
            if layout is not None:
                archive.writestr("layout.json", json.dumps(layout, ensure_ascii=False))
            for image_path, content in (images or {"images/chart.png": MINIMAL_PNG_BYTES}).items():
                archive.writestr(image_path, content)


if __name__ == "__main__":
    unittest.main()

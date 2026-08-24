from __future__ import annotations

# tests/test_document_ir_word_renderer_layout.py: 覆盖 DocumentIRWordRenderer 的 v2 layout plan 路径。

import hashlib
import io
import re
import tempfile
import unittest
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from xml.etree import ElementTree

from docx import Document
from docx.enum.section import WD_ORIENTATION, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu

from app.conversion.docx_rules.text_semantics import (
    replace_simple_table_math_identifiers,
    simple_table_math_identifier_html,
)
from app.conversion.runtime.table_html_parser import parse_html_table
from app.services.document_ir.models import (
    DocumentIR,
    DocumentIRAsset,
    DocumentIRBlock,
    DocumentIRBlockLayout,
    DocumentIRMetadata,
    DocumentIRPage,
)
from app.services.document_ir.integrity import DocumentIntegrityError
from app.services.renderer.evidence import build_renderer_evidence, ensure_renderer_integrity
from app.services.renderer.docx_quality import analyze_docx_quality_bytes
from app.services.renderer.layout_plan import build_layout_plan
from app.services.renderer.table_layout import (
    _unbreakable_units,
    choose_table_fit_decision,
    insert_engineering_token_breaks,
)
from app.services.renderer.template_config import TemplateConfigPackage
from app.services.renderer.word_renderer import DocumentIRRenderOptions, DocumentIRWordRenderer, resolve_template_config


MINIMAL_PNG_BYTES = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
    b"\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
    b"\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01"
    b"\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)


class DocumentIRWordRendererLayoutTests(unittest.TestCase):
    def test_seed_config_declares_western_and_complex_fonts_for_all_text_groups(self) -> None:
        config = TemplateConfigPackage.load(
            Path(__file__).resolve().parents[1] / "resources" / "template_configs" / "datasheet_template.yaml"
        )

        for group_name in (
            "part_number",
            "category",
            "heading1",
            "heading2",
            "heading3",
            "toc_title",
            "toc_level1",
            "toc_level2",
            "toc_level3",
            "normal",
            "table_content",
            "table_note",
            "caption",
            "long_table_caption",
        ):
            style = getattr(config, group_name)
            self.assertEqual(style.western_font_name, "Arial", group_name)
            self.assertEqual(style.complex_script_font_name, "Times New Roman", group_name)
        self.assertEqual(config.header.title_western_font_name, "Arial")
        self.assertEqual(config.header.title_complex_script_font_name, "Times New Roman")
        self.assertEqual(config.footer.western_font_name, "Arial")
        self.assertEqual(config.footer.complex_script_font_name, "Times New Roman")

    def test_build_layout_plan_separates_decorations_and_two_column_regions(self) -> None:
        document = self._make_v2_two_column_document()

        plan = build_layout_plan(document)

        self.assertEqual(plan.mode, "layout_v2")
        self.assertEqual(len(plan.pages), 1)
        self.assertEqual([block.text for block in plan.pages[0].header_blocks], ["Header"])
        self.assertEqual([block.text for block in plan.pages[0].footer_blocks], ["Footer"])
        self.assertEqual([region.kind for region in plan.pages[0].regions], ["two_col", "one_col"])

    def test_build_layout_plan_prefers_column_hint_for_stable_non_overlapping_regions(self) -> None:
        document = self._make_v2_stable_hint_document()

        plan = build_layout_plan(document)

        self.assertEqual(plan.mode, "layout_v2")
        self.assertEqual([region.kind for region in plan.pages[0].regions], ["one_col", "two_col", "one_col"])
        self.assertEqual([block.text for block in plan.pages[0].regions[0].blocks], ["Left A"])
        self.assertEqual([block.text for block in plan.pages[0].regions[1].blocks], ["Left B", "Right A"])
        self.assertEqual(plan.pages[0].regions[1].column_break_index, 1)
        self.assertEqual([block.text for block in plan.pages[0].regions[2].blocks], ["Right B", "Wide after hints"])

    def test_build_layout_plan_keeps_explicit_wide_blocks_out_of_stable_two_column_regions(self) -> None:
        document = self._make_v2_wide_between_columns_document()

        plan = build_layout_plan(document)

        self.assertEqual(plan.mode, "layout_v2")
        self.assertEqual([region.kind for region in plan.pages[0].regions], ["two_col", "one_col", "two_col"])
        self.assertEqual([block.text for block in plan.pages[0].regions[0].blocks], ["Left A", "Right A"])
        self.assertEqual([block.text for block in plan.pages[0].regions[1].blocks], ["Wide callout"])
        self.assertEqual([block.text for block in plan.pages[0].regions[2].blocks], ["Left B", "Right B"])

    def test_build_layout_plan_keeps_single_column_gap_between_real_two_column_bands(self) -> None:
        def block(block_id: str, text: str, bbox: tuple[float, float, float, float], hint: str) -> DocumentIRBlock:
            return DocumentIRBlock(
                id=block_id,
                type="paragraph",
                order=len(text),
                text=text,
                layout=DocumentIRBlockLayout(page_index=0, bbox=bbox, role="body", column_hint=hint),
            )

        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_local_bands", source_file_id="file_local_bands", document_language="en"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                block("upper_left", "Upper left", (0, 60, 170, 120), "left"),
                block("upper_right", "Upper right", (220, 65, 390, 125), "right"),
                block("middle_left", "Middle only", (0, 170, 170, 230), "left"),
                block("lower_left", "Lower left", (0, 310, 170, 400), "left"),
                block("lower_right", "Lower right", (220, 315, 390, 400), "right"),
            ],
        )

        plan = build_layout_plan(document)

        self.assertEqual([region.kind for region in plan.pages[0].regions], ["two_col", "one_col", "two_col"])
        self.assertEqual([block.text for block in plan.pages[0].regions[1].blocks], ["Middle only"])

    def test_renderer_evidence_summarizes_layout_and_docx_quality(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_evidence", source_file_id="file_evidence", document_language="en"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                DocumentIRBlock(
                    id="block_0001",
                    type="paragraph",
                    order=0,
                    page_index=0,
                    text="Formal renderer evidence should remain editable in Word output.",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(72, 80, 320, 140), role="body"),
                )
            ],
            quality_flags=["layout_sample:block_0001"],
        )
        layout_plan = build_layout_plan(document)
        word_content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )

        evidence = build_renderer_evidence(document, word_content, layout_plan)

        self.assertEqual(evidence["schema_version"], "document-ir/v2")
        self.assertEqual(evidence["renderer_mode"], "layout_v2")
        self.assertEqual(evidence["layout_page_count"], 1)
        self.assertEqual(evidence["layout_block_count"], 1)
        self.assertEqual(evidence["expected_layout_block_count"], 1)
        self.assertEqual(evidence["layout_block_count_delta"], 0)
        self.assertEqual(evidence["layout_quality_flags"], ["layout_sample:block_0001"])
        self.assertTrue(evidence["docx_quality"]["editable_content_available"])
        self.assertFalse(evidence["docx_quality"]["image_only_suspected"])
        self.assertEqual(evidence["quality_flags"], [])
        self.assertEqual(evidence["blocking_quality_flags"], [])

    def test_renderer_evidence_blocks_missing_layout_drawing_and_table(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_blocking", source_file_id="file_blocking", document_language="zh"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            assets=[DocumentIRAsset(id="asset_chart", kind="image", source_path="chart.png")],
            blocks=[
                DocumentIRBlock(id="missing_page", type="paragraph", order=0, text="不应静默丢失的正文"),
                DocumentIRBlock(
                    id="missing_drawing",
                    type="image",
                    order=1,
                    asset_ref="asset_chart",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(10, 80, 180, 180), role="body"),
                ),
                DocumentIRBlock(
                    id="missing_table",
                    type="table",
                    order=2,
                    html="not-a-table",
                    text="表格回退文本",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(10, 200, 390, 300), role="body"),
                ),
            ],
        )
        word_content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )

        evidence = build_renderer_evidence(document, word_content)

        self.assertEqual(evidence["expected_layout_block_count"], 3)
        self.assertEqual(evidence["layout_block_count"], 2)
        self.assertEqual(evidence["expected_inline_shape_count"], 1)
        self.assertEqual(evidence["inline_shape_count"], 0)
        self.assertEqual(evidence["expected_table_count"], 1)
        self.assertEqual(evidence["table_count"], 0)
        self.assertIn("ir_layout_integrity:missing_blocks", evidence["blocking_quality_flags"])
        self.assertIn("ir_docx_integrity:inline_shape_count_mismatch", evidence["blocking_quality_flags"])
        self.assertIn("ir_docx_integrity:table_count_mismatch", evidence["blocking_quality_flags"])
        with self.assertRaises(DocumentIntegrityError) as raised:
            ensure_renderer_integrity(evidence)
        self.assertIn("expected_inline_shape_count=1", str(raised.exception))
        self.assertIn("inline_shape_count=0", str(raised.exception))

    def test_renderer_evidence_proves_visible_payloads_and_image_asset_identity(self) -> None:
        checksum = f"sha256:{hashlib.sha256(MINIMAL_PNG_BYTES).hexdigest()}"
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_payload", source_file_id="file_payload", document_language="en"),
            assets=[
                DocumentIRAsset(
                    id="asset_chart",
                    kind="image",
                    source_path="chart.png",
                    checksum_sha256=checksum,
                )
            ],
            blocks=[
                DocumentIRBlock(id="heading", type="heading", order=0, level=1, text="Electrical data"),
                DocumentIRBlock(id="body", type="paragraph", order=1, text="Visible body payload"),
                DocumentIRBlock(id="list", type="list", order=2, list_items=["• First item", "• Second item"]),
                DocumentIRBlock(
                    id="table",
                    type="table",
                    order=3,
                    html=r"<table><tr><th>Parameter</th></tr><tr><td>\frac{1}{2}</td></tr></table>",
                    captions=["Table 1 Limits"],
                    footnotes=["Note: tested at 25°C"],
                ),
                DocumentIRBlock(id="formula", type="formula", order=4, latex=r"V_{OUT}=1"),
                DocumentIRBlock(
                    id="chart",
                    type="image",
                    order=5,
                    asset_ref="asset_chart",
                    captions=["Figure 1 Response"],
                    footnotes=["Typical condition"],
                ),
            ],
        )
        config = TemplateConfigPackage.load(None)
        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: MINIMAL_PNG_BYTES,
            template_config=config,
        )

        evidence = build_renderer_evidence(document, content, template_config=config)

        self.assertEqual(evidence["blocking_quality_flags"], [])
        self.assertEqual(evidence["visible_payload_evidence"]["missing_payload_ids"], [])
        self.assertEqual(evidence["visible_payload_evidence"]["counts"]["table_cell"], {"expected": 2, "covered": 2})
        self.assertEqual(evidence["visible_payload_evidence"]["counts"]["annotation"], {"expected": 4, "covered": 4})
        self.assertEqual(evidence["image_evidence"]["expected_asset_ids"], ["asset_chart"])
        self.assertEqual(evidence["image_evidence"]["rendered_assets"][0]["checksum_sha256"], checksum)

    def test_renderer_evidence_blocks_missing_visible_payload_and_asset_identity_mismatch(self) -> None:
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_tamper", source_file_id="file_tamper", document_language="en"),
            assets=[DocumentIRAsset(id="asset_chart", kind="image", source_path="chart.png")],
            blocks=[
                DocumentIRBlock(id="body", type="paragraph", order=0, text="Critical visible payload"),
                DocumentIRBlock(id="chart", type="image", order=1, asset_ref="asset_chart"),
            ],
        )
        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: MINIMAL_PNG_BYTES,
        )
        tampered = self._replace_document_xml(
            content,
            {
                "Critical visible payload": "Removed payload",
                "md2word:asset:asset_chart": "md2word:asset:wrong_asset",
            },
        )

        evidence = build_renderer_evidence(document, tampered)

        self.assertIn("body", evidence["visible_payload_evidence"]["missing_payload_ids"])
        self.assertIn("ir_docx_integrity:visible_payload_missing", evidence["blocking_quality_flags"])
        self.assertIn("ir_docx_integrity:image_asset_identity_mismatch", evidence["blocking_quality_flags"])

    def test_renderer_evidence_blocks_duplicate_or_wrong_crop_fragments(self) -> None:
        normalization = {
            "reason": "editable_table_extracted_from_containing_image",
            "table_block_id": "dimensions",
            "image_block_id": "drawing",
            "masked_rect": [0.3, 0.2, 0.7, 0.8],
            "crop_fragment_count": 4,
        }
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_crop", source_file_id="file_crop", document_language="en"),
            pages=[DocumentIRPage(index=0, width=800.0, height=800.0)],
            assets=[DocumentIRAsset(id="asset_drawing", kind="image", source_path="drawing.png")],
            blocks=[
                DocumentIRBlock(
                    id="dimensions",
                    type="table",
                    order=0,
                    html="<table><tr><td>DIM</td><td>20</td></tr></table>",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(300, 200, 700, 800), role="body"),
                    metadata={
                        "visual_overlap_group_id": "visual_overlap:dimensions:drawing",
                        "visual_overlap_normalization": normalization,
                    },
                ),
                DocumentIRBlock(
                    id="drawing",
                    type="image",
                    order=1,
                    asset_ref="asset_drawing",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 0, 800, 800), role="body"),
                    metadata={
                        "visual_overlap_group_id": "visual_overlap:dimensions:drawing",
                        "visual_overlap_normalization": normalization,
                        "visual_crop_rects": [
                            [0.0, 0.0, 0.3, 1.0],
                            [0.3, 0.0, 0.7, 0.2],
                            [0.7, 0.0, 1.0, 1.0],
                            [0.3, 0.8, 0.7, 1.0],
                        ]
                    },
                )
            ],
        )
        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: MINIMAL_PNG_BYTES,
        )
        evidence = build_renderer_evidence(document, content)
        self.assertEqual(evidence["blocking_quality_flags"], [])
        self.assertEqual(
            evidence["visual_table_composite_evidence"]["rendered_group_ids"],
            ["visual_overlap:dimensions:drawing"],
        )

        tampered = self._replace_document_xml(
            content,
            {
                '<a:srcRect l="0" t="0" r="70000" b="0"/>': (
                    '<a:srcRect l="30000" t="0" r="30000" b="80000"/>'
                )
            },
        )
        tampered_evidence = build_renderer_evidence(document, tampered)

        self.assertEqual(tampered_evidence["image_evidence"]["missing_asset_ids"], [])
        self.assertEqual(tampered_evidence["image_evidence"]["unexpected_asset_ids"], [])
        self.assertEqual(len(tampered_evidence["image_evidence"]["missing_fragments"]), 1)
        self.assertEqual(len(tampered_evidence["image_evidence"]["unexpected_fragments"]), 1)
        self.assertIn(
            "ir_docx_integrity:image_fragment_identity_mismatch",
            tampered_evidence["blocking_quality_flags"],
        )

        reordered = self._replace_document_xml(
            content,
            {
                '<a:srcRect l="0" t="0" r="70000" b="0"/>': (
                    '<a:srcRect l="99999" t="99999" r="99999" b="99999"/>'
                ),
                '<a:srcRect l="70000" t="0" r="0" b="0"/>': (
                    '<a:srcRect l="0" t="0" r="70000" b="0"/>'
                ),
                '<a:srcRect l="99999" t="99999" r="99999" b="99999"/>': (
                    '<a:srcRect l="70000" t="0" r="0" b="0"/>'
                ),
            },
        )
        reordered_evidence = build_renderer_evidence(document, reordered)

        self.assertEqual(reordered_evidence["image_evidence"]["missing_fragments"], [])
        self.assertEqual(reordered_evidence["image_evidence"]["unexpected_fragments"], [])
        self.assertEqual(
            reordered_evidence["image_evidence"]["order_mismatch_asset_ids"],
            ["asset_drawing"],
        )
        self.assertIn(
            "ir_docx_integrity:image_fragment_identity_mismatch",
            reordered_evidence["blocking_quality_flags"],
        )

        wrong_group = self._replace_document_xml(
            content,
            {
                "md2word:visual_table_composite:visual_overlap:dimensions:drawing": (
                    "md2word:visual_table_composite:visual_overlap:wrong:group"
                )
            },
        )
        wrong_group_evidence = build_renderer_evidence(document, wrong_group)

        self.assertEqual(
            wrong_group_evidence["visual_table_composite_evidence"]["missing_group_ids"],
            ["visual_overlap:dimensions:drawing"],
        )
        self.assertEqual(
            wrong_group_evidence["visual_table_composite_evidence"]["unexpected_group_ids"],
            ["visual_overlap:wrong:group"],
        )
        self.assertIn(
            "ir_docx_integrity:visual_table_composite_identity_mismatch",
            wrong_group_evidence["blocking_quality_flags"],
        )

    def test_visual_table_composite_keeps_portrait_mask_for_semantically_breakable_tokens(self) -> None:
        mask = [0.3, 0.2, 0.7, 0.8]
        normalization = {
            "reason": "editable_table_extracted_from_containing_image",
            "table_block_id": "dimensions",
            "image_block_id": "drawing",
            "masked_rect": mask,
            "crop_fragment_count": 4,
        }
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(
                job_id="docjob_token_composite",
                source_file_id="file_token_composite",
                document_language="en",
            ),
            pages=[DocumentIRPage(index=0, width=800.0, height=800.0)],
            assets=[DocumentIRAsset(id="asset_drawing", kind="image", source_path="drawing.png")],
            blocks=[
                DocumentIRBlock(
                    id="dimensions",
                    type="table",
                    order=0,
                    html="<table><tr><td>MS-001</td><td>0.775(19,69)</td></tr></table>",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(240, 160, 560, 640), role="body"),
                    metadata={
                        "visual_overlap_group_id": "visual_overlap:dimensions:drawing",
                        "visual_overlap_normalization": normalization,
                    },
                ),
                DocumentIRBlock(
                    id="drawing",
                    type="image",
                    order=1,
                    asset_ref="asset_drawing",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(80, 0, 720, 800), role="body"),
                    metadata={
                        "visual_overlap_group_id": "visual_overlap:dimensions:drawing",
                        "visual_overlap_normalization": normalization,
                        "visual_crop_rects": [
                            [0.0, 0.0, 0.3, 1.0],
                            [0.3, 0.0, 0.7, 0.2],
                            [0.7, 0.0, 1.0, 1.0],
                            [0.3, 0.8, 0.7, 1.0],
                        ],
                    },
                ),
            ],
        )
        config = TemplateConfigPackage.load(
            Path(__file__).resolve().parents[1] / "resources" / "template_configs" / "datasheet_template.yaml"
        )

        plan = build_layout_plan(document, config)
        decision = plan.table_fit_decisions["dimensions"]
        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: MINIMAL_PNG_BYTES,
            layout_plan=plan,
            template_config=config,
        )
        rendered = Document(io.BytesIO(content))
        outer_table = rendered.tables[0]
        outer_widths = [int(column.width) for column in outer_table.columns]
        inner_table = outer_table.cell(1, 1).tables[0]
        inner_width = sum(int(column.width) for column in inner_table.columns)
        content_width = int(
            rendered.sections[0].page_width
            - rendered.sections[0].left_margin
            - rendered.sections[0].right_margin
        )

        self.assertEqual(plan.pages[0].regions[0].kind, "visual_table_composite")
        self.assertEqual(decision.container, "visual_table_composite")
        self.assertTrue(decision.fits)
        self.assertGreaterEqual(decision.selected_font_size, decision.minimum_font_size)
        self.assertAlmostEqual(sum(outer_widths), content_width, delta=635)
        self.assertAlmostEqual(outer_widths[1], decision.available_width, delta=635)
        self.assertLessEqual(inner_width, outer_widths[1] + 635)
        self.assertEqual(sum(decision.column_widths), decision.available_width)
        self.assertEqual(rendered.sections[0].orientation, WD_ORIENTATION.PORTRAIT)
        self.assertEqual(build_renderer_evidence(document, content, plan, config)["blocking_quality_flags"], [])

    def test_renderer_evidence_matches_reordered_prefix_payloads_without_false_block(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(
                job_id="docjob_prefix_payload",
                source_file_id="file_prefix_payload",
                document_language="en",
            ),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                DocumentIRBlock(
                    id="right_short",
                    type="paragraph",
                    order=0,
                    text="A",
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(230, 80, 380, 100),
                        role="body",
                        column_hint="right",
                    ),
                ),
                DocumentIRBlock(
                    id="left_long",
                    type="paragraph",
                    order=1,
                    text="AB",
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(20, 80, 170, 100),
                        role="body",
                        column_hint="left",
                    ),
                ),
            ],
        )
        config = TemplateConfigPackage.load(None)
        plan = build_layout_plan(document, config)

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
            layout_plan=plan,
            template_config=config,
        )
        evidence = build_renderer_evidence(document, content, plan, config)

        self.assertEqual([block.id for block in plan.pages[0].regions[0].blocks], ["left_long", "right_short"])
        self.assertEqual(evidence["visible_payload_evidence"]["missing_payload_ids"], [])
        self.assertNotIn("ir_docx_integrity:visible_payload_missing", evidence["blocking_quality_flags"])

    def test_renderer_evidence_matches_renderer_visible_text_normalization(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(
                job_id="docjob_visible_normalization",
                source_file_id="file_visible_normalization",
                document_language="en",
            ),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                DocumentIRBlock(id="bullet", type="paragraph", order=0, text="• Native list item"),
                DocumentIRBlock(
                    id="overline",
                    type="paragraph",
                    order=1,
                    text=r"Output $\overline{OE}$ remains editable",
                ),
                DocumentIRBlock(
                    id="table",
                    type="table",
                    order=2,
                    html=r"<table><tr><td>2000 | LARGE T&R</td><td>$\overline{OE}$</td></tr></table>",
                    captions=[r"Table 8 Register $\overline{SAMPLE}$"],
                ),
            ],
        )
        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )

        evidence = build_renderer_evidence(document, content)

        self.assertEqual(evidence["visible_payload_evidence"]["missing_payload_ids"], [])
        self.assertNotIn("ir_docx_integrity:visible_payload_missing", evidence["blocking_quality_flags"])
        self.assertIn("<m:bar>", self._document_xml(content))

    def test_render_to_bytes_uses_native_columns_for_v2_region(self) -> None:
        renderer = DocumentIRWordRenderer()

        content = renderer.render_to_bytes(
            self._make_v2_two_column_document(),
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        rendered = Document(io.BytesIO(content))
        xml = self._document_xml(content)

        self.assertEqual(len(rendered.tables), 0)
        self.assertTrue(any(paragraph.text == "Left column" for paragraph in rendered.paragraphs))
        self.assertTrue(any(paragraph.text == "Right column" for paragraph in rendered.paragraphs))
        self.assertTrue(any(paragraph.text == "Wide block" for paragraph in rendered.paragraphs))
        self.assertIn('<w:cols w:space="720" w:num="2"', xml)
        self.assertIn('<w:br w:type="column"', xml)
        self.assertGreaterEqual(xml.count('<w:type w:val="continuous"'), 2)
        self.assertEqual(rendered.sections[0].page_width, rendered.sections[1].page_width)
        self.assertEqual(rendered.sections[0].left_margin, rendered.sections[1].left_margin)
        self.assertTrue(rendered.sections[1].header.is_linked_to_previous)

    def test_render_to_bytes_prefers_column_hint_when_ordering_native_columns(self) -> None:
        renderer = DocumentIRWordRenderer()

        content = renderer.render_to_bytes(
            self._make_v2_hint_override_document(),
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        rendered = Document(io.BytesIO(content))

        self.assertEqual(len(rendered.tables), 0)
        visible_paragraphs = [paragraph.text for paragraph in rendered.paragraphs if paragraph.text]
        self.assertEqual(visible_paragraphs, ["Left by hint", "Right by hint"])
        self.assertIn('<w:br w:type="column"', self._document_xml(content))
        self.assertEqual(rendered.sections[-1].start_type, WD_SECTION_START.CONTINUOUS)
        self.assertNotIn('<w:br w:type="page"', self._document_xml(content))

    def test_render_to_bytes_keeps_paired_image_caption_fields_in_row_order(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_image_rows", source_file_id="file_image_rows", document_language="en"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            assets=[DocumentIRAsset(id="asset_chart", kind="image", source_path="images/chart.png")],
            blocks=[
                DocumentIRBlock(
                    id=f"block_{index}",
                    type="image",
                    order=index,
                    asset_ref="asset_chart",
                    captions=[f"Figure {index}: {tail}"],
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=bbox,
                        role="body",
                        column_hint="left" if index % 2 else "right",
                    ),
                )
                for index, tail, bbox in (
                    (1, "Left first", (10, 80, 170, 180)),
                    (2, "Right first", (220, 82, 380, 180)),
                    (3, "Left second", (10, 220, 170, 320)),
                    (4, "Right second", (220, 218, 380, 320)),
                )
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: MINIMAL_PNG_BYTES,
        )
        xml = self._document_xml(content)
        rendered = Document(io.BytesIO(content))

        self.assertLess(xml.index("Left first"), xml.index("Right first"))
        self.assertLess(xml.index("Right first"), xml.index("Left second"))
        self.assertLess(xml.index("Left second"), xml.index("Right second"))
        self.assertNotIn('<w:br w:type="column"', xml)
        self.assertEqual(len(rendered.sections), 1)
        self.assertEqual(len(rendered.tables), 2)
        self.assertIn("Left first", rendered.tables[0].cell(0, 0).text)
        self.assertIn("Right first", rendered.tables[0].cell(0, 1).text)
        self.assertIn("Left second", rendered.tables[1].cell(0, 0).text)
        self.assertIn("Right second", rendered.tables[1].cell(0, 1).text)
        self.assertEqual(xml.count('w:val="md2word:paired_visual_row"'), 2)
        self.assertNotIn("<w:tabs>", xml)
        for source_number in range(1, 5):
            self.assertIn(f" SEQ Figure \\r {source_number} \\* ARABIC ", xml)
        quality = analyze_docx_quality_bytes(content)
        self.assertEqual(quality["table_count"], 0)
        self.assertEqual(quality["layout_table_count"], 2)

    def test_render_to_bytes_keeps_multi_annotation_images_in_native_columns(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_image_notes", source_file_id="file_image_notes", document_language="en"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            assets=[DocumentIRAsset(id="asset_chart", kind="image", source_path="images/chart.png")],
            blocks=[
                DocumentIRBlock(
                    id=f"block_{index}",
                    type="image",
                    order=index,
                    asset_ref="asset_chart",
                    captions=[f"Figure {index}: chart"],
                    footnotes=["Source note"] if index == 1 else [],
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=bbox,
                        role="body",
                        column_hint="left" if index == 1 else "right",
                    ),
                )
                for index, bbox in ((1, (10, 80, 170, 180)), (2, (220, 82, 380, 180)))
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: MINIMAL_PNG_BYTES,
        )
        xml = self._document_xml(content)

        self.assertIn('<w:cols w:space="720" w:num="2"', xml)
        self.assertIn('<w:br w:type="column"', xml)
        self.assertIn("Source note", xml)
        self.assertNotIn("<w:tabs>", xml)

    def test_render_to_bytes_does_not_force_break_between_same_orientation_source_pages(self) -> None:
        document = self._make_v2_hint_override_document()
        document.pages.append(DocumentIRPage(index=1, width=400.0, height=600.0))
        document.blocks.append(
            DocumentIRBlock(
                id="block_next_page",
                type="heading",
                order=2,
                level=2,
                text="Next source page",
                layout=DocumentIRBlockLayout(page_index=1, bbox=(0, 60, 400, 100), role="body", column_hint="wide"),
            )
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        rendered = Document(io.BytesIO(content))
        xml = self._document_xml(content)

        self.assertEqual(rendered.sections[-1].start_type, WD_SECTION_START.CONTINUOUS)
        self.assertNotIn('<w:br w:type="page"', xml)
        self.assertIn("<w:pageBreakBefore/>", xml)
        self.assertTrue(any(paragraph.text == "Next source page" for paragraph in rendered.paragraphs))

    def test_render_to_bytes_anchors_semantic_page_breaks_to_real_content(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_semantic_pages", source_file_id="file_semantic_pages", document_language="zh"),
            pages=[DocumentIRPage(index=index, width=400.0, height=600.0) for index in range(5)],
            blocks=[
                DocumentIRBlock(
                    id="intro",
                    type="paragraph",
                    order=0,
                    text="Intro",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 60, 400, 100), role="body", column_hint="wide"),
                ),
                DocumentIRBlock(
                    id="working_principle",
                    type="heading",
                    order=1,
                    level=2,
                    text="工作原理",
                    layout=DocumentIRBlockLayout(page_index=1, bbox=(0, 60, 400, 100), role="body", column_hint="wide"),
                ),
                DocumentIRBlock(
                    id="captioned_table",
                    type="table",
                    order=2,
                    captions=["绝对最大额定值"],
                    html="<table><tr><th>参数</th></tr><tr><td>值</td></tr></table>",
                    layout=DocumentIRBlockLayout(page_index=2, bbox=(0, 80, 400, 180), role="body", column_hint="wide"),
                ),
                DocumentIRBlock(
                    id="table_context",
                    type="paragraph",
                    order=3,
                    text="默认测试条件：VS = ±18V、VCM = 0V、TA = 25°C。",
                    metadata={"annotation_role": "table_context", "source_block_id": "context_table"},
                    layout=DocumentIRBlockLayout(page_index=3, bbox=(0, 60, 400, 80), role="body", column_hint="wide"),
                ),
                DocumentIRBlock(
                    id="context_table",
                    type="table",
                    order=4,
                    html="<table><tr><th>参数</th></tr><tr><td>值</td></tr></table>",
                    layout=DocumentIRBlockLayout(page_index=3, bbox=(0, 90, 400, 190), role="body", column_hint="wide"),
                ),
                DocumentIRBlock(
                    id="flowing_page",
                    type="paragraph",
                    order=5,
                    text="Natural continuation",
                    layout=DocumentIRBlockLayout(page_index=4, bbox=(0, 60, 400, 100), role="body", column_hint="wide"),
                ),
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        rendered = Document(io.BytesIO(content))
        xml = self._document_xml(content)
        paragraphs = {paragraph.text: paragraph for paragraph in rendered.paragraphs if paragraph.text}

        self.assertTrue(paragraphs["工作原理"].paragraph_format.page_break_before)
        self.assertTrue(paragraphs["工作原理"].paragraph_format.keep_with_next)
        self.assertTrue(paragraphs["绝对最大额定值"].paragraph_format.page_break_before)
        self.assertTrue(paragraphs["默认测试条件：VS = ±18V、VCM = 0V、TA = 25°C。"].paragraph_format.page_break_before)
        self.assertTrue(paragraphs["Natural continuation"].paragraph_format.page_break_before)
        self.assertEqual(xml.count("<w:pageBreakBefore/>"), 4)
        self.assertNotIn('<w:br w:type="page"', xml)

    def test_soft_page_group_uses_final_left_to_right_anchor_and_explicit_continuation_reason(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_soft_groups", source_file_id="file_soft_groups", document_language="zh"),
            pages=[DocumentIRPage(index=index, width=400.0, height=600.0) for index in range(3)],
            blocks=[
                DocumentIRBlock(
                    id="intro",
                    type="paragraph",
                    order=0,
                    text="前言",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 60, 400, 100), role="body", column_hint="wide"),
                ),
                DocumentIRBlock(
                    id="right_first_by_y",
                    type="heading",
                    order=1,
                    level=2,
                    text="右栏标题",
                    metadata={"content_group_id": "continued_body"},
                    layout=DocumentIRBlockLayout(page_index=1, bbox=(220, 40, 390, 130), role="body", column_hint="right"),
                ),
                DocumentIRBlock(
                    id="left_final_first",
                    type="paragraph",
                    order=2,
                    text="左栏正文",
                    layout=DocumentIRBlockLayout(page_index=1, bbox=(10, 80, 180, 150), role="body", column_hint="left"),
                ),
                DocumentIRBlock(
                    id="continued_body",
                    type="paragraph",
                    order=3,
                    text="跨源页续接正文",
                    metadata={"content_group_id": "continued_body"},
                    layout=DocumentIRBlockLayout(page_index=2, bbox=(0, 50, 400, 100), role="body", column_hint="wide"),
                ),
            ],
        )

        plan = build_layout_plan(document)

        self.assertEqual(len(plan.pages), 2)
        self.assertEqual(plan.pages[1].source_page_indices, (1, 2))
        self.assertEqual(plan.pages[1].break_before_block_id, "left_final_first")
        self.assertEqual(plan.pages[1].break_reason, "source_page_boundary")
        self.assertEqual(plan.pages[1].continuation_reasons, ("content_group:continued_body",))

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
            layout_plan=plan,
        )
        rendered = Document(io.BytesIO(content))
        paragraphs = {paragraph.text: paragraph for paragraph in rendered.paragraphs if paragraph.text}
        self.assertTrue(paragraphs["左栏正文"].paragraph_format.page_break_before)
        self.assertIsNone(paragraphs["右栏标题"].paragraph_format.page_break_before)
        self.assertIsNone(paragraphs["跨源页续接正文"].paragraph_format.page_break_before)
        evidence = build_renderer_evidence(document, content, plan)
        self.assertEqual(evidence["soft_page_groups"][1]["source_page_indices"], [1, 2])
        self.assertEqual(evidence["soft_page_groups"][1]["break_before_block_id"], "left_final_first")
        self.assertEqual(
            evidence["soft_page_groups"][1]["continuation_reasons"],
            ["content_group:continued_body"],
        )

    def test_soft_page_group_continues_after_promoted_page_edge_footer(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(
                job_id="docjob_promoted_footer",
                source_file_id="file_promoted_footer",
                document_language="en",
            ),
            pages=[DocumentIRPage(index=index, width=400.0, height=600.0) for index in range(2)],
            blocks=[
                DocumentIRBlock(
                    id="body",
                    type="paragraph",
                    order=0,
                    text="Page body",
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(20, 80, 380, 120),
                        role="body",
                        column_hint="wide",
                    ),
                ),
                DocumentIRBlock(
                    id="promoted_footer",
                    type="paragraph",
                    order=1,
                    text="Important source notice",
                    metadata={
                        "source_type": "page_footer",
                        "semantic_decision_reason": "decoration_candidate_promoted_to_body",
                    },
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(20, 500, 380, 530),
                        role="body",
                        column_hint="wide",
                    ),
                ),
                DocumentIRBlock(
                    id="next_page",
                    type="heading",
                    order=2,
                    level=1,
                    text="Next source group",
                    layout=DocumentIRBlockLayout(
                        page_index=1,
                        bbox=(20, 60, 380, 90),
                        role="body",
                        column_hint="wide",
                    ),
                ),
            ],
        )

        plan = build_layout_plan(document)

        self.assertEqual(len(plan.pages), 1)
        self.assertEqual(plan.pages[0].source_page_indices, (0, 1))
        self.assertEqual(
            plan.pages[0].continuation_reasons,
            ("promoted_page_edge_decoration:promoted_footer",),
        )
        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
            layout_plan=plan,
        )
        self.assertNotIn("<w:pageBreakBefore/>", self._document_xml(content))

    def test_short_two_column_visual_prelude_degrades_to_kept_single_column_flow(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(
                job_id="docjob_visual_prelude",
                source_file_id="file_visual_prelude",
                document_language="en",
            ),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            assets=[DocumentIRAsset(id="asset_subject", kind="image", source_path="subject.png")],
            blocks=[
                DocumentIRBlock(
                    id="left_label",
                    type="paragraph",
                    order=0,
                    text="Package code",
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(10, 40, 150, 60),
                        role="body",
                        column_hint="left",
                    ),
                ),
                DocumentIRBlock(
                    id="right_title",
                    type="paragraph",
                    order=1,
                    text="Package outline",
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(250, 35, 390, 55),
                        role="body",
                        column_hint="right",
                    ),
                ),
                DocumentIRBlock(
                    id="right_label",
                    type="paragraph",
                    order=2,
                    text="Package family",
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(250, 60, 390, 80),
                        role="body",
                        column_hint="right",
                    ),
                ),
                DocumentIRBlock(
                    id="subject",
                    type="image",
                    order=3,
                    asset_ref="asset_subject",
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(10, 90, 390, 390),
                        role="body",
                        column_hint="wide",
                    ),
                ),
            ],
        )

        plan = build_layout_plan(document)

        self.assertEqual([region.kind for region in plan.pages[0].regions], ["visual_prelude"])
        self.assertEqual(
            [block.id for block in plan.pages[0].regions[0].blocks],
            ["right_title", "left_label", "right_label", "subject"],
        )
        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: MINIMAL_PNG_BYTES,
            layout_plan=plan,
        )
        rendered = Document(io.BytesIO(content))
        paragraphs = {paragraph.text: paragraph for paragraph in rendered.paragraphs if paragraph.text}
        for label in ("Package outline", "Package code", "Package family"):
            self.assertTrue(paragraphs[label].paragraph_format.keep_with_next)
        xml = self._document_xml(content)
        self.assertNotRegex(xml, r'<w:cols[^>]*w:num="2"')
        self.assertLess(xml.index("Package family"), xml.index("<w:drawing>"))
        evidence = build_renderer_evidence(document, content, plan)
        self.assertEqual(
            evidence["column_downgrades"],
            [
                {
                    "page_index": 0,
                    "region_kind": "visual_prelude",
                    "reason": "short_text_prelude_before_wide_visual",
                    "block_ids": ["right_title", "left_label", "right_label", "subject"],
                }
            ],
        )

    def test_renderer_evidence_records_explicit_missing_page_width_column_downgrade(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_downgrade", source_file_id="file_downgrade", document_language="en"),
            pages=[DocumentIRPage(index=0, height=600.0)],
            blocks=[
                DocumentIRBlock(
                    id="body",
                    type="paragraph",
                    order=0,
                    text="Safe single-column fallback",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(10, 80, 180, 120), role="body"),
                )
            ],
        )
        plan = build_layout_plan(document)
        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
            layout_plan=plan,
        )

        evidence = build_renderer_evidence(document, content, plan)

        self.assertEqual(
            evidence["column_downgrades"],
            [{"page_index": 0, "region_kind": "one_col", "reason": "missing_page_width", "block_ids": ["body"]}],
        )
        self.assertEqual(evidence["blocking_quality_flags"], [])

    def test_layout_keeps_tall_right_column_table_in_its_column(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_tall_table", source_file_id="file_tall_table", document_language="zh"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                DocumentIRBlock(
                    id="left_a",
                    type="paragraph",
                    order=0,
                    text="左栏第一段",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(10, 80, 180, 130), role="body", column_hint="left"),
                ),
                DocumentIRBlock(
                    id="right_table",
                    type="table",
                    order=1,
                    html="<table><tr><th>参数</th><th>值</th></tr><tr><td>A</td><td>1</td></tr></table>",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(220, 70, 390, 430), role="body", column_hint="right"),
                ),
                DocumentIRBlock(
                    id="left_b",
                    type="paragraph",
                    order=2,
                    text="左栏第二段",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(10, 150, 180, 200), role="body", column_hint="left"),
                ),
            ],
        )

        plan = build_layout_plan(document)

        self.assertEqual([region.kind for region in plan.pages[0].regions], ["two_col"])
        region = plan.pages[0].regions[0]
        self.assertEqual(region.column_break_index, 2)
        self.assertEqual([block.id for block in region.blocks], ["left_a", "left_b", "right_table"])

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
            layout_plan=plan,
        )
        xml = self._document_xml(content)
        self.assertIn('<w:cols w:space="720" w:num="2"', xml)
        self.assertIn('<w:br w:type="column"', xml)
        self.assertEqual(analyze_docx_quality_bytes(content)["table_count"], 1)

    def test_table_fit_uses_bundled_minimum_to_keep_right_column_membership(self) -> None:
        document = self._make_table_fit_two_column_document()
        config = TemplateConfigPackage.load(
            Path(__file__).resolve().parents[1] / "resources" / "template_configs" / "datasheet_template.yaml"
        )

        plan = build_layout_plan(document, config)

        decision = plan.table_fit_decisions["right_table"]
        self.assertEqual([region.kind for region in plan.pages[0].regions], ["two_col"])
        self.assertEqual(decision.container, "column")
        self.assertEqual(decision.selected_font_size, 8.5)
        self.assertGreaterEqual(decision.selected_font_size, config.table_content.min_font_size)
        self.assertIn("table_font_reduced:right_table:9->8.5", decision.quality_flags)

    def test_table_fit_without_minimum_moves_bound_context_and_table_to_full_width(self) -> None:
        document = self._make_table_fit_two_column_document()
        config = TemplateConfigPackage.load(
            Path(__file__).resolve().parents[1] / "resources" / "template_configs" / "datasheet_template.yaml"
        )
        config = config.model_copy(
            update={"table_content": config.table_content.model_copy(update={"min_font_size": None})}
        )

        plan = build_layout_plan(document, config)

        decision = plan.table_fit_decisions["right_table"]
        self.assertEqual([region.kind for region in plan.pages[0].regions], ["two_col", "one_col"])
        self.assertEqual([block.id for block in plan.pages[0].regions[1].blocks], ["table_context", "right_table"])
        self.assertEqual(decision.container, "full_width")
        self.assertEqual(decision.selected_font_size, 9)
        self.assertNotIn("table_font_reduced:right_table:9->8.5", decision.quality_flags)
        self.assertIn("table_container_degraded:right_table:column->full_width", decision.quality_flags)

    def test_two_column_landscape_fallback_keeps_context_bound_and_records_each_step(self) -> None:
        document = self._make_table_fit_two_column_document()
        long_value = "A" * 50
        document.blocks[2].html = (
            "<table><tr>"
            f"<td>{long_value}</td><td>{long_value}</td>"
            f"<td>{long_value}</td><td>{long_value}</td>"
            "</tr></table>"
        )
        config = TemplateConfigPackage.load(
            Path(__file__).resolve().parents[1] / "resources" / "template_configs" / "datasheet_template.yaml"
        )

        plan = build_layout_plan(document, config)

        decision = plan.table_fit_decisions["right_table"]
        self.assertEqual([region.kind for region in plan.pages[0].regions], ["two_col", "landscape_table"])
        self.assertEqual([block.id for block in plan.pages[0].regions[1].blocks], ["table_context", "right_table"])
        self.assertEqual(decision.reason, "column_and_full_width_unreadable_landscape_unresolved")
        self.assertIn("table_container_degraded:right_table:column->full_width", decision.quality_flags)
        self.assertIn("table_container_degraded:right_table:full_width->landscape", decision.quality_flags)

    def test_landscape_fallback_moves_single_column_context_with_table(self) -> None:
        long_value = "A" * 50
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(
                job_id="docjob_landscape_context",
                source_file_id="file_landscape_context",
                document_language="en",
            ),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                DocumentIRBlock(
                    id="table_context",
                    type="paragraph",
                    order=0,
                    text="Absolute maximum ratings",
                    metadata={"annotation_role": "table_context", "source_block_id": "wide_table"},
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(10, 60, 390, 75),
                        role="body",
                        column_hint="wide",
                    ),
                ),
                DocumentIRBlock(
                    id="wide_table",
                    type="table",
                    order=1,
                    html=(
                        "<table><tr>"
                        f"<td>{long_value}</td><td>{long_value}</td>"
                        f"<td>{long_value}</td><td>{long_value}</td>"
                        "</tr></table>"
                    ),
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(10, 80, 390, 180),
                        role="body",
                        column_hint="wide",
                    ),
                ),
            ],
        )
        config = TemplateConfigPackage.load(
            Path(__file__).resolve().parents[1] / "resources" / "template_configs" / "datasheet_template.yaml"
        )

        plan = build_layout_plan(document, config)

        decision = plan.table_fit_decisions["wide_table"]
        self.assertEqual([region.kind for region in plan.pages[0].regions], ["landscape_table"])
        self.assertEqual([block.id for block in plan.pages[0].regions[0].blocks], ["table_context", "wide_table"])
        self.assertEqual(decision.container, "landscape")
        self.assertEqual(decision.selected_font_size, 7)
        self.assertFalse(decision.fits)
        self.assertEqual(decision.source_width_ratio, 0.95)
        self.assertIn("table_container_degraded:wide_table:full_width->landscape", decision.quality_flags)
        self.assertIn("table_fit_unresolved:wide_table:landscape", decision.quality_flags)

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
            layout_plan=plan,
            template_config=config,
        )
        rendered = Document(io.BytesIO(content))
        self.assertEqual(len(rendered.sections), 1)
        self.assertLess(self._document_text(content).index("Absolute maximum ratings"), self._document_text(content).index(long_value))
        self.assertIn('w:orient="landscape"', self._document_xml(content))

    def test_landscape_fallback_does_not_insert_empty_source_orientation_section(self) -> None:
        long_value = "A" * 50
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(
                job_id="docjob_landscape_transition",
                source_file_id="file_landscape_transition",
                document_language="en",
            ),
            pages=[
                DocumentIRPage(index=0, width=792.0, height=612.0),
                DocumentIRPage(index=1, width=612.0, height=792.0),
            ],
            blocks=[
                DocumentIRBlock(
                    id="landscape_intro",
                    type="paragraph",
                    order=0,
                    text="Landscape intro",
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(40, 60, 300, 90),
                        role="body",
                        column_hint="wide",
                    ),
                ),
                DocumentIRBlock(
                    id="wide_table",
                    type="table",
                    order=1,
                    html=(
                        "<table><tr>"
                        + "".join(f"<td>{long_value}</td>" for _index in range(4))
                        + "</tr></table>"
                    ),
                    layout=DocumentIRBlockLayout(
                        page_index=1,
                        bbox=(40, 80, 572, 200),
                        role="body",
                        column_hint="wide",
                    ),
                ),
            ],
        )
        config = TemplateConfigPackage.load(
            Path(__file__).resolve().parents[1] / "resources" / "template_configs" / "datasheet_template.yaml"
        )
        plan = build_layout_plan(document, config)

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
            layout_plan=plan,
            template_config=config,
        )
        rendered = Document(io.BytesIO(content))

        self.assertEqual([region.kind for region in plan.pages[1].regions], ["landscape_table"])
        self.assertEqual(len(rendered.sections), 2)
        self.assertTrue(all(section.orientation == WD_ORIENTATION.LANDSCAPE for section in rendered.sections))

    def test_render_to_bytes_keeps_continuous_restore_when_wide_content_follows_columns(self) -> None:
        document = self._make_v2_two_column_document()
        document.pages.append(DocumentIRPage(index=1, width=400.0, height=600.0))
        document.blocks.append(
            DocumentIRBlock(
                id="block_next_page",
                type="heading",
                order=5,
                level=2,
                text="Next source page",
                layout=DocumentIRBlockLayout(page_index=1, bbox=(0, 60, 400, 100), role="body", column_hint="wide"),
            )
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        rendered = Document(io.BytesIO(content))
        xml = self._document_xml(content)

        self.assertEqual(rendered.sections[-1].start_type, WD_SECTION_START.CONTINUOUS)
        self.assertTrue(any(section.start_type == WD_SECTION_START.CONTINUOUS for section in rendered.sections))
        self.assertNotIn('<w:br w:type="page"', xml)
        self.assertTrue(any(paragraph.text == "Wide block" for paragraph in rendered.paragraphs))

    def test_render_to_bytes_starts_new_section_when_page_orientation_changes(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_orientation", source_file_id="file_orientation", document_language="en"),
            pages=[
                DocumentIRPage(index=0, width=612.0, height=792.0),
                DocumentIRPage(index=1, width=792.0, height=612.0),
            ],
            blocks=[
                DocumentIRBlock(
                    id="block_portrait",
                    type="paragraph",
                    order=0,
                    text="Portrait",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(72, 80, 300, 120), role="body"),
                ),
                DocumentIRBlock(
                    id="block_landscape",
                    type="paragraph",
                    order=1,
                    text="Landscape",
                    layout=DocumentIRBlockLayout(page_index=1, bbox=(72, 80, 300, 120), role="body"),
                ),
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        rendered = Document(io.BytesIO(content))

        self.assertEqual(rendered.sections[-1].start_type, WD_SECTION_START.NEW_PAGE)
        self.assertEqual(rendered.sections[-1].orientation, WD_ORIENTATION.LANDSCAPE)
        self.assertNotIn('<w:br w:type="page"', self._document_xml(content))

    def test_render_to_bytes_preserves_template_decorations_when_config_disabled(self) -> None:
        renderer = DocumentIRWordRenderer()
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            template_path = temp_path / "template.docx"
            config_path = temp_path / "template.yaml"
            template = Document()
            template.sections[0].header.paragraphs[0].text = "Template Header"
            template.sections[0].footer.paragraphs[0].text = "Template Footer"
            template.save(template_path)
            config_path.write_text(
                (
                    "schema_version: template-config/v1\n"
                    "header:\n"
                    "  enabled: false\n"
                    "footer:\n"
                    "  enabled: false\n"
                ),
                encoding="utf-8",
            )

            content = renderer.render_to_bytes(
                self._make_v2_two_column_document(),
                template_path=template_path,
                config_path=config_path,
                image_loader=lambda _asset: None,
            )
            rendered = Document(io.BytesIO(content))

        self.assertEqual(rendered.sections[0].header.paragraphs[0].text, "Template Header")
        self.assertEqual(rendered.sections[0].footer.paragraphs[0].text, "Template Footer")
        xml = self._document_xml(content)
        self.assertEqual(xml.count("<w:headerReference"), 1)
        self.assertEqual(xml.count("<w:footerReference"), 1)
        self.assertTrue(all(section.header.is_linked_to_previous for section in rendered.sections[1:]))
        self.assertTrue(all(section.footer.is_linked_to_previous for section in rendered.sections[1:]))

    def test_render_to_bytes_writes_configured_body_header_and_footer_font_slots(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_fonts", source_file_id="file_fonts", document_language="zh"),
            blocks=[
                DocumentIRBlock(
                    id="block_fonts",
                    type="paragraph",
                    order=0,
                    text="中文 Latin العربية",
                )
            ],
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "template.yaml"
            config_path.write_text(
                "\n".join(
                    [
                        "schema_version: template-config/v1",
                        "normal:",
                        "  font_name: SimSun",
                        "  western_font_name: Arial",
                        "  complex_script_font_name: Times New Roman",
                        "header:",
                        "  enabled: true",
                        "  title_text: 中文 Header العربية",
                        "  title_font_name: Microsoft YaHei",
                        "  title_western_font_name: Calibri",
                        "  title_complex_script_font_name: Arial",
                        "footer:",
                        "  enabled: true",
                        "  show_page_number: true",
                        "  font_name: FangSong",
                        "  western_font_name: Arial",
                        "  complex_script_font_name: Times New Roman",
                    ]
                ),
                encoding="utf-8",
            )

            content = renderer.render_to_bytes(
                document,
                template_path=(
                    Path(__file__).resolve().parents[1]
                    / "resources"
                    / "templates"
                    / "default"
                    / "datasheet_template.docx"
                ),
                config_path=config_path,
                image_loader=lambda _asset: None,
            )

        footer_xml = self._related_xml(content, "word/footer")
        self.assertIn("<w:fldChar", footer_xml)
        self.assertIn(" PAGE ", footer_xml)
        self.assertNotIn("第  页", footer_xml)
        expected_parts = (
            (self._document_xml(content), ("SimSun", "Arial", "Times New Roman")),
            (self._related_xml(content, "word/header"), ("Microsoft YaHei", "Calibri", "Arial")),
            (footer_xml, ("FangSong", "Arial", "Times New Roman")),
        )
        for xml, (east_asia, western, complex_script) in expected_parts:
            run_fonts = re.search(r"<w:rFonts\b[^>]*/>", xml)
            self.assertIsNotNone(run_fonts)
            font_attributes = run_fonts.group(0)
            self.assertIn(f'w:eastAsia="{east_asia}"', font_attributes)
            self.assertIn(f'w:ascii="{western}"', font_attributes)
            self.assertIn(f'w:hAnsi="{western}"', font_attributes)
            self.assertIn(f'w:cs="{complex_script}"', font_attributes)

    def test_render_to_bytes_writes_native_custom_footer_and_ignores_source_footer(self) -> None:
        renderer = DocumentIRWordRenderer()
        repo_root = Path(__file__).resolve().parents[1]
        content = renderer.render_to_bytes(
            self._make_v2_two_column_document(),
            template_path=repo_root / "resources" / "templates" / "default" / "datasheet_template.docx",
            config_path=repo_root / "resources" / "template_configs" / "datasheet_template.yaml",
            image_loader=lambda _asset: None,
        )

        rendered = Document(io.BytesIO(content))
        footer_xml = self._related_xml(content, "word/footer")
        self.assertIn("v1.0", footer_xml)
        self.assertIn("info@ichelps.com", footer_xml)
        self.assertIn(" PAGE ", footer_xml)
        self.assertNotIn("Footer", footer_xml)
        self.assertNotIn("<w:tbl", footer_xml)
        self.assertIn('w:ptab w:alignment="center" w:relativeTo="margin" w:leader="none"', footer_xml)
        self.assertIn('w:ptab w:alignment="right" w:relativeTo="margin" w:leader="none"', footer_xml)
        self.assertTrue(all(section.footer.is_linked_to_previous for section in rendered.sections[1:]))

    def test_enabled_config_without_content_preserves_template_decorations_and_ignores_source_decorations(self) -> None:
        renderer = DocumentIRWordRenderer()
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template_path = root / "template.docx"
            config_path = root / "template.yaml"
            template = Document()
            template.sections[0].header.paragraphs[0].text = "Template Header"
            template.sections[0].footer.paragraphs[0].text = "Template Footer"
            template.save(template_path)
            config_path.write_text(
                "\n".join(
                    [
                        "schema_version: template-config/v1",
                        "header:",
                        "  enabled: true",
                        "footer:",
                        "  enabled: true",
                        "  show_page_number: false",
                    ]
                ),
                encoding="utf-8",
            )

            content = renderer.render_to_bytes(
                self._make_v2_two_column_document(),
                template_path=template_path,
                config_path=config_path,
                image_loader=lambda _asset: None,
            )
            rendered = Document(io.BytesIO(content))

        self.assertEqual(rendered.sections[0].header.paragraphs[0].text, "Template Header")
        self.assertEqual(rendered.sections[0].footer.paragraphs[0].text, "Template Footer")
        self.assertNotIn("Header", "\n".join(paragraph.text for paragraph in rendered.paragraphs))
        self.assertNotIn("Footer", "\n".join(paragraph.text for paragraph in rendered.paragraphs))

    def test_render_to_bytes_keeps_v1_sequential_fallback(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_v1", source_file_id="file_v1", document_language="en"),
            blocks=[
                DocumentIRBlock(id="block_0001", type="paragraph", order=0, text="First"),
                DocumentIRBlock(id="block_0002", type="paragraph", order=1, text="Second"),
            ],
        )

        content = renderer.render_to_bytes(document, template_path=None, config_path=None, image_loader=lambda _asset: None)
        rendered = Document(io.BytesIO(content))

        self.assertEqual(len(rendered.tables), 0)
        self.assertEqual([paragraph.text for paragraph in rendered.paragraphs if paragraph.text], ["First", "Second"])

    def test_render_to_bytes_preserves_html_table_spans_and_annotations(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_table", source_file_id="file_table", document_language="en"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                DocumentIRBlock(
                    id="block_table",
                    type="table",
                    order=0,
                    html=(
                        "<table><tr><th colspan=\"2\">Header</th><th>Value</th></tr>"
                        "<tr><td rowspan=\"2\">A</td><td>B1</td><td>C1</td></tr>"
                        "<tr><td>B2</td><td>C2</td></tr></table>"
                    ),
                    captions=["Table 1: merged cells"],
                    footnotes=["Note: kept editable"],
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 80, 400, 180), role="body"),
                )
            ],
        )

        content = renderer.render_to_bytes(document, template_path=None, config_path=None, image_loader=lambda _asset: None)
        rendered = Document(io.BytesIO(content))
        xml = self._document_xml(content)

        self.assertEqual(len(rendered.tables), 1)
        self.assertIn("Header", rendered.tables[0].cell(0, 0).text)
        self.assertIn("A", rendered.tables[0].cell(1, 0).text)
        self.assertIn('w:gridSpan', xml)
        self.assertIn('w:vMerge', xml)
        self.assertIn("<w:tblHeader", xml)
        self.assertGreaterEqual(xml.count("<w:cantSplit"), 3)
        self.assertIn('<w:jc w:val="center"', xml)
        self.assertIn('<w:jc w:val="left"', xml)
        self.assertIn('<w:spacing w:before="0" w:after="0"', xml)
        self.assertLess(xml.index(" SEQ Table "), xml.index("<w:tbl>"))
        self.assertLess(xml.index("<w:tbl>"), xml.index("Note: kept editable"))
        self.assertIn("<w:keepNext/>", xml[: xml.index("<w:tbl>")])
        self.assertTrue(any(paragraph.text == "Table 1: merged cells" for paragraph in rendered.paragraphs))
        self.assertTrue(any(paragraph.text == "Note: kept editable" for paragraph in rendered.paragraphs))

    def test_invalid_table_body_deduplicates_caption_but_keeps_distinct_fallback(self) -> None:
        document = DocumentIR(
            metadata=DocumentIRMetadata(
                job_id="docjob_invalid_tables",
                source_file_id="file_invalid_tables",
                document_language="zh",
            ),
            blocks=[
                DocumentIRBlock(
                    id="duplicate_caption",
                    type="table",
                    order=0,
                    text="参数表 A",
                    html="<div>并非表格</div>",
                    captions=["参数表 A"],
                ),
                DocumentIRBlock(
                    id="distinct_fallback",
                    type="table",
                    order=1,
                    text="表体解析失败时保留的独立说明",
                    html="<table></table>",
                    captions=["参数表 B"],
                ),
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        rendered = Document(io.BytesIO(content))
        visible_paragraphs = [paragraph.text for paragraph in rendered.paragraphs if paragraph.text]
        evidence = build_renderer_evidence(document, content)

        self.assertEqual(visible_paragraphs.count("参数表 A"), 1)
        self.assertEqual(visible_paragraphs.count("参数表 B"), 1)
        self.assertEqual(visible_paragraphs.count("表体解析失败时保留的独立说明"), 1)
        self.assertEqual(len(rendered.tables), 0)
        self.assertIn("ir_docx_integrity:table_count_mismatch", evidence["blocking_quality_flags"])
        with self.assertRaisesRegex(DocumentIntegrityError, "table_count_mismatch"):
            ensure_renderer_integrity(evidence)

    def test_render_to_bytes_writes_fraction_omml_in_formula_and_table_cells(self) -> None:
        formula = (
            r"\mathrm{激励频率} = \frac {(\mathrm{FCW} \times \mathrm{f} _ {\mathrm{CLKIN}})}"
            r"{2 ^ {1 5}}"
        )
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_fraction", source_file_id="file_fraction", document_language="zh"),
            blocks=[
                DocumentIRBlock(id="block_formula", type="formula", order=0, latex=formula),
                DocumentIRBlock(
                    id="block_table",
                    type="table",
                    order=1,
                    html=f"<table><tr><td>{formula}</td></tr></table>",
                ),
                DocumentIRBlock(id="block_inline_times", type="paragraph", order=2, text=r"A \times B"),
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        xml = self._document_xml(content)
        visible_text = self._document_text(content)
        root = ElementTree.fromstring(xml)
        namespaces = {
            "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        }

        def script_text(element: ElementTree.Element, name: str) -> str:
            script_part = element.find(f"m:{name}", namespaces)
            self.assertIsNotNone(script_part)
            return "".join(script_part.itertext())

        superscripts = [
            (script_text(element, "e"), script_text(element, "sup"))
            for element in root.findall(".//m:sSup", namespaces)
        ]
        subscripts = [
            (script_text(element, "e"), script_text(element, "sub"))
            for element in root.findall(".//m:sSub", namespaces)
        ]
        math_runs = root.findall(".//m:r", namespaces)
        math_style_values = [
            style.get("{http://schemas.openxmlformats.org/officeDocument/2006/math}val") if style is not None else None
            for run in math_runs
            for style in [run.find("m:rPr/m:sty", namespaces)]
        ]
        settings_root = ElementTree.fromstring(self._settings_xml(content))
        template_math_font = settings_root.find(".//m:mathFont", namespaces)
        self.assertIsNotNone(template_math_font)
        template_math_font_name = template_math_font.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/math}val"
        )

        self.assertGreaterEqual(xml.count("<m:oMath>"), 2)
        self.assertGreaterEqual(xml.count("<m:f>"), 2)
        self.assertGreaterEqual(xml.count("<m:sSup>"), 2)
        self.assertGreaterEqual(xml.count("<m:sSub>"), 2)
        self.assertGreaterEqual(superscripts.count(("2", "15")), 2)
        self.assertGreaterEqual(subscripts.count(("f", "CLKIN")), 2)
        self.assertTrue(math_runs)
        self.assertIn("p", math_style_values)
        self.assertIn(None, math_style_values)
        for run, style_value in zip(math_runs, math_style_values, strict=True):
            fonts = run.find("w:rPr/w:rFonts", namespaces)
            if style_value == "p":
                self.assertIsNotNone(fonts)
                self.assertEqual(fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii"), template_math_font_name)
                self.assertEqual(fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi"), template_math_font_name)
            else:
                self.assertIsNone(fonts)
        self.assertIn("FCW ×fCLKIN", visible_text)
        self.assertIn("215", visible_text)
        self.assertIn("A × B", visible_text)
        self.assertNotIn(r"\frac", xml)
        self.assertNotIn(r"\times", xml)

    def test_thin_space_tex_command_renders_as_omml_without_latex_residue(self) -> None:
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_thin_space", source_file_id="file_thin_space", document_language="en"),
            blocks=[DocumentIRBlock(id="block_formula", type="formula", order=0, latex=r"x\,y")],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        quality = analyze_docx_quality_bytes(content)

        self.assertIn("<m:oMath>", self._document_xml(content))
        self.assertIn("x\u2009y", self._document_text(content))
        self.assertEqual(quality["latex_residue_count"], 0)

    def test_malformed_fraction_remains_visible_and_is_reported_as_latex_residue(self) -> None:
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_bad_fraction", source_file_id="file_bad_fraction", document_language="en"),
            blocks=[DocumentIRBlock(id="block_formula", type="formula", order=0, latex=r"\frac{1}{2")],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        quality = analyze_docx_quality_bytes(content)
        evidence = build_renderer_evidence(document, content)

        self.assertIn(r"\frac", self._document_xml(content))
        self.assertEqual(quality["latex_residue_count"], 1)
        self.assertIn("docx_latex_residue:1", evidence["quality_flags"])
        self.assertEqual(evidence["blocking_quality_flags"], [])

    def test_render_to_bytes_writes_extended_omml_in_all_shared_text_paths(self) -> None:
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_extended_omml", source_file_id="file_extended_omml", document_language="zh"),
            blocks=[
                DocumentIRBlock(
                    id="block_formula",
                    type="formula",
                    order=0,
                    latex=r"\overline{\mathbf{SOE}} + \boxed{\mathrm{X}} + \bigotimes + \mathbf{t}_{\text{TRACK}}",
                ),
                DocumentIRBlock(
                    id="block_heading",
                    type="heading",
                    order=1,
                    level=2,
                    text=r"状态 $\overline{\mathbf{SAMPLE}}$",
                ),
                DocumentIRBlock(
                    id="block_text",
                    type="paragraph",
                    order=2,
                    text=r"电源 AV $_{DD}$ 与 $\bigotimes$",
                ),
                DocumentIRBlock(
                    id="block_table",
                    type="table",
                    order=3,
                    html=r"<table><tr><td>V<sub>OS</sub>，A<sub>V</sub> $_{DD}$，DV $_{DD}$ 与 $\boxed{\mathrm{X}}$</td></tr></table>",
                ),
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        xml = self._document_xml(content)
        root = ElementTree.fromstring(xml)
        namespace = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}

        def script_text(element: ElementTree.Element, name: str) -> str:
            part = element.find(f"m:{name}", namespace)
            self.assertIsNotNone(part)
            return "".join(part.itertext())

        subscripts = [
            (script_text(element, "e"), script_text(element, "sub"))
            for element in root.findall(".//m:sSub", namespace)
        ]
        self.assertGreaterEqual(len(root.findall(".//m:bar", namespace)), 2)
        self.assertTrue(all(item.get("{http://schemas.openxmlformats.org/officeDocument/2006/math}val") == "top" for item in root.findall(".//m:barPr/m:pos", namespace)))
        self.assertGreaterEqual(len(root.findall(".//m:rPr/m:sty", namespace)), 3)
        self.assertGreaterEqual(len(root.findall(".//m:borderBox", namespace)), 2)
        self.assertGreaterEqual(self._document_text(content).count("⨂"), 2)
        self.assertTrue({("t", "TRACK"), ("AV", "DD"), ("DV", "DD")} <= set(subscripts))
        self.assertGreaterEqual(subscripts.count(("AV", "DD")), 2)
        self.assertIn('w:vertAlign w:val="subscript"', xml)
        self.assertNotIn("&lt;/td&gt;", xml)
        for residue in (r"\overline", r"\mathbf", r"\bigotimes", r"\boxed", "_{DD}", "_{TRACK}"):
            self.assertNotIn(residue, xml)
        self.assertEqual(analyze_docx_quality_bytes(content)["latex_residue_count"], 0)

    def test_render_to_bytes_preprocesses_cm4112_formulas_without_tex_residue(self) -> None:
        fraction = r"\frac {\left(V _ {I N} - V _ {S}\right)}{R _ {S} + 5 0 0 \Omega} \leq 5 \mathrm{mA}"
        total_delay = (
            r"\mathrm{总} t _ {s} = \sqrt {\left(t _ {s , C M 7 5 0 1}\right) ^ {2} + "
            r"\left(t _ {s , C M 4 1 1 X}\right) ^ {2}}"
        )
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_cm4112_math", source_file_id="file_cm4112_math", document_language="zh"),
            blocks=[
                DocumentIRBlock(id="fraction", type="formula", order=0, latex=fraction),
                DocumentIRBlock(id="delay", type="formula", order=1, latex=total_delay),
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        xml = self._document_xml(content)
        visible_text = self._document_text(content)

        self.assertIn("<m:f>", xml)
        self.assertIn("<m:rad>", xml)
        self.assertIn('<m:degHide m:val="1"/>', xml)
        self.assertGreaterEqual(xml.count("<m:sSub>"), 6)
        self.assertGreaterEqual(xml.count("<m:sSup>"), 2)
        for expected in ("≤", "Ω", "总", "CM7501", "CM411X"):
            self.assertIn(expected, visible_text)
        for residue in (r"\frac", r"\left", r"\right", r"\leq", "I N", "5 0 0", "C M 7 5 0 1"):
            self.assertNotIn(residue, xml)
        self.assertEqual(analyze_docx_quality_bytes(content)["latex_residue_count"], 0)

    def test_render_to_bytes_supports_datasheet_tex_matrix_in_all_math_entries(self) -> None:
        expression = (
            r"\dfrac{\mathbf{V}_{I N}^{2}}{\sqrt[3]{\mathit{x}}} + "
            r"\left(\boldsymbol{\Phi} \approx \varphi\right) + \text{总} + "
            r"\pm + \mu + \circ + \times + \theta + \Omega + \Delta + \leq + \sim + \ldots + \rightarrow \tag{1}"
        )
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_tex_matrix", source_file_id="file_tex_matrix", document_language="zh"),
            blocks=[
                DocumentIRBlock(id="formula", type="formula", order=0, latex=expression),
                DocumentIRBlock(id="inline", type="paragraph", order=1, text=f"正文 ${expression}$"),
                DocumentIRBlock(id="table", type="table", order=2, html=f"<table><tr><td>${expression}$</td></tr></table>"),
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        xml = self._document_xml(content)
        visible_text = self._document_text(content)
        root = ElementTree.fromstring(xml)
        word_namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

        def element_text(element: ElementTree.Element) -> str:
            return "".join(element.itertext())

        body_paragraphs = [element_text(paragraph) for paragraph in root.findall("./w:body/w:p", word_namespace)]
        table_cells = [element_text(cell) for cell in root.findall(".//w:tbl//w:tc", word_namespace)]

        self.assertGreaterEqual(xml.count("<m:f>"), 3)
        self.assertGreaterEqual(xml.count("<m:rad>"), 3)
        self.assertGreaterEqual(xml.count("<m:deg>"), 3)
        self.assertNotIn("<m:degHide", xml)
        self.assertGreaterEqual(xml.count("<m:sSub>"), 3)
        self.assertGreaterEqual(xml.count("<m:sSup>"), 3)
        self.assertTrue(any("VIN" in text and "(1)" in text for text in body_paragraphs))
        self.assertTrue(any("正文" in text and "VIN" in text and "(1)" in text for text in body_paragraphs))
        self.assertTrue(any("VIN" in text and "(1)" in text for text in table_cells))
        for expected in ("VIN", "Φ", "≈", "ϕ", "总", "±", "μ", "°", "×", "θ", "Ω", "Δ", "≤", "∼", "…", "→"):
            self.assertIn(expected, visible_text)
        self.assertEqual(visible_text.count("(1)"), 3)
        for residue in (r"\dfrac", r"\sqrt", r"\left", r"\right", r"\boldsymbol", r"\tag"):
            self.assertNotIn(residue, xml)
        self.assertEqual(analyze_docx_quality_bytes(content)["latex_residue_count"], 0)

    def test_invalid_formula_tag_remains_visible_with_quality_signal(self) -> None:
        invalid_expressions = (r"x + \tag{}", r"x + \tag{\unsupported}", r"x + \tag*{1}")
        for index, expression in enumerate(invalid_expressions):
            with self.subTest(expression=expression):
                document = DocumentIR(
                    metadata=DocumentIRMetadata(
                        job_id=f"docjob_bad_tag_{index}",
                        source_file_id=f"file_bad_tag_{index}",
                        document_language="en",
                    ),
                    blocks=[
                        DocumentIRBlock(id="formula", type="formula", order=0, latex=expression),
                        DocumentIRBlock(id="inline", type="paragraph", order=1, text=f"Inline ${expression}$"),
                        DocumentIRBlock(
                            id="table",
                            type="table",
                            order=2,
                            html=f"<table><tr><td>${expression}$</td></tr></table>",
                        ),
                    ],
                )

                content = DocumentIRWordRenderer().render_to_bytes(
                    document,
                    template_path=None,
                    config_path=None,
                    image_loader=lambda _asset: None,
                )
                xml = self._document_xml(content)
                quality = analyze_docx_quality_bytes(content)
                evidence = build_renderer_evidence(document, content)
                rendered = Document(io.BytesIO(content))
                visible_paragraphs = [paragraph.text for paragraph in rendered.paragraphs if paragraph.text]

                self.assertIn(expression, visible_paragraphs)
                self.assertIn(f"Inline {expression}", visible_paragraphs)
                self.assertEqual(rendered.tables[0].cell(0, 0).text, expression)
                self.assertNotIn("<m:oMath>", xml)
                self.assertGreater(quality["latex_residue_count"], 0)
                self.assertTrue(any(flag.startswith("docx_latex_residue:") for flag in evidence["quality_flags"]))
                self.assertEqual(evidence["blocking_quality_flags"], [])

    def test_render_to_bytes_preserves_template_math_font_for_tagged_formulas(self) -> None:
        expression = (
            r"x + \mathrm{R} + \text{T} + \mathsf{S} + \mathbf{B} + \mathit{I}"
            r" + \sqrt[\mathrm{D}]{z} \tag{1}"
        )
        document = DocumentIR(
            metadata=DocumentIRMetadata(
                job_id="docjob_template_math_font",
                source_file_id="file_template_math_font",
                document_language="en",
            ),
            blocks=[
                DocumentIRBlock(id="formula", type="formula", order=0, latex=expression),
                DocumentIRBlock(id="inline", type="paragraph", order=1, text=f"Inline ${expression}$"),
                DocumentIRBlock(
                    id="table",
                    type="table",
                    order=2,
                    html=f"<table><tr><td>${expression}$</td></tr></table>",
                ),
            ],
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            template_path = Path(temp_dir) / "stix-template.docx"
            template = Document()
            math_properties = template.settings.element.find(qn("m:mathPr"))
            if math_properties is None:
                math_properties = OxmlElement("m:mathPr")
                template.settings.element.append(math_properties)
            math_font = math_properties.find(qn("m:mathFont"))
            if math_font is None:
                math_font = OxmlElement("m:mathFont")
                math_properties.insert(0, math_font)
            math_font.set(qn("m:val"), "STIX Two Math")
            template.save(template_path)

            content = DocumentIRWordRenderer().render_to_bytes(
                document,
                template_path=template_path,
                config_path=None,
                image_loader=lambda _asset: None,
            )

        settings_xml = self._settings_xml(content)
        document_xml = self._document_xml(content)
        root = ElementTree.fromstring(document_xml)
        namespaces = {
            "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        }

        def element_text(element: ElementTree.Element) -> str:
            return "".join(element.itertext())

        body_paragraphs = [element_text(paragraph) for paragraph in root.findall("./w:body/w:p", namespaces)]
        table_cells = [element_text(cell) for cell in root.findall(".//w:tbl//w:tc", namespaces)]
        math_runs = root.findall(".//m:r", namespaces)
        quality = analyze_docx_quality_bytes(content)
        evidence = build_renderer_evidence(document, content)

        self.assertIn('<m:mathFont m:val="STIX Two Math"', settings_xml)
        self.assertNotIn('<m:mathFont m:val="Cambria Math"', settings_xml)
        self.assertEqual(document_xml.count("<m:oMath>"), 3)
        self.assertTrue(math_runs)
        roman_runs = []
        for run in math_runs:
            text = "".join(run.itertext())
            style = run.find("m:rPr/m:sty", namespaces)
            fonts = run.find("w:rPr/w:rFonts", namespaces)
            if text in {"R", "T", "S", "D"}:
                roman_runs.append(run)
                self.assertIsNotNone(style)
                self.assertEqual(style.get("{http://schemas.openxmlformats.org/officeDocument/2006/math}val"), "p")
                self.assertIsNotNone(fonts)
                self.assertEqual(
                    fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii"),
                    "STIX Two Math",
                )
                self.assertEqual(
                    fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi"),
                    "STIX Two Math",
                )
                continue
            if text == "B":
                self.assertIsNone(fonts)
                self.assertEqual(style.get("{http://schemas.openxmlformats.org/officeDocument/2006/math}val"), "b")
            elif text == "I":
                self.assertIsNone(fonts)
                self.assertEqual(style.get("{http://schemas.openxmlformats.org/officeDocument/2006/math}val"), "i")
            elif style is not None and style.get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/math}val"
            ) == "p":
                self.assertIsNotNone(fonts)
                self.assertEqual(
                    fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii"),
                    "STIX Two Math",
                )
                self.assertEqual(
                    fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi"),
                    "STIX Two Math",
                )
            else:
                self.assertIsNone(style)
                self.assertIsNone(fonts)
        self.assertEqual(len(roman_runs), 12)
        expected_visible = "x + R + T + S + B + I + Dz (1)"
        self.assertTrue(any(expected_visible in text for text in body_paragraphs))
        self.assertTrue(any(f"Inline {expected_visible}" in text for text in body_paragraphs))
        self.assertTrue(any(expected_visible in text for text in table_cells))
        self.assertNotIn(r"\tag", document_xml)
        self.assertEqual(quality["latex_residue_count"], 0)
        self.assertEqual(evidence["blocking_quality_flags"], [])

    def test_render_to_bytes_adds_default_math_font_when_template_omits_it(self) -> None:
        document = DocumentIR(
            metadata=DocumentIRMetadata(
                job_id="docjob_missing_template_math_font",
                source_file_id="file_missing_template_math_font",
                document_language="en",
            ),
            blocks=[
                DocumentIRBlock(
                    id="formula",
                    type="formula",
                    order=0,
                    latex=r"x + \mathrm{R} + \mathsf{S} + \text{T}",
                )
            ],
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            template_path = Path(temp_dir) / "no-math-font-template.docx"
            template = Document()
            math_properties = template.settings.element.find(qn("m:mathPr"))
            math_font = math_properties.find(qn("m:mathFont")) if math_properties is not None else None
            if math_font is not None:
                math_properties.remove(math_font)
            template.save(template_path)
            content = DocumentIRWordRenderer().render_to_bytes(
                document,
                template_path=template_path,
                config_path=None,
                image_loader=lambda _asset: None,
            )

        settings_xml = self._settings_xml(content)
        document_xml = self._document_xml(content)
        root = ElementTree.fromstring(document_xml)
        namespaces = {
            "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        }

        self.assertIn('<m:mathFont m:val="Cambria Math"', settings_xml)
        self.assertIn("x + R + S + T", self._document_text(content))
        for run in root.findall(".//m:r", namespaces):
            text = "".join(run.itertext())
            style = run.find("m:rPr/m:sty", namespaces)
            fonts = run.find("w:rPr/w:rFonts", namespaces)
            if text in {"R", "S", "T"} or (
                style is not None
                and style.get("{http://schemas.openxmlformats.org/officeDocument/2006/math}val") == "p"
            ):
                self.assertIsNotNone(style)
                self.assertEqual(style.get("{http://schemas.openxmlformats.org/officeDocument/2006/math}val"), "p")
                self.assertIsNotNone(fonts)
                self.assertEqual(
                    fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii"),
                    "Cambria Math",
                )
                self.assertEqual(
                    fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi"),
                    "Cambria Math",
                )
            else:
                self.assertIsNone(style)
                self.assertIsNone(fonts)

    def test_math_font_context_is_scoped_to_each_concurrent_document_part(self) -> None:
        document = DocumentIR(
            metadata=DocumentIRMetadata(
                job_id="docjob_concurrent_math_font",
                source_file_id="file_concurrent_math_font",
                document_language="en",
            ),
            blocks=[DocumentIRBlock(id="formula", type="formula", order=0, latex=r"x+1+\mathrm{R}")],
        )
        renderer = DocumentIRWordRenderer()

        with tempfile.TemporaryDirectory() as temp_dir:
            template_paths: dict[str, Path] = {}
            for name, value in (("custom", "STIX Two Math"), ("missing", None), ("empty", "  ")):
                path = Path(temp_dir) / f"{name}.docx"
                template = Document()
                math_properties = template.settings.element.find(qn("m:mathPr"))
                if math_properties is None:
                    math_properties = OxmlElement("m:mathPr")
                    template.settings.element.append(math_properties)
                math_font = math_properties.find(qn("m:mathFont"))
                if value is None and math_font is not None:
                    math_properties.remove(math_font)
                elif value is not None:
                    if math_font is None:
                        math_font = OxmlElement("m:mathFont")
                        math_properties.insert(0, math_font)
                    math_font.set(qn("m:val"), value)
                template.save(path)
                template_paths[name] = path

            def render(name: str) -> tuple[str, bytes]:
                return name, renderer.render_to_bytes(
                    document,
                    template_path=template_paths[name],
                    config_path=None,
                    image_loader=lambda _asset: None,
                )

            names = ("custom", "missing", "empty") * 10
            with ThreadPoolExecutor(max_workers=6) as executor:
                outputs = list(executor.map(render, names))

        namespaces = {
            "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        }
        for name, content in outputs:
            expected = "STIX Two Math" if name == "custom" else "Cambria Math"
            settings = ElementTree.fromstring(self._settings_xml(content))
            self.assertEqual(
                settings.find(".//m:mathFont", namespaces).get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/math}val"
                ),
                expected,
            )
            root = ElementTree.fromstring(self._document_xml(content))
            explicit_fonts = root.findall(".//m:r/w:rPr/w:rFonts", namespaces)
            self.assertTrue(explicit_fonts)
            self.assertTrue(
                all(
                    fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii") == expected
                    and fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi") == expected
                    for fonts in explicit_fonts
                )
            )

    def test_render_to_bytes_uses_word_runs_for_simple_table_math_identifiers(self) -> None:
        document = DocumentIR(
            metadata=DocumentIRMetadata(
                job_id="docjob_simple_table_math",
                source_file_id="file_simple_table_math",
                document_language="en",
            ),
            blocks=[
                DocumentIRBlock(id="body", type="paragraph", order=0, text=r"Body $V_{CC}$"),
                DocumentIRBlock(
                    id="table",
                    type="table",
                    order=1,
                    html=(
                        r"<table><tr><th>$V_{CC}$</th><th>$I^2C$</th></tr>"
                        r"<tr><td>$R_{\theta JA}$</td><td>$t_{ON}$</td></tr>"
                        r"<tr><td>$V_{OUT}=1$</td><td>$\frac{1}{2}$</td></tr>"
                        r"<tr><td>Input low voltage $V_{IL}$; limit $V_{IN}=1$</td>"
                        r"<td>Plain <b>HTML</b></td></tr></table>"
                    ),
                ),
            ],
        )
        config = TemplateConfigPackage.load(
            Path(__file__).resolve().parents[1] / "resources" / "template_configs" / "datasheet_template.yaml"
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
            template_config=config,
        )
        evidence = build_renderer_evidence(document, content, template_config=config)
        root = ElementTree.fromstring(self._document_xml(content))
        namespaces = {
            "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        }
        cells = root.findall(".//w:tbl//w:tc", namespaces)
        self.assertEqual(len(cells), 8)

        expected_runs = (
            (("V", None), ("CC", "subscript")),
            (("I", None), ("2", "superscript"), ("C", None)),
            (("R", None), ("θJA", "subscript")),
            (("t", None), ("ON", "subscript")),
        )
        for cell, expected in zip(cells[:4], expected_runs, strict=True):
            self.assertIsNone(cell.find(".//m:oMath", namespaces))
            actual: list[tuple[str, str | None]] = []
            for run in cell.findall(".//w:r", namespaces):
                text = "".join(run.itertext())
                if not text:
                    continue
                vertical = run.find("w:rPr/w:vertAlign", namespaces)
                actual.append(
                    (
                        text,
                        vertical.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
                        if vertical is not None
                        else None,
                    )
                )
                fonts = run.find("w:rPr/w:rFonts", namespaces)
                self.assertIsNotNone(fonts)
                self.assertEqual(
                    fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia"),
                    "宋体",
                )
                self.assertEqual(
                    fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii"),
                    "Arial",
                )
                self.assertEqual(
                    run.find("w:rPr/w:sz", namespaces).get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"
                    ),
                    "18",
                )
            self.assertEqual(tuple(actual), expected)

        for cell in cells[4:6]:
            self.assertIsNotNone(cell.find(".//m:oMath", namespaces))
        mixed_cell = cells[6]
        self.assertIsNotNone(mixed_cell.find(".//m:oMath", namespaces))
        mixed_word_runs = [
            (
                "".join(run.itertext()),
                vertical.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
                if (vertical := run.find("w:rPr/w:vertAlign", namespaces)) is not None
                else None,
            )
            for run in mixed_cell.findall(".//w:r", namespaces)
            if "".join(run.itertext())
        ]
        self.assertTrue(any(text.endswith("V") and vertical is None for text, vertical in mixed_word_runs))
        self.assertIn(("IL", "subscript"), mixed_word_runs)
        self.assertIsNone(cells[7].find(".//m:oMath", namespaces))
        self.assertIsNotNone(cells[7].find(".//w:r/w:rPr/w:b", namespaces))
        body_paragraph = root.find("./w:body/w:p", namespaces)
        self.assertIsNotNone(body_paragraph)
        self.assertIsNotNone(body_paragraph.find(".//m:oMath", namespaces))
        self.assertEqual(
            cells[0].find("w:p/w:pPr/w:jc", namespaces).get(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"
            ),
            "center",
        )
        self.assertEqual(
            cells[2].find("w:p/w:pPr/w:jc", namespaces).get(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"
            ),
            "left",
        )
        self.assertEqual(evidence["blocking_quality_flags"], [])
        self.assertEqual(analyze_docx_quality_bytes(content)["latex_residue_count"], 0)

    def test_simple_table_math_classifier_rejects_complex_or_unsupported_script_order(self) -> None:
        expected = {
            r"$V_{CC}$": "V<sub>CC</sub>",
            r"$R_{\theta JA}$": "R<sub>θJA</sub>",
            r"$t_{ON}$": "t<sub>ON</sub>",
            r"$I^2C$": "I<sup>2</sup>C",
        }
        for source, rendered in expected.items():
            with self.subTest(source=source):
                self.assertEqual(simple_table_math_identifier_html(source), rendered)

        for source in (
            r"$A^2_B$",
            r"$V_{OUT}=1$",
            r"$x^2+y^2$",
            r"$\frac{1}{2}$",
            r"$\sqrt{x}$",
            r"$\overline{AB}$",
            r"$\boxed{x}$",
            r"$x\tag{1}$",
            r"$V_{A_{B}}$",
        ):
            with self.subTest(source=source):
                self.assertIsNone(simple_table_math_identifier_html(source))
                self.assertEqual(
                    replace_simple_table_math_identifiers(f"before {source} after"),
                    f"before {source} after",
                )

    def test_render_to_bytes_styles_only_upright_math_tokens_for_wps_visibility(self) -> None:
        expressions = (r"V_{OUT}=1", r"\frac{1}{2}", r"x^2+y^2", r"\mathbf{x+1}", r"\mathit{x+1}")
        document = DocumentIR(
            metadata=DocumentIRMetadata(
                job_id="docjob_wps_upright_math",
                source_file_id="file_wps_upright_math",
                document_language="en",
            ),
            blocks=[
                *(
                    DocumentIRBlock(id=f"formula_{index}", type="formula", order=index, latex=expression)
                    for index, expression in enumerate(expressions)
                ),
                DocumentIRBlock(
                    id="table",
                    type="table",
                    order=len(expressions),
                    html=r"<table><tr><td>$V_{OUT}=1$</td><td>$\frac{1}{2}$</td></tr></table>",
                ),
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        root = ElementTree.fromstring(self._document_xml(content))
        namespaces = {
            "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        }
        style_by_text: dict[str, set[tuple[str | None, str | None]]] = {}
        for run in root.findall(".//m:r", namespaces):
            text = "".join(run.itertext())
            style = run.find("m:rPr/m:sty", namespaces)
            fonts = run.find("w:rPr/w:rFonts", namespaces)
            style_by_text.setdefault(text, set()).add(
                (
                    style.get("{http://schemas.openxmlformats.org/officeDocument/2006/math}val")
                    if style is not None
                    else None,
                    fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii")
                    if fonts is not None
                    else None,
                )
            )

        for text in ("=1", "1", "2", "+"):
            self.assertIn(("p", "Cambria Math"), style_by_text[text])
        for text in ("V", "OUT", "x", "y"):
            self.assertIn((None, None), style_by_text[text])
        self.assertIn(("b", None), style_by_text["+1"])
        self.assertIn(("i", None), style_by_text["+1"])
        self.assertGreaterEqual(self._document_xml(content).count("<m:sSup>"), 2)
        self.assertEqual(analyze_docx_quality_bytes(content)["latex_residue_count"], 0)
        self.assertEqual(build_renderer_evidence(document, content)["blocking_quality_flags"], [])

    def test_render_to_bytes_styles_later_th_cells_as_table_headers(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_later_th", source_file_id="file_later_th", document_language="en"),
            blocks=[
                DocumentIRBlock(
                    id="block_later_th",
                    type="table",
                    order=0,
                    html=(
                        "<table><tr><td colspan=\"2\">Electrical Characteristics</td></tr>"
                        "<tr><th>PARAMETER</th><th>TEST CONDITIONS</th></tr>"
                        "<tr><td>VOS</td><td>TA = +25°C</td></tr></table>"
                    ),
                )
            ],
        )

        content = renderer.render_to_bytes(document, template_path=None, config_path=None, image_loader=lambda _asset: None)
        xml = self._document_xml(content)

        self.assertIn("PARAMETER", xml)
        self.assertIn("TEST CONDITIONS", xml)
        self.assertGreaterEqual(xml.count("<w:shd"), 3)

    def test_render_to_bytes_normalizes_formula_text_in_table_cells_and_notes(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_table_formula", source_file_id="file_table_formula", document_language="en"),
            blocks=[
                DocumentIRBlock(
                    id="block_table_formula",
                    type="table",
                    order=0,
                    html=(
                        r"<table><tr><th>Parameter</th><th>Value</th></tr>"
                        r"<tr><td>$R_{\theta JA}$</td><td>nV/ \sqrt{Hz}</td></tr></table>"
                    ),
                    captions=[r"Table 1: $R_{\theta JA}$ formulas"],
                    footnotes=[r"Note: E=mc^2"],
                )
            ],
        )

        content = renderer.render_to_bytes(document, template_path=None, config_path=None, image_loader=lambda _asset: None)
        xml = self._document_xml(content)
        visible_text = self._document_text(content)

        self.assertIn("R", visible_text)
        self.assertIn("θJA", visible_text)
        self.assertIn("Table 1: RθJA formulas", visible_text)
        self.assertIn("nV/√Hz", visible_text)
        self.assertIn("Note: E=mc2", visible_text)
        self.assertIn(" SEQ Table ", xml)
        self.assertIn("<m:oMath>", xml)
        self.assertIn("<m:sSub>", xml)
        self.assertIn('w:val="superscript"', xml)
        self.assertNotIn(r"\theta", xml)
        self.assertNotIn(r"\sqrt", xml)

    def test_render_to_bytes_applies_subscript_config_to_formal_renderer_outputs(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_terms", source_file_id="file_terms", document_language="en"),
            blocks=[
                DocumentIRBlock(id="block_text", type="paragraph", order=0, text="VOS and IOS are specified."),
                DocumentIRBlock(
                    id="block_table",
                    type="table",
                    order=1,
                    html="<table><tr><th>Parameter</th></tr><tr><td>VOS</td></tr></table>",
                    captions=["Table 1: VOS limits"],
                ),
            ],
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            subscript_config_path = Path(temp_dir) / "subscripts.yaml"
            subscript_config_path.write_text(
                "\n".join(
                    [
                        "subscript_rules:",
                        "  electrical:",
                        "    VOS: V<sub>OS</sub>",
                        "    IOS: I<sub>OS</sub>",
                    ]
                ),
                encoding="utf-8",
            )

            content = renderer.render_to_bytes(
                document,
                template_path=None,
                config_path=None,
                image_loader=lambda _asset: None,
                options=DocumentIRRenderOptions(subscript_config_path=subscript_config_path, document_language="en"),
            )

        xml = self._document_xml(content)
        self.assertGreaterEqual(xml.count('w:val="subscript"'), 4)
        self.assertIn("V</w:t>", xml)
        self.assertIn("OS", xml)
        self.assertNotIn("&lt;sub&gt;", xml)

    def test_datasheet_five_column_table_widths_prioritize_description_column(self) -> None:
        table = parse_html_table(
            "<table><tr><th>Pin</th><th>DescriptionWithLongIdentifier</th>"
            "<th>Min</th><th>Typ</th><th>Max</th></tr></table>"
        )
        assert table is not None

        decision = choose_table_fit_decision(
            block_id="five_columns",
            table=table,
            container="full_width",
            available_width=10_000_000,
            preferred_font_size=9,
            minimum_font_size=7,
            source_width_ratio=0.8,
        )

        self.assertEqual(sum(decision.column_widths), 10_000_000)
        self.assertGreater(decision.column_widths[1], decision.column_widths[0] * 3)
        self.assertGreater(decision.column_widths[1], decision.column_widths[4] * 3)
        self.assertEqual(decision.source_width_ratio, 0.8)

    def test_short_six_column_table_keeps_preferred_font_size(self) -> None:
        table = parse_html_table(
            "<table><tr><td>A</td><td>B</td><td>C</td><td>D</td><td>E</td><td>F</td></tr></table>"
        )
        assert table is not None

        decision = choose_table_fit_decision(
            block_id="short_six_columns",
            table=table,
            container="column",
            available_width=int(Cm(8.4)),
            preferred_font_size=9,
            minimum_font_size=7,
        )

        self.assertTrue(decision.fits)
        self.assertEqual(decision.selected_font_size, 9)
        self.assertEqual(decision.reason, "preferred_font_fits")
        self.assertEqual(decision.quality_flags, ())

    def test_table_fit_uses_semantic_boundaries_for_engineering_tokens(self) -> None:
        table = parse_html_table("<table><tr><td>0.775(19,69)</td></tr></table>")
        assert table is not None

        decision = choose_table_fit_decision(
            block_id="engineering_token",
            table=table,
            container="column",
            available_width=600_000,
            preferred_font_size=9,
            minimum_font_size=9,
        )

        self.assertAlmostEqual(_unbreakable_units("MS-001"), 5.7)
        self.assertAlmostEqual(_unbreakable_units("MS-001VARIATION"), 9.0)
        self.assertAlmostEqual(_unbreakable_units("0.775(19,69)"), 6.1)
        self.assertEqual(insert_engineering_token_breaks("MS-001VARIATION"), "MS-001 VARIATION")
        self.assertEqual(
            insert_engineering_token_breaks('<span data-code="MS-001VARIATION">MS-001VARIATION</span>'),
            '<span data-code="MS-001VARIATION">MS-001 VARIATION</span>',
        )
        self.assertTrue(decision.fits)
        self.assertEqual(decision.reason, "preferred_font_fits")

    def test_render_to_bytes_normalizes_datasheet_temperature_and_microsecond_symbols(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_formula_edge", source_file_id="file_formula_edge", document_language="en"),
            blocks=[
                DocumentIRBlock(
                    id="block_formula_edge",
                    type="formula",
                    order=0,
                    latex=r"$T_A = +25^\circ C$, $5 \mu s$",
                )
            ],
        )

        content = renderer.render_to_bytes(document, template_path=None, config_path=None, image_loader=lambda _asset: None)
        xml = self._document_xml(content)
        visible_text = self._document_text(content)

        self.assertIn("TA = +25°C, 5 μs", visible_text)
        self.assertIn("<m:oMath>", xml)
        self.assertIn("<m:sSub>", xml)
        self.assertNotIn("^°", xml)
        self.assertNotIn(r"\circ", xml)
        self.assertNotIn(r"\mu", xml)

    def test_render_to_bytes_normalizes_manual_quality_record_formula_residue(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_manual_formula_residue", source_file_id="file_manual_formula_residue", document_language="en"),
            blocks=[
                DocumentIRBlock(
                    id="block_manual_formula_residue",
                    type="paragraph",
                    order=0,
                    text=(
                        r"$\\sqrt{Hz}$ $5 \\mu s$ $PACKAGE^{(1)}$ $-40^{\\circ}C$ "
                        r"$R_{\\theta JC(top)}$ $\\psi_{JT}$ $V_S = \\pm 5$ "
                        r"$V_{CM}$ $I_B$ $R_L=5 kΩ$ nV/\\sqrt{Hz} "
                        r"$\\mathrm { \\mu V / ^ { \\circ } C }$"
                    ),
                )
            ],
        )

        content = renderer.render_to_bytes(document, template_path=None, config_path=None, image_loader=lambda _asset: None)
        xml = self._document_xml(content)
        visible_text = " ".join(self._document_text(content).split())

        for expected in ("Hz", "5 μs", "PACKAGE", "-40°C", "θJC(top)", "ψJT", "±5", "CM", "nV/√Hz", "μV / ° C"):
            self.assertIn(expected, visible_text)
        self.assertIn("<m:sSub>", xml)
        self.assertIn("<m:sSup>", xml)
        self.assertIn("<m:rad>", xml)
        self.assertNotIn(r"\mathrm", xml)
        self.assertNotIn(r"\sqrt", xml)
        self.assertNotIn(r"\theta", xml)
        self.assertNotIn(r"\psi", xml)
        self.assertNotIn(r"\circ", xml)
        self.assertNotIn(r"\pm", xml)
        self.assertNotIn(r"\mu", xml)

    def test_render_to_bytes_uses_bbox_width_for_images_and_annotations(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_image", source_file_id="file_image", document_language="en"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            assets=[DocumentIRAsset(id="asset_image", kind="image", source_path="images/chart.png")],
            blocks=[
                DocumentIRBlock(
                    id="block_image",
                    type="image",
                    order=0,
                    text="Ignored fallback",
                    asset_ref="asset_image",
                    captions=[r"Figure 1: $V_{OS}$ bbox-sized chart"],
                    footnotes=[r"Source: $5 \mu s$ MinerU"],
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 80, 100, 180), role="body"),
                )
            ],
        )

        content = renderer.render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: MINIMAL_PNG_BYTES,
        )
        xml = self._document_xml(content)
        settings_xml = self._settings_xml(content)
        extent_match = re.search(r'<wp:extent cx="(\d+)" cy="(\d+)"', xml)

        self.assertIsNotNone(extent_match)
        self.assertLess(int(extent_match.group(1)), 1_500_000)
        self.assertLess(xml.index("<w:keepNext/>"), xml.index("<w:drawing>"))
        self.assertLess(xml.index("<w:drawing>"), xml.index(" SEQ Figure "))
        self.assertGreaterEqual(xml.count("<w:keepNext/>"), 2)
        self.assertIn("<m:sSub>", xml)
        self.assertNotIn(r"\mu", xml)
        self.assertIn("Figure 1: VOS bbox-sized chart", self._document_text(content))
        self.assertIn("Source: 5 μs MinerU", self._document_text(content))
        self.assertIn('<w:doNotAutoCompressPictures w:val="true"', settings_xml)
        self.assertEqual(self._first_media_bytes(content), MINIMAL_PNG_BYTES)

    def test_render_to_bytes_caps_normal_and_paired_images_by_source_relative_height(self) -> None:
        group_id = "image_group:paired:caption"
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(
                job_id="docjob_image_height",
                source_file_id="file_image_height",
                document_language="en",
            ),
            pages=[
                DocumentIRPage(index=0, width=400.0, height=600.0),
                DocumentIRPage(index=1, width=400.0, height=600.0),
            ],
            assets=[
                DocumentIRAsset(id="asset_left", kind="image", source_path="left.png"),
                DocumentIRAsset(id="asset_right", kind="image", source_path="right.png"),
                DocumentIRAsset(id="asset_subject", kind="image", source_path="subject.png"),
            ],
            blocks=[
                DocumentIRBlock(
                    id="left_image",
                    type="image",
                    order=0,
                    asset_ref="asset_left",
                    metadata={"content_group_id": group_id},
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(10, 300, 180, 400),
                        role="body",
                        column_hint="left",
                    ),
                ),
                DocumentIRBlock(
                    id="right_image",
                    type="image",
                    order=1,
                    asset_ref="asset_right",
                    metadata={"content_group_id": group_id},
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(220, 300, 390, 400),
                        role="body",
                        column_hint="right",
                    ),
                ),
                DocumentIRBlock(
                    id="shared_caption",
                    type="paragraph",
                    order=2,
                    text="Figure 1. Shared logic diagram",
                    metadata={"annotation_role": "standalone_caption", "content_group_id": group_id},
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(80, 410, 320, 435),
                        role="body",
                        column_hint="wide",
                    ),
                ),
                DocumentIRBlock(
                    id="left_label",
                    type="paragraph",
                    order=3,
                    text="Package code",
                    layout=DocumentIRBlockLayout(
                        page_index=1,
                        bbox=(10, 45, 170, 70),
                        role="body",
                        column_hint="left",
                    ),
                ),
                DocumentIRBlock(
                    id="right_label",
                    type="paragraph",
                    order=4,
                    text="Package outline",
                    layout=DocumentIRBlockLayout(
                        page_index=1,
                        bbox=(230, 45, 390, 70),
                        role="body",
                        column_hint="right",
                    ),
                ),
                DocumentIRBlock(
                    id="subject",
                    type="image",
                    order=5,
                    asset_ref="asset_subject",
                    layout=DocumentIRBlockLayout(
                        page_index=1,
                        bbox=(10, 85, 390, 385),
                        role="body",
                        column_hint="wide",
                    ),
                ),
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: MINIMAL_PNG_BYTES,
        )
        rendered = Document(io.BytesIO(content))
        content_height = int(
            rendered.sections[-1].page_height
            - rendered.sections[-1].top_margin
            - rendered.sections[-1].bottom_margin
        )
        expected_paired_height = int(content_height * (100 / 600))
        expected_subject_height = int(content_height * (300 / 600))
        shapes = list(rendered.inline_shapes)

        self.assertEqual(len(shapes), 3)
        for shape in shapes[:2]:
            self.assertAlmostEqual(shape.height, expected_paired_height, delta=1)
            self.assertEqual(shape.width, shape.height)
        self.assertAlmostEqual(shapes[2].height, expected_subject_height, delta=1)
        self.assertEqual(shapes[2].width, shapes[2].height)
        paired_table_text = " ".join(cell.text for row in rendered.tables[0].rows for cell in row.cells)
        self.assertNotIn("Shared logic diagram", paired_table_text)
        self.assertIn("Figure 1. Shared logic diagram", [paragraph.text for paragraph in rendered.paragraphs])
        xml = self._document_xml(content)
        self.assertEqual(xml.count("<w:pageBreakBefore/>"), 1)
        self.assertNotIn('<w:br w:type="page"', xml)

    def test_render_to_bytes_preserves_inline_rich_text_markup(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_rich", source_file_id="file_rich", document_language="en"),
            blocks=[
                DocumentIRBlock(
                    id="block_rich",
                    type="paragraph",
                    order=0,
                    text="V<sub>OS</sub> = 5<sup>2</sup><br/>next",
                )
            ],
        )

        content = renderer.render_to_bytes(document, template_path=None, config_path=None, image_loader=lambda _asset: None)
        xml = self._document_xml(content)

        self.assertIn('w:val="subscript"', xml)
        self.assertIn('w:val="superscript"', xml)
        self.assertIn("<w:br", xml)

    def test_render_to_bytes_writes_heading_list_and_formula_ooxml_semantics(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_semantic", source_file_id="file_semantic", document_language="en"),
            blocks=[
                DocumentIRBlock(
                    id="block_heading",
                    type="heading",
                    order=0,
                    level=2,
                    text="6.5 Thermal Information",
                ),
                DocumentIRBlock(
                    id="block_list",
                    type="list",
                    order=1,
                    list_items=["• Low noise", "- Stable supply"],
                ),
                DocumentIRBlock(
                    id="block_formula",
                    type="formula",
                    order=2,
                    latex=r"$R_{\theta JA}$, E=mc^2, nV/ \sqrt{Hz}",
                ),
            ],
        )

        content = renderer.render_to_bytes(document, template_path=None, config_path=None, image_loader=lambda _asset: None)
        rendered = Document(io.BytesIO(content))
        xml = self._document_xml(content)
        numbering_xml = self._numbering_xml(content)

        self.assertEqual(rendered.paragraphs[0].style.name, "Heading 2")
        self.assertEqual([paragraph.text for paragraph in rendered.paragraphs[1:3]], ["Low noise", "Stable supply"])
        self.assertIn('w:val="Heading2"', xml)
        self.assertIn("<w:numPr>", xml)
        self.assertIn('w:val="bullet"', numbering_xml)
        self.assertIn("<m:sSub>", xml)
        self.assertIn("<m:sSup>", xml)
        self.assertIn("<m:rad>", xml)
        self.assertIn("nV/", self._document_text(content))
        self.assertNotIn("• Low noise", xml)
        self.assertNotIn(r"\sqrt", xml)

    def test_render_to_bytes_promotes_bullet_like_paragraphs_to_word_lists(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_auto_list", source_file_id="file_auto_list", document_language="en"),
            blocks=[
                DocumentIRBlock(id="block_bullet", type="paragraph", order=0, text="• Paragraph bullet"),
                DocumentIRBlock(id="block_dash", type="paragraph", order=1, text="- Nested paragraph bullet"),
                DocumentIRBlock(id="block_temperature", type="paragraph", order=2, text="-40°C operation"),
            ],
        )

        content = renderer.render_to_bytes(document, template_path=None, config_path=None, image_loader=lambda _asset: None)
        rendered = Document(io.BytesIO(content))
        xml = self._document_xml(content)
        numbering_xml = self._numbering_xml(content)

        self.assertEqual(
            [paragraph.text for paragraph in rendered.paragraphs if paragraph.text],
            ["Paragraph bullet", "Nested paragraph bullet", "-40°C operation"],
        )
        self.assertEqual(xml.count("<w:numPr>"), 2)
        self.assertIn('w:val="bullet"', numbering_xml)
        self.assertNotIn("• Paragraph bullet", xml)
        self.assertNotIn("- Nested paragraph bullet", xml)
        self.assertIn("-40°C operation", xml)

    def test_render_to_bytes_uses_native_decimal_numbering_and_restarts_sequences(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(
                job_id="docjob_decimal_lists",
                source_file_id="file_decimal_lists",
                document_language="en",
            ),
            blocks=[
                DocumentIRBlock(
                    id="first_sequence",
                    type="list",
                    order=0,
                    list_items=[
                        "1. First step",
                        "2. Second step",
                        "3. Third step",
                        "5. Jumped step",
                        "6. Continued after jump",
                    ],
                ),
                DocumentIRBlock(id="separator", type="paragraph", order=1, text="Sequence separator"),
                DocumentIRBlock(
                    id="second_sequence",
                    type="list",
                    order=2,
                    list_items=["1. Restarted step", "2. Continued step"],
                ),
                DocumentIRBlock(id="isolated", type="paragraph", order=3, text="1. Isolated sentence."),
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        rendered = Document(io.BytesIO(content))
        paragraphs = {paragraph.text: paragraph for paragraph in rendered.paragraphs if paragraph.text}
        first_num_ids = [
            str(paragraphs[text]._p.pPr.numPr.numId.val)
            for text in ("First step", "Second step", "Third step")
        ]
        second_num_ids = [
            str(paragraphs[text]._p.pPr.numPr.numId.val)
            for text in ("Restarted step", "Continued step")
        ]
        jump_num_ids = [
            str(paragraphs[text]._p.pPr.numPr.numId.val)
            for text in ("Jumped step", "Continued after jump")
        ]

        self.assertEqual(len(set(first_num_ids)), 1)
        self.assertEqual(len(set(jump_num_ids)), 1)
        self.assertEqual(len(set(second_num_ids)), 1)
        self.assertNotEqual(first_num_ids[0], jump_num_ids[0])
        self.assertNotEqual(jump_num_ids[0], second_num_ids[0])
        self.assertNotEqual(first_num_ids[0], second_num_ids[0])
        for text in ("Sequence separator", "1. Isolated sentence."):
            paragraph_properties = paragraphs[text]._p.pPr
            self.assertTrue(paragraph_properties is None or paragraph_properties.numPr is None)
        document_xml = self._document_xml(content)
        self.assertNotIn("1. First step", document_xml)
        self.assertNotIn("1. Restarted step", document_xml)
        self.assertIn("1. Isolated sentence.", document_xml)

        numbering_root = ElementTree.fromstring(self._numbering_xml(content))
        word_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        num_to_abstract = {
            num.get(f"{word_ns}numId"): num.find(f"{word_ns}abstractNumId").get(f"{word_ns}val")
            for num in numbering_root.findall(f"{word_ns}num")
        }
        abstract_by_id = {
            abstract.get(f"{word_ns}abstractNumId"): abstract
            for abstract in numbering_root.findall(f"{word_ns}abstractNum")
        }
        for num_id, start in (
            (first_num_ids[0], "1"),
            (jump_num_ids[0], "5"),
            (second_num_ids[0], "1"),
        ):
            level = abstract_by_id[num_to_abstract[num_id]].find(f"{word_ns}lvl")
            self.assertEqual(level.find(f"{word_ns}numFmt").get(f"{word_ns}val"), "decimal")
            self.assertEqual(level.find(f"{word_ns}start").get(f"{word_ns}val"), start)
            self.assertEqual(level.find(f"{word_ns}suff").get(f"{word_ns}val"), "space")

        evidence = build_renderer_evidence(document, content)
        self.assertEqual(evidence["visible_payload_evidence"]["counts"]["list_item"], {"expected": 7, "covered": 7})
        self.assertEqual(evidence["blocking_quality_flags"], [])

    def test_render_to_bytes_keeps_decimal_values_inside_explicit_list_items(self) -> None:
        document = DocumentIR(
            metadata=DocumentIRMetadata(
                job_id="docjob_decimal_values",
                source_file_id="file_decimal_values",
                document_language="en",
            ),
            blocks=[
                DocumentIRBlock(
                    id="voltage_values",
                    type="list",
                    order=0,
                    list_items=["2.2V operation", "3.3V operation"],
                )
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        rendered = Document(io.BytesIO(content))

        self.assertEqual([paragraph.text for paragraph in rendered.paragraphs], ["2.2V operation", "3.3V operation"])
        self.assertIn('w:val="bullet"', self._numbering_xml(content))
        evidence = build_renderer_evidence(document, content)
        self.assertEqual(evidence["visible_payload_evidence"]["counts"]["list_item"], {"expected": 2, "covered": 2})
        self.assertEqual(evidence["blocking_quality_flags"], [])

    def test_render_to_bytes_keeps_v1_heading_like_paragraph_compatibility(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_auto_heading", source_file_id="file_auto_heading", document_language="en"),
            blocks=[
                DocumentIRBlock(id="block_numbered_heading", type="paragraph", order=0, text="6.5 Thermal Information"),
                DocumentIRBlock(id="block_title_case_heading", type="paragraph", order=1, text="Absolute Maximum Ratings"),
                DocumentIRBlock(
                    id="block_body",
                    type="paragraph",
                    order=2,
                    text="This section describes the device behavior",
                ),
            ],
        )

        content = renderer.render_to_bytes(document, template_path=None, config_path=None, image_loader=lambda _asset: None)
        rendered = Document(io.BytesIO(content))
        xml = self._document_xml(content)

        self.assertEqual([paragraph.text for paragraph in rendered.paragraphs], [
            "6.5 Thermal Information",
            "Absolute Maximum Ratings",
            "This section describes the device behavior",
        ])
        self.assertEqual(rendered.paragraphs[0].style.name, "Heading 2")
        self.assertEqual(rendered.paragraphs[1].style.name, "Heading 2")
        self.assertEqual(rendered.paragraphs[2].style.name, "Normal")
        self.assertGreaterEqual(xml.count('w:val="Heading2"'), 2)

    def test_render_to_bytes_uses_explicit_v2_semantics_without_heading_guessing(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_v2_semantics", source_file_id="file_v2_semantics", document_language="zh"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                DocumentIRBlock(
                    id="caption",
                    type="paragraph",
                    order=0,
                    text="图 1 输出电压",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(10, 30, 180, 55), role="body"),
                ),
                DocumentIRBlock(
                    id="reference",
                    type="paragraph",
                    order=1,
                    text="图 13 显示了输出响应",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(10, 65, 180, 90), role="body"),
                ),
                DocumentIRBlock(
                    id="explicit_heading",
                    type="heading",
                    order=2,
                    level=1,
                    text="工作原理",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(10, 100, 390, 130), role="body", column_hint="wide"),
                ),
                DocumentIRBlock(
                    id="bullet",
                    type="paragraph",
                    order=3,
                    text="• 应用说明",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(10, 140, 180, 165), role="body"),
                ),
                DocumentIRBlock(
                    id="toc",
                    type="paragraph",
                    order=4,
                    text="工作原理 .... 10",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(10, 175, 180, 200), role="body"),
                ),
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        rendered = Document(io.BytesIO(content))
        paragraphs = {paragraph.text: paragraph for paragraph in rendered.paragraphs if paragraph.text}

        self.assertEqual(paragraphs["图 1 输出电压"].style.name, "Normal")
        self.assertEqual(paragraphs["图 13 显示了输出响应"].style.name, "Normal")
        self.assertEqual(paragraphs["工作原理"].style.name, "Heading 1")
        self.assertIn("应用说明", paragraphs)
        self.assertIn("工作原理 .... 10", paragraphs)
        self.assertEqual(paragraphs["工作原理 .... 10"].style.name, "Normal")
        self.assertEqual(self._document_xml(content).count(" TOC \\o"), 0)

    def test_layout_merges_single_side_tail_before_paired_images_into_one_column_flow(self) -> None:
        def body(
            block_id: str,
            block_type: str,
            text: str | None,
            bbox: tuple[float, float, float, float],
            hint: str,
            *,
            asset_ref: str | None = None,
        ) -> DocumentIRBlock:
            return DocumentIRBlock(
                id=block_id,
                type=block_type,
                order=len(block_id),
                text=text,
                asset_ref=asset_ref,
                layout=DocumentIRBlockLayout(page_index=0, bbox=bbox, role="body", column_hint=hint),
            )

        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_tail", source_file_id="file_tail", document_language="zh"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            assets=[
                DocumentIRAsset(id="asset_left", kind="image", source_path="left.png"),
                DocumentIRAsset(id="asset_right", kind="image", source_path="right.png"),
            ],
            blocks=[
                body("left", "paragraph", "左栏正文", (10, 50, 180, 100), "left"),
                body("right", "paragraph", "右栏正文", (220, 50, 390, 100), "right"),
                body("tail", "paragraph", "左栏尾部", (10, 120, 180, 155), "left"),
                body("image_left", "image", None, (10, 220, 180, 330), "left", asset_ref="asset_left"),
                body("image_right", "image", None, (220, 220, 390, 330), "right", asset_ref="asset_right"),
            ],
        )

        plan = build_layout_plan(document)
        self.assertEqual([region.kind for region in plan.pages[0].regions], ["two_col", "paired_visual_row"])
        self.assertEqual([block.id for block in plan.pages[0].regions[0].blocks], ["left", "tail", "right"])

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: MINIMAL_PNG_BYTES,
        )
        xml = self._document_xml(content)
        self.assertEqual(xml.count('w:num="2"'), 1)
        self.assertEqual(analyze_docx_quality_bytes(content)["inline_shape_count"], 2)

    def test_layout_keeps_text_after_image_only_region_outside_image_flow(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_image_tail", source_file_id="file_image_tail", document_language="zh"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                DocumentIRBlock(
                    id="image_left",
                    type="image",
                    order=0,
                    asset_ref="asset_left",
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(10, 50, 180, 180),
                        role="body",
                        column_hint="left",
                    ),
                ),
                DocumentIRBlock(
                    id="image_right",
                    type="image",
                    order=1,
                    asset_ref="asset_right",
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(220, 50, 390, 180),
                        role="body",
                        column_hint="right",
                    ),
                ),
                DocumentIRBlock(
                    id="text_tail",
                    type="paragraph",
                    order=2,
                    text="下一段正文",
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(10, 220, 180, 260),
                        role="body",
                        column_hint="left",
                    ),
                ),
            ],
        )

        plan = build_layout_plan(document)

        self.assertEqual([region.kind for region in plan.pages[0].regions], ["paired_visual_row", "one_col"])
        self.assertEqual([block.id for block in plan.pages[0].regions[1].blocks], ["text_tail"])

    def test_render_to_bytes_uses_one_native_toc_field_for_toc_like_entries(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_toc", source_file_id="file_toc", document_language="en"),
            blocks=[
                DocumentIRBlock(id="block_toc_title", type="paragraph", order=0, text="目录"),
                DocumentIRBlock(id="block_toc_entry", type="paragraph", order=1, text="Features........................ 1"),
                DocumentIRBlock(
                    id="block_toc_list",
                    type="list",
                    order=2,
                    list_items=["特性 .... 1", "目录 .... 3", "Appendix A Device Pinout 120", "• Updated numbering"],
                ),
            ],
        )

        content = renderer.render_to_bytes(document, template_path=None, config_path=None, image_loader=lambda _asset: None)
        rendered = Document(io.BytesIO(content))
        xml = self._document_xml(content)

        self.assertEqual([paragraph.text for paragraph in rendered.paragraphs], ["目录", "更新目录", "Updated numbering"])
        self.assertEqual(xml.count(" TOC \\o"), 1)
        self.assertNotIn("Features", xml)
        self.assertNotIn("Appendix A Device Pinout", xml)
        self.assertNotIn("<w:tab/>", xml)
        self.assertEqual(xml.count("<w:numPr>"), 1)
        self.assertIn('<w:updateFields w:val="true"', self._settings_xml(content))
        evidence = build_renderer_evidence(document, content)
        self.assertIsNone(evidence["toc_evidence"]["expected_field_count"])
        self.assertEqual(evidence["blocking_quality_flags"], [])

    def test_v2_toc_anchor_uses_final_headings_as_cache_without_source_page_numbers(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_toc_anchor", source_file_id="file_toc_anchor", document_language="zh"),
            pages=[DocumentIRPage(index=0, width=612.0, height=792.0)],
            blocks=[
                DocumentIRBlock(
                    id="toc_title",
                    type="paragraph",
                    order=0,
                    text="目录",
                    layout=DocumentIRBlockLayout(page_index=0, role="body"),
                    metadata={"annotation_role": "toc_title"},
                ),
                DocumentIRBlock(
                    id="toc_anchor",
                    type="paragraph",
                    order=0,
                    text="",
                    layout=DocumentIRBlockLayout(page_index=0, role="body"),
                    metadata={"annotation_role": "toc_anchor"},
                ),
                DocumentIRBlock(
                    id="toc_source",
                    type="list",
                    order=1,
                    list_items=["Product Overview................ 2"],
                    layout=DocumentIRBlockLayout(page_index=0, role="toc_source"),
                ),
                DocumentIRBlock(
                    id="heading",
                    type="heading",
                    order=2,
                    level=1,
                    text="产品概述",
                    layout=DocumentIRBlockLayout(page_index=0, role="body"),
                ),
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        xml = self._document_xml(content)
        visible_text = self._document_text(content)

        self.assertEqual(xml.count(" TOC \\o"), 1)
        self.assertIn("更新目录（在 Word/WPS 中更新字段以生成页码）", visible_text)
        self.assertIn("产品概述", visible_text)
        self.assertNotIn("Product Overview", visible_text)
        evidence = build_renderer_evidence(document, content)
        self.assertEqual(evidence["layout_block_count_delta"], 0)
        self.assertEqual(evidence["blocking_quality_flags"], [])

    def test_render_to_bytes_links_unique_caption_references_with_native_fields(self) -> None:
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_refs", source_file_id="file_refs", document_language="zh"),
            assets=[DocumentIRAsset(id="asset_figure", kind="image", source_path="figure.png")],
            blocks=[
                DocumentIRBlock(id="block_ref", type="paragraph", order=0, text="参见图 3、表 2；图 99 保持原文，表 7 也保持原文。"),
                DocumentIRBlock(
                    id="block_figure",
                    type="image",
                    order=1,
                    asset_ref="asset_figure",
                    captions=["图 3. 典型波形"],
                    footnotes=["测试条件说明"],
                ),
                DocumentIRBlock(
                    id="block_table",
                    type="table",
                    order=2,
                    html="<table><tr><td>见图 3</td></tr></table>",
                    captions=["表 2 参数"],
                ),
                DocumentIRBlock(
                    id="block_duplicate_a",
                    type="table",
                    order=3,
                    html="<table><tr><td>A</td></tr></table>",
                    captions=["Table 7: duplicate A"],
                ),
                DocumentIRBlock(
                    id="block_duplicate_b",
                    type="table",
                    order=4,
                    html="<table><tr><td>B</td></tr></table>",
                    captions=["Table 7: duplicate B"],
                ),
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: MINIMAL_PNG_BYTES,
        )
        xml = self._document_xml(content)

        self.assertIn(" SEQ 图 \\r 3 \\* ARABIC ", xml)
        self.assertIn(" SEQ 表 \\r 2 \\* ARABIC ", xml)
        self.assertEqual(xml.count(" SEQ Table \\r 7 \\* ARABIC "), 2)
        self.assertIn('w:name="_md2word_figure_3"', xml)
        self.assertIn('w:name="_md2word_table_2"', xml)
        self.assertGreaterEqual(xml.count(" REF _md2word_figure_3 \\h "), 2)
        self.assertIn(" REF _md2word_table_2 \\h ", xml)
        self.assertNotIn("REF _md2word_figure_99", xml)
        self.assertNotIn("REF _md2word_table_7", xml)
        self.assertNotIn('w:name="_md2word_table_7"', xml)
        self.assertIn("测试条件说明", xml)
        self.assertIn('<w:updateFields w:val="true"', self._settings_xml(content))

    def test_caption_sequence_fields_preserve_source_integer_numbers(self) -> None:
        document = DocumentIR(
            metadata=DocumentIRMetadata(
                job_id="docjob_caption_source_numbers",
                source_file_id="file_caption_source_numbers",
                document_language="en",
            ),
            assets=[DocumentIRAsset(id="asset_figure", kind="image", source_path="figure.png")],
            blocks=[
                DocumentIRBlock(
                    id="references",
                    type="paragraph",
                    order=0,
                    text="See Figure 5, Figure 9, Table 3, and Table 7.",
                ),
                *[
                    DocumentIRBlock(
                        id=f"figure_{index}",
                        type="image",
                        order=index,
                        asset_ref="asset_figure",
                        captions=[caption],
                    )
                    for index, caption in enumerate(
                        (
                            "Figure 5 start",
                            "Figure 6 consecutive",
                            "Figure 9 skipped",
                            "Figure 9 repeated",
                            "Figure 1.1 compound",
                        ),
                        start=1,
                    )
                ],
                *[
                    DocumentIRBlock(
                        id=f"table_{index}",
                        type="table",
                        order=index,
                        html="<table><tr><td>value</td></tr></table>",
                        captions=[caption],
                    )
                    for index, caption in enumerate(
                        (
                            "Table 3 start",
                            "Table 4 consecutive",
                            "Table 7 skipped",
                            "Table 7 repeated",
                            "Table 2-1 compound",
                        ),
                        start=6,
                    )
                ],
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: MINIMAL_PNG_BYTES,
        )
        xml = self._document_xml(content)
        visible_text = self._document_text(content)

        for label, number in (("Figure", 5), ("Figure", 6), ("Table", 3), ("Table", 4)):
            self.assertIn(f" SEQ {label} \\r {number} \\* ARABIC ", xml)
        self.assertEqual(xml.count(" SEQ Figure \\r 9 \\* ARABIC "), 2)
        self.assertEqual(xml.count(" SEQ Table \\r 7 \\* ARABIC "), 2)
        self.assertNotIn(" SEQ Figure \\* ARABIC ", xml)
        self.assertNotIn(" SEQ Table \\* ARABIC ", xml)
        self.assertIn('w:name="_md2word_figure_5"', xml)
        self.assertIn('w:name="_md2word_table_3"', xml)
        self.assertNotIn('w:name="_md2word_figure_9"', xml)
        self.assertNotIn('w:name="_md2word_table_7"', xml)
        self.assertIn(" REF _md2word_figure_5 \\h ", xml)
        self.assertIn(" REF _md2word_table_3 \\h ", xml)
        self.assertNotIn("REF _md2word_figure_9", xml)
        self.assertNotIn("REF _md2word_table_7", xml)
        for caption in (
            "Figure 5 start",
            "Figure 6 consecutive",
            "Figure 9 skipped",
            "Figure 9 repeated",
            "Table 3 start",
            "Table 4 consecutive",
            "Table 7 skipped",
            "Table 7 repeated",
        ):
            self.assertIn(caption, visible_text)
        self.assertIn("Figure 1.1 compound", visible_text)
        self.assertIn("Table 2-1 compound", visible_text)

        root = ElementTree.fromstring(xml)
        word_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        sequence_fields = []
        for paragraph in root.iter(f"{word_ns}p"):
            runs = paragraph.findall(f"{word_ns}r")
            for index, run in enumerate(runs):
                instruction = run.find(f"{word_ns}instrText")
                if instruction is None or not (instruction.text or "").strip().startswith("SEQ "):
                    continue
                sequence_fields.append(instruction.text or "")
                source_number = re.search(r"\\r (\d+)", instruction.text or "")
                self.assertIsNotNone(source_number)
                self.assertGreaterEqual(index, 1)
                self.assertLess(index + 3, len(runs))
                self.assertEqual(
                    runs[index - 1].find(f"{word_ns}fldChar").get(f"{word_ns}fldCharType"),
                    "begin",
                )
                self.assertEqual(
                    runs[index + 1].find(f"{word_ns}fldChar").get(f"{word_ns}fldCharType"),
                    "separate",
                )
                self.assertEqual(runs[index + 2].find(f"{word_ns}t").text, source_number.group(1))
                self.assertEqual(
                    runs[index + 3].find(f"{word_ns}fldChar").get(f"{word_ns}fldCharType"),
                    "end",
                )
        self.assertEqual(len(sequence_fields), 8)

    def test_caption_bookmarks_follow_only_actual_caption_fields(self) -> None:
        document = DocumentIR(
            metadata=DocumentIRMetadata(
                job_id="docjob_caption_field_targets",
                source_file_id="file_caption_field_targets",
                document_language="zh",
            ),
            blocks=[
                DocumentIRBlock(id="references", type="paragraph", order=0, text="参见图 8、表 4 和 Table 12。"),
                DocumentIRBlock(
                    id="standalone_figure",
                    type="paragraph",
                    order=1,
                    text="图 8 独立题注",
                    metadata={"annotation_role": "standalone_caption"},
                ),
                DocumentIRBlock(
                    id="real_table",
                    type="table",
                    order=2,
                    html="<table><tr><td>value</td></tr></table>",
                    captions=["表 4 参数"],
                ),
                DocumentIRBlock(
                    id="table_context",
                    type="table",
                    order=3,
                    html="<table><tr><td>context</td></tr></table>",
                    captions=["Table 12 at VS = 5 V, RL = 10 kΩ"],
                ),
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        xml = self._document_xml(content)

        self.assertIn(" SEQ 图 \\r 8 \\* ARABIC ", xml)
        self.assertIn(" SEQ 表 \\r 4 \\* ARABIC ", xml)
        self.assertNotIn(" SEQ Table \\r 12 \\* ARABIC ", xml)
        self.assertIn('w:name="_md2word_figure_8"', xml)
        self.assertIn('w:name="_md2word_table_4"', xml)
        self.assertNotIn('w:name="_md2word_table_12"', xml)
        self.assertIn(" REF _md2word_figure_8 \\h ", xml)
        self.assertIn(" REF _md2word_table_4 \\h ", xml)
        self.assertNotIn(" REF _md2word_table_12 \\h ", xml)

    def test_render_to_bytes_keeps_compound_caption_numbers_as_plain_text(self) -> None:
        document = DocumentIR(
            metadata=DocumentIRMetadata(
                job_id="docjob_compound_captions",
                source_file_id="file_compound_captions",
                document_language="zh",
            ),
            blocks=[
                DocumentIRBlock(id="block_refs", type="paragraph", order=0, text="参见表1-1、表 1.1、表1–1、Figure 1A，以及表 1。"),
                DocumentIRBlock(
                    id="block_valid",
                    type="table",
                    order=1,
                    html="<table><tr><td>valid</td></tr></table>",
                    captions=["表 1 有效题注"],
                ),
                DocumentIRBlock(
                    id="block_hyphen",
                    type="table",
                    order=2,
                    html="<table><tr><td>hyphen</td></tr></table>",
                    captions=["表1-1 复合编号"],
                ),
                DocumentIRBlock(
                    id="block_decimal",
                    type="table",
                    order=3,
                    html="<table><tr><td>decimal</td></tr></table>",
                    captions=["表 1.1 复合编号"],
                ),
                DocumentIRBlock(
                    id="block_en_dash",
                    type="table",
                    order=4,
                    html="<table><tr><td>en dash</td></tr></table>",
                    captions=["表1–1 复合编号"],
                ),
                DocumentIRBlock(
                    id="block_suffix",
                    type="table",
                    order=5,
                    html="<table><tr><td>suffix</td></tr></table>",
                    captions=["Figure 1A compound number"],
                ),
            ],
        )

        content = DocumentIRWordRenderer().render_to_bytes(
            document,
            template_path=None,
            config_path=None,
            image_loader=lambda _asset: None,
        )
        xml = self._document_xml(content)
        visible_text = self._document_text(content)

        self.assertEqual(xml.count(" SEQ 表 \\r 1 \\* ARABIC "), 1)
        self.assertNotIn(" SEQ Figure ", xml)
        self.assertEqual(xml.count(" REF _md2word_table_1 \\h "), 1)
        self.assertIn("表1-1", visible_text)
        self.assertIn("表 1.1", visible_text)
        self.assertIn("表1–1", visible_text)
        self.assertIn("Figure 1A", visible_text)

    def test_render_to_bytes_uses_configured_toc_styles(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_toc_style", source_file_id="file_toc_style", document_language="en"),
            blocks=[
                DocumentIRBlock(id="block_toc_title", type="paragraph", order=0, text="Table of Contents"),
                DocumentIRBlock(id="block_toc_level1", type="paragraph", order=1, text="1 Features........................ 1"),
                DocumentIRBlock(id="block_toc_level2", type="paragraph", order=2, text="1.1 Device Details................ 2"),
            ],
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "template.yaml"
            config_path.write_text(
                "\n".join(
                    [
                        "schema_version: template-config/v1",
                        "toc_title:",
                        "  font_size: 15",
                        "  color: '#112233'",
                        "  alignment: center",
                        "toc_level1:",
                        "  font_size: 11",
                        "  color: '#445566'",
                        "  alignment: left",
                        "toc_level2:",
                        "  font_name: SimSun",
                        "  western_font_name: Arial",
                        "  complex_script_font_name: Noto Sans Arabic",
                        "  font_size: 10",
                        "  bold: true",
                        "  italic: true",
                        "  color: '#778899'",
                        "  alignment: right",
                        "  line_spacing: 1.2",
                        "  space_before: 2",
                        "  space_after: 3",
                        "  first_line_indent_pt: 4",
                        "  left_indent_pt: 5",
                        "  right_indent_pt: 6",
                        "toc_level3:",
                        "  font_size: 9",
                        "  color: '#AABBCC'",
                        "  alignment: justify",
                    ]
                ),
                encoding="utf-8",
            )

            content = renderer.render_to_bytes(document, template_path=None, config_path=config_path, image_loader=lambda _asset: None)

        xml = self._document_xml(content)
        styles_root = ElementTree.fromstring(self._styles_xml(content))
        word_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        toc2 = styles_root.find(f".//{word_ns}style[@{word_ns}styleId='TOC2']")
        toc3 = styles_root.find(f".//{word_ns}style[@{word_ns}styleId='TOC3']")

        self.assertIsNotNone(toc2)
        self.assertIsNotNone(toc3)
        assert toc2 is not None
        assert toc3 is not None
        toc2_fonts = toc2.find(f"./{word_ns}rPr/{word_ns}rFonts")
        self.assertIsNotNone(toc2_fonts)
        assert toc2_fonts is not None
        self.assertIn('w:val="112233"', xml)
        self.assertIn('w:val="445566"', xml)
        self.assertEqual(toc2_fonts.get(f"{word_ns}eastAsia"), "SimSun")
        self.assertEqual(toc2_fonts.get(f"{word_ns}ascii"), "Arial")
        self.assertEqual(toc2_fonts.get(f"{word_ns}hAnsi"), "Arial")
        self.assertEqual(toc2_fonts.get(f"{word_ns}cs"), "Noto Sans Arabic")
        self.assertEqual(toc2.find(f"./{word_ns}rPr/{word_ns}color").get(f"{word_ns}val"), "778899")
        self.assertIsNotNone(toc2.find(f"./{word_ns}rPr/{word_ns}b"))
        self.assertIsNotNone(toc2.find(f"./{word_ns}rPr/{word_ns}i"))
        self.assertEqual(toc2.find(f"./{word_ns}pPr/{word_ns}jc").get(f"{word_ns}val"), "right")
        self.assertEqual(toc2.find(f"./{word_ns}pPr/{word_ns}ind").get(f"{word_ns}left"), "100")
        self.assertEqual(toc3.find(f"./{word_ns}rPr/{word_ns}color").get(f"{word_ns}val"), "AABBCC")
        self.assertEqual(toc3.find(f"./{word_ns}pPr/{word_ns}jc").get(f"{word_ns}val"), "both")
        self.assertEqual(xml.count(" TOC \\o"), 1)

    def test_render_to_bytes_uses_configured_long_table_caption_and_border_color(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_table_style", source_file_id="file_table_style", document_language="en"),
            blocks=[
                DocumentIRBlock(
                    id="block_table",
                    type="table",
                    order=0,
                    html="<table><tr><th>Parameter</th></tr><tr><td>Gain</td></tr></table>",
                    captions=["A.1 Continued electrical characteristics"],
                )
            ],
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "template.yaml"
            config_path.write_text(
                "\n".join(
                    [
                        "schema_version: template-config/v1",
                        "table_style:",
                        "  style_name: Table Grid",
                        "  header_fill: AABBCC",
                        "  border_color: '123456'",
                        "caption:",
                        "  color: '#111111'",
                        "long_table_caption:",
                        "  color: '#ABCDEF'",
                    ]
                ),
                encoding="utf-8",
            )

            content = renderer.render_to_bytes(document, template_path=None, config_path=config_path, image_loader=lambda _asset: None)

        xml = self._document_xml(content)
        self.assertIn('w:val="ABCDEF"', xml)
        self.assertIn('w:fill="AABBCC"', xml)
        self.assertIn('w:color="123456"', xml)

    def test_render_to_bytes_uses_configured_page_size_and_header_bold_flag(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_page_config", source_file_id="file_page_config", document_language="en"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                DocumentIRBlock(
                    id="block_table",
                    type="table",
                    order=0,
                    html="<table><tr><th>Parameter</th></tr><tr><td>Gain</td></tr></table>",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(72, 80, 300, 160), role="body"),
                )
            ],
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "template.yaml"
            config_path.write_text(
                "\n".join(
                    [
                        "schema_version: template-config/v1",
                        "page_setup:",
                        "  page_width_cm: 10",
                        "  page_height_cm: 20",
                        "table_style:",
                        "  header_bold: false",
                        "  header_fill: ''",
                    ]
                ),
                encoding="utf-8",
            )

            content = renderer.render_to_bytes(document, template_path=None, config_path=config_path, image_loader=lambda _asset: None)
            rendered = Document(io.BytesIO(content))

        xml = self._document_xml(content)
        self.assertAlmostEqual(rendered.sections[0].page_width.cm, Cm(10).cm, places=1)
        self.assertAlmostEqual(rendered.sections[0].page_height.cm, Cm(20).cm, places=1)
        self.assertNotIn("<w:b/>", xml)
        self.assertNotIn("<w:b ", xml)
        self.assertNotIn("<w:shd", xml)

    def test_template_page_size_defaults_feed_layout_plan_and_renderer(self) -> None:
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(
                job_id="docjob_template_page_size",
                source_file_id="file_template_page_size",
                document_language="en",
            ),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                DocumentIRBlock(
                    id="block_table",
                    type="table",
                    order=0,
                    html="<table><tr><th>Parameter</th></tr><tr><td>Gain</td></tr></table>",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(40, 80, 360, 160), role="body"),
                )
            ],
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            template_path = Path(temp_dir) / "non-a4-template.docx"
            template = Document()
            template.sections[-1].page_width = Cm(18)
            template.sections[-1].page_height = Cm(25)
            template.save(template_path)

            config = resolve_template_config(document, None, None, template_path=template_path)
            plan = build_layout_plan(document, config)
            content = DocumentIRWordRenderer().render_to_bytes(
                document,
                template_path=template_path,
                config_path=None,
                image_loader=lambda _asset: None,
                layout_plan=plan,
                template_config=config,
            )
            rendered = Document(io.BytesIO(content))

        self.assertAlmostEqual(config.page_setup.page_width_cm or 0, 18, places=2)
        self.assertAlmostEqual(config.page_setup.page_height_cm or 0, 25, places=2)
        self.assertAlmostEqual(
            plan.table_fit_decisions["block_table"].available_width,
            int(Cm((config.page_setup.page_width_cm or 0) - 2 * 3.175)),
            delta=2,
        )
        self.assertAlmostEqual(rendered.sections[-1].page_width.cm, 18, places=1)
        self.assertAlmostEqual(rendered.sections[-1].page_height.cm, 25, places=1)

    def test_render_to_bytes_uses_configured_heading_word_style_when_available(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_heading_style", source_file_id="file_heading_style", document_language="en"),
            blocks=[
                DocumentIRBlock(
                    id="block_heading",
                    type="heading",
                    order=0,
                    level=2,
                    text="Configured title style",
                )
            ],
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "template.yaml"
            config_path.write_text(
                (
                    "schema_version: template-config/v1\n"
                    "heading2:\n"
                    "  word_style_name: Title\n"
                    "  font_size: 14\n"
                ),
                encoding="utf-8",
            )

            content = renderer.render_to_bytes(
                document,
                template_path=None,
                config_path=config_path,
                image_loader=lambda _asset: None,
            )
            rendered = Document(io.BytesIO(content))

        self.assertEqual(rendered.paragraphs[0].style.name, "Title")

    def test_render_to_bytes_preserves_portrait_section_for_wide_table_page(self) -> None:
        renderer = DocumentIRWordRenderer()
        wide_row = "".join(f"<th>H{index}</th>" for index in range(8))
        wide_values = "".join(f"<td>V{index}</td>" for index in range(8))
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_wide", source_file_id="file_wide", document_language="en"),
            pages=[
                DocumentIRPage(index=0, width=612.0, height=792.0),
                DocumentIRPage(index=1, width=612.0, height=792.0),
            ],
            blocks=[
                DocumentIRBlock(
                    id="block_intro",
                    type="paragraph",
                    order=0,
                    text="Portrait intro",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(72, 80, 260, 120), role="body", column_hint="left"),
                ),
                DocumentIRBlock(
                    id="block_wide_table",
                    type="table",
                    order=1,
                    html=f"<table><tr>{wide_row}</tr><tr>{wide_values}</tr></table>",
                    layout=DocumentIRBlockLayout(page_index=1, bbox=(40, 80, 572, 260), role="body", column_hint="wide"),
                ),
            ],
        )

        content = renderer.render_to_bytes(document, template_path=None, config_path=None, image_loader=lambda _asset: None)
        rendered = Document(io.BytesIO(content))
        xml = self._document_xml(content)
        grid_widths = [int(value) for value in re.findall(r'<w:gridCol w:w="(\d+)"', xml)]
        available_width = Emu(
            rendered.sections[-1].page_width
            - rendered.sections[-1].left_margin
            - rendered.sections[-1].right_margin
        )
        available_width_twips = available_width.twips

        self.assertEqual(len(rendered.sections), 1)
        self.assertEqual(rendered.sections[0].orientation, WD_ORIENTATION.PORTRAIT)
        self.assertEqual(rendered.sections[-1].orientation, WD_ORIENTATION.PORTRAIT)
        self.assertLess(rendered.sections[-1].page_width, rendered.sections[-1].page_height)
        self.assertNotIn('w:orient="landscape"', xml)
        self.assertEqual(len(rendered.tables[-1].columns), 8)
        self.assertEqual(len(grid_widths), 8)
        self.assertLessEqual(sum(grid_widths), available_width_twips + 8)

    def test_render_to_bytes_keeps_source_landscape_page_landscape(self) -> None:
        renderer = DocumentIRWordRenderer()
        document = DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_landscape", source_file_id="file_landscape", document_language="en"),
            pages=[DocumentIRPage(index=0, width=792.0, height=612.0)],
            blocks=[
                DocumentIRBlock(
                    id="block_body",
                    type="paragraph",
                    order=0,
                    text="Landscape source content",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(72, 80, 400, 120), role="body"),
                ),
            ],
        )

        content = renderer.render_to_bytes(document, template_path=None, config_path=None, image_loader=lambda _asset: None)
        rendered = Document(io.BytesIO(content))

        self.assertEqual(rendered.sections[-1].orientation, WD_ORIENTATION.LANDSCAPE)
        self.assertGreater(rendered.sections[-1].page_width, rendered.sections[-1].page_height)

    def test_render_to_bytes_renders_context_like_table_caption_as_body_text(self) -> None:
        renderer = DocumentIRWordRenderer()
        condition = (
            "at TA = +25°C, VS = ±5 V to ±15 V, VCM = VOUT = VS / 2, "
            "and RL = 10 kΩ connected to VS / 2 (unless otherwise noted)"
        )
        document = DocumentIR(
            metadata=DocumentIRMetadata(job_id="docjob_table_context", source_file_id="file_table_context", document_language="en"),
            blocks=[
                DocumentIRBlock(
                    id="block_table",
                    type="table",
                    order=0,
                    html="<table><tr><th>PARAMETER</th><th>VALUE</th></tr><tr><td>VOS</td><td>5</td></tr></table>",
                    captions=[condition],
                )
            ],
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "template.yaml"
            config_path.write_text(
                "\n".join(
                    [
                        "schema_version: template-config/v1",
                        "normal:",
                        "  alignment: left",
                        "  first_line_indent_pt: 21",
                        "caption:",
                        "  alignment: center",
                    ]
                ),
                encoding="utf-8",
            )

            content = renderer.render_to_bytes(document, template_path=None, config_path=config_path, image_loader=lambda _asset: None)
            rendered = Document(io.BytesIO(content))

        xml = self._document_xml(content)
        before_table_xml = xml[: xml.index("<w:tbl>")]
        self.assertEqual(rendered.paragraphs[0].text, condition)
        self.assertEqual(rendered.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertIn('<w:jc w:val="left"', before_table_xml)
        self.assertNotIn('<w:jc w:val="center"', before_table_xml)

    def _document_xml(self, content: bytes) -> str:
        with zipfile.ZipFile(io.BytesIO(content)) as docx_zip:
            return docx_zip.read("word/document.xml").decode("utf-8")

    def _numbering_xml(self, content: bytes) -> str:
        with zipfile.ZipFile(io.BytesIO(content)) as docx_zip:
            return docx_zip.read("word/numbering.xml").decode("utf-8")

    def _settings_xml(self, content: bytes) -> str:
        with zipfile.ZipFile(io.BytesIO(content)) as docx_zip:
            return docx_zip.read("word/settings.xml").decode("utf-8")

    def _styles_xml(self, content: bytes) -> str:
        with zipfile.ZipFile(io.BytesIO(content)) as docx_zip:
            return docx_zip.read("word/styles.xml").decode("utf-8")

    def _first_media_bytes(self, content: bytes) -> bytes:
        with zipfile.ZipFile(io.BytesIO(content)) as docx_zip:
            media_name = next(name for name in docx_zip.namelist() if name.startswith("word/media/"))
            return docx_zip.read(media_name)

    def _document_text(self, content: bytes) -> str:
        root = ElementTree.fromstring(self._document_xml(content))
        return "".join(element.text or "" for element in root.iter() if element.tag.endswith("}t"))

    def _related_xml(self, content: bytes, prefix: str) -> str:
        with zipfile.ZipFile(io.BytesIO(content)) as docx_zip:
            part_name = next(name for name in docx_zip.namelist() if name.startswith(prefix) and name.endswith(".xml"))
            return docx_zip.read(part_name).decode("utf-8")

    def _replace_document_xml(self, content: bytes, replacements: dict[str, str]) -> bytes:
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(
            output,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as target:
            for item in source.infolist():
                payload = source.read(item.filename)
                if item.filename == "word/document.xml":
                    text = payload.decode("utf-8")
                    for old, new in replacements.items():
                        text = text.replace(old, new)
                    payload = text.encode("utf-8")
                target.writestr(item, payload)
        return output.getvalue()

    def _make_v2_two_column_document(self) -> DocumentIR:
        return DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_v2", source_file_id="file_v2", document_language="en"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                DocumentIRBlock(
                    id="block_header",
                    type="paragraph",
                    order=0,
                    text="Header",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 0, 400, 30), role="page_header"),
                ),
                DocumentIRBlock(
                    id="block_left",
                    type="paragraph",
                    order=1,
                    text="Left column",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 60, 150, 100), role="body"),
                ),
                DocumentIRBlock(
                    id="block_right",
                    type="paragraph",
                    order=2,
                    text="Right column",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(220, 60, 380, 100), role="body"),
                ),
                DocumentIRBlock(
                    id="block_wide",
                    type="paragraph",
                    order=3,
                    text="Wide block",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 130, 400, 170), role="body"),
                ),
                DocumentIRBlock(
                    id="block_footer",
                    type="paragraph",
                    order=4,
                    text="Footer",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 560, 400, 590), role="page_footer"),
                ),
            ],
        )

    def _make_table_fit_two_column_document(self) -> DocumentIR:
        cells = "".join(f"<td>{value}</td>" for value in ("ABCDEFGHIJK", "LMNOPQRSTUV", "WXYZABCDEFG", "HIJKLMNOPQR"))
        return DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_table_fit", source_file_id="file_table_fit", document_language="en"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                DocumentIRBlock(
                    id="left_a",
                    type="paragraph",
                    order=0,
                    text="Left A",
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(10, 80, 180, 140),
                        role="body",
                        column_hint="left",
                    ),
                ),
                DocumentIRBlock(
                    id="table_context",
                    type="paragraph",
                    order=1,
                    text="Right table context",
                    metadata={"annotation_role": "table_context", "source_block_id": "right_table"},
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(220, 55, 390, 75),
                        role="body",
                        column_hint="right",
                    ),
                ),
                DocumentIRBlock(
                    id="right_table",
                    type="table",
                    order=2,
                    html=f"<table><tr>{cells}</tr></table>",
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(220, 80, 390, 150),
                        role="body",
                        column_hint="right",
                    ),
                ),
                DocumentIRBlock(
                    id="left_b",
                    type="paragraph",
                    order=3,
                    text="Left B",
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(10, 165, 180, 230),
                        role="body",
                        column_hint="left",
                    ),
                ),
                DocumentIRBlock(
                    id="right_body",
                    type="paragraph",
                    order=4,
                    text="Right body",
                    layout=DocumentIRBlockLayout(
                        page_index=0,
                        bbox=(220, 165, 390, 230),
                        role="body",
                        column_hint="right",
                    ),
                ),
            ],
        )

    def _make_v2_stable_hint_document(self) -> DocumentIR:
        return DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_hint", source_file_id="file_hint", document_language="en"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                DocumentIRBlock(
                    id="block_left_a",
                    type="paragraph",
                    order=0,
                    text="Left A",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 60, 150, 80), role="body", column_hint="left"),
                ),
                DocumentIRBlock(
                    id="block_right_a",
                    type="paragraph",
                    order=1,
                    text="Right A",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(220, 90, 380, 110), role="body", column_hint="right"),
                ),
                DocumentIRBlock(
                    id="block_left_b",
                    type="paragraph",
                    order=2,
                    text="Left B",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 200, 150, 220), role="body", column_hint="left"),
                ),
                DocumentIRBlock(
                    id="block_right_b",
                    type="paragraph",
                    order=3,
                    text="Right B",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(220, 240, 380, 260), role="body", column_hint="right"),
                ),
                DocumentIRBlock(
                    id="block_wide",
                    type="paragraph",
                    order=4,
                    text="Wide after hints",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 300, 400, 330), role="body", column_hint="wide"),
                ),
            ],
        )

    def _make_v2_wide_between_columns_document(self) -> DocumentIR:
        return DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_wide_between", source_file_id="file_wide_between", document_language="en"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                DocumentIRBlock(
                    id="block_left_a",
                    type="paragraph",
                    order=0,
                    text="Left A",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 60, 150, 80), role="body", column_hint="left"),
                ),
                DocumentIRBlock(
                    id="block_right_a",
                    type="paragraph",
                    order=1,
                    text="Right A",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(220, 70, 380, 90), role="body", column_hint="right"),
                ),
                DocumentIRBlock(
                    id="block_wide_callout",
                    type="paragraph",
                    order=2,
                    text="Wide callout",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 180, 400, 220), role="body", column_hint="wide"),
                ),
                DocumentIRBlock(
                    id="block_left_b",
                    type="paragraph",
                    order=3,
                    text="Left B",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 300, 150, 320), role="body", column_hint="left"),
                ),
                DocumentIRBlock(
                    id="block_right_b",
                    type="paragraph",
                    order=4,
                    text="Right B",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(220, 310, 380, 330), role="body", column_hint="right"),
                ),
            ],
        )

    def _make_v2_hint_override_document(self) -> DocumentIR:
        return DocumentIR(
            schema_version="document-ir/v2",
            metadata=DocumentIRMetadata(job_id="docjob_override", source_file_id="file_override", document_language="en"),
            pages=[DocumentIRPage(index=0, width=400.0, height=600.0)],
            blocks=[
                DocumentIRBlock(
                    id="block_left_hint",
                    type="paragraph",
                    order=0,
                    text="Left by hint",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(220, 60, 380, 100), role="body", column_hint="left"),
                ),
                DocumentIRBlock(
                    id="block_right_hint",
                    type="paragraph",
                    order=1,
                    text="Right by hint",
                    layout=DocumentIRBlockLayout(page_index=0, bbox=(0, 70, 150, 110), role="body", column_hint="right"),
                ),
            ],
        )


if __name__ == "__main__":
    unittest.main()

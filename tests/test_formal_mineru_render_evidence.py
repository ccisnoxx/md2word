"""正式本地转换入口必须消费同次渲染的 v2 内容、目标和锚点证明。"""

from __future__ import annotations

import hashlib
import json
from dataclasses import replace
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase, mock

from docx import Document

from app.services.document_ir.integrity import DocumentIntegrityError
from app.services.document_ir.models import DocumentIR
from app.services.renderer.render_plan import document_ir_hash
from scripts.dev import formal_mineru_conversion


class FormalMinerURenderEvidenceTests(TestCase):
    def setUp(self) -> None:
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        root = Path(directory.name)
        package = root / "mineru-package"
        package.mkdir()
        (package / "sample_content_list_v2.json").write_text(
            json.dumps([
                {"type": "title", "text": "1 Overview", "text_level": 1,
                 "bbox": [72, 80, 260, 100], "page_idx": 0},
                {"type": "text", "text": "Precision amplifier characteristics",
                 "bbox": [72, 120, 480, 150], "page_idx": 0},
                {"type": "table", "table_body": "<table><tr><th>Parameter</th></tr><tr><td>Gain</td></tr></table>",
                 "bbox": [72, 170, 540, 260], "page_idx": 0},
                {"type": "equation", "latex": r"V_{OUT}=1",
                 "bbox": [72, 280, 260, 310], "page_idx": 0},
            ]),
            encoding="utf-8",
        )
        self.request = formal_mineru_conversion.FormalMinerUConversionRequest(
            input_path=package,
            output_path=root / "out.docx",
            no_template=True,
            document_language="en",
        )

    def test_formal_conversion_returns_bound_v2_evidence(self) -> None:
        result = formal_mineru_conversion.convert_formal_mineru_package(self.request)

        metadata = result.renderer_metadata
        self.assertEqual(metadata["renderer_content_contract_version"], "renderer-content/v2")
        self.assertEqual(metadata["render_trace_status"], "complete")
        self.assertEqual(metadata["blocking_quality_flags"], [])
        plan = metadata["render_plan"]
        trace = metadata["block_render_evidence"]
        anchors = metadata["block_anchor_evidence"]
        coverage = metadata["formula_source_coverage"]
        self.assertEqual(plan["schema_version"], "render-plan/v2")
        self.assertEqual(trace["schema_version"], "block-render-evidence/v2")
        self.assertEqual(anchors["schema_version"], "block-anchor-evidence/v2")
        self.assertEqual(coverage["schema_version"], "formula-source-coverage/v1")
        self.assertTrue(trace["complete"])

        document_ir = DocumentIR.model_validate_json(result.document_ir_path.read_text(encoding="utf-8"))
        ir_sha256 = document_ir_hash(document_ir)
        docx_sha256 = hashlib.sha256(result.output_docx_path.read_bytes()).hexdigest()
        for proof in (plan, trace, anchors, coverage):
            self.assertEqual(proof["input_ir_sha256"], ir_sha256)
        for proof in (trace, anchors, coverage):
            self.assertEqual(proof["docx_sha256"], docx_sha256)
        for proof in (trace, anchors):
            self.assertEqual(proof["plan_sha256"], plan["plan_sha256"])
        self.assertEqual(trace["resource_identity"], plan["resource_identity"])
        block_ids = [block.id for block in document_ir.blocks]
        self.assertEqual([decision["block_id"] for decision in plan["decisions"]], block_ids)
        self.assertEqual([record["block_id"] for record in trace["block_render_map"]], block_ids)
        self.assertEqual({entry["block_id"] for entry in anchors["entries"]}, set(block_ids))
        self.assertEqual(len(plan["table_units"]), 1)
        self.assertEqual(len(trace["content_scopes"]), 1)
        self.assertEqual(len(coverage["entries"]), 1)
        self.assertEqual(coverage["entries"][0]["source"]["block_id"], document_ir.blocks[-1].id)

        document = Document(result.output_docx_path)
        self.assertIn("1 Overview", [paragraph.text for paragraph in document.paragraphs])
        self.assertIn("Precision amplifier characteristics", [paragraph.text for paragraph in document.paragraphs])
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(document.tables[0].cell(0, 0).text, "Parameter")
        self.assertEqual(document.tables[0].cell(1, 0).text, "Gain")

    def test_formal_conversion_rejects_missing_same_render_proofs(self) -> None:
        render_document = formal_mineru_conversion._render_document
        for field in ("block_anchor_evidence", "_anchor_ownership", "_ownership"):
            with self.subTest(field=field):
                request = replace(self.request, output_path=self.request.output_path.with_name(f"{field}.docx"))

                def without_proof(*args, **kwargs):
                    rendered = render_document(*args, **kwargs)
                    return replace(rendered, **{field: None})

                with mock.patch.object(formal_mineru_conversion, "_render_document", side_effect=without_proof):
                    with self.assertRaises(DocumentIntegrityError) as rejected:
                        formal_mineru_conversion.convert_formal_mineru_package(request)

                self.assertEqual(rejected.exception.boundary, "DocumentIR→DOCX")
                self.assertEqual(rejected.exception.evidence["blocking_quality_flags"],
                                 ["ir_render_evidence:target_mismatch"])
                self.assertFalse(request.output_path.exists())
                self.assertFalse(request.output_path.with_suffix(".document-ir.json").exists())

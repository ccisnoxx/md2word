"""TAB-301 兼容证据：正式 converter 来源投影不改写 IR 或既有原生表格输出。"""
from __future__ import annotations

import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from docx import Document

from app.services.document_ir import MinerUToDocumentIRService
from app.services.document_ir.models import DocumentIR, RepeatedHeaderHint, SourceSpan
from app.services.renderer import DocumentIRWordRenderer
from app.services.tables.models import (
    CellRun, TableCell, TableGrid, TableSourceFragment, cell_content_sha256,
    make_cell_id, make_run_id, raw_html_set_sha256,
)


def render_signature(document: DocumentIR) -> tuple:
    """比较正式输出的可见内容和既有结构，不以 ZIP 时间戳作行为证据。"""
    content = DocumentIRWordRenderer().render_to_bytes(
        document, template_path=None, config_path=None, image_loader=lambda _asset: None)
    rendered = Document(io.BytesIO(content))
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    structures = tuple(tuple(ElementTree.tostring(element) for element in root.findall(f".//w:{name}", namespace))
                       for name in ("gridSpan", "vMerge", "tblHeader"))
    return (tuple(tuple(cell.text for row in table.rows for cell in row.cells) for table in rendered.tables),
            tuple(paragraph.text for paragraph in rendered.paragraphs if paragraph.text), structures)


class CanonicalTableGridCompatibilityTests(unittest.TestCase):
    def test_legacy_v1_v2_html_readers_keep_schema_source_metadata_and_native_output(self):
        for version in ("document-ir/v1", "document-ir/v2"):
            with self.subTest(version=version):
                payload = {
                    "schema_version": version,
                    "metadata": {"job_id": "legacy", "source_file_id": "legacy", "document_language": "zh"},
                    "blocks": [{"id": "legacy-table", "type": "table", "order": 0,
                                "html": "<table><tr><td>旧表值</td></tr></table>",
                                "metadata": {"source_type": "table", "source_spans": [{"legacy": True}]}}],
                }
                original = DocumentIR.model_validate(payload)
                decoded = DocumentIR.model_validate_json(original.model_dump_json())
                self.assertEqual(decoded, original)
                self.assertEqual(decoded.schema_version, version)
                self.assertEqual(decoded.blocks[0].metadata, payload["blocks"][0]["metadata"])
                self.assertIsNone(decoded.metadata.logical_table_groups_version)
                self.assertEqual(decoded.metadata.logical_table_groups, [])
                self.assertNotIn("table_grid", decoded.model_dump())
                self.assertNotIn("table_grid", decoded.blocks[0].model_dump())
                before = render_signature(original)
                self.assertEqual(render_signature(decoded), before)
                self.assertEqual(before[0], (("旧表值",),))

    def test_converter_sources_project_to_independent_grid_without_ledger_or_renderer_writes(self):
        # 这里只手写已知语义 cell 输入，不使用 parser 或伪造未来 builder。
        html = (' \n<table><tr><th colspan="2">源表头 中文 Русский 😀 &amp;</th></tr>'
                '<tr><td rowspan="2">源值</td><td>B1</td></tr><tr><td>B2</td></tr></table>\n ')
        payload = [
            {"type": "table", "page_idx": 0, "bbox": [10, 60, 90, 90], "html": html,
             "_mineru_source": {"source_index": 999}, "private_raw": "原始字段保留"},
            {"type": "table", "page_idx": 1, "html": " \t\n", "img_path": ""},
            {"type": "table_placeholder", "page_idx": 2},
        ]
        with tempfile.TemporaryDirectory(prefix="tab301_contract_") as directory:
            path = Path(directory) / "mineru.zip"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("content_list_v2.json", json.dumps(payload, ensure_ascii=False))
                archive.writestr("layout.json", json.dumps({"pdf_info": [
                    {"page_idx": index, "page_size": [100, 100]} for index in range(3)
                ]}))
            result = MinerUToDocumentIRService().convert_zip(
                path, job_id="grid-contract", source_file_id="source", document_language="zh")
        document = DocumentIR.model_validate_json(result.document.model_dump_json())
        original_json = document.model_dump_json()
        original_signature = render_signature(document)
        group, = document.metadata.logical_table_groups
        fragments = [TableSourceFragment(
            fragment_id=fragment.fragment_id, source_span=fragment.source_span,
            html_state=fragment.html_state, raw_html=fragment.raw_html, raw_html_sha256=fragment.raw_html_sha256,
            repeated_header_hint=fragment.repeated_header_hint,
        ) for fragment in group.fragments]
        identity = make_cell_id(fragments[0].fragment_id, 1, 0)
        run = CellRun(id=make_run_id(identity, 0), source_run_index=0, type="text", source="源值", target="译文")
        cell = TableCell(id=identity, source_fragment_id=fragments[0].fragment_id,
                         source_row_index=1, source_cell_index=0, row=1, column=0,
                         rowspan=2, colspan=1, is_header=False, runs=[run], content_sha256=cell_content_sha256([run]))
        grid = TableGrid(schema_version="canonical-table/v1", table_id=group.logical_group_id,
                         source_content_sha256=result.validation_report.content_list_sha256,
                         raw_html_sha256=raw_html_set_sha256(fragments), rows=3, columns=2,
                         source_fragments=fragments, cells=[cell], header_rows=[0])
        decoded = TableGrid.model_validate_json(grid.model_dump_json())
        self.assertEqual(decoded, grid)
        self.assertEqual(decoded.table_id, group.logical_group_id)
        self.assertEqual(decoded.source_content_sha256, document.metadata.mineru_validation.content_list_sha256)
        self.assertEqual([fragment.raw_html for fragment in decoded.source_fragments], [html, " \t\n", None])
        self.assertEqual([fragment.html_state for fragment in decoded.source_fragments], ["present", "empty", "missing"])
        self.assertEqual([fragment.source_span.source_index for fragment in decoded.source_fragments], [0, 1, 2])
        self.assertEqual(group.fragments[0].raw_fragment, payload[0])
        self.assertIsInstance(grid.source_fragments[0].source_span, SourceSpan)
        grid.source_fragments[0].source_span.page_index = 99
        self.assertEqual(document.model_dump_json(), original_json)
        self.assertEqual(render_signature(document), original_signature)
        self.assertEqual(len(original_signature[0]), 1)
        self.assertIn("源值", original_signature[0][0])
        self.assertTrue(all(original_signature[2]))
        self.assertNotIn("译文", " ".join(original_signature[0][0]))

    def test_verified_multifragment_header_hints_keep_upstream_types_and_independent_copies(self):
        def raw_table(page: int, value: int) -> dict:
            return {"type": "table", "page_idx": page, "bbox": [10, 60 if page == 0 else 10, 90, 90],
                    "html": ('<table><tr><th>编号</th><th>参数</th></tr>'
                             f'<tr><td>{value}</td><td>源值</td></tr></table>')}
        with tempfile.TemporaryDirectory(prefix="tab301_hint_") as directory:
            path = Path(directory) / "mineru.zip"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("content_list_v2.json", json.dumps([raw_table(0, 1), raw_table(1, 2)]))
                archive.writestr("layout.json", json.dumps({"pdf_info": [
                    {"page_idx": page, "page_size": [100, 100]} for page in range(2)
                ]}))
            document = MinerUToDocumentIRService().convert_zip(
                path, job_id="hint", source_file_id="source", document_language="zh").document
        group, = document.metadata.logical_table_groups
        original = document.model_dump_json()
        snapshots = [TableSourceFragment(
            fragment_id=fragment.fragment_id, source_span=fragment.source_span,
            html_state=fragment.html_state, raw_html=fragment.raw_html,
            raw_html_sha256=fragment.raw_html_sha256, repeated_header_hint=fragment.repeated_header_hint,
        ) for fragment in group.fragments]
        hint = snapshots[1].repeated_header_hint
        self.assertIsInstance(hint, RepeatedHeaderHint)
        self.assertEqual(hint, group.fragments[1].repeated_header_hint)
        self.assertIsNot(hint, group.fragments[1].repeated_header_hint)
        reread = TableSourceFragment.model_validate_json(snapshots[1].model_dump_json())
        self.assertEqual(reread.repeated_header_hint, hint)
        hint.row_count = 2
        self.assertEqual(document.model_dump_json(), original)


if __name__ == "__main__":
    unittest.main()

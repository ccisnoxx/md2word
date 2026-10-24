"""RND-505 的同次创建、最终配对、内容守恒和兼容 reader 合同。"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import zipfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
import yaml

from app.services.document_ir.integrity import DocumentIntegrityError
from app.services.document_ir.models import DocumentIRAsset, DocumentIRBlock, DocumentIRBlockLayout, DocumentIRPage
from app.services.renderer import DocumentIRWordRenderer
from app.services.renderer.block_anchors import (
    ANCHOR_NAME_RE, BlockAnchorEvidence, block_bookmark_name, validate_block_anchor_evidence,
)
from app.services.renderer.block_evidence import PackageIndex
from app.services.renderer.evidence import (
    build_renderer_evidence, build_renderer_event_summary, ensure_renderer_integrity,
    validate_persisted_render_evidence,
)
from app.services.renderer.render_plan import content_hash
from app.services.renderer.template_config import TemplateConfigPackage
from tests.test_document_ir_logical_table_group import convert_payload, table
from tests.test_document_ir_word_renderer_layout import MINIMAL_PNG_BYTES
from tests.test_image_group_rendering import image_ir, image_loader
from tests.test_renderer_render_plan import make_ir, render


def _images(positions, *, captions):
    ir = image_ir(len(positions))
    for block, caption in zip(ir.blocks, captions, strict=True):
        block.captions, block.footnotes = [caption], []
    return ir


def _bookmark(parent, name, number):
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:name"), name)
    for node in (start, end):
        node.set(qn("w:id"), str(number))
        parent.append(node)


def _rewrite(content, change):
    source, target = io.BytesIO(content), io.BytesIO()
    with zipfile.ZipFile(source) as old, zipfile.ZipFile(target, "w") as new:
        for info in old.infolist():
            data = old.read(info.filename)
            if info.filename == "word/document.xml":
                root = etree.fromstring(data)
                change(root)
                data = etree.tostring(root)
            new.writestr(info, data)
    return target.getvalue()


def _rehash(proof):
    proof["evidence_sha256"] = content_hash({key: value for key, value in proof.items() if key != "evidence_sha256"})
    return proof


def _xml_parts(content, *, remove_generic=False):
    result = {}
    with zipfile.ZipFile(io.BytesIO(content)) as package:
        for name in package.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                root = etree.fromstring(package.read(name))
                if remove_generic:
                    ids = {node.get(qn("w:id")) for node in root.iter(qn("w:bookmarkStart"))
                           if ANCHOR_NAME_RE.fullmatch(node.get(qn("w:name"), ""))}
                    for node in list(root.iter()):
                        if node.tag in {qn("w:bookmarkStart"), qn("w:bookmarkEnd")} and node.get(qn("w:id")) in ids:
                            node.getparent().remove(node)
                result[name] = etree.tostring(root, method="c14n")
    return result


class BlockAnchorTests(TestCase):
    def authenticate(self, ir, result):
        evidence = build_renderer_evidence(ir, result.content, render_result=result)
        ensure_renderer_integrity(evidence)
        proof = evidence.get("block_anchor_evidence")
        self.assertIsNotNone(proof)
        self.assertEqual(BlockAnchorEvidence.model_validate_json(json.dumps(proof)), result.block_anchor_evidence)
        self.assertIsNotNone(validate_block_anchor_evidence(proof, plan=result.plan,
            block_evidence=result.block_evidence, word_content=result.content, ownership=result._anchor_ownership))
        return proof

    def test_body_objects_empty_and_table_cell_have_final_paired_anchors(self):
        ir = make_ir([
            DocumentIRBlock(id="heading", type="heading", order=0, text="1 Overview"),
            DocumentIRBlock(id="body", type="paragraph", order=1, text="Editable retained content."),
            DocumentIRBlock(id="list", type="list", order=2, list_items=["First", "Second"]),
            DocumentIRBlock(id="formula", type="formula", order=3, latex=r"\frac{a}{b}"),
            DocumentIRBlock(id="table", type="table", order=4, html="<table><tr><td>5 V</td></tr></table>"),
            DocumentIRBlock(id="captioned", type="table", order=5, html="<table><tr><td>X</td></tr></table>",
                captions=["Table 1 Parameters"], footnotes=["Retained note"]),
            DocumentIRBlock(id="image", type="image", order=6, asset_ref="asset", captions=["Figure 1 Example"]),
            DocumentIRBlock(id="empty", type="paragraph", order=7),
            DocumentIRBlock(id="empty_list", type="list", order=8, list_items=[]),
        ])
        ir.assets = [DocumentIRAsset(id="asset", kind="image", source_path="image.png")]
        before = ir.model_dump_json()
        result = render(ir, image_loader=lambda _: MINIMAL_PNG_BYTES)
        proof = self.authenticate(ir, result)
        self.assertEqual(ir.model_dump_json(), before)
        entries = {entry["block_id"]: entry for entry in proof["entries"]}
        self.assertEqual(set(entries), {block.id for block in ir.blocks} - {"empty_list"})
        index = PackageIndex(result.content)
        for entry in entries.values():
            start, end, target = [index.nodes[value["part"], tuple(value["child_path"])]
                                  for value in (entry["start"], entry["end"], entry["target"])]
            self.assertEqual(start.get(qn("w:name")), block_bookmark_name(entry["anchor_ref"]))
            self.assertEqual(start.get(qn("w:id")), end.get(qn("w:id")))
            self.assertIs(start.getparent(), end.getparent())
            if entry["block_id"] == "table":
                self.assertEqual(target.tag, qn("w:tbl"))
                self.assertIs(next(target.iter(qn("w:p"))), start.getparent())
        self.assertEqual(entries["captioned"]["target"]["role"], "caption")
        self.assertTrue(all(record.word_page is None for record in result.block_evidence.block_render_map))

    def test_three_image_cells_keep_drawing_owner_and_original_first_bookmark(self):
        ir = _images([10, 130, 250], captions=["Figure 1 A $x+y$", "Figure 2 B", "Figure 3 C"])
        result = render(ir, image_loader=image_loader)
        proof = self.authenticate(ir, result)
        index = PackageIndex(result.content)
        cells = []
        for entry, record in zip(proof["entries"], result.block_evidence.block_render_map, strict=True):
            drawing = index.nodes[entry["target"]["part"], tuple(entry["target"]["child_path"])]
            self.assertEqual(drawing.tag, qn("w:drawing"))
            self.assertEqual([target.role for target in record.targets], ["image", "caption"])
            self.assertFalse(any(ANCHOR_NAME_RE.fullmatch(name) for name in record.targets[0].bookmark_names))
            self.assertTrue(record.bookmark_name.startswith("md2w_img_"))
            cells.append(next(drawing.iterancestors(qn("w:tc"))))
        self.assertEqual(len(set(cells)), 3)
        restored = build_renderer_evidence(ir, result.content, render_plan=result.plan.model_dump(mode="json"),
            block_render_evidence=result.block_evidence.model_dump(mode="json"), block_anchor_evidence=proof,
            formula_report=result.formula_report, image_layout_report=result.image_layout_report)
        self.assertEqual(restored["block_anchor_evidence"], proof)

    def test_marker_removal_proves_all_word_xml_and_reports_unchanged(self):
        cases = [make_ir([
            DocumentIRBlock(id="body", type="paragraph", order=0, text="Value $x+y$ remains."),
            DocumentIRBlock(id="table", type="table", order=1, html="<table><tr><td>Gain $a+b$</td></tr></table>", captions=["Table 1 Parameters"]),
        ]), _images([10, 130, 250], captions=["Figure 1 A $x+y$", "Figure 2 B", "Figure 3 C"])]
        for ir in cases:
            with self.subTest(blocks=len(ir.blocks)):
                enabled = render(ir, image_loader=image_loader)
                with patch("app.services.renderer.block_anchors.BLOCK_ANCHOR_GENERATION_POLICY", "read-only/v1"):
                    compatible = render(ir, image_loader=image_loader)
                self.assertIsNone(compatible.block_anchor_evidence)
                self.assertEqual(_xml_parts(enabled.content, remove_generic=True), _xml_parts(compatible.content))
                self.assertEqual(enabled.formula_report, compatible.formula_report)
                self.assertEqual(enabled.image_layout_report, compatible.image_layout_report)
                # B1 reader 继续接受先前写出的 marker 和原证明。
                with patch("app.services.renderer.block_anchors.BLOCK_ANCHOR_GENERATION_POLICY", "read-only/v1"):
                    self.authenticate(ir, enabled)

    def test_name_stability_concurrency_ir_binding_and_high_reserved_id(self):
        renderer = DocumentIRWordRenderer()
        ir = make_ir()
        def run(value):
            return renderer.render_to_result(value, template_path=None, config_path=None, image_loader=lambda _: None)
        with ThreadPoolExecutor(max_workers=3) as pool:
            results = list(pool.map(run, [ir, ir, ir]))
        self.assertEqual([result.block_anchor_evidence.entries for result in results], [results[0].block_anchor_evidence.entries] * 3)
        other = make_ir([DocumentIRBlock(id="body", type="paragraph", order=0, text="Different fixed IR.")])
        different = run(other)
        self.assertNotEqual(results[0].block_anchor_evidence.input_ir_sha256, different.block_anchor_evidence.input_ir_sha256)
        self.assertIsNone(validate_block_anchor_evidence(results[0].block_anchor_evidence, plan=different.plan,
            block_evidence=different.block_evidence, word_content=different.content))
        with TemporaryDirectory() as directory:
            path = Path(directory) / "template.docx"
            document = Document()
            _bookmark(document.sections[0].header.paragraphs[0]._p, "reserved_template", 9000000)
            document.save(path)
            result = render(ir, template_path=path)
        index = PackageIndex(result.content)
        entry = result.block_anchor_evidence.entries[0]
        self.assertGreater(int(index.nodes[entry.start.part, tuple(entry.start.child_path)].get(qn("w:id"))), 9000000)

    def test_template_and_inside_target_name_collision_never_claim_creation(self):
        ir = make_ir()
        expected = block_bookmark_name(render(ir).block_evidence.block_render_map[0].anchor_ref)
        with TemporaryDirectory() as directory:
            path = Path(directory) / "template.docx"
            document = Document()
            _bookmark(document.add_paragraph()._p, expected, 6)
            document.save(path)
            result = render(ir, template_path=path)
        self.assertEqual(self.authenticate(ir, result)["entries"], [])
        self.assertEqual(result.block_anchor_counts["name_conflict"], 1)
        original = DocumentIRWordRenderer._fill_planned_text
        def preoccupied(renderer, paragraph, *args, **kwargs):
            original(renderer, paragraph, *args, **kwargs)
            _bookmark(paragraph._p, expected, 17)
        with patch.object(DocumentIRWordRenderer, "_fill_planned_text", preoccupied):
            result = render(ir)
        self.assertEqual(result.block_evidence.block_render_map[0].bookmark_name, expected)
        self.assertEqual(self.authenticate(ir, result)["entries"], [])
        self.assertEqual(result.block_anchor_counts["name_conflict"], 1)
        forged = render(ir).block_anchor_evidence
        self.assertNotIn("block_anchor_evidence", build_renderer_evidence(ir, result.content,
            render_result=replace(result, block_anchor_evidence=forged, _anchor_ownership=None)))

    def test_toc_fragment_owners_and_suppressed_blocks_have_no_fabricated_entries(self):
        ir = make_ir([
            DocumentIRBlock(id="first", type="paragraph", order=0, text="1 Overview ........ 3"),
            DocumentIRBlock(id="second", type="paragraph", order=1, text="1 Overview ........ 3"),
            DocumentIRBlock(id="empty", type="list", order=2, list_items=[]),
        ])
        result = render(ir)
        proof = self.authenticate(ir, result)
        self.assertEqual([entry["block_id"] for entry in proof["entries"]], ["first"])
        self.assertEqual(result.block_evidence.block_render_map[1].representative_block_id, "first")
        ir = convert_payload([table(0, "1"), {"type": "table_placeholder", "page_idx": 1, "bbox": [10, 10, 90, 50]}])
        result = render(ir)
        proof = self.authenticate(ir, result)
        self.assertEqual(len(proof["entries"]), 1)
        self.assertEqual(len(result.block_evidence.block_render_map[0].provenance.source_spans), 2)
        decorations = make_ir([DocumentIRBlock(id="header", type="paragraph", order=0, text="Source header",
            layout=DocumentIRBlockLayout(page_index=0, role="page_header"))])
        decorations.schema_version = "document-ir/v2"
        decorations.pages = [DocumentIRPage(index=0)]
        result = render(decorations)
        self.assertEqual(self.authenticate(decorations, result)["entries"], [])

    def test_shared_header_footer_markers_keep_physical_part_and_block_owners(self):
        ir = make_ir([DocumentIRBlock(id=role + str(i), type="paragraph", order=i, text="Retained decoration",
            layout=DocumentIRBlockLayout(page_index=0, role=role))
            for i, role in enumerate(("page_header", "page_header", "page_footer", "page_number"))])
        ir.blocks.append(DocumentIRBlock(id="body", type="paragraph", order=4, text="Retained body",
            layout=DocumentIRBlockLayout(page_index=0, role="body")))
        ir.schema_version = "document-ir/v2"
        ir.pages = [DocumentIRPage(index=0)]
        with TemporaryDirectory() as directory:
            path = Path(directory) / "config.yaml"
            config = TemplateConfigPackage(header={"enabled": True}, footer={"enabled": True})
            path.write_text(yaml.safe_dump(config.model_dump(mode="json")), encoding="utf-8")
            result = render(ir, config_path=path)
        proof = self.authenticate(ir, result)
        entries = {entry["block_id"]: entry for entry in proof["entries"]}
        self.assertEqual(len(entries), 5)
        for first, second, role in (("page_header0", "page_header1", "header"),
                                    ("page_footer2", "page_number3", "footer")):
            self.assertEqual(entries[first]["target"], entries[second]["target"])
            self.assertEqual(entries[first]["target"]["role"], role)
            self.assertIn(role, entries[first]["target"]["part"])
            self.assertNotEqual(entries[first]["start"], entries[second]["start"])

    def test_optional_proof_bad_version_binding_limits_and_creation_ownership_are_local(self):
        ir, result = make_ir(), render(make_ir())
        proof = self.authenticate(ir, result)
        cases = []
        for key, value in (("schema_version", "unknown/v2"), ("docx_sha256", "0" * 64),
                           ("plan_sha256", "0" * 64), ("input_ir_sha256", "0" * 64)):
            changed = copy.deepcopy(proof)
            changed[key] = value
            cases.append(changed)
        duplicate = copy.deepcopy(proof)
        duplicate["entries"].append(duplicate["entries"][0])
        cases.append(_rehash(duplicate))
        for value in cases:
            with self.subTest(version=value["schema_version"]):
                restored = build_renderer_evidence(ir, result.content, render_plan=result.plan,
                    block_render_evidence=result.block_evidence, block_anchor_evidence=value,
                    formula_report=result.formula_report, image_layout_report=result.image_layout_report)
                self.assertEqual(restored["render_trace_status"], "complete")
                self.assertNotIn("block_anchor_evidence", restored)
        with patch("app.services.renderer.block_anchors.MAX_ANCHOR_BYTES", 100):
            limited = render(ir)
            self.assertIsNone(limited.block_anchor_evidence)
            self.assertEqual(limited.block_anchor_counts["limit_exceeded"], 1)
            safe = build_renderer_event_summary(build_renderer_evidence(ir, limited.content, render_result=limited))
            self.assertEqual(safe["block_anchor_summary"]["limit_exceeded"], 1)
            self.assertIsNone(validate_block_anchor_evidence(proof, plan=result.plan,
                block_evidence=result.block_evidence, word_content=result.content))
        with patch("app.services.renderer.block_anchors.MAX_ANCHOR_ENTRIES", 0):
            self.assertIsNone(render(ir).block_anchor_evidence)
        self.assertNotIn("block_anchor_evidence", build_renderer_evidence(ir, result.content,
            render_result=replace(result, _anchor_ownership=None)))
        damaged_trace = result.block_evidence.model_copy(deep=True)
        damaged_trace.block_render_map[0].targets[0].child_path = [999]
        self.assertIsNone(validate_persisted_render_evidence(ir, result.content,
            render_plan=result.plan, block_render_evidence=damaged_trace))
        with self.assertRaises(DocumentIntegrityError):
            build_renderer_evidence(ir, result.content, render_result=replace(result, block_evidence=damaged_trace))
        metadata = build_renderer_evidence(ir, result.content, render_result=result)
        metadata["block_anchor_evidence"]["entries"][0]["block_id"] = "ANCHOR_CANARY"
        metadata["block_anchor_summary"]["arbitrary"] = "ANCHOR_CANARY"
        summary = json.dumps(build_renderer_event_summary(metadata))
        self.assertNotIn("block_anchor_evidence", summary)
        self.assertNotIn("ANCHOR_CANARY", summary)

    def test_deep_optional_proof_does_not_escape_original_content_validation(self):
        import sys

        ir = make_ir()
        result = render(ir)
        nested = 0
        for _ in range(sys.getrecursionlimit() + 1):
            nested = [nested]
        proof = {"schema_version": "block-anchor-evidence/v1", "entries": [], "extra": nested}
        restored = build_renderer_evidence(ir, result.content, render_plan=result.plan,
            block_render_evidence=result.block_evidence, block_anchor_evidence=proof,
            formula_report=result.formula_report, image_layout_report=result.image_layout_report)
        ensure_renderer_integrity(restored)
        self.assertEqual(restored["render_trace_status"], "complete")
        self.assertNotIn("block_anchor_evidence", restored)

    def test_wrong_run_part_path_pair_and_template_duplicates_are_not_certified(self):
        ir = _images([10, 130], captions=["Figure 1 A", "Figure 2 B"])
        result = render(ir, image_loader=image_loader)
        original = self.authenticate(ir, result)
        for failure in ("other_image", "wrong_part", "stale_path", "wrong_end", "wrong_name"):
            value = copy.deepcopy(original)
            entry = value["entries"][0]
            if failure == "other_image":
                entry["target"] = value["entries"][1]["target"]
            elif failure == "wrong_part":
                entry["end"]["part"] = "word/header1.xml"
            elif failure == "stale_path":
                entry["start"]["child_path"][-1] += 1
            elif failure == "wrong_end":
                entry["end"] = value["entries"][1]["end"]
            else:
                entry["bookmark_name"] = value["entries"][1]["bookmark_name"]
            self.assertIsNone(validate_block_anchor_evidence(_rehash(value), plan=result.plan,
                block_evidence=result.block_evidence, word_content=result.content))

    def test_same_paragraph_other_run_and_same_cell_other_paragraph_are_not_owner(self):
        ir = _images([10, 130], captions=["Figure 1 A", "Figure 2 B"])
        result = render(ir, image_loader=image_loader)
        proof = self.authenticate(ir, result)
        for other_paragraph in (False, True):
            def add_foreign_run(root):
                drawing = next(root.iter(qn("w:drawing")))
                run = next(drawing.iterancestors(qn("w:r")))
                paragraph = run.getparent()
                if other_paragraph:
                    new_paragraph = OxmlElement("w:p")
                    new_paragraph.append(copy.deepcopy(run))
                    paragraph.getparent().append(new_paragraph)
                else:
                    paragraph.append(copy.deepcopy(run))
            content = _rewrite(result.content, add_foreign_run)
            index = PackageIndex(content)
            cell = next(index.roots["word/document.xml"].iter(qn("w:tc")))
            foreign = list(cell.iter(qn("w:drawing")))[-1]
            part, path = index.paths[foreign]
            trace = result.block_evidence.model_copy(deep=True)
            trace.docx_sha256 = hashlib.sha256(content).hexdigest()
            target = trace.block_render_map[0].targets[0]
            target.part, target.child_path = part, list(path)
            target.element_kind, target.actual, target.bookmark_names = index.project(part, path)
            changed = copy.deepcopy(proof)
            changed["docx_sha256"] = trace.docx_sha256
            changed["entries"][0]["target"].update(part=part, child_path=list(path))
            with self.subTest(other_paragraph=other_paragraph):
                self.assertIsNone(validate_block_anchor_evidence(_rehash(changed), plan=result.plan,
                    block_evidence=trace, word_content=content))

    def test_short_name_collision_and_partial_insertion_leave_no_half_marker(self):
        ir = make_ir([DocumentIRBlock(id=str(i), type="paragraph", order=i, text="Retained text") for i in range(2)])
        with patch("app.services.renderer.block_anchors.block_bookmark_name", return_value="md2w_blk_" + "0" * 30):
            result = render(ir)
        self.assertEqual(result.block_anchor_evidence.entries, [])
        self.assertEqual(result.block_anchor_counts["name_conflict"], 2)
        class FailingParagraph:
            def __init__(self, paragraph):
                self.paragraph = paragraph
            def find(self, tag):
                return self.paragraph.find(tag)
            def insert(self, position, node):
                self.paragraph.insert(position, node)
            def append(self, _node):
                raise ValueError("模拟第二个 marker 插入失败")
        with patch("app.services.renderer.block_anchors._placement",
                   side_effect=lambda target: (FailingParagraph(target.element), None)):
            failed = render(ir)
        self.assertEqual(failed.block_anchor_evidence.entries, [])
        self.assertEqual(failed.block_anchor_counts["target_unavailable"], 2)
        self.assertFalse(any(ANCHOR_NAME_RE.fullmatch(node.get(qn("w:name"), ""))
            for node in PackageIndex(failed.content).roots["word/document.xml"].iter(qn("w:bookmarkStart"))))
        with patch("app.services.renderer.block_anchors.BLOCK_ANCHOR_GENERATION_POLICY", "read-only/v1"):
            compatible = render(ir)
        self.assertEqual(_xml_parts(failed.content), _xml_parts(compatible.content))

    def test_image_reader_rejects_half_unknown_reversed_marker_and_changed_content(self):
        ir = _images([10, 130], captions=["Figure 1 A", "Figure 2 B"])
        result = render(ir, image_loader=image_loader)
        self.authenticate(ir, result)
        for failure in ("missing_end", "duplicate_id", "unknown_name", "reverse", "extra_run", "changed_dimension", "marker_content", "marker_attribute"):
            def change(root):
                start = next(node for node in root.iter(qn("w:bookmarkStart"))
                    if ANCHOR_NAME_RE.fullmatch(node.get(qn("w:name"), "")))
                end = next(node for node in root.iter(qn("w:bookmarkEnd")) if node.get(qn("w:id")) == start.get(qn("w:id")))
                paragraph = start.getparent()
                if failure == "missing_end":
                    paragraph.remove(end)
                elif failure == "duplicate_id":
                    start.set(qn("w:id"), "2000000")
                elif failure == "unknown_name":
                    start.set(qn("w:name"), "md2w_blk_unknown")
                elif failure == "reverse":
                    start.addprevious(end)
                elif failure == "extra_run":
                    paragraph.append(OxmlElement("w:r"))
                elif failure == "marker_content":
                    start.append(OxmlElement("w:r"))
                elif failure == "marker_attribute":
                    start.set(qn("w:colFirst"), "0")
                else:
                    next(root.iter(qn("wp:extent"))).set("cx", "1")
            with self.subTest(failure=failure), self.assertRaises(ValueError):
                build_renderer_evidence(ir, _rewrite(result.content, change),
                    formula_report=result.formula_report, image_layout_report=result.image_layout_report)

"""IR-205 合成合同；表号用于场景命名，不代表真实 PDF 或 Word 视觉验收。"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document

from app.services.document_ir import MinerUToDocumentIRService
from app.services.document_ir.artifact_validator import MinerUArtifactValidationError
from app.services.document_ir.integrity import DocumentIntegrityError
from app.services.document_ir.logical_table_group import attach_logical_table_groups, capture_table_sources
from app.services.document_ir.models import DocumentIR
from app.services.renderer import DocumentIRWordRenderer


def table_html(first: str, last: str | None = None, *, header: str = "编号", columns: int = 2) -> str:
    rows = [f"<tr><th>{header}</th>" + "<th>参数</th>" * (columns - 1) + "</tr>"] if header else []
    for number in [first] + ([last] if last is not None else []):
        rows.append(f"<tr><td>{number}</td>" + "<td>源值</td>" * (columns - 1) + "</tr>")
    return "<table>" + "".join(rows) + "</table>"


def table(page: int, first: str, last: str | None = None, **changes) -> dict:
    return {"type": "table", "page_idx": page, "bbox": [10, 10 if page else 60, 90, 90],
            "html": table_html(first, last), **changes}


def convert_payload(payload, *, geometry: bool = True):
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "mineru.zip"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("content_list_v2.json", json.dumps(payload, ensure_ascii=False))
            if geometry:
                archive.writestr("layout.json", json.dumps({"pdf_info": [
                    {"page_idx": page, "page_size": [100, 100]} for page in range(12)
                ]}))
        before = path.read_bytes()
        result = MinerUToDocumentIRService().convert_zip(path, job_id="job", source_file_id="source", document_language="zh")
        if path.read_bytes() != before:
            raise AssertionError("转换改写了原始 ZIP。")
        return result.document


class LogicalTableGroupTests(unittest.TestCase):
    def test_filtered_promotional_tables_require_exact_source_receipts(self):
        """推广封底只凭既有过滤回执移除来源；保留表和归并空片仍须完整对账。"""
        def payload(html=""):
            return [table(0, "1"), {"type": "table_placeholder", "page_idx": 1},
                    {"type": "text", "page_idx": 11, "text": "可使用微信扫码进入芯片助手小程序"},
                    {"type": "image", "page_idx": 11, "img_path": "images/missing.png"},
                    {"type": "table", "page_idx": 11, "html": html}]

        for geometry in (False, True):
            for html in ("", "<span>推广封底占位</span>"):
                with self.subTest(geometry=geometry, html=html):
                    document = convert_payload(payload(html), geometry=geometry)
                    report = document.metadata.mineru_validation
                    self.assertTrue(report.can_convert)
                    self.assertEqual(report.inventory["raw_block_count"], 5)
                    self.assertEqual(report.normalization["filtered_block_count"], 3)
                    self.assertEqual(report.normalization["merged_table_placeholder_count"], 1)
                    self.assertEqual([item["json_pointer"] for item in report.normalization["filtered_blocks"]],
                                     ["/2", "/3", "/4"])
                    fragments = document.metadata.logical_table_groups[0].fragments
                    self.assertEqual([fragment.source_span.json_pointer for fragment in fragments], ["/0", "/1"])
                    self.assertEqual(fragments[0].raw_fragment, payload(html)[0])
                    self.assertEqual(len(document.blocks), 1)
                    self.assertEqual(DocumentIR.model_validate_json(document.model_dump_json()), document)

        def corrupt(document, ledger, mutation):
            report = document.metadata.mineru_validation
            receipts = report.normalization["filtered_blocks"]
            if mutation == "missing":
                receipts.pop()
            elif mutation == "locator":
                receipts[-1]["source_index"] = 0
            elif mutation == "rule":
                receipts[-1]["rule_id"] = "untrusted_filter/v1"
            elif mutation == "reason":
                receipts[-1]["reason"] = "untrusted_filter"
            elif mutation == "page":
                report.normalization["filtered_promotional_page_index"] = 0
            elif mutation == "missing_owner":
                document.blocks.clear()
            elif mutation == "duplicate":
                receipts.append(dict(receipts[-1]))
            else:
                source = ledger.tables[0 if mutation == "retained" else 1]
                receipts.append({**source.span.model_dump(exclude={"bbox"}),
                                 "rule_id": "final_promotional_page/v1", "reason": "final_promotional_page"})
            attach_logical_table_groups(document, ledger)

        for mutation in ("missing", "locator", "rule", "reason", "page", "missing_owner", "duplicate", "retained", "merged"):
            with self.subTest(mutation=mutation), patch(
                "app.services.document_ir.mineru_converter.attach_logical_table_groups",
                side_effect=lambda document, ledger: corrupt(document, ledger, mutation),
            ), self.assertRaisesRegex(ValueError, "逻辑表格来源合同无效"):
                convert_payload(payload(), geometry=False)

    def test_original_internal_fields_do_not_override_consumed_source_evidence(self):
        raw = table(0, "99")
        raw.pop("html")
        raw["_mineru_source"] = {"html": table_html("1")}
        raw["alternate"] = table_html("99")
        document = convert_payload([raw])
        fragment = document.metadata.logical_table_groups[0].fragments[0]
        self.assertEqual(fragment.raw_fragment, raw)
        self.assertEqual(fragment.raw_html_pointer, "/alternate")
        self.assertEqual(fragment.raw_html, document.blocks[0].html)
        self.assertEqual(DocumentIR.model_validate_json(document.model_dump_json()), document)

        left = table(0, "1", html=table_html("1", header=""),
                     _mineru_source={"table_caption": ["Table 9 Parameters"]})
        right = table(1, "99", html=table_html("99", header=""),
                      _mineru_source={"table_caption": ["Table 9 (continued)"]})
        document = convert_payload([left, right])
        self.assertEqual([block.captions for block in document.blocks], [[], []])
        self.assertEqual(len(document.metadata.logical_table_groups), 2)

        document = convert_payload([table(0, "1", _mineru_source={"table_footnote": ["伪结束标记"]}), table(1, "2")])
        self.assertEqual(len(document.metadata.logical_table_groups), 1)

    def test_ambiguous_table_numbers_cannot_supply_any_continuation_evidence(self):
        captions = [
            ("Table 9 - 1 Parameters", "Table 9 - 2 (continued)"),
            ("Table 9(a) Parameters", "Table 9(b) (continued)"),
            ("Table 9 Parameters", "Table 9 continued; Table 10 continued"),
            ("表 9-1 参数", "表 9-2（续）"),
            ("Table 9A Parameters", "Table 9B (continued)"),
        ]
        for left_caption, right_caption in captions:
            # 既检查错误表号单独误判，也检查不能回落到表头+序号正证据掩盖歧义。
            for header, following in (("", "99"), ("编号", "2")):
                with self.subTest(captions=(left_caption, right_caption), header=header):
                    document = convert_payload([
                        table(0, "1", html=table_html("1", header=header), table_caption=[left_caption]),
                        table(1, following, html=table_html(following, header=header), table_caption=[right_caption]),
                    ])
                    self.assertEqual(len(document.metadata.logical_table_groups), 2)
                    self.assertFalse(any(group.continuation_relations for group in document.metadata.logical_table_groups))

    def test_reloaded_verified_relations_require_direct_source_and_real_evidence(self):
        document = convert_payload([table(0, "1"), table(1, "2"), table(2, "3")])
        valid = document.model_dump(mode="json")
        for corruption in ("skip_source", "false_evidence"):
            with self.subTest(corruption=corruption):
                payload = copy.deepcopy(valid)
                group = payload["metadata"]["logical_table_groups"][0]
                relation = group["continuation_relations"][1]
                if corruption == "skip_source":
                    relation["from_fragment_id"] = group["fragments"][0]["fragment_id"]
                else:
                    relation["evidence_codes"][-1] = "same_table_number_continued"
                with self.assertRaisesRegex(ValueError, "verified_relation"):
                    DocumentIR.model_validate(payload)

    def test_table3_keeps_all_empty_raw_fragments_without_adding_render_blocks(self):
        first = table(8, "1", table_caption=["表 3 参数"], html=" \n" + table_html("1") + "\t ")
        empty = {"type": "table", "page_idx": 9, "img_path": "", "html": " \t\n", "bbox": [10, 10, 90, 60]}
        missing = {"type": "table_placeholder", "page_idx": 10, "bbox": [10, 10, 90, 70],
                   "_mineru_source": {"json_pointer": "/forged"}, "_mineru_custom": {"keep": "原始字段"}}
        payload = [first, empty, missing]
        original = copy.deepcopy(payload)
        document = convert_payload(payload)
        self.assertEqual(payload, original)
        self.assertEqual(document.metadata.source_block_count, 1)
        self.assertEqual(document.metadata.source_type_counts, {"table": 1})
        self.assertEqual(len(document.blocks), 1)
        self.assertEqual(document.blocks[0].id, "block_0001")
        self.assertEqual(document.blocks[0].layout.source_index, 0)
        self.assertEqual(document.blocks[0].layout.continuation_heights, [50, 60])
        group = document.metadata.logical_table_groups[0]
        self.assertEqual([fragment.raw_fragment for fragment in group.fragments], payload)
        self.assertEqual([fragment.raw_html for fragment in group.fragments], [first["html"], empty["html"], None])
        self.assertEqual([fragment.html_state for fragment in group.fragments], ["present", "empty", "missing"])
        self.assertEqual([fragment.block_id for fragment in group.fragments], ["block_0001", None, None])
        self.assertEqual([fragment.source_span.page_index for fragment in group.fragments], [8, 9, 10])
        self.assertEqual([fragment.source_span.json_pointer for fragment in group.fragments], ["/0", "/1", "/2"])
        self.assertEqual([fragment.normalization_reason for fragment in group.fragments],
                         [None, "strict_empty_table_placeholder", "explicit_table_placeholder"])
        self.assertEqual([relation.kind for relation in group.continuation_relations], ["normalized_empty_placeholder"] * 2)
        self.assertEqual(group.fragments[0].raw_html_sha256, hashlib.sha256(first["html"].encode()).hexdigest())
        self.assertIsNone(group.fragments[2].raw_html_sha256)
        self.assertEqual(len(document.blocks[0].metadata["source_spans"]), 3)
        self.assertNotIn("原始字段", document.metadata.mineru_validation.model_dump_json())
        rendered = Document(io.BytesIO(DocumentIRWordRenderer().render_to_bytes(
            document, template_path=None, config_path=None, image_loader=lambda _asset: None)))
        self.assertEqual(len(rendered.tables), 1)
        self.assertIn("源值", " ".join(cell.text for row in rendered.tables[0].rows for cell in row.cells))

    def test_reloaded_empty_fragments_cannot_detach_from_receipt_owner_group(self):
        document = convert_payload([table(0, "1"), {"type": "table_placeholder", "page_idx": 1}])
        payload = document.model_dump(mode="json")
        group = payload["metadata"]["logical_table_groups"][0]
        empty = group["fragments"].pop()
        group["continuation_relations"] = []
        # 连同派生计数和 owner spans 一起伪造，仍必须由 receipt 的来源合同拒绝。
        identity = hashlib.sha256(json.dumps(["logical-table-groups/v1", empty["fragment_id"]],
                                             ensure_ascii=False, sort_keys=True,
                                             separators=(",", ":")).encode()).hexdigest()
        payload["metadata"]["logical_table_groups"].append({
            "logical_group_id": "logical-table:" + identity, "fragments": [empty],
            "continuation_relations": [], "diagnostics": [],
        })
        payload["metadata"]["logical_table_group_counts"]["group_count"] = 2
        payload["blocks"][0]["metadata"]["source_spans"].pop()
        with self.assertRaisesRegex(ValueError, "normalization_group"):
            DocumentIR.model_validate(payload)

    def test_table9_continued_number_is_positive_evidence_even_without_th(self):
        document = convert_payload([table(0, "A", table_caption=["Table 9 Parameters"], html=table_html("A", header="")),
                                    table(1, "B", table_caption=["Table 9 (continued)"], html=table_html("B", header=""))])
        group, = document.metadata.logical_table_groups
        self.assertEqual(len(group.fragments), 2)
        self.assertEqual(group.continuation_relations[0].evidence_codes[-1], "same_table_number_continued")
        self.assertIsNone(group.fragments[1].header_signature)
        self.assertIsNone(group.fragments[1].repeated_header_hint)

    def test_table9_and_table10_same_header_remain_distinct(self):
        document = convert_payload([table(0, "1", table_caption=["表 9 参数"]),
                                    table(1, "2", table_caption=["表 10 参数（续）"])])
        self.assertEqual(len(document.metadata.logical_table_groups), 2)
        self.assertFalse(any(group.diagnostics for group in document.metadata.logical_table_groups))

    def test_table10_empty_and_missing_pages_never_invent_sources(self):
        document = convert_payload([table(0, "1", table_caption=["表 10 参数"]),
                                    {"type": "table", "page_idx": 1, "img_path": "", "html": ""},
                                    table(3, "2")])
        self.assertEqual(document.metadata.logical_table_group_counts["fragment_count"], 3)
        self.assertEqual([fragment.source_span.page_index for group in document.metadata.logical_table_groups
                          for fragment in group.fragments], [0, 1, 3])
        self.assertEqual(len(document.blocks), 2)
        with self.assertRaises(DocumentIntegrityError):
            convert_payload([{"type": "table", "html": "", "img_path": "", "page_idx": 0}])

    def test_table24_multifragment_keeps_repeated_headers_and_local_spans(self):
        payload = [table(0, "0x00", "0x01", table_caption=["表 24 寄存器"]),
                   table(1, "0x02", "0x03"), table(2, "0x04", "0x05")]
        document = convert_payload(payload)
        group, = document.metadata.logical_table_groups
        self.assertEqual([block.html for block in document.blocks], [item["html"] for item in payload])
        self.assertEqual(len(group.fragments), 3)
        self.assertEqual(len(group.continuation_relations), 2)
        self.assertIsNone(group.fragments[0].repeated_header_hint)
        self.assertEqual([fragment.repeated_header_hint.row_count for fragment in group.fragments[1:]], [1, 1])
        for block, fragment in zip(document.blocks, group.fragments):
            self.assertEqual(block.metadata["source_spans"], [fragment.source_span.model_dump(mode="json")])
            self.assertEqual(fragment.render_owner_block_id, block.id)
        # 去掉派生 metadata 的正式 renderer 输出应保持同样的表数和可见内容。
        old = document.model_dump(mode="json")
        for key in ("logical_table_groups_version", "logical_table_groups", "logical_table_group_counts"):
            old["metadata"].pop(key)
        for block in old["blocks"]:
            for key in ("logical_group_id", "table_fragment_id", "source_spans"):
                block["metadata"].pop(key)
        before = Document(io.BytesIO(DocumentIRWordRenderer().render_to_bytes(
            DocumentIR.model_validate(old), template_path=None, config_path=None, image_loader=lambda _asset: None)))
        after = Document(io.BytesIO(DocumentIRWordRenderer().render_to_bytes(
            document, template_path=None, config_path=None, image_loader=lambda _asset: None)))
        self.assertEqual(len(after.tables), 3)
        self.assertEqual([[cell.text for row in item.rows for cell in row.cells] for item in after.tables],
                         [[cell.text for row in item.rows for cell in row.cells] for item in before.tables])

    def test_stable_ids_roundtrip_and_rebuilding_replace_derived_fields(self):
        payload = [table(0, "1"), table(1, "2"), {"type": "table_placeholder", "page_idx": 2}]
        captured = []
        def capture(*args, **kwargs):
            ledger = capture_table_sources(*args, **kwargs)
            captured.append(ledger)
            return ledger
        with patch("app.services.document_ir.mineru_converter.capture_table_sources", side_effect=capture):
            first = convert_payload(payload)
            second = convert_payload(payload)
        self.assertEqual(first.metadata.logical_table_groups, second.metadata.logical_table_groups)
        decoded = DocumentIR.model_validate_json(first.model_dump_json())
        self.assertEqual(decoded, first)
        before = decoded.model_dump(mode="json")
        attach_logical_table_groups(decoded, captured[0])
        self.assertEqual(decoded.model_dump(mode="json"), before)
        with self.assertRaisesRegex(ValueError, "source_inventory"):
            attach_logical_table_groups(decoded, type(captured[0])(
                captured[0].checksum, captured[0].tables[:-1], captured[0].body_order))
        encoded = json.loads(decoded.model_dump_json())
        self.assertNotIn("raw_html", encoded["metadata"]["logical_table_groups"][0]["fragments"][0])

    def test_nested_flat_and_wrapper_preserve_trusted_locator_and_all_raw_carriers(self):
        raw_html = " \n" + table_html("1") + "\n "
        base = {"block_type": "text", "table_html": raw_html, "html": table_html("9"),
                "content": {"html": ""}, "_mineru_source": {"source_index": 999},
                "_mineru_validated_asset": {"raw": "合法输入字段"}}
        for shape in ("flat", "nested", "wrapper_flat", "wrapper_nested"):
            with self.subTest(shape=shape):
                raw = copy.deepcopy(base)
                nested = "nested" in shape
                if not nested:
                    raw["page_idx"] = 4
                payload = [[raw]] if nested else [raw]
                if shape.startswith("wrapper"):
                    payload = {"content_list": payload}
                document = convert_payload(payload)
                fragment = document.metadata.logical_table_groups[0].fragments[0]
                expected_pointer = ("/content_list" if shape.startswith("wrapper") else "") + ("/0/0" if nested else "/0")
                self.assertEqual(fragment.raw_fragment, raw)
                self.assertEqual(fragment.raw_html, raw_html)
                self.assertEqual(fragment.raw_html_pointer, "/table_html")
                self.assertEqual(fragment.source_span.json_pointer, expected_pointer)
                self.assertEqual(fragment.source_span.source_index, 0)
                self.assertEqual(fragment.source_span.source_index_in_page, 0)
                self.assertEqual(fragment.source_span.page_index, 0 if nested else 4)
                self.assertIsNone(fragment.source_span.bbox)

    def test_unknown_page_and_nested_html_path_are_preserved(self):
        raw = {"type": "table", "content": {"payload/a~b": [{"html": " \n" + table_html("1") + " "}]}}
        document = convert_payload([raw])
        fragment = document.metadata.logical_table_groups[0].fragments[0]
        self.assertIsNone(fragment.source_span.page_index)
        self.assertIsNone(fragment.source_span.source_index_in_page)
        self.assertEqual(fragment.raw_html_pointer, "/content/payload~1a~0b/0/html")
        self.assertEqual(fragment.raw_html, raw["content"]["payload/a~b"][0]["html"])

    def test_conservative_candidates_do_not_merge_or_raise_fatal(self):
        cases = {
            "page_gap": ({"page_idx": 2}, "page_not_adjacent"),
            "nonmonotonic": ({"page_idx": 0}, None),
            "no_ordinal": ({"html": table_html("address 2")}, "insufficient_evidence"),
            "stride": ({"html": table_html("3")}, "insufficient_evidence"),
            "radix_change": ({"html": table_html("0x02")}, "insufficient_evidence"),
            "header_case": ({"html": table_html("2", header="NUMBER")}, "header_mismatch"),
            "columns": ({"html": table_html("2", header="", columns=3)}, "column_mismatch"),
            "no_header": ({"html": table_html("2", header="")}, "insufficient_evidence"),
            "rowspan": ({"html": '<table><tr><th rowspan="2">编号</th><th>参数</th></tr><tr><td>2</td><td>源值</td></tr></table>'}, "header_mismatch"),
            "bbox_missing": ({"bbox": None}, "missing_geometry"),
            "high_top": ({"bbox": [10, 25.01, 90, 90]}, "geometry_mismatch"),
            "horizontal": ({"bbox": [15.01, 10, 90, 90]}, "geometry_mismatch"),
        }
        for name, (changes, reason) in cases.items():
            with self.subTest(name=name):
                following = table(1, "2", **changes)
                if following.get("bbox", "absent") is None:
                    following.pop("bbox")
                document = convert_payload([table(0, "1"), following])
                self.assertEqual(len(document.metadata.logical_table_groups), 2)
                diagnostics = [item for group in document.metadata.logical_table_groups for item in group.diagnostics]
                self.assertEqual([item.reason for item in diagnostics], [reason] if reason else [])
                self.assertTrue(document.metadata.mineru_validation.can_convert)
                self.assertEqual(len([flag for flag in document.quality_flags if flag.startswith("table_continuation_uncertain:")]),
                                 1 if reason else 0)

    def test_geometry_thresholds_and_trusted_page_size_are_required(self):
        boundary = [table(0, "1", bbox=[10, 60, 90, 75]), table(1, "2", bbox=[15, 25, 95, 90])]
        self.assertEqual(len(convert_payload(boundary).metadata.logical_table_groups), 1)
        before_boundary = [table(0, "1", bbox=[10, 60, 90, 74.99]), table(1, "2")]
        document = convert_payload(before_boundary)
        self.assertEqual(document.metadata.logical_table_groups[-1].diagnostics[0].reason, "geometry_mismatch")
        document = convert_payload(boundary, geometry=False)
        self.assertEqual(document.metadata.logical_table_groups[-1].diagnostics[0].reason, "missing_geometry")

    def test_complex_spans_header_units_and_page_body_boundaries_stay_conservative(self):
        complex_html = '<table><tr><th rowspan="2">编号</th><th>参数</th></tr><tr><td>1</td><td>源值</td></tr></table>'
        document = convert_payload([table(0, "1", html=complex_html, table_caption=["表 24 参数"]),
                                    table(1, "2", html=complex_html, table_caption=["表 24（续）"])])
        self.assertEqual(document.metadata.logical_table_groups[-1].diagnostics[0].reason, "unreliable_columns")
        for header in ("编号 (V)", "编号 2"):
            document = convert_payload([table(0, "1"), table(1, "2", html=table_html("2", header=header))])
            self.assertEqual(document.metadata.logical_table_groups[-1].diagnostics[0].reason, "header_mismatch")
        document = convert_payload([table(2, "1"), table(1, "2")])
        self.assertEqual(document.metadata.logical_table_groups[-1].diagnostics[0].reason, "page_not_adjacent")
        document = convert_payload([table(0, "1"), table(1, "2"), {"type": "text", "text": "页内后续正文", "page_idx": 0}])
        self.assertEqual(document.metadata.logical_table_groups[-1].diagnostics[0].reason, "page_body_boundary")

    def test_explicit_margin_roles_can_be_skipped_but_body_heading_caption_and_footnote_cut_groups(self):
        margin = {"type": "footer", "text": "页脚", "page_idx": 0}
        document = convert_payload([table(0, "1"), margin, table(1, "2")])
        self.assertEqual(len(document.metadata.logical_table_groups), 1)
        cases = [
            [table(0, "1"), {"type": "paragraph", "text": "正文", "page_idx": 0}, table(1, "2")],
            [table(0, "1"), {"type": "heading", "text": "新节", "page_idx": 1}, table(1, "2")],
            [table(0, "1"), table(1, "2", table_caption=["独立参数表"])],
            [table(0, "1", table_footnote=["注：本表结束"]), table(1, "2")],
        ]
        for payload in cases:
            with self.subTest(payload=payload):
                document = convert_payload(payload)
                self.assertEqual(len(document.metadata.logical_table_groups), 2)
                self.assertFalse(any(group.diagnostics for group in document.metadata.logical_table_groups))

    def test_old_v1_and_v2_do_not_fabricate_groups(self):
        for version in ("document-ir/v1", "document-ir/v2"):
            document = DocumentIR.model_validate({"schema_version": version,
                "metadata": {"job_id": "old", "source_file_id": "old", "document_language": "en"}})
            self.assertIsNone(document.metadata.logical_table_groups_version)
            self.assertEqual(document.metadata.logical_table_groups, [])

    def test_legacy_independent_source_spans_do_not_declare_logical_groups(self):
        """任意块的独立 spans 沿用 metadata 开放合同，不由 IR-205 声明或接管。"""
        for block_type in ("paragraph", "table"):
            payload = {"schema_version": "document-ir/v2",
                       "metadata": {"job_id": "old", "source_file_id": "old", "document_language": "en"},
                       "blocks": [{"id": "legacy", "type": block_type, "order": 0,
                                   "metadata": {"source_spans": [{"page_index": 0, "legacy": True}]}}]}
            document = DocumentIR.model_validate(payload)
            self.assertIsNone(document.metadata.logical_table_groups_version)
            self.assertEqual(DocumentIR.model_validate_json(document.model_dump_json()).blocks[0].metadata,
                             payload["blocks"][0]["metadata"])

    def test_corrupt_persisted_contracts_fail_explicitly(self):
        document = convert_payload([table(0, "1"), table(1, "2"), {"type": "table_placeholder", "page_idx": 2}])
        valid = document.model_dump(mode="json")
        def fragment(payload):
            return payload["metadata"]["logical_table_groups"][0]["fragments"][0]
        mutations = {
            "duplicate": lambda p: p["metadata"]["logical_table_groups"].append(copy.deepcopy(p["metadata"]["logical_table_groups"][0])),
            "raw_hash": lambda p: fragment(p).update(raw_html_sha256="0" * 64),
            "bad_pointer": lambda p: fragment(p).update(raw_html_pointer="/missing"),
            "owner": lambda p: fragment(p).update(render_owner_block_id="missing"),
            "spans": lambda p: p["blocks"][0]["metadata"].update(source_spans=[]),
            "cycle": lambda p: p["metadata"]["logical_table_groups"][0]["continuation_relations"][0].update(to_fragment_id=fragment(p)["fragment_id"]),
            "version": lambda p: p["metadata"].update(logical_table_groups_version=None),
            "counts": lambda p: p["metadata"]["logical_table_group_counts"].update(fragment_count=99),
            "missing_fragment": lambda p: p["metadata"]["logical_table_groups"][0]["fragments"].pop(),
        }
        for name, mutate in mutations.items():
            with self.subTest(name=name):
                payload = copy.deepcopy(valid)
                mutate(payload)
                with self.assertRaisesRegex(ValueError, "逻辑表格来源合同无效"):
                    DocumentIR.model_validate(payload)

    def test_contentful_empty_and_nonempty_bad_html_keep_fail_closed(self):
        for empty in ({"type": "table", "img_path": "", "html": "", "text": "正文", "page_idx": 1},
                      {"type": "table", "img_path": "", "html": "<table><tr><td>损坏", "page_idx": 1}):
            with self.subTest(empty=empty), self.assertRaises((DocumentIntegrityError, MinerUArtifactValidationError)):
                convert_payload([table(0, "1"), empty])


if __name__ == "__main__":
    unittest.main()

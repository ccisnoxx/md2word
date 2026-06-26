"""Validate the local MinerU phase-2 quality matrix runner."""

from __future__ import annotations

import os
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase, mock

from docx import Document

from scripts.dev import run_local_mineru_quality_matrix as quality_matrix


class LocalMinerUQualityMatrixTests(TestCase):
    def test_phase2_core_sample_set_covers_required_samples(self) -> None:
        sample_ids = {sample.sample_id for sample in quality_matrix.PHASE2_CORE_SAMPLES}

        self.assertEqual(
            sample_ids,
            {"OPA228_P1", "OPA228_P2", "OPA228_P6", "OPA228_P7", "RS2G04"},
        )
        p0_ids = {sample.sample_id for sample in quality_matrix.PHASE2_CORE_SAMPLES if sample.priority == "P0"}
        self.assertEqual(p0_ids, {"OPA228_P1", "OPA228_P2", "OPA228_P6", "OPA228_P7"})

    def test_ru2zh_acceptance_sample_set_targets_russian_mineru_package(self) -> None:
        samples = quality_matrix.SAMPLE_SETS[quality_matrix.RU2ZH_ACCEPTANCE_SAMPLE_SET]

        self.assertEqual([sample.sample_id for sample in samples], ["RU_DIODE_ANODE_TOP"])
        self.assertEqual(samples[0].document_language, "ru")
        self.assertIn("диод анод сверху", samples[0].input_glob)

    def test_dry_run_writes_report_with_mineru_stats_and_qa_command(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            package = self._make_mineru_package(root)
            output_dir = root / "runs"

            with mock.patch.object(quality_matrix, "_find_sample_matches", return_value=[package]):
                report = quality_matrix.run_matrix(
                    quality_matrix.MatrixConfig(
                        output_dir=output_dir,
                        sample_ids=("OPA228_P1",),
                        modes=("direct",),
                        timestamped=False,
                        dry_run=True,
                    )
                )

            result = report["results"][0]
            self.assertEqual(report["summary"]["dry_run"], 1)
            self.assertEqual(result["status"], "dry_run")
            self.assertEqual(result["mineru_stats"]["page_count"], 1)
            self.assertEqual(result["mineru_stats"]["block_type_counts"]["title"], 1)
            self.assertIn("render_docx.py", " ".join(result["qa"]["render_command"]))
            self.assertTrue((output_dir / "report.json").exists())
            self.assertTrue((output_dir / "report.md").exists())
            self.assertTrue((output_dir / "manual-signoff-template.yaml").exists())

    def test_en2zh_without_key_or_response_is_reported_as_blocker(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            package = self._make_mineru_package(root)

            with (
                mock.patch.object(quality_matrix, "_find_sample_matches", return_value=[package]),
                mock.patch.dict(os.environ, {quality_matrix.OPENAI_API_KEY_ENV: ""}, clear=False),
                mock.patch.object(quality_matrix.formal_mineru_conversion, "run_conversion") as run_conversion,
            ):
                report = quality_matrix.run_matrix(
                    quality_matrix.MatrixConfig(
                        output_dir=root / "runs",
                        sample_ids=("OPA228_P1",),
                        modes=("en2zh",),
                        timestamped=False,
                    )
                )

            result = report["results"][0]
            self.assertEqual(result["status"], "blocked")
            self.assertEqual(result["error_code"], "translation_api_key_missing")
            self.assertEqual(result["issues"][0]["severity"], "P1")
            self.assertEqual(report["summary"]["blocked"], 1)
            self.assertEqual(report["summary"]["phase3_entry_blocker"], "environment_or_input_blockers_must_be_resolved")
            run_conversion.assert_not_called()

    def test_ru2zh_without_key_or_response_is_reported_as_blocker(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            package = self._make_mineru_package(root)

            with (
                mock.patch.object(quality_matrix, "_find_sample_matches", return_value=[package]),
                mock.patch.dict(os.environ, {quality_matrix.OPENAI_API_KEY_ENV: ""}, clear=False),
                mock.patch.object(quality_matrix.formal_mineru_conversion, "run_conversion") as run_conversion,
            ):
                report = quality_matrix.run_matrix(
                    quality_matrix.MatrixConfig(
                        output_dir=root / "runs",
                        sample_set=quality_matrix.RU2ZH_ACCEPTANCE_SAMPLE_SET,
                        sample_ids=("RU_DIODE_ANODE_TOP",),
                        modes=("ru2zh",),
                        timestamped=False,
                    )
                )

            result = report["results"][0]
            self.assertEqual(result["status"], "blocked")
            self.assertEqual(result["error_code"], "translation_api_key_missing")
            self.assertIn("ru2zh requires", result["error_message"])
            run_conversion.assert_not_called()

    def test_successful_direct_run_records_docx_quality_flags(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            package = self._make_mineru_package(root)

            def fake_run_conversion(**kwargs):
                docx_path = Path(kwargs["output_path"])
                docx_path.parent.mkdir(parents=True, exist_ok=True)
                document = Document()
                document.add_heading("1 Features", level=1)
                document.add_paragraph("The OPA228 input offset voltage remains editable.")
                table = document.add_table(rows=2, cols=2)
                table.cell(0, 0).text = "Parameter"
                table.cell(0, 1).text = "Value"
                table.cell(1, 0).text = "Noise"
                table.cell(1, 1).text = "3 nV/√Hz"
                document.save(docx_path)

            with (
                mock.patch.object(quality_matrix, "_find_sample_matches", return_value=[package]),
                mock.patch.object(quality_matrix.formal_mineru_conversion, "run_conversion", side_effect=fake_run_conversion),
            ):
                report = quality_matrix.run_matrix(
                    quality_matrix.MatrixConfig(
                        output_dir=root / "runs",
                        sample_ids=("OPA228_P1",),
                        modes=("direct",),
                        timestamped=False,
                    )
                )

            result = report["results"][0]
            self.assertEqual(result["status"], "succeeded")
            self.assertTrue(result["docx_openable"])
            self.assertTrue(result["docx_sha256"])
            self.assertGreaterEqual(result["docx_quality"]["table_count"], 1)
            self.assertIn("manual_review_required", result["qa"])
            self.assertTrue((root / "runs" / "manual-signoff-template.yaml").exists())

    def _make_mineru_package(self, root: Path) -> Path:
        package = root / "mineru-package"
        image_dir = package / "images"
        image_dir.mkdir(parents=True)
        (image_dir / "plot.jpg").write_bytes(b"fake-image")
        (package / "sample_origin.pdf").write_bytes(b"%PDF-1.4\n")
        (package / "sample_content_list_v2.json").write_text(
            """
[
  [
    {"type": "title", "content": {"title_content": [{"type": "text", "content": "1 Features"}]}, "bbox": [0, 0, 10, 10]},
    {"type": "paragraph", "content": {"paragraph_content": [{"type": "text", "content": "Body"}]}, "bbox": [0, 20, 10, 30]},
    {"type": "table", "content": {"html": "<table><tr><td>A</td></tr></table>"}, "bbox": [0, 40, 10, 50]}
  ]
]
""",
            encoding="utf-8",
        )
        return package

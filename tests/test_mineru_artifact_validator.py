"""覆盖原始 MinerU 输入在正式转换边界上的观察合同。"""
from __future__ import annotations

import io
import random
import json
import tempfile
import hashlib
import stat
import warnings
from contextlib import contextmanager
from unittest.mock import patch
import unittest
import zipfile
from pathlib import Path

from PIL import Image, EpsImagePlugin

from app.services.document_ir import MinerUToDocumentIRService
from app.services.document_ir.artifact_validator import (
    MinerUArtifactValidationError, MinerUArtifactValidator, MinerUValidationReport,
)
from app.services.document_ir.mineru_package import (
    MinerUPackageError, MinerUValidationPolicy, extract_zip, resolve_asset_path,
)
from app.services.document_ir.integrity import DocumentIntegrityError
from app.services.renderer import DocumentIRWordRenderer
from app.services.renderer.evidence import (
    build_renderer_event_summary, build_renderer_evidence, build_renderer_metadata, ensure_renderer_integrity,
)
from docx import Document
from docx.image.image import Image as WordImage


def png_bytes() -> bytes:
    buffer = io.BytesIO()
    with Image.new("RGB", (2, 2), "white") as image:
        image.save(buffer, format="PNG")
    return buffer.getvalue()


@contextmanager
def warnings_suppressed_duplicate_zip():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        yield


class MinerUArtifactValidatorTests(unittest.TestCase):
    def convert(self, payload, members=None, *, policy=None):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "result.zip"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("content_list_v2.json", json.dumps(payload))
                for name, data in (members or {}).items():
                    archive.writestr(name, data)
            return MinerUToDocumentIRService(**({"policy": policy} if policy else {})).convert_zip(
                path, job_id="job", source_file_id="source", document_language="en"
            )

    def test_invalid_page_is_not_silently_dropped(self):
        with self.assertRaises(ValueError):
            self.convert([[{"type": "text", "text": "Body"}], 42])

    def test_traversal_is_not_repaired_with_package_basename(self):
        with self.assertRaises(ValueError):
            self.convert([{"type": "image", "img_path": "../images/a.png", "page_idx": 0}],
                         {"images/a.png": png_bytes()})

    def test_nonfinite_bbox_stops_before_mapping_without_changing_input(self):
        block = {"type": "text", "text": "12 µV", "page_idx": 0,
                 "bbox": [0, 0, float("nan"), 30]}
        with patch.object(MinerUToDocumentIRService, "_convert_blocks") as mapping:
            with self.assertRaises(MinerUArtifactValidationError) as caught:
                self.convert([block])
        mapping.assert_not_called()
        self.assertEqual(caught.exception.report.first_fatal.code, "mineru_bbox_invalid")
        self.assertEqual(block["text"], "12 µV")
        self.assertNotEqual(block["bbox"][2], block["bbox"][2])
        json.dumps(caught.exception.report.model_dump(mode="json"), allow_nan=False)

    def test_raw_nested_wrapper_locator_inventory_and_empty_page(self):
        payload = {"content_list": [[], [
            {"type": "paragraph", "content": {"paragraph_content": [
                {"type": "text", "content": "12 "}, {"type": "equation_inline", "content": "V_{OS}"}]}},
            {"block_type": "text", "text": "Legacy"},
        ]]}
        result = self.convert(payload)
        report = result.validation_report
        self.assertEqual(report.inventory["input_shape"], "wrapper_nested")
        self.assertEqual(report.inventory["raw_page_count"], 2)
        self.assertEqual(report.inventory["raw_block_count"], 2)
        empty_page = next(issue for issue in report.warnings if issue.code == "mineru_page_empty")
        self.assertEqual(empty_page.json_pointer, "/content_list/0")
        self.assertEqual(empty_page.page_index, 0)
        self.assertFalse(empty_page.stops_conversion)
        self.assertEqual([block.metadata["mineru_source"]["json_pointer"] for block in result.document.blocks],
                         ["/content_list/1/0", "/content_list/1/1"])
        self.assertEqual([block.metadata["mineru_source"]["source_index"] for block in result.document.blocks], [0, 1])
        self.assertEqual([block.page_index for block in result.document.blocks], [1, 1])
        self.assertEqual(result.document.blocks[0].text, "12 $V_{OS}$")
        self.assertEqual(result.document.blocks[1].text, "Legacy")

    def test_real_empty_margin_runs_are_warnings_and_empty_body_still_fails(self):
        blocks = [{"type": "text", "text": "12 µV"}]
        blocks.extend({"type": source_type, "content": {field: []}} for source_type, field in (
            ("page_header", "page_header_content"), ("page_footer", "page_footer_content"), ("page_number", "page_number_content")))
        result = self.convert([blocks])
        self.assertEqual(len(result.document.blocks), 4)
        self.assertEqual(result.validation_report.issue_counts["warning"]["mineru_margin_payload_empty"], 3)
        self.assertEqual(result.validation_report.fatal_count, 0)
        self.assertEqual([block.layout.role for block in result.document.blocks], ["body", "page_header", "page_footer", "page_number"])
        for block in ({"type": "paragraph", "content": {"paragraph_content": []}},
                      {"type": "title", "content": {"title_content": []}},
                      {"type": "formula", "content": {"math_content": ""}},
                      {"type": "page_header", "content": {"page_header_content": False}}):
            with self.subTest(block=block), self.assertRaises(MinerUArtifactValidationError):
                self.convert([[block]])

    def test_page_and_block_failures_have_source_pointer(self):
        cases = [
            ([[{"type": "text", "text": "Body"}], False], "mineru_page_invalid", "/1"),
            ([[False, {"type": "text", "text": "Body"}]], "mineru_block_invalid", "/0/0"),
            ([{"type": "text", "text": "Body", "page_idx": True}], "mineru_page_invalid", "/0"),
            ([{"type": "text", "text": "Body", "page_idx": 1.5}], "mineru_page_invalid", "/0"),
            ([{"type": "text", "text": "Body", "page_idx": -1}], "mineru_page_invalid", "/0"),
            ([{"type": "text", "text": "Body", "page_idx": "1.0"}], "mineru_page_invalid", "/0"),
            ([[{"type": "text", "text": "Body", "page_idx": 1}]], "mineru_page_invalid", "/0/0"),
            ([{"type": "text", "text": "Body", "page_idx": 0, "page_index": 1}], "mineru_page_invalid", "/0"),
            ([{"type": 5, "text": "Body", "page_idx": 0}], "mineru_block_invalid", "/0"),
        ]
        for payload, code, pointer in cases:
            with self.subTest(payload=payload), self.assertRaises(MinerUArtifactValidationError) as caught:
                self.convert(payload)
            report = caught.exception.report
            self.assertFalse(report.can_convert)
            self.assertEqual(report.first_fatal.code, code)
            self.assertEqual(report.first_fatal.json_pointer, pointer)

    def test_flat_decimal_legacy_page_preserves_order(self):
        result = self.convert([{"type": "text", "text": "First", "page_idx": "2"},
                               {"type": "text", "text": "Second", "page_index": 0},
                               {"type": "text", "text": "Third", "page_idx": 2}])
        self.assertEqual([block.text for block in result.document.blocks], ["First", "Second", "Third"])
        self.assertEqual([block.metadata["mineru_source"]["source_index_in_page"] for block in result.document.blocks], [0, 0, 1])
        self.assertIn("legacy_page_index_coerced", result.validation_report.issue_counts["warning"])
        self.assertIn("mineru_page_order_nonmonotonic", result.validation_report.issue_counts["warning"])
        self.assertEqual(result.validation_report.issue_counts["warning"]["mineru_page_gap"], 1)
        gap = next(issue for issue in result.validation_report.warnings if issue.code == "mineru_page_gap")
        self.assertEqual(gap.json_pointer, "/2")
        self.assertEqual(gap.page_index, 2)
        self.assertEqual(gap.facts, {"previous_page_index": 0, "index_gap": 1})
        self.assertFalse(gap.stops_conversion)

    def test_flat_missing_page_warns_without_guessing_and_renderer_still_blocks(self):
        payload = {"content_list": [{"type": "text", "text": "First", "page_idx": 2},
                                    {"type": "text", "text": "12 µV"},
                                    {"type": "text", "text": "Third", "page_idx": 0}]}
        result = self.convert(payload)
        self.assertEqual([block.text for block in result.document.blocks], ["First", "12 µV", "Third"])
        self.assertEqual([block.page_index for block in result.document.blocks], [2, None, 0])
        issue = next(issue for issue in result.validation_report.warnings if issue.code == "mineru_page_missing")
        self.assertEqual(issue.json_pointer, "/content_list/1")
        self.assertEqual(issue.source_index, 1)
        self.assertIsNone(issue.page_index)
        self.assertFalse(issue.stops_conversion)
        self.assertNotIn("page_idx", result.document.blocks[1].metadata)
        word = DocumentIRWordRenderer().render_to_bytes(result.document, template_path=None, config_path=None,
                                                        image_loader=lambda _asset: None)
        with self.assertRaises(DocumentIntegrityError):
            ensure_renderer_integrity(build_renderer_evidence(result.document, word))

    def test_invalid_bbox_reasons_and_json_are_finite(self):
        values = [(None, "wrong_length_or_type"), (False, "wrong_length_or_type"),
                  ([0, 1, 2], "wrong_length_or_type"), ([False, 0, 20, 30], "non_numeric"),
                  (["0", 0, 20, 30], "non_numeric"), ([0, 0, float("inf"), 30], "non_finite"),
                  ([0, 0, 0, 30], "non_positive_area"), ([20, 0, 0, 30], "non_positive_area"),
                  ([0, 10, 20, 0], "non_positive_area")]
        for bbox, reason in values:
            with self.subTest(reason=reason):
                with self.assertRaises(MinerUArtifactValidationError) as caught:
                    self.convert([{"type": "text", "text": "±15 V", "bbox": bbox, "page_idx": 0}])
                issue = caught.exception.report.first_fatal
                self.assertEqual(issue.code, "mineru_bbox_invalid")
                self.assertEqual(issue.severity, "blocker")
                self.assertTrue(issue.stops_conversion)
                self.assertEqual(issue.facts["reason"], reason)
                json.dumps(caught.exception.report.model_dump(mode="json"), allow_nan=False)
        result = self.convert([{"type": "text", "text": "Negative origin", "bbox": [-20, -10, 1, 2], "page_idx": 0}])
        self.assertEqual(result.document.blocks[0].layout.bbox, (-20., -10., 1., 2.))

    def test_rich_payload_rejects_silent_stringification_and_schema_only_content(self):
        cases = [
            {"type": "paragraph", "content": {"paragraph_content": [{"type": "text", "content": 12}]}},
            {"type": "paragraph", "content": {"paragraph_content": "Body"}},
            {"type": "paragraph", "content": {"schema_type": "paragraph"}},
            {"type": "list", "items": [False]},
            {"type": "equation_interline", "content": {"math_type": "latex"}},
            {"type": "equation_interline", "content": {"math_type": "latex", "description": "必须保留的公式说明"}},
            {"type": "image", "img_path": "images/a.png", "caption": [{"type": "text", "content": {"wrong": "value"}}]},
        ]
        for block in cases:
            with self.subTest(block=block), self.assertRaises(MinerUArtifactValidationError) as caught:
                self.convert([[block]], {"images/a.png": png_bytes()})
            self.assertIn("mineru_block_payload_invalid", caught.exception.report.issue_counts["blocker"])
        legacy = self.convert([[{"block_type": "list", "items": [{"text": "12 µV"}, "±15 V"]}]])
        self.assertEqual(legacy.document.blocks[0].list_items, ["12 µV", "±15 V"])
        unknown = self.convert([[{"type": "custom", "text": "可保留正文"}]])
        self.assertEqual(unknown.document.blocks[0].text, "可保留正文")
        self.assertIn("unsupported_block_type", unknown.validation_report.issue_counts["warning"])
        for block in ({"type": "custom", "content": {"schema_type": "custom"}},
                      {"type": "custom", "description": "不能静默丢失的正文"}):
            with self.subTest(block=block), self.assertRaises(MinerUArtifactValidationError) as caught:
                self.convert([[block]])
            self.assertEqual(caught.exception.report.first_fatal.code, "unsupported_block_type")
        with self.assertRaises(MinerUArtifactValidationError):
            self.convert([[{"type": "paragraph", "description": "不能静默丢失的正文"}]])

    def test_invalid_selected_v2_and_aux_json_do_not_fall_back(self):
        for member_name in ("content_list_v2.json", "layout.json"):
            with self.subTest(member=member_name), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "result.zip"
                with zipfile.ZipFile(path, "w") as archive:
                    archive.writestr("content_list.json", json.dumps([{ "type": "text", "text": "Legacy", "page_idx": 0}]))
                    if member_name == "layout.json":
                        archive.writestr("content_list_v2.json", json.dumps([[{"type": "text", "text": "Body"}]]))
                    archive.writestr(member_name, b"invalid json")
                with self.assertRaises(MinerUArtifactValidationError) as caught:
                    MinerUToDocumentIRService().convert_zip(path, job_id="job", source_file_id="source", document_language="en")
                self.assertEqual(caught.exception.report.first_fatal.code, "mineru_json_invalid")
                self.assertEqual(caught.exception.report.content_list_path, "content_list_v2.json")
                self.assertEqual(caught.exception.report.package_sha256, hashlib.sha256(path.read_bytes()).hexdigest())

    def test_empty_or_unsupported_document_fails(self):
        for payload in ([], [[]], {}, "Body"):
            with self.subTest(payload=payload), self.assertRaises(MinerUArtifactValidationError):
                self.convert(payload)

    def test_zip_paths_and_special_entries_are_rejected_before_any_write(self):
        names = ["../outside", "/outside", "C:/outside", "C:\\outside", "\\\\server\\outside", "a/../../outside", "a//file", "a//"]
        for name in names:
            with self.subTest(name=name), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                output = root / "output"
                output.mkdir()
                path = root / "result.zip"
                with zipfile.ZipFile(path, "w") as archive:
                    archive.writestr("safe.txt", b"safe")
                    archive.writestr(name, b"unsafe")
                with self.assertRaises(MinerUPackageError):
                    extract_zip(path, output)
                self.assertEqual(list(output.iterdir()), [])
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for kind in (stat.S_IFLNK, stat.S_IFIFO):
                path = root / "result.zip"
                info = zipfile.ZipInfo("special")
                info.create_system = 3
                info.external_attr = (kind | 0o777) << 16
                with zipfile.ZipFile(path, "w") as archive:
                    archive.writestr(info, "../outside")
                with self.assertRaises(MinerUPackageError):
                    extract_zip(path, root / "output")

    def test_zip_file_directory_conflicts_are_preflighted_in_both_orders(self):
        for names in (("a", "a/child"), ("a/child", "a"), ("a", "a/sub/"), ("a/sub/", "a")):
            with self.subTest(names=names), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                output = root / "output"
                output.mkdir()
                path = root / "result.zip"
                with zipfile.ZipFile(path, "w") as archive:
                    archive.writestr("safe.txt", b"safe")
                    for name in names:
                        archive.writestr(name, b"")
                with self.assertRaises(MinerUPackageError) as caught:
                    extract_zip(path, output)
                self.assertEqual(caught.exception.code, "mineru_path_invalid")
                self.assertEqual(caught.exception.facts, {"reason": "file_directory_conflict"})
                self.assertEqual(list(output.iterdir()), [])

    def test_zip_collisions_follow_target_filesystem_before_extraction(self):
        for names in (("a", "A/child"), ("A/child", "a"), ("a", "A/sub/"),
                      ("A/sub/", "a"), ("a", "A")):
            with self.subTest(names=names), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                probe = root / "case-probe"
                probe.touch()
                case_sensitive = not (root / "CASE-PROBE").exists()
                probe.unlink()
                output = root / "output"
                output.mkdir()
                archive_path = root / "result.zip"
                with zipfile.ZipFile(archive_path, "w") as archive:
                    archive.writestr("safe.txt", b"safe")
                    for name in names:
                        archive.writestr(name, b"")
                if case_sensitive:
                    extract_zip(archive_path, output)
                    self.assertEqual((output / "safe.txt").read_bytes(), b"safe")
                    self.assertTrue(all((output / name).exists() for name in names))
                else:
                    with self.assertRaises(MinerUPackageError) as caught:
                        extract_zip(archive_path, output)
                    self.assertEqual(caught.exception.code, "mineru_path_invalid")
                    self.assertEqual(caught.exception.facts["reason"],
                                     "file_directory_conflict" if any("/" in name for name in names)
                                     else "duplicate_normalized_path")
                    self.assertEqual(list(output.iterdir()), [])

    def test_zip_duplicate_corrupt_crc_and_limits_have_reports(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "result.zip"
            with warnings_suppressed_duplicate_zip():
                with zipfile.ZipFile(path, "w") as archive:
                    archive.writestr("content_list.json", b"[]")
                    archive.writestr("content_list.json", b"[]")
            with self.assertRaises(MinerUArtifactValidationError) as caught:
                MinerUToDocumentIRService().convert_zip(path, job_id="j", source_file_id="s", document_language="en")
            self.assertEqual(caught.exception.report.first_fatal.facts["reason"], "duplicate_normalized_path")
            path.write_bytes(b"not a zip")
            with self.assertRaises(MinerUArtifactValidationError) as caught:
                MinerUToDocumentIRService().convert_zip(path, job_id="j", source_file_id="s", document_language="en")
            self.assertEqual(caught.exception.report.first_fatal.code, "mineru_package_invalid")
            with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_STORED) as archive:
                archive.writestr("content_list.json", b"[]")
            data = bytearray(path.read_bytes())
            with zipfile.ZipFile(path) as archive:
                info = archive.infolist()[0]
                offset = info.header_offset + 30 + len(info.filename.encode()) + len(info.extra)
            data[offset] ^= 1
            path.write_bytes(data)
            with self.assertRaises(MinerUArtifactValidationError) as caught:
                MinerUToDocumentIRService().convert_zip(path, job_id="j", source_file_id="s", document_language="en")
            self.assertEqual(caught.exception.report.first_fatal.code, "mineru_package_invalid")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "result.zip"
            with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("content_list.json", b"[]")
            data = bytearray(path.read_bytes())
            with zipfile.ZipFile(path) as archive:
                info = archive.infolist()[0]
                offset = info.header_offset + 30 + len(info.filename.encode()) + len(info.extra)
            data[offset] = 7
            path.write_bytes(data)
            with self.assertRaises(MinerUArtifactValidationError) as caught:
                MinerUToDocumentIRService().convert_zip(path, job_id="j", source_file_id="s", document_language="en")
            self.assertEqual(caught.exception.report.first_fatal.code, "mineru_package_invalid")
        for policy in (MinerUValidationPolicy(max_entries=1), MinerUValidationPolicy(max_file_bytes=1),
                       MinerUValidationPolicy(max_total_bytes=1), MinerUValidationPolicy(max_json_bytes=1)):
            with self.subTest(policy=policy), self.assertRaises(MinerUArtifactValidationError) as caught:
                self.convert([[{"type": "text", "text": "Body"}]], {"extra.txt": b"extra"}, policy=policy)
            self.assertEqual(caught.exception.report.first_fatal.code, "mineru_package_limit_exceeded")

    def test_asset_exact_resolution_ambiguity_and_symlink_never_reads_outside(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory).resolve()
            package = root / "package"
            (package / "nested/images").mkdir(parents=True)
            (package / "images").mkdir()
            (package / "images/a.png").write_bytes(png_bytes())
            (package / "nested/images/a.png").write_bytes(png_bytes())
            with self.assertRaises(MinerUPackageError):
                resolve_asset_path("images/a.png", package, package / "nested")
            self.assertEqual(resolve_asset_path("missing/a.png", package)[0], None)
            sentinel = root / "sentinel.png"
            sentinel.write_bytes(png_bytes())
            (package / "images/link.png").symlink_to(sentinel)
            with patch.object(Path, "read_bytes", side_effect=AssertionError("不得读取包外文件")):
                with self.assertRaises(MinerUPackageError):
                    resolve_asset_path("images/link.png", package)
            for raw in ("../images/a.png", "/images/a.png", "C:/images/a.png", "C:\\images\\a.png", "\\\\server\\a.png", "images/../a.png"):
                with self.subTest(raw=raw), self.assertRaises(MinerUPackageError):
                    resolve_asset_path(raw, package)
            self.assertEqual(resolve_asset_path("images/", package)[1], "not_regular_file")
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert([[{"type": "image", "img_path": "images/a.png"}]], {"nested/images/a.png": png_bytes()})
        self.assertEqual(caught.exception.report.first_fatal.code, "missing_image_asset")

    def test_content_list_directory_asset_resolution_is_exact_and_recorded(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "result.zip"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("nested/content_list_v2.json", json.dumps([[{"type": "chart", "content": {"image_source": {"path": "images/a.png"}}}]]))
                archive.writestr("nested/images/a.png", png_bytes())
            result = MinerUToDocumentIRService().convert_zip(path, job_id="j", source_file_id="s", document_language="en")
        self.assertEqual(result.image_assets[0].source_path, "nested/images/a.png")
        self.assertEqual(result.document.blocks[0].metadata["mineru_asset_resolution"], "content_list_directory")

    def test_images_empty_invalid_word_unsupported_and_local_bomb(self):
        for data in (b"", b"fake-image", b"<svg/> "):
            with self.subTest(data=data), self.assertRaises(MinerUArtifactValidationError) as caught:
                self.convert([[{"type": "image", "img_path": "images/a.png"}]], {"images/a.png": data})
            self.assertEqual(caught.exception.report.first_fatal.code, "invalid_image_asset")
        buffer = io.BytesIO()
        with Image.new("RGB", (2, 2)) as image:
            image.save(buffer, format="WEBP")
        # Word 不识别的 WebP 不再失败，而是无损转 PNG（IMAGE-001 例外）。
        result = self.convert([[{"type": "image", "img_path": "images/a.webp"}]], {"images/a.webp": buffer.getvalue()})
        self.assertEqual(result.image_assets[0].mime_type, "image/png")
        with patch.object(Image, "MAX_IMAGE_PIXELS", 3), self.assertRaises(MinerUArtifactValidationError):
            self.convert([[{"type": "image", "img_path": "images/a.png"}]], {"images/a.png": png_bytes()})
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert([[{"type": "image", "content": {"image_source": {"path": "images/a.png"}},
                            "img_path": "images/missing.png"}]], {"images/a.png": png_bytes()})
        self.assertEqual(caught.exception.report.first_fatal.code, "missing_image_asset")
        result = self.convert([[{"type": "chart", "img_path": "images/a.bin"}]], {"images/a.bin": png_bytes()})
        self.assertEqual(result.image_assets[0].mime_type, "image/png")
        self.assertEqual(result.image_assets[0].content, png_bytes())

    def test_allowed_raster_formats_keep_word_and_reference_boundaries(self):
        for image_format in ("PNG", "JPEG", "GIF", "BMP", "TIFF", "WEBP"):
            with self.subTest(image_format=image_format):
                buffer = io.BytesIO()
                with Image.new("RGB", (2, 2), "white") as image:
                    image.save(buffer, format=image_format)
                image_bytes = buffer.getvalue()
                table = self.convert([[{"type": "table", "html": "<table><tr><td>12 V</td></tr></table>", "img_path": "images/a.png"}]],
                                     {"images/a.png": image_bytes})
                self.assertEqual(table.validation_report.table_checks[0].reference_status, "valid")
                result = self.convert([[{"type": "image", "img_path": "images/a.png"}]], {"images/a.png": image_bytes})
                normalized = result.validation_report.issue_counts["warning"].get("image_asset_normalized", 0)
                if image_format == "WEBP":
                    self.assertEqual(result.image_assets[0].mime_type, "image/png")
                    self.assertEqual(result.image_assets[0].normalized_from[0], "WEBP")
                    self.assertEqual(normalized, 1)
                else:
                    # Word 可识别的格式保持原字节，不产生规范化 warning。
                    self.assertEqual(result.image_assets[0].content, image_bytes)
                    self.assertIsNone(result.image_assets[0].normalized_from)
                    self.assertEqual(normalized, 0)

    def test_eps_never_invokes_ghostscript_for_image_or_table_reference(self):
        buffer = io.BytesIO()
        with Image.new("RGB", (2, 2), "white") as image:
            image.save(buffer, format="EPS")
        eps = buffer.getvalue()
        with patch.object(EpsImagePlugin, "Ghostscript", side_effect=AssertionError("不得调用 Ghostscript")) as ghostscript:
            with self.assertRaises(MinerUArtifactValidationError) as caught:
                self.convert([[{"type": "image", "img_path": "images/a.png"}]], {"images/a.png": eps})
            self.assertEqual(caught.exception.report.first_fatal.code, "invalid_image_asset")
            self.assertEqual(caught.exception.report.first_fatal.facts["reason"], "not_allowed_raster_format")
            result = self.convert([[{"type": "table", "html": "<table><tr><td>12 V</td></tr></table>", "img_path": "images/a.png"}]],
                                  {"images/a.png": eps})
            self.assertEqual(result.validation_report.issue_counts["warning"]["invalid_table_reference"], 1)
            self.assertTrue(result.validation_report.can_convert)
            ghostscript.assert_not_called()

    def test_html_four_classes_and_quality_blockers_remain_separate(self):
        payload = [[
            {"type": "table", "html": "<table><tr><td>12 µV</td></tr></table>"},
            {"type": "table", "html": '<table><tr><td colspan="bad">±15 V</td></tr></table>'},
            {"type": "table", "html": "", "img_path": "images/a.png"},
            {"type": "table", "html": "", "img_path": "images/"},
        ]]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory).resolve()
            (root / "images").mkdir()
            (root / "images/a.png").write_bytes(png_bytes())
            report = MinerUValidationReport(package_sha256="0" * 64)
            validated = MinerUArtifactValidator(report).validate(payload, package_root=root, content_list_path=root / "content_list_v2.json")
        self.assertEqual([check.classification for check in report.table_checks], [
            "html_present_and_parseable", "html_present_but_suspicious", "html_missing_reference_present", "html_missing_reference_missing"])
        self.assertEqual(report.table_checks[-1].reference_status, "invalid")
        self.assertEqual(report.table_checks[2].reference_resolution, "package_root")
        self.assertEqual(report.blocker_count, 2)
        self.assertEqual(report.fatal_count, 0)
        self.assertTrue(report.can_convert)
        self.assertEqual(len(validated.blocks), 4)
        result = self.convert([payload[0][:2]])
        word = DocumentIRWordRenderer().render_to_bytes(result.document, template_path=None, config_path=None, image_loader=lambda _asset: None)
        document = Document(io.BytesIO(word))
        self.assertEqual([table.cell(0, 0).text for table in document.tables], ["12 µV", "±15 V"])
        evidence = build_renderer_evidence(result.document, word)
        self.assertEqual(evidence["blocking_quality_flags"], [])
        self.assertEqual(evidence["mineru_validation"], result.validation_report.model_dump(mode="json"))
        for html in ("text without table", 42):
            with self.subTest(html=html), self.assertRaises((DocumentIntegrityError, MinerUArtifactValidationError)) as caught:
                self.convert([[{"type": "table", "html": html}]])
            if isinstance(caught.exception, DocumentIntegrityError):
                report_data = caught.exception.evidence["mineru_validation"]
                self.assertIn("invalid_table_html", report_data["issue_counts"]["blocker"])

    def test_unconsumable_html_stops_before_ir_construction_with_report(self):
        with patch("app.services.document_ir.mineru_converter.DocumentIR", side_effect=AssertionError("不得构建无效 IR")), \
             self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert([[{"type": "table", "html": "<table><tr><td>unclosed"}]])
        report = caught.exception.report
        self.assertEqual(report.blocker_count, 1)
        self.assertFalse(report.can_convert)
        self.assertEqual(report.first_fatal.code, "invalid_table_html")
        self.assertEqual(report.first_fatal.json_pointer, "/0/0")
        self.assertFalse(report.table_checks[0].parseable)

    def test_leading_dot_asset_compatibility_and_large_span_are_reported(self):
        result = self.convert([[{"type": "image", "img_path": "./images/a.png"}]], {"images/a.png": png_bytes()})
        self.assertEqual(result.image_assets[0].source_path, "images/a.png")
        result = self.convert([[{"type": "table", "html": '<table><tr><td colspan="' + "9" * 5000 + '">12 V</td></tr></table>'}]])
        self.assertIn("suspicious_table_html", result.validation_report.issue_counts["warning"])
        self.assertTrue(result.validation_report.can_convert)
        result = self.convert([[{"type": "table", "html": "<table><tr><td>12 V</td></tr></table>"},
                                {"type": "table", "html": "", "img_path": "./images/"}]])
        self.assertEqual(len(result.document.blocks), 1)
        self.assertEqual(result.validation_report.issue_counts["blocker"]["empty_table_html"], 1)

    def test_selected_v2_list_runs_cannot_silently_drop_string_item(self):
        payload = [[{"type": "list", "content": {"list_items": [
            {"item_content": [{"type": "text", "content": "First"}]},
            {"item_content": "12 V"},
        ]}}]]
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert(payload)
        issue = caught.exception.report.first_fatal
        self.assertEqual(issue.code, "mineru_block_payload_invalid")
        self.assertEqual(issue.field, "/0/0/content/list_items/1/item_content")

    def test_empty_authoritative_rich_runs_do_not_consume_sibling_description(self):
        for source_type in ("paragraph", "custom"):
            with self.subTest(source_type=source_type):
                with self.assertRaises(MinerUArtifactValidationError) as caught:
                    self.convert([[{"type": source_type, "content": {"paragraph_content": [], "description": "12 V"}}]])
                self.assertEqual(caught.exception.report.first_fatal.code, "mineru_block_payload_invalid")
                self.assertEqual(caught.exception.report.first_fatal.field, "/0/0/content/description")

    def test_v2_list_run_nodes_fail_with_report_before_consumer_getters(self):
        for runs in ([False], [[{"type": "text", "content": "12 V"}]]):
            with self.subTest(runs=runs):
                with self.assertRaises(MinerUArtifactValidationError) as caught:
                    self.convert([[{"type": "list", "content": {"list_items": [{"item_content": runs}]}}]])
                self.assertEqual(caught.exception.report.first_fatal.code, "mineru_block_payload_invalid")
                self.assertEqual(caught.exception.report.first_fatal.field, "/0/0/content/list_items/0/item_content/0")

    def test_selected_list_and_inferred_list_do_not_drop_semantic_siblings(self):
        cases = [
            ({"type": "list", "content": {"list_items": [
                {"item_content": [{"type": "text", "content": "First"}]},
                {"item_content": [], "description": "12 V"},
            ]}}, "/0/0/content/list_items/1/description"),
            ({"type": "paragraph", "content": {"paragraph_content": [{"type": "text", "content": "- First\n- Second"}],
                                                     "description": "12 V"}}, "/0/0/content/description"),
        ]
        for block, field in cases:
            with self.subTest(field=field):
                with self.assertRaises(MinerUArtifactValidationError) as caught:
                    self.convert([[block]])
                self.assertEqual(caught.exception.report.first_fatal.field, field)

    def test_repeated_canonical_images_reuse_reads_checks_and_bytes(self):
        buffer = io.BytesIO()
        with Image.frombytes("RGB", (37, 37), random.Random(201).randbytes(37 * 37 * 3)) as image:
            image.save(buffer, format="PNG")
        image_bytes = buffer.getvalue()
        self.assertGreater(len(image_bytes) * 10, 5_000)
        blocks = [{"type": "image", "img_path": "images/a.png"} for _ in range(10)]
        original_read = Path.read_bytes
        image_reads = []
        def traced_read(path):
            if path.suffix == ".png":
                image_reads.append(path.resolve())
            return original_read(path)
        with patch.object(Path, "read_bytes", traced_read), patch.object(Image, "open", wraps=Image.open) as opened, \
             patch.object(WordImage, "from_blob", wraps=WordImage.from_blob) as word_header:
            result = self.convert([blocks], {"images/a.png": image_bytes}, policy=MinerUValidationPolicy(max_total_bytes=5_000))
        self.assertEqual(word_header.call_count, 1)
        self.assertEqual(len(image_reads), 1)
        self.assertEqual(opened.call_count, 2)
        self.assertEqual(result.validation_report.inventory["image_check_count"], 10)
        self.assertEqual(len(result.image_assets), 10)
        self.assertTrue(all(asset.content is result.image_assets[0].content for asset in result.image_assets))
        self.assertEqual([block.metadata["mineru_source"]["json_pointer"] for block in result.document.blocks],
                         [f"/0/{index}" for index in range(10)])

    def test_cached_reference_raster_normalizes_once_for_all_references(self):
        buffer = io.BytesIO()
        with Image.new("RGB", (2, 2), "white") as image:
            image.save(buffer, format="WEBP")
        blocks = [{"type": "table", "html": "<table><tr><td>12 V</td></tr></table>", "img_path": "images/a.webp"},
                  {"type": "image", "img_path": "images/a.webp"},
                  {"type": "image", "img_path": "./images/a.webp"},
                  {"type": "table", "html": "<table><tr><td>±15 V</td></tr></table>", "img_path": "images/a.webp"}]
        with patch.object(Image, "open", wraps=Image.open) as opened, patch.object(WordImage, "from_blob", wraps=WordImage.from_blob) as word_header:
            result = self.convert([blocks], {"images/a.webp": buffer.getvalue()})
        report = result.validation_report
        # 同一文件：安全探测两次 + 规范化一次；Word 检查原字节与新字节各一次。
        self.assertEqual(opened.call_count, 3)
        self.assertEqual(word_header.call_count, 2)
        self.assertEqual([check.reference_status for check in report.table_checks], ["valid", "valid"])
        self.assertEqual(report.blocker_count, 0)
        self.assertEqual(report.issue_counts["warning"]["image_asset_normalized"], 2)
        self.assertEqual([issue.json_pointer for issue in report.warnings if issue.code == "image_asset_normalized"],
                         ["/0/1", "/0/2"])
        self.assertTrue(all(asset.content is result.image_assets[0].content for asset in result.image_assets))

    def test_word_rejected_rasters_are_normalized_to_png_losslessly(self):
        def encoded(mode, image_format, **options):
            buffer = io.BytesIO()
            with Image.new(mode, (6, 4), "red" if mode in {"RGB", "RGBA"} else 0) as image:
                if mode == "RGBA":
                    image.putpixel((0, 0), (0, 0, 255, 0))
                image.save(buffer, format=image_format, **options)
            return buffer.getvalue()

        jfif = encoded("RGB", "JPEG")
        no_app0 = jfif[:2] + jfif[4 + int.from_bytes(jfif[4:6], "big"):]
        cases = {"cmyk_jpeg": (encoded("CMYK", "JPEG"), "JPEG", "RGB"),
                 "jpeg_without_app0": (no_app0, "JPEG", "RGB"),
                 "webp_with_alpha": (encoded("RGBA", "WEBP", lossless=True), "WEBP", "RGBA")}
        for name, (data, source_format, mode) in cases.items():
            with self.subTest(name):
                first = self.convert([[{"type": "image", "img_path": "images/a.bin", "page_idx": 0}]], {"images/a.bin": data})
                second = self.convert([[{"type": "image", "img_path": "images/a.bin", "page_idx": 0}]], {"images/a.bin": data})
                asset = first.image_assets[0]
                self.assertEqual((asset.mime_type, asset.normalized_from),
                                 ("image/png", (source_format, hashlib.sha256(data).hexdigest())))
                self.assertEqual(asset.content, second.image_assets[0].content)
                self.assertEqual(asset.stored_name, "a.png")
                WordImage.from_blob(asset.content)
                with Image.open(io.BytesIO(asset.content)) as png:
                    self.assertEqual((png.format, png.mode, png.size), ("PNG", mode, (6, 4)))
                    if name == "webp_with_alpha":
                        self.assertEqual(png.getpixel((0, 0))[3], 0)
                warning = next(issue for issue in first.validation_report.warnings if issue.code == "image_asset_normalized")
                self.assertEqual((warning.page_index, warning.facts), (0, {"reason": "word_incompatible_raster"}))
                self.assertEqual(first.document.blocks[0].metadata["mineru_asset_normalized_from"], source_format)
                summary = build_renderer_event_summary({"mineru_validation": first.validation_report.model_dump(mode="json")})
                self.assertEqual(summary["mineru_validation"]["issue_counts"]["warning"]["image_asset_normalized"], 1)

    def test_normalization_keeps_dpi_and_skips_word_accepted_cmyk_jpeg(self):
        from app.services.document_ir.artifact_validator import _word_compatible_png
        buffer = io.BytesIO()
        with Image.new("RGB", (6, 4), "red") as image:
            image.save(buffer, format="PNG", dpi=(300, 300))
        with Image.open(io.BytesIO(_word_compatible_png(buffer.getvalue(), max_bytes=10**6))) as png:
            self.assertEqual(tuple(round(value) for value in png.info["dpi"]), (300, 300))
        self.assertIsNone(_word_compatible_png(buffer.getvalue(), max_bytes=10))
        # 指定 DPI 时 Pillow 为 CMYK JPEG 写入 JFIF 头，Word 可识别，必须保持原字节。
        cmyk = io.BytesIO()
        with Image.new("CMYK", (6, 4), 0) as image:
            image.save(cmyk, format="JPEG", dpi=(300, 300))
        result = self.convert([[{"type": "image", "img_path": "images/a.jpg"}]], {"images/a.jpg": cmyk.getvalue()})
        self.assertEqual((result.image_assets[0].content, result.image_assets[0].normalized_from), (cmyk.getvalue(), None))

    def test_unnormalizable_word_rejected_raster_keeps_original_failure(self):
        buffer = io.BytesIO()
        with Image.new("RGB", (2, 2)) as image:
            image.save(buffer, format="WEBP")
        with patch("app.services.document_ir.artifact_validator._word_compatible_png", return_value=None), \
                self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert([[{"type": "image", "img_path": "images/a.webp"}]], {"images/a.webp": buffer.getvalue()})
        self.assertEqual((caught.exception.report.first_fatal.code, caught.exception.report.first_fatal.facts),
                         ("invalid_image_asset", {"reason": "unreadable_or_unsupported_image"}))

    def test_table_reference_raster_is_normalized_without_quality_warning(self):
        buffer = io.BytesIO()
        with Image.new("RGB", (2, 2), "white") as image:
            image.save(buffer, format="WEBP")
        result = self.convert([[{"type": "table", "html": "<table><tr><td>12 V</td></tr></table>", "img_path": "images/a.webp"}]],
                              {"images/a.webp": buffer.getvalue()})
        self.assertEqual(result.validation_report.table_checks[0].reference_status, "valid")
        self.assertEqual([(asset.mime_type, asset.content[:8]) for asset in result.table_reference_assets],
                         [("image/png", b"\x89PNG\r\n\x1a\n")])
        self.assertNotIn("image_asset_normalized", result.validation_report.issue_counts["warning"])

    def test_whitespace_source_type_keeps_v2_image_payload_and_caption(self):
        result = self.convert([[{"type": " image ", "content": {"image_source": {"path": "images/a.png"},
                                  "image_caption": [{"type": "text", "content": "Figure 1. 12 V"}]}}]],
                              {"images/a.png": png_bytes()})
        self.assertEqual(result.document.blocks[0].type, "image")
        self.assertEqual(result.document.blocks[0].captions, ["Figure 1. 12 V"])
        self.assertEqual(result.document.blocks[0].metadata["source_type"], "image")
        self.assertEqual(result.validation_report.inventory["type_counts"], {"image": 1})
        self.assertEqual(result.image_assets[0].content, png_bytes())
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert([[{"type": " ", "text": "12 V"}]])
        self.assertEqual(caught.exception.report.first_fatal.code, "mineru_block_invalid")

    def test_mapping_preserves_existing_header_heading_table_and_image_priority(self):
        payload = [[
            {"type": "header", "text": "Header", "img_path": "images/a.png"},
            {"type": "title", "text": "Title", "table_body": "<table><tr><td>Not selected</td></tr></table>", "img_path": "images/a.png"},
            {"block_type": "text", "table_body": "<table><tr><td>12 V</td></tr></table>", "img_path": "images/a.png"},
            {"block_type": "text", "text": "Figure", "img_path": "images/a.png", "items": ["Not selected"]},
        ]]
        result = self.convert(payload, {"images/a.png": png_bytes()})
        self.assertEqual([block.type for block in result.document.blocks], ["paragraph", "heading", "table", "image"])
        self.assertEqual(result.validation_report.inventory["image_check_count"], 1)
        self.assertEqual(result.validation_report.inventory["table_reference_check_count"], 1)
        self.assertEqual(len(result.image_assets), 1)

    def test_inferred_table_cannot_publish_an_invisible_missing_payload(self):
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert([[{"block_type": "text", "table_body": "12 V"}]])
        self.assertEqual(caught.exception.report.first_fatal.code, "mineru_block_payload_invalid")
        self.assertIn("invalid_table_html", caught.exception.report.issue_counts["blocker"])

    def test_legacy_inferred_image_and_table_use_the_same_validation_as_conversion(self):
        for source_type in ("text", "custom"):
            with self.subTest(source_type=source_type):
                image = self.convert([[{"block_type": source_type, "text": "Figure 1", "img_path": "images/a.png"}]],
                                     {"images/a.png": png_bytes()})
                self.assertEqual(image.document.blocks[0].type, "image")
                self.assertEqual(image.image_assets[0].content, png_bytes())
                self.assertEqual(image.validation_report.inventory["image_check_count"], 1)
                table = self.convert([[{"block_type": source_type, "table_body": "<table><tr><td>12 V</td></tr></table>"}]])
                word = DocumentIRWordRenderer().render_to_bytes(table.document, template_path=None, config_path=None, image_loader=lambda _asset: None)
                self.assertEqual(Document(io.BytesIO(word)).tables[0].cell(0, 0).text, "12 V")
                self.assertEqual(table.validation_report.inventory["table_reference_check_count"], 1)
        for path, data, code in (("../images/a.png", png_bytes(), "mineru_path_invalid"),
                                 ("images/a.png", b"bad", "invalid_image_asset")):
            with self.subTest(path=path):
                with self.assertRaises(MinerUArtifactValidationError) as caught:
                    self.convert([[{"block_type": "custom", "text": "Figure 1", "img_path": path}]], {"images/a.png": data})
                self.assertEqual(caught.exception.report.first_fatal.code, code)

    def test_list_containers_and_typed_runs_reject_unconsumed_semantic_leaves(self):
        cases = [
            ({"type": "list", "content": {"list_items": [{"item_content": [{"type": "text", "content": "First"}]}],
                                           "description": "12 V"}}, "/0/0/content/description"),
            ({"type": "list", "items": ["First"], "text": "12 V"}, "/0/0/text"),
            ({"type": "paragraph", "content": {"paragraph_content": [{"type": "text", "content": "First", "description": "12 V"}]}},
             "/0/0/content/paragraph_content/0/description"),
            ({"type": "list", "content": {"list_items": [{"item_content": [{"type": "text", "content": "First", "description": "12 V"}]}]}},
             "/0/0/content/list_items/0/item_content/0/description"),
        ]
        for block, field in cases:
            with self.subTest(field=field):
                with self.assertRaises(MinerUArtifactValidationError) as caught:
                    self.convert([[block]])
                issue = caught.exception.report.first_fatal
                self.assertEqual(issue.code, "mineru_block_payload_invalid")
                self.assertEqual(issue.field, field)
        valid = self.convert([[{"type": "list", "items": [{"text": "First", "description": "12 V"}]}]])
        self.assertEqual(valid.document.blocks[0].list_items, ["First", "12 V"])

    def test_zero_word_dpi_is_fatal_for_required_image_but_valid_for_reference(self):
        buffer = io.BytesIO()
        with Image.new("RGB", (2, 2), "white") as image:
            image.save(buffer, format="PNG", dpi=(0.1, 0.1))
        content = buffer.getvalue()
        self.assertEqual(WordImage.from_blob(content).horz_dpi, 0)
        table = {"type": "table", "html": "<table><tr><td>12 V</td></tr></table>", "img_path": "images/a.png"}
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert([[table, {"type": "image", "img_path": "images/a.png"}]], {"images/a.png": content})
        report = caught.exception.report
        self.assertEqual(report.first_fatal.code, "invalid_image_asset")
        self.assertEqual(report.first_fatal.json_pointer, "/0/1")
        self.assertEqual(report.table_checks[0].reference_status, "valid")

    def test_selected_list_items_validate_all_known_semantic_containers(self):
        first = [{"type": "text", "content": "First"}]
        extra = [{"type": "text", "content": "12 V"}]
        cases = [
            ({"type": "list", "content": {"list_items": [{"item_content": first, "paragraph_content": extra}]}},
             "/0/0/content/list_items/0/paragraph_content"),
            ({"type": "list", "items": [{"type": "text", "content": "First", "item_content": extra}]},
             "/0/0/items/0/item_content"),
            ({"type": "paragraph", "content": {"paragraph_content": [{"type": "text", "content": "First", "item_content": extra}]}},
             "/0/0/content/paragraph_content/0/item_content"),
        ]
        for block, field in cases:
            with self.subTest(field=field):
                with self.assertRaises(MinerUArtifactValidationError) as caught:
                    self.convert([[block]])
                self.assertEqual(caught.exception.report.first_fatal.code, "mineru_block_payload_invalid")
                self.assertEqual(caught.exception.report.first_fatal.field, field)
        preserved = self.convert([[{"type": "list", "items": [{"text": "First", "description": "12 V"},
                                                               {"item_content": extra}]}]])
        self.assertEqual(preserved.document.blocks[0].list_items, ["First", "12 V", "12 V"])
        empty = self.convert([[{"type": "list", "content": {"list_items": [{"item_content": first, "paragraph_content": []}]}}]])
        self.assertEqual(empty.document.blocks[0].list_items, ["First"])

    def test_table_report_uses_the_html_selected_by_the_converter(self):
        good = "<table><tr><td>12 V</td></tr></table>"
        bad = "<table><tr><td>unclosed"
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert([[{"type": "table", "html": good, "table_body": bad}]])
        check = caught.exception.report.table_checks[0]
        self.assertFalse(check.parseable)
        self.assertEqual(check.classification, "html_present_but_suspicious")
        self.assertIn("invalid_table_html", caught.exception.report.issue_counts["blocker"])
        self.assertEqual(caught.exception.report.first_fatal.code, "invalid_table_html")

    def test_report_limit_stops_at_total_details_and_remains_deterministic(self):
        blocks = [[{"type": "text", "text": f"Body {index}"} for index in range(250)]]
        policy = MinerUValidationPolicy(max_issue_details=2)
        reports = []
        for _ in range(2):
            with self.assertRaises(MinerUArtifactValidationError) as caught:
                self.convert(blocks, policy=policy)
            reports.append(caught.exception.report)
        first, second = reports
        self.assertEqual(first, second)
        self.assertEqual(first.inventory["raw_block_count"], 2)
        self.assertEqual(first.warning_count, 2)
        self.assertEqual(first.issue_counts["warning"]["mineru_bbox_missing"], 2)
        self.assertLessEqual(len(first.warnings) + len(first.blockers), 2)
        self.assertTrue(first.details_truncated)
        self.assertFalse(first.complete)
        self.assertFalse(first.can_convert)
        self.assertEqual(first.first_fatal.code, "report_limit_exceeded")
        self.assertEqual(first.first_fatal.facts["reason"], "issue_details")
        self.assertEqual(first.issue_counts["blocker"]["report_limit_exceeded"], 1)
        mixed = [[{"type": "custom", "content": {"schema_type": "custom"}}]]
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert(mixed, policy=policy)
        self.assertEqual(caught.exception.report.warning_count, 1)
        self.assertIn("unsupported_block_type", caught.exception.report.issue_counts["blocker"])
        self.assertIn("report_limit_exceeded", caught.exception.report.issue_counts["blocker"])

    def test_policy_defaults_are_versioned_and_keep_json_limit(self):
        policy = MinerUValidationPolicy()
        self.assertEqual(policy.version, "mineru-input-policy/v2")
        self.assertEqual(policy.max_entries, 10_000)
        self.assertEqual(policy.max_file_bytes, 128 * 1024 * 1024)
        self.assertEqual(policy.max_total_bytes, 1024 * 1024 * 1024)
        self.assertEqual(policy.max_json_bytes, 64 * 1024 * 1024)
        self.assertEqual(policy.max_issue_details, 1000)
        self.assertEqual(policy.max_report_bytes, 512 * 1024)
        with self.assertRaisesRegex(ValueError, "至少为 4096"):
            MinerUValidationPolicy(max_report_bytes=4095)

    def test_policy_rejects_non_integer_or_non_positive_limits(self):
        limits = ("max_entries", "max_file_bytes", "max_total_bytes", "max_json_bytes",
                  "max_issue_details", "max_report_bytes")
        for field in limits:
            for value in (float("nan"), float("inf"), True, False, 1.5, 1.0, "1", None, 0, -1):
                with self.subTest(field=field, value=value), self.assertRaises(ValueError):
                    MinerUValidationPolicy(**{field: value})

    def test_escaped_isolated_surrogates_stop_with_safe_report_and_input_hash(self):
        cases = [
            [{"type": "\ud800", "text": "body", "page_idx": 0}],
            [[{"type": "text", "text": "body\udfff"}]],
            {"content_list": [[{"type": "text", "text": "body", "\ud800": "metadata"}]]},
            [[{"type": "text", "text": "body", "metadata": {"nested": ["safe", "\ud800x\udfff"]}}]],
        ]
        for payload in cases:
            with self.subTest(payload=payload):
                with self.assertRaises(MinerUArtifactValidationError) as caught:
                    self.convert(payload)
                report = caught.exception.report
                self.assertEqual(report.first_fatal.code, "mineru_json_invalid")
                self.assertEqual(report.first_fatal.facts, {"reason": "isolated_unicode_surrogate"})
                self.assertFalse(report.can_convert)
                self.assertEqual(report.content_list_sha256, hashlib.sha256(json.dumps(payload).encode("utf-8")).hexdigest())
                self.assertEqual(len(report.package_sha256), 64)
                encoded = json.dumps(report.model_dump(mode="json"), ensure_ascii=False, allow_nan=False).encode("utf-8")
                self.assertLess(len(encoded), MinerUValidationPolicy().max_report_bytes)

    def test_direct_validator_rejects_surrogate_before_report_fields(self):
        report = MinerUValidationReport(package_sha256="0" * 64)
        validator = MinerUArtifactValidator(report)
        with tempfile.TemporaryDirectory() as directory, self.assertRaises(MinerUArtifactValidationError):
            root = Path(directory)
            validator.validate([{"type": "\ud800", "text": "body", "page_idx": 0}],
                               package_root=root, content_list_path=root / "content_list.json")
        self.assertEqual(report.first_fatal.code, "mineru_json_invalid")
        self.assertEqual(report.inventory, {})
        json.dumps(report.model_dump(mode="json"), ensure_ascii=False).encode("utf-8")

    def test_valid_surrogate_pairs_chinese_and_emoji_keep_text(self):
        for text in ("中文 \ud83d\ude00 12 µV", "中文 😀 12 µV"):
            with self.subTest(text=text):
                result = self.convert([[{"type": "text", "text": text, "bbox": [0, 0, 10, 10]}]])
                self.assertEqual(result.document.blocks[0].text, "中文 😀 12 µV")
                self.assertTrue(result.validation_report.can_convert)
                result.document.model_dump_json().encode("utf-8")

    def test_report_byte_limit_covers_table_checks_and_package_errors(self):
        policy = MinerUValidationPolicy(max_report_bytes=4096)
        blocks = [[{"type": "table", "html": "<table><tr><td>12 V</td></tr></table>",
                    "img_path": "images/a.png", "bbox": [0, 0, 10, 10]} for _ in range(50)]]
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert(blocks, {"images/a.png": png_bytes()}, policy=policy)
        report = caught.exception.report
        self.assertFalse(report.complete)
        self.assertFalse(report.can_convert)
        self.assertEqual(report.first_fatal.facts["reason"], "report_bytes")
        self.assertLess(report.inventory["raw_block_count"], 50)
        self.assertGreater(report.table_check_count, 0)
        self.assertEqual(report.warning_count, 0)
        self.assertGreater(report.omitted_table_check_count, 0)
        self.assertLess(len(json.dumps(report.model_dump(mode="json"), ensure_ascii=False).encode("utf-8")), 4096)
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert([[{"type": "text", "text": "Body"}]], {"../outside": b"bad"}, policy=policy)
        self.assertEqual(caught.exception.report.first_fatal.code, "mineru_path_invalid")
        self.assertLess(len(json.dumps(caught.exception.report.model_dump(mode="json")).encode("utf-8")), 4096)

    def test_selected_path_and_large_page_metadata_have_bounded_failure_reports(self):
        policy = MinerUValidationPolicy(max_report_bytes=4096)
        report = MinerUValidationReport(package_sha256="0" * 64, content_list_path="目录/" * 1000)
        with self.assertRaises(MinerUArtifactValidationError):
            MinerUArtifactValidator(report, policy=policy)
        self.assertIsNone(report.content_list_path)
        self.assertFalse(report.complete)
        self.assertLess(len(json.dumps(report.model_dump(mode="json"), ensure_ascii=False).encode("utf-8")), 4096)
        report = MinerUValidationReport(package_sha256="0" * 64)
        validator = MinerUArtifactValidator(report, policy=policy)
        with self.assertRaises(MinerUArtifactValidationError):
            validator.record_normalization([], [], filtered_page_index=10 ** 4200, filtered_block_count=2)
        self.assertFalse(report.complete)
        self.assertEqual(report.normalization["filtered_block_count"], 2)
        self.assertLess(len(json.dumps(report.model_dump(mode="json"), ensure_ascii=False).encode("utf-8")), 4096)

    def test_normalization_keeps_raw_issues_and_distinguishes_resolved_tables(self):
        real = {"type": "table", "html": "<table><tr><td>12 V</td></tr></table>", "bbox": [0, 0, 10, 10]}
        empty = {"type": "table", "html": "", "img_path": "images/", "bbox": [0, 0, 10, 10]}
        result = self.convert([[real], [empty]])
        report = result.validation_report
        self.assertEqual(report.inventory["raw_block_count"], 2)
        self.assertEqual(report.effective_inventory["block_count"], 1)
        self.assertEqual(report.effective_inventory["type_counts"], {"table": 1})
        self.assertEqual(report.issue_counts["blocker"]["empty_table_html"], 1)
        self.assertEqual(report.normalization["resolved_issue_counts"], {"blocker": {"empty_table_html": 1}})
        self.assertEqual([entry["json_pointer"] for entry in report.normalization["merged_table_placeholders"]], ["/1/0"])
        receipt = report.normalization["merged_table_placeholders"][0]
        self.assertEqual(receipt["target"]["json_pointer"], "/0/0")
        self.assertEqual(receipt["rule_id"], "mineru_empty_table_continuation/v1")
        self.assertEqual(receipt["reason"], "strict_empty_table_placeholder")
        for payload in ([[empty]], [[real], [{**empty, "text": "12 µV"}]]):
            with self.subTest(payload=payload), self.assertRaises(DocumentIntegrityError) as caught:
                self.convert(payload)
            failure = caught.exception.evidence["mineru_validation"]
            self.assertEqual(failure["normalization"]["resolved_issue_counts"], {"blocker": {}})
            self.assertEqual(failure["effective_inventory"]["block_count"], len(payload))
            self.assertIn("source_ir_integrity:payload_missing:table", caught.exception.evidence["blocking_quality_flags"])

    def test_normalization_metadata_cannot_cross_report_byte_limit(self):
        real = {"type": "table", "html": "<table><tr><td>12 V</td></tr></table>", "bbox": [0, 0, 10, 10]}
        empty = {"type": "table", "html": "", "img_path": "images/", "bbox": [0, 0, 10, 10]}
        blocks = [[real, *[dict(empty) for _ in range(8)]]]
        completed = self.convert(blocks).validation_report
        before_normalization = completed.model_copy(deep=True)
        before_normalization.effective_inventory = {}
        before_normalization.normalization = {}
        raw_size = len(json.dumps(before_normalization.model_dump(mode="json"), ensure_ascii=False).encode("utf-8"))
        final_size = len(json.dumps(completed.model_dump(mode="json"), ensure_ascii=False).encode("utf-8"))
        self.assertGreater(final_size, raw_size)
        byte_limit = (raw_size + final_size) // 2
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert(blocks, policy=MinerUValidationPolicy(max_report_bytes=byte_limit))
        report = caught.exception.report
        self.assertFalse(report.complete)
        self.assertIn("report_limit_exceeded", report.issue_counts["blocker"])
        self.assertEqual(report.normalization["merged_table_placeholder_count"], 8)
        self.assertLess(len(json.dumps(report.model_dump(mode="json"), ensure_ascii=False).encode("utf-8")), byte_limit)

    def test_formula_optional_reference_still_obeys_raw_path_safety(self):
        formula = {"type": "formula", "content": {"math_content": "V_{OS}", "image_source": {"path": "images/missing.png"}}}
        result = self.convert([[formula]])
        self.assertEqual(result.document.blocks[0].latex, "V_{OS}")
        self.assertEqual(result.validation_report.fatal_count, 0)
        self.assertEqual(result.image_assets, [])
        formula["content"]["image_source"]["path"] = "../private.png"
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert([[formula]])
        self.assertEqual(caught.exception.report.first_fatal.code, "mineru_path_invalid")
        self.assertEqual(caught.exception.report.first_fatal.json_pointer, "/0/0")

    def test_all_legacy_asset_strings_are_checked_before_mapping(self):
        cases = [
            ({"type": "formula", "content": {"math_content": "V_{OS}", "image_source": [
                "images/missing.png", "../private.png"]}}, "/0/0/content/image_source/1"),
            ({"type": "formula", "content": {"math_content": "V_{OS}", "image_source": [
                "images/missing.png", "../private"]}}, "/0/0/content/image_source/1"),
            ({"type": "formula", "content": {"math_content": "V_{OS}", "image_source": {
                "path": "images/missing.png", "content": [["images/also-missing.png", "../private.png"]]}}},
             "/0/0/content/image_source/content/0/1"),
            ({"type": "image", "content": ["images/missing.png", "../private.png"]}, "/0/0/content/1"),
        ]
        for block, field in cases:
            with self.subTest(block=block):
                with patch.object(MinerUToDocumentIRService, "_convert_blocks") as mapping:
                    with self.assertRaises(MinerUArtifactValidationError) as caught:
                        self.convert([[block]])
                mapping.assert_not_called()
                issue = caught.exception.report.first_fatal
                self.assertEqual(issue.code, "mineru_path_invalid")
                self.assertEqual(issue.field, field)
                self.assertNotIn("../private.png", caught.exception.report.model_dump_json())
        with patch.object(MinerUToDocumentIRService, "_convert_blocks") as mapping:
            with self.assertRaises(MinerUArtifactValidationError) as caught:
                self.convert([[{"type": "formula", "content": {"math_content": "V_{OS}", "image_source": {
                    "path": ["images/missing.png", "../private.png"]}}}]])
        mapping.assert_not_called()
        self.assertEqual(caught.exception.report.first_fatal.code, "mineru_block_payload_invalid")
        self.assertEqual(caught.exception.report.first_fatal.facts["reason"], "expected_asset_path_string")

    def test_legacy_asset_strings_keep_optional_and_first_reference_semantics(self):
        result = self.convert([[{"type": "formula", "content": {"math_content": "V_{OS}", "image_source": [
            "images/missing.png", "images/also-missing.png"]}}]])
        self.assertEqual(result.document.blocks[0].latex, "V_{OS}")
        self.assertTrue(result.validation_report.can_convert)
        self.assertEqual(result.image_assets, [])
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert([[{"type": "image", "content": {"image_source": ["images/missing.png", "images/a.png"]}}]],
                         {"images/a.png": png_bytes()})
        self.assertEqual(caught.exception.report.first_fatal.code, "missing_image_asset")
        result = self.convert([[{"type": "image", "content": {"image_source": ["images/a.png", "images/missing.png"]}}]],
                              {"images/a.png": png_bytes()})
        self.assertEqual(result.image_assets[0].source_path, "images/a.png")
        self.assertEqual(result.image_assets[0].content, png_bytes())
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert([[{"type": "image", "content": {"image_source": "assets/a.bin"}}]],
                         {"assets/a.bin": png_bytes()})
        self.assertEqual(caught.exception.report.first_fatal.code, "missing_image_asset")

    def test_asset_string_collection_does_not_scan_body_runs(self):
        result = self.convert([[{"type": "paragraph", "content": {"paragraph_content": [
            {"type": "text", "content": "../private.png"}]}}]])
        self.assertEqual(result.document.blocks[0].text, "../private.png")
        self.assertTrue(result.validation_report.can_convert)

    def test_filtered_media_can_be_missing_but_cannot_bypass_raw_safety(self):
        payload = [[{"type": "text", "text": "12 µV", "bbox": [0, 0, 20, 20]}], [
            {"type": "text", "text": "可使用微信扫码进入芯片助手小程序"},
            {"type": "image", "img_path": "images/missing.png"},
        ]]
        result = self.convert(payload)
        self.assertEqual([block.text for block in result.document.blocks], ["12 µV"])
        report = result.validation_report
        self.assertTrue(report.can_convert)
        self.assertEqual(report.normalization["filtered_block_count"], 2)
        self.assertEqual([source["json_pointer"] for source in report.normalization["filtered_blocks"]], ["/1/0", "/1/1"])
        self.assertTrue(all(source["rule_id"] == "final_promotional_page/v1" for source in report.normalization["filtered_blocks"]))
        payload[1][1]["img_path"] = "../private.png"
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert(payload)
        self.assertEqual(caught.exception.report.first_fatal.code, "mineru_path_invalid")
        payload[1][1]["img_path"] = "images/missing.png"
        payload[1][1]["bbox"] = [1, 1, 0, 0]
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert(payload)
        self.assertEqual(caught.exception.report.first_fatal.code, "mineru_bbox_invalid")

    def test_contentful_explicit_placeholder_cannot_discard_payload(self):
        real = {"type": "table", "html": "<table><tr><td>12 V</td></tr></table>"}
        for content in ({"text": "12 µV"}, {"caption": ["Table 2"]}, {"footnote": ["1. 25°C"]},
                        {"html": "<table><tr><td>±15 V</td></tr></table>"}):
            with self.subTest(content=content), self.assertRaises(MinerUArtifactValidationError) as caught:
                self.convert([[real, {"type": "table_placeholder", **content}]])
            self.assertEqual(caught.exception.report.first_fatal.code, "mineru_block_payload_invalid")
            self.assertEqual(caught.exception.report.first_fatal.facts["reason"], "contentful_table_placeholder")

    def test_raw_html_and_typed_content_in_explicit_placeholder_stop_before_mapping(self):
        real = {"type": "table", "html": "<table><tr><td>12 V</td></tr></table>"}
        payloads = [{key: "12 µV"} for key in ("html", "table_html", "table_body")]
        payloads.extend({"content": {key: "12 µV"}} for key in ("html", "table_html", "table_body"))
        payloads.extend([
            {"html": ["12 µV"]},
            {"html": {"content": ["12 µV"]}},
            {"content": [{"html": "12 µV"}]},
            {"content": {"type": "text", "content": "12 µV"}},
            {"content": {"content": "12 µV"}},
            {"content": {"type": "text", "content": "../private.png"}},
        ])
        for content in payloads:
            with self.subTest(content=content):
                payload = [[real, {"type": "table_placeholder", **content}]]
                original = json.dumps(payload)
                with patch.object(MinerUToDocumentIRService, "_convert_blocks") as mapping:
                    with self.assertRaises(MinerUArtifactValidationError) as caught:
                        self.convert(payload)
                mapping.assert_not_called()
                issue = caught.exception.report.first_fatal
                self.assertEqual(issue.code, "mineru_block_payload_invalid")
                self.assertEqual(issue.facts["reason"], "contentful_table_placeholder")
                self.assertEqual(issue.json_pointer, "/0/1")
                self.assertEqual(json.dumps(payload), original)
                self.assertNotIn("12 µV", caught.exception.report.model_dump_json())

    def test_empty_and_structural_explicit_placeholders_remain_mergeable(self):
        html = "<table><tr><td>12 µV</td></tr></table>"
        for content in ({}, {"html": " \n"}, {"content": {"type": "text", "content": ""}},
                        {"content": {"image_source": {"path": "images/"}, "table_type": "simple_table",
                                     "table_nest_level": 0, "html": "", "table_caption": [], "table_footnote": []}}):
            with self.subTest(content=content):
                result = self.convert([[{"type": "table", "html": html}, {"type": "table_placeholder", **content}]])
                self.assertEqual(len(result.document.blocks), 1)
                self.assertEqual(result.document.blocks[0].html, html)
                self.assertTrue(result.validation_report.can_convert)
                self.assertEqual(result.validation_report.normalization["merged_table_placeholder_count"], 1)
                self.assertEqual(result.validation_report.normalization["merged_table_placeholders"][0]["reason"],
                                 "explicit_table_placeholder")

    def test_merged_ordinary_table_cannot_discard_supported_body_content(self):
        real = {"type": "table", "html": "<table><tr><td>12 V</td></tr></table>"}
        for content in ("12 µV", {"type": "text", "content": "12 µV"},
                        [{"type": "text", "content": "12 µV"}]):
            with self.subTest(content=content):
                with self.assertRaises(MinerUArtifactValidationError) as caught:
                    self.convert([[real, {"type": "table", "img_path": "images/", "content": content}]])
                report = caught.exception.report
                self.assertEqual(report.first_fatal.code, "mineru_block_payload_invalid")
                self.assertEqual(report.first_fatal.facts["reason"], "contentful_table_placeholder")
                self.assertEqual(report.first_fatal.json_pointer, "/0/1")
                self.assertFalse(report.can_convert)
                self.assertEqual(report.normalization["resolved_issue_counts"], {"blocker": {}})
                self.assertNotIn("12 µV", report.model_dump_json())

        result = self.convert([[real, {"type": "table", "img_path": "images/", "content": {
            "table_type": "simple_table", "table_nest_level": 0,
            "image_source": {"path": "images/"}, "table_caption": [], "table_footnote": [],
        }}]])
        self.assertTrue(result.validation_report.can_convert)
        self.assertEqual(len(result.document.blocks), 1)
        self.assertEqual(result.document.blocks[0].html, real["html"])
        self.assertEqual(result.validation_report.normalization["merged_table_placeholder_count"], 1)

    def test_nested_content_arrays_cannot_disappear_as_table_continuations(self):
        real = {"type": "table", "html": "<table><tr><td>12 V</td></tr></table>"}
        contents = ({"content": ["12 µV"]}, {"content": {"content": ["12 µV"]}},
                    {"content": [{"content": ["12 µV"]}]})
        for source_type in ("table", "table_placeholder"):
            for content in contents:
                with self.subTest(source_type=source_type, content=content):
                    with self.assertRaises(MinerUArtifactValidationError) as caught:
                        self.convert([[real, {"type": source_type, "img_path": "images/", "content": content}]])
                    self.assertEqual(caught.exception.report.first_fatal.facts["reason"], "contentful_table_placeholder")
                    self.assertFalse(caught.exception.report.can_convert)
                    self.assertNotIn("12 µV", caught.exception.report.model_dump_json())
            for content in ({"content": []}, {"content": {"content": []}},
                            {"table_type": "simple_table", "table_nest_level": 0,
                             "metadata": {"labels": ["height only"]}}):
                with self.subTest(source_type=source_type, structure=content):
                    result = self.convert([[real, {"type": source_type, "img_path": "images/", "content": content}]])
                    self.assertTrue(result.validation_report.can_convert)
                    self.assertEqual(len(result.document.blocks), 1)
                    self.assertEqual(result.document.blocks[0].html, real["html"])

    def test_placeholder_semantic_nodes_reject_invalid_required_field_types(self):
        real = {"type": "table", "html": "<table><tr><td>12 V</td></tr></table>"}
        for source_type in ("table", "table_placeholder"):
            for node_type in ("text", "equation_inline", "inline_formula"):
                for value in (12, None, False, [], {}):
                    for content in ({"type": node_type, "content": value},
                                    {"content": [{"type": node_type, "content": value}]}):
                        with self.subTest(source_type=source_type, node_type=node_type, value=value, content=content):
                            with patch.object(MinerUToDocumentIRService, "_convert_blocks") as mapping:
                                with self.assertRaises(MinerUArtifactValidationError) as caught:
                                    self.convert([[real, {"type": source_type, "img_path": "images/", "content": content}]])
                            mapping.assert_not_called()
                            report = caught.exception.report
                            self.assertEqual(report.first_fatal.code, "mineru_block_payload_invalid")
                            self.assertEqual(report.first_fatal.json_pointer, "/0/1")
                            self.assertFalse(report.can_convert)
        for field in ("html", "table_html", "table_body"):
            for value in (12, None, False, [], {}):
                with self.subTest(field=field, value=value):
                    with patch.object(MinerUToDocumentIRService, "_convert_blocks") as mapping:
                        with self.assertRaises(MinerUArtifactValidationError) as caught:
                            self.convert([[real, {"type": "table_placeholder", "content": {field: value}}]])
                    mapping.assert_not_called()
                    self.assertEqual(caught.exception.report.first_fatal.code, "mineru_block_payload_invalid")
                    self.assertEqual(caught.exception.report.first_fatal.json_pointer, "/0/1")

    def test_unhashable_nested_type_has_a_structured_input_failure(self):
        for value in ([], {}):
            for payload in ([{"type": "table_placeholder", "content": {"metadata": {"type": value}}}],
                            [{"type": "text", "content": {"paragraph_content": [
                                {"type": value, "content": "12 µV"}]}}]):
                with self.subTest(value=value, payload=payload):
                    with self.assertRaises(MinerUArtifactValidationError) as caught:
                        self.convert([payload])
                    self.assertEqual(caught.exception.report.first_fatal.code, "mineru_block_payload_invalid")
                    self.assertEqual(caught.exception.report.first_fatal.json_pointer, "/0/0")
                    self.assertEqual(caught.exception.report.first_fatal.facts["reason"], "expected_type_string")

    def test_nonempty_invalid_html_cannot_use_empty_continuation_receipt(self):
        real = {"type": "table", "html": "<table><tr><td>12 V</td></tr></table>"}
        with self.assertRaises(MinerUArtifactValidationError) as caught:
            self.convert([[real], [{"type": "table", "html": "12 µV", "img_path": "images/"}]])
        self.assertEqual(caught.exception.report.first_fatal.code, "invalid_table_html")
        self.assertEqual(caught.exception.report.first_fatal.json_pointer, "/1/0")
        self.assertEqual(caught.exception.report.normalization["resolved_issue_counts"], {"blocker": {}})

    def test_report_uses_controlled_labels_and_natural_source_order(self):
        secret = "/private/IR201_private_token"
        unknown = self.convert([[{"type": secret, "text": "12 µV"}]])
        serialized = json.dumps(unknown.validation_report.model_dump(mode="json"), allow_nan=False)
        self.assertNotIn(secret, serialized)
        self.assertEqual(unknown.validation_report.inventory["type_counts"], {"unsupported": 1})
        self.assertNotIn(secret, json.dumps(build_renderer_metadata(unknown.document), allow_nan=False))
        self.assertEqual(unknown.document.metadata.source_type_counts, {"unsupported": 1})
        payload = [[{"type": "paragraph", "text": "12 µV", "content": {secret: {"text": False}}}]]
        reports = []
        for _ in range(2):
            with self.assertRaises(MinerUArtifactValidationError) as caught:
                self.convert(payload)
            reports.append(caught.exception.report.model_dump(mode="json"))
        self.assertEqual(reports[0], reports[1])
        self.assertNotIn(secret, json.dumps(reports[0], allow_nan=False))
        self.assertIn("unknown_", reports[0]["first_fatal"]["field"])
        blocks = [[{"type": "text", "text": "12 µV"} for _ in range(12)]]
        report = self.convert(blocks).validation_report
        self.assertEqual([issue.json_pointer for issue in report.warnings], [f"/0/{index}" for index in range(12)])


if __name__ == "__main__":
    unittest.main()

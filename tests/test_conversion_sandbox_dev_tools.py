"""验证本地转换沙箱辅助脚本。"""

from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase, mock
from xml.etree import ElementTree

import json
import os
import zipfile

from docx import Document

from app.conversion import runtime_facade
from app.services.document_ir.integrity import DocumentIntegrityError
from scripts.dev import (
    audit_conversion_sandbox,
    formal_mineru_conversion,
    migrate_md2word_main_sandbox,
    run_local_mineru_conversion,
    run_local_mineru_debug_config,
)


class LocalMinerUConversionCliTests(TestCase):
    def test_formal_adapter_saves_same_image_report_and_rejects_missing_inline(self) -> None:
        from tests.test_image_group_rendering import write_numbered_image_zip, change_docx, NS
        from app.services.renderer.word_renderer import DocumentIRWordRenderer
        from dataclasses import replace
        with TemporaryDirectory() as directory:
            root = Path(directory)
            package = root / "images.zip"
            write_numbered_image_zip(package, (24, 25))
            kwargs = dict(input_path=str(package), output_path=str(root / "images.docx"), template_path=None,
                layout_config_path=None, subscript_config=None, no_template=True, use_mineru_header_footer=False,
                document_language="zh", translation_mode="direct")
            result = formal_mineru_conversion.run_conversion(**kwargs)
            report = result.renderer_metadata["image_layout_report"]
            self.assertEqual(report["schema_version"], "image-layout/v1")
            self.assertEqual(report["render_inputs"]["schema_version"], "image-render-inputs/v1")
            self.assertEqual(result.renderer_metadata["layout_table_count"], 1)
            self.assertEqual(result.renderer_metadata["content_table_count"], 0)
            self.assertIn("图 24", Document(result.output_docx_path).tables[0].cell(0, 0).text)
            original = DocumentIRWordRenderer.render_to_result
            def remove_inline(renderer, *args, **options):
                rendered = original(renderer, *args, **options)
                def remove(root):
                    parent = next(node for node in root.iter() if node.find("./wp:inline", NS) is not None)
                    parent.remove(parent.find("./wp:inline", NS))
                return replace(rendered, content=change_docx(rendered.content, remove))
            kwargs["output_path"] = str(root / "rejected.docx")
            with mock.patch.object(DocumentIRWordRenderer, "render_to_result", remove_inline), self.assertRaises(ValueError):
                formal_mineru_conversion.run_conversion(**kwargs)
            self.assertFalse(Path(kwargs["output_path"]).exists())

    def test_build_conversion_kwargs_uses_official_runtime_contract(self) -> None:
        args = run_local_mineru_conversion.parse_args(
            [
                "--output",
                "out.docx",
                "--layout-config",
                "layout.yaml",
                "--translation-mode",
                "llm_en_to_zh",
                "--translation-trace-output",
                "trace.json",
                "mineru-package",
            ]
        )

        kwargs = run_local_mineru_conversion.build_conversion_kwargs(args)

        self.assertEqual(kwargs["input_path"], "mineru-package")
        self.assertEqual(kwargs["output_path"], "out.docx")
        self.assertEqual(kwargs["layout_config_path"], "layout.yaml")
        self.assertEqual(kwargs["translation_mode"], "en2zh")
        self.assertEqual(kwargs["translation_trace_output"], "trace.json")

    def test_build_conversion_kwargs_accepts_ru2zh_alias(self) -> None:
        args = run_local_mineru_conversion.parse_args(
            [
                "--translation-mode",
                "llm_ru_to_zh",
                "--document-language",
                "ru",
                "mineru-package",
            ]
        )

        kwargs = run_local_mineru_conversion.build_conversion_kwargs(args)

        self.assertEqual(kwargs["translation_mode"], "ru2zh")
        self.assertEqual(kwargs["document_language"], "ru")

    def test_build_conversion_kwargs_accepts_input_from_environment(self) -> None:
        args = run_local_mineru_conversion.parse_args([])

        kwargs = run_local_mineru_conversion.build_conversion_kwargs(
            args,
            environ={run_local_mineru_conversion.DEFAULT_INPUT_ENV: "/tmp/mineru-package"},
        )

        self.assertEqual(kwargs["input_path"], "/tmp/mineru-package")

    def test_missing_input_fails_before_runtime_call(self) -> None:
        args = run_local_mineru_conversion.parse_args([])

        with self.assertRaises(ValueError):
            run_local_mineru_conversion.build_conversion_kwargs(args, environ={})

    def test_main_calls_formal_conversion_adapter(self) -> None:
        with mock.patch.object(run_local_mineru_conversion, "configure_logging") as configure_logging, mock.patch.object(
            run_local_mineru_conversion.formal_mineru_conversion,
            "run_conversion",
        ) as run_conversion:
            run_local_mineru_conversion.main(["/tmp/mineru-package", "--document-language", "en"])

        configure_logging.assert_called_once_with("INFO")
        run_conversion.assert_called_once()
        self.assertEqual(run_conversion.call_args.kwargs["input_path"], "/tmp/mineru-package")
        self.assertEqual(run_conversion.call_args.kwargs["document_language"], "en")

    def test_formal_adapter_converts_minimal_mineru_package_to_openable_docx(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            package = root / "mineru-package"
            package.mkdir()
            (package / "sample_content_list_v2.json").write_text(
                json.dumps(
                    [
                        {"type": "title", "text": "1 Features", "text_level": 1, "bbox": [72, 80, 180, 100], "page_idx": 0},
                        {"type": "text", "text": "Low noise precision amplifier", "bbox": [72, 120, 260, 150], "page_idx": 0},
                        {
                            "type": "table",
                            "table_caption": ["Table 1 Electrical characteristics"],
                            "table_body": "<table><tr><th>PARAMETER</th></tr><tr><td>VOS</td></tr></table>",
                            "bbox": [72, 170, 540, 260],
                            "page_idx": 0,
                        },
                        {"type": "equation", "latex": r"V_{OUT}=1", "bbox": [72, 280, 260, 310], "page_idx": 0},
                    ]
                ),
                encoding="utf-8",
            )
            output_path = root / "out.docx"

            result = formal_mineru_conversion.run_conversion(
                input_path=str(package),
                output_path=str(output_path),
                template_path=None,
                layout_config_path=None,
                subscript_config=None,
                no_template=True,
                use_mineru_header_footer=False,
                document_language="en",
                translation_mode="direct",
            )

            self.assertTrue(output_path.exists())
            self.assertTrue(result.document_ir_path.exists())
            self.assertEqual(result.renderer_metadata["blocking_quality_flags"], [])
            self.assertEqual(result.renderer_metadata["expected_layout_block_count"], 4)
            self.assertEqual(result.renderer_metadata["layout_block_count"], 4)
            rendered = Document(output_path)
            self.assertIn("Features", "\n".join(paragraph.text for paragraph in rendered.paragraphs))
            with zipfile.ZipFile(output_path) as docx_zip:
                document_xml = ElementTree.fromstring(docx_zip.read("word/document.xml"))
            word_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            math_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
            visible_runs = [run for run in document_xml.iter(f"{word_ns}r") if run.find(f".//{word_ns}t") is not None]
            self.assertTrue(visible_runs)
            for run in visible_runs:
                fonts = run.find(f"./{word_ns}rPr/{word_ns}rFonts")
                self.assertIsNotNone(fonts)
                self.assertEqual(fonts.get(f"{word_ns}eastAsia"), "宋体")
                self.assertEqual(fonts.get(f"{word_ns}ascii"), "Arial")
                self.assertEqual(fonts.get(f"{word_ns}hAnsi"), "Arial")
                self.assertEqual(fonts.get(f"{word_ns}cs"), "Times New Roman")
            math_runs = list(document_xml.iter(f"{math_ns}r"))
            self.assertTrue(math_runs)
            self.assertTrue(all(run.find(f".//{word_ns}rFonts") is None for run in math_runs))

    def test_formal_adapter_resolves_default_explicit_and_no_template_resources(self) -> None:
        default_template = Path("/tmp/default-template.docx")
        default_layout = Path("/tmp/default-layout.yaml")
        explicit_template = Path("/tmp/explicit-template.docx")
        explicit_layout = Path("/tmp/explicit-layout.yaml")
        converted = object()
        settings = mock.Mock(default_template_path=default_template, layout_config_path=default_layout)

        with mock.patch.object(formal_mineru_conversion, "Settings", return_value=settings), mock.patch.object(
            formal_mineru_conversion,
            "convert_formal_mineru_package",
            return_value=converted,
        ) as convert:
            default_result = formal_mineru_conversion.run_conversion(
                input_path="mineru-package",
                output_path="out.docx",
                template_path=None,
                layout_config_path=None,
                subscript_config=None,
                no_template=False,
                use_mineru_header_footer=False,
            )
            explicit_result = formal_mineru_conversion.run_conversion(
                input_path="mineru-package",
                output_path="out.docx",
                template_path=str(explicit_template),
                layout_config_path=str(explicit_layout),
                subscript_config=None,
                no_template=False,
                use_mineru_header_footer=False,
            )
            no_template_result = formal_mineru_conversion.run_conversion(
                input_path="mineru-package",
                output_path="out.docx",
                template_path=str(explicit_template),
                layout_config_path=None,
                subscript_config=None,
                no_template=True,
                use_mineru_header_footer=False,
            )

        self.assertIs(default_result, converted)
        self.assertIs(explicit_result, converted)
        self.assertIs(no_template_result, converted)
        default_request, explicit_request, no_template_request = [call.args[0] for call in convert.call_args_list]
        self.assertEqual(default_request.template_path, default_template)
        self.assertEqual(default_request.layout_config_path, default_layout)
        self.assertEqual(explicit_request.template_path, explicit_template.resolve())
        self.assertEqual(explicit_request.layout_config_path, explicit_layout.resolve())
        self.assertIsNone(no_template_request.template_path)
        self.assertEqual(no_template_request.layout_config_path, default_layout)

        missing_template_settings = mock.Mock(default_template_path=None, layout_config_path=default_layout)
        with mock.patch.object(formal_mineru_conversion, "Settings", return_value=missing_template_settings):
            with self.assertRaisesRegex(FileNotFoundError, "默认 Word 模板不存在"):
                formal_mineru_conversion.run_conversion(
                    input_path="mineru-package",
                    output_path="out.docx",
                    template_path=None,
                    layout_config_path=None,
                    subscript_config=None,
                    no_template=False,
                    use_mineru_header_footer=False,
                )

        missing_settings = mock.Mock(default_template_path=None, layout_config_path=None)
        with mock.patch.object(formal_mineru_conversion, "Settings", return_value=missing_settings):
            with self.assertRaisesRegex(FileNotFoundError, "默认版式配置文件不存在"):
                formal_mineru_conversion.run_conversion(
                    input_path="mineru-package",
                    output_path="out.docx",
                    template_path=None,
                    layout_config_path=None,
                    subscript_config=None,
                    no_template=True,
                    use_mineru_header_footer=False,
                )

    def test_formal_adapter_does_not_publish_docx_when_renderer_integrity_blocks(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            package = root / "mineru-package"
            package.mkdir()
            (package / "sample_content_list_v2.json").write_text(
                json.dumps([{"type": "text", "text": "Visible body", "bbox": [10, 20, 180, 60], "page_idx": 0}]),
                encoding="utf-8",
            )
            output_path = root / "blocked.docx"
            blocking_evidence = {
                "blocking_quality_flags": ["ir_docx_integrity:inline_shape_count_mismatch"],
                "docx_quality": {},
                "quality_flags": [],
            }

            with mock.patch.object(
                formal_mineru_conversion,
                "build_renderer_evidence",
                return_value=blocking_evidence,
            ), self.assertRaises(DocumentIntegrityError):
                formal_mineru_conversion.run_conversion(
                    input_path=str(package),
                    output_path=str(output_path),
                    template_path=None,
                    layout_config_path=None,
                    subscript_config=None,
                    no_template=True,
                    use_mineru_header_footer=False,
                    document_language="en",
                    translation_mode="direct",
                )

            self.assertFalse(output_path.exists())

    def test_formal_adapter_writes_translation_request_before_missing_key_failure(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            package = root / "mineru-package"
            package.mkdir()
            (package / "sample_content_list_v2.json").write_text(
                json.dumps([{"type": "text", "text": "Low noise amplifier", "bbox": [72, 80, 260, 110], "page_idx": 0}]),
                encoding="utf-8",
            )
            request_path = root / "translation-request.json"

            with mock.patch.dict(os.environ, {"OPENAI_API_KEY": ""}, clear=False):
                with self.assertRaises(RuntimeError):
                    formal_mineru_conversion.run_conversion(
                        input_path=str(package),
                        output_path=str(root / "out.docx"),
                        template_path=None,
                        layout_config_path=None,
                        subscript_config=None,
                        no_template=True,
                        use_mineru_header_footer=False,
                        document_language="en",
                        translation_mode="en2zh",
                        translation_request_output=str(request_path),
                        translation_api_key="",
                    )

            payload = json.loads(request_path.read_text(encoding="utf-8"))
            self.assertEqual(payload["schema_version"], "document-ir-local-translation-request/v2")
            self.assertEqual(payload["language_pair"], "en2zh")
            self.assertEqual(payload["items"][0]["text"], "Low noise amplifier")


class LocalMinerUDebugConfigTests(TestCase):
    def test_repo_debug_configs_target_project_local_mineru_result(self) -> None:
        repo_root = Path.cwd()

        default_kwargs = run_local_mineru_debug_config.build_conversion_kwargs(
            repo_root / "scripts/dev/local_mineru_debug_config.json",
            repo_root=repo_root,
        )
        direct_kwargs = run_local_mineru_debug_config.build_conversion_kwargs(
            repo_root / "scripts/dev/local_mineru_debug_direct_config.json",
            repo_root=repo_root,
        )
        en2zh_kwargs = run_local_mineru_debug_config.build_conversion_kwargs(
            repo_root / "scripts/dev/local_mineru_debug_en2zh_config.json",
            repo_root=repo_root,
        )
        ru2zh_kwargs = run_local_mineru_debug_config.build_conversion_kwargs(
            repo_root / "scripts/dev/local_mineru_debug_ru2zh_config.json",
            repo_root=repo_root,
        )

        self.assertEqual(default_kwargs["translation_mode"], "direct")
        self.assertEqual(direct_kwargs["translation_mode"], "direct")
        self.assertEqual(en2zh_kwargs["translation_mode"], "en2zh")
        self.assertEqual(ru2zh_kwargs["translation_mode"], "ru2zh")
        self.assertEqual(default_kwargs["input_path"], direct_kwargs["input_path"])
        self.assertEqual(Path(direct_kwargs["input_path"]).relative_to(repo_root).parts[0], "mineru_result")
        self.assertEqual(Path(en2zh_kwargs["input_path"]).relative_to(repo_root).parts[0], "mineru_result")
        self.assertEqual(Path(ru2zh_kwargs["input_path"]).relative_to(repo_root).parts[0], "mineru_result")
        self.assertIn("tests/tmp/local_mineru_debug/direct", direct_kwargs["output_path"])
        self.assertIn("tests/tmp/local_mineru_debug/en2zh", en2zh_kwargs["output_path"])
        self.assertIn("tests/tmp/local_mineru_debug/ru2zh", ru2zh_kwargs["output_path"])
        self.assertIsNone(en2zh_kwargs["translation_api_key"])
        self.assertIsNone(ru2zh_kwargs["translation_api_key"])
        self.assertIsNotNone(en2zh_kwargs["translation_request_output"])
        self.assertIsNotNone(en2zh_kwargs["translation_trace_output"])
        self.assertIsNotNone(ru2zh_kwargs["translation_request_output"])
        self.assertIsNotNone(ru2zh_kwargs["translation_trace_output"])

    def test_pycharm_run_configs_reference_local_debug_configs(self) -> None:
        repo_root = Path.cwd()
        expectations = {
            "Local MinerU Direct.run.xml": "scripts/dev/local_mineru_debug_direct_config.json",
            "Local MinerU En2Zh.run.xml": "scripts/dev/local_mineru_debug_en2zh_config.json",
        }

        for filename, config_path in expectations.items():
            with self.subTest(filename=filename):
                tree = ElementTree.parse(repo_root / ".run" / filename)
                options = {option.attrib["name"]: option.attrib.get("value", "") for option in tree.findall(".//option[@name]")}

                self.assertEqual(options["SCRIPT_NAME"], "$PROJECT_DIR$/scripts/dev/run_local_mineru_debug_config.py")
                self.assertEqual(options["WORKING_DIRECTORY"], "$PROJECT_DIR$")
                self.assertIn(config_path, options["PARAMETERS"])

    def test_build_conversion_kwargs_resolves_relative_paths_from_official_repo(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            repo_root = Path(tmp_dir)
            config_path = repo_root / "local_mineru_debug_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "input_path": "tests/fixtures/mineru-package",
                        "output_path": "tests/tmp/local_mineru_debug/out.docx",
                        "template": "resources/templates/default/datasheet_template.docx",
                        "layout_config": "resources/template_configs/datasheet_template.yaml",
                        "subscript_config": "resources/term_rules/subscripts_dictionary.yaml",
                        "translation_mode": "llm_en_to_zh",
                        "translation_trace_output": "tests/tmp/local_mineru_debug/trace.json",
                    }
                ),
                encoding="utf-8",
            )

            kwargs = run_local_mineru_debug_config.build_conversion_kwargs(config_path, repo_root=repo_root)

        self.assertEqual(kwargs["input_path"], str((repo_root / "tests/fixtures/mineru-package").resolve()))
        self.assertEqual(kwargs["output_path"], str((repo_root / "tests/tmp/local_mineru_debug/out.docx").resolve()))
        self.assertEqual(kwargs["template_path"], str((repo_root / "resources/templates/default/datasheet_template.docx").resolve()))
        self.assertEqual(kwargs["layout_config_path"], str((repo_root / "resources/template_configs/datasheet_template.yaml").resolve()))
        self.assertEqual(kwargs["translation_mode"], "en2zh")
        self.assertEqual(kwargs["translation_trace_output"], str((repo_root / "tests/tmp/local_mineru_debug/trace.json").resolve()))

    def test_build_conversion_kwargs_keeps_absolute_mineru_asset_path(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            repo_root = Path(tmp_dir)
            absolute_package = Path(tmp_dir) / "md2word_main" / "mineru_result" / "OPA228_P1.pdf-uuid"
            config_path = repo_root / "local_mineru_debug_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "input_path": str(absolute_package),
                        "output_path": "tests/tmp/out.docx",
                        "document_language": "en",
                        "translation_mode": "direct",
                    }
                ),
                encoding="utf-8",
            )

            kwargs = run_local_mineru_debug_config.build_conversion_kwargs(config_path, repo_root=repo_root)

        self.assertEqual(kwargs["input_path"], str(absolute_package.resolve()))
        self.assertEqual(kwargs["document_language"], "en")
        self.assertEqual(kwargs["translation_mode"], "direct")

    def test_unknown_config_key_fails_before_runtime_call(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            config_path = Path(tmp_dir) / "local_mineru_debug_config.json"
            config_path.write_text(json.dumps({"input_path": "mineru", "surprise": True}), encoding="utf-8")

            with self.assertRaises(ValueError):
                run_local_mineru_debug_config.build_conversion_kwargs(config_path)

    def test_main_calls_formal_conversion_adapter_from_config(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            config_path = Path(tmp_dir) / "local_mineru_debug_config.json"
            config_path.write_text(json.dumps({"input_path": "/tmp/mineru-package", "log_level": "DEBUG"}), encoding="utf-8")

            with mock.patch.object(run_local_mineru_debug_config, "configure_logging") as configure_logging, mock.patch.object(
                run_local_mineru_debug_config.formal_mineru_conversion,
                "run_conversion",
            ) as run_conversion:
                run_local_mineru_debug_config.main(["--config", str(config_path)])

        configure_logging.assert_called_once_with("DEBUG")
        run_conversion.assert_called_once()
        self.assertEqual(run_conversion.call_args.kwargs["input_path"], str(Path("/tmp/mineru-package").resolve()))


class ConversionSandboxAuditTests(TestCase):
    def test_drift_report_classifies_changed_and_one_sided_files(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            official = root / "official"
            sandbox = root / "sandbox"
            (official / "app" / "conversion").mkdir(parents=True)
            (sandbox / "app" / "conversion").mkdir(parents=True)
            (official / "app" / "conversion" / "same.py").write_text("same\n", encoding="utf-8")
            (sandbox / "app" / "conversion" / "same.py").write_text("same\n", encoding="utf-8")
            (official / "app" / "conversion" / "changed.py").write_text("official\n", encoding="utf-8")
            (sandbox / "app" / "conversion" / "changed.py").write_text("sandbox\n", encoding="utf-8")
            (official / "app" / "conversion" / "official_only.py").write_text("official\n", encoding="utf-8")
            (sandbox / "app" / "conversion" / "sandbox_only.py").write_text("sandbox\n", encoding="utf-8")
            (sandbox / "app" / "conversion" / "ignored.pyc").write_bytes(b"bytecode")

            report = audit_conversion_sandbox.build_drift_report(
                official,
                sandbox,
                ["app/conversion"],
            )

        self.assertTrue(report.has_drift)
        self.assertEqual(report.same_files, ["app/conversion/same.py"])
        self.assertEqual(report.changed_files, ["app/conversion/changed.py"])
        self.assertEqual(report.official_only_files, ["app/conversion/official_only.py"])
        self.assertEqual(report.sandbox_only_files, ["app/conversion/sandbox_only.py"])

    def test_sample_inventory_reports_pdfs_and_mineru_packages(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            sandbox = Path(tmp_dir)
            (sandbox / "test_pdf").mkdir()
            (sandbox / "test_pdf" / "OPA228_P1.pdf").write_bytes(b"%PDF-1.4\n")
            package = sandbox / "mineru_result" / "OPA228_P1.pdf-uuid"
            (package / "images").mkdir(parents=True)
            (package / "abc_content_list_v2.json").write_text("[]", encoding="utf-8")
            (package / "abc_content_list.json").write_text("[]", encoding="utf-8")
            (package / "abc_model.json").write_text("{}", encoding="utf-8")
            (package / "layout.json").write_text("{}", encoding="utf-8")
            (package / "block_list.json").write_text("[]", encoding="utf-8")
            (package / "full.md").write_text("# title\n", encoding="utf-8")
            (package / "images" / "plot.jpg").write_bytes(b"image")

            report = audit_conversion_sandbox.build_sample_inventory(sandbox)

        self.assertEqual(report.pdfs, ["OPA228_P1.pdf"])
        self.assertEqual(len(report.mineru_packages), 1)
        package_report = report.mineru_packages[0]
        self.assertEqual(package_report.candidate_pdf, "OPA228_P1.pdf")
        self.assertTrue(package_report.has_content_list_v2)
        self.assertTrue(package_report.has_content_list)
        self.assertTrue(package_report.has_model_json)
        self.assertTrue(package_report.has_layout_json)
        self.assertTrue(package_report.has_block_list_json)
        self.assertTrue(package_report.has_full_markdown)
        self.assertEqual(package_report.image_count, 1)

    def test_thin_shell_report_blocks_active_duplicate_sources(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            sandbox = Path(tmp_dir)
            for relative_path in ("mineru_to_docx.py", "sandbox_config.json", "README.md"):
                (sandbox / relative_path).write_text("", encoding="utf-8")
            for relative_path in ("test_pdf", "mineru_result", "out", "app"):
                (sandbox / relative_path).mkdir(parents=True)

            report = audit_conversion_sandbox.build_thin_shell_report(sandbox)

        self.assertFalse(report.ok)
        self.assertEqual(report.active_duplicate_paths, ["app"])
        self.assertEqual(report.missing_expected_paths, [])


class Md2wordMainMigrationTests(TestCase):
    def test_migration_archives_duplicate_sources_and_writes_thin_shell(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            official = root / "official"
            sandbox = root / "md2word_main"
            (official / "app" / "conversion").mkdir(parents=True)
            for relative_path in ("app/conversion", "cli_core", "resources", "tests"):
                (sandbox / relative_path).mkdir(parents=True)
                (sandbox / relative_path / "old.py").write_text("old\n", encoding="utf-8")
            for relative_path in ("test_pdf", "mineru_result", "out"):
                (sandbox / relative_path).mkdir(parents=True)
            (sandbox / "mineru_result" / "OPA228_P1.pdf-d1491424-9662-411d-8e39-fa6bcbd785a4").mkdir()

            result = migrate_md2word_main_sandbox.migrate_sandbox(
                sandbox_root=sandbox,
                official_root=official,
                apply=True,
            )

            thin_report = audit_conversion_sandbox.build_thin_shell_report(sandbox)

            self.assertTrue(thin_report.ok)
            self.assertFalse((sandbox / "app").exists())
            self.assertFalse((sandbox / "cli_core").exists())
            self.assertFalse((sandbox / "resources").exists())
            self.assertTrue((sandbox / "_legacy_duplicate_core" / "app" / "conversion" / "old.py").exists())
            self.assertTrue((sandbox / "_legacy_duplicate_core" / "cli_core" / "old.py").exists())
            self.assertTrue((sandbox / "_legacy_duplicate_core" / "resources" / "old.py").exists())
            self.assertTrue((sandbox / "mineru_to_docx.py").exists())
            self.assertTrue((sandbox / "config.py").exists())
            self.assertTrue((sandbox / "sandbox_config.json").exists())
            self.assertTrue((sandbox / "tests" / "test_thin_shell.py").exists())
            self.assertIn("app -> _legacy_duplicate_core/app", result.moved_paths)

            second_result = migrate_md2word_main_sandbox.migrate_sandbox(
                sandbox_root=sandbox,
                official_root=official,
                apply=True,
            )
            readme = (sandbox / "README.md").read_text(encoding="utf-8")
            entrypoint = (sandbox / "mineru_to_docx.py").read_text(encoding="utf-8")

            self.assertEqual(second_result.moved_paths, [])
            self.assertFalse((sandbox / "_legacy_duplicate_core" / "tests.2").exists())
            self.assertIn("run_local_mineru_debug_config.py", readme)
            self.assertIn("已标记为 deprecated", readme)
            self.assertIn("已废弃", entrypoint)
            self.assertIn("run_local_mineru_debug_config.py", entrypoint)
            self.assertNotIn("run_local_mineru_conversion", entrypoint)


class OfficialConversionRuntimeBindingTests(TestCase):
    def test_runtime_facade_load_pages_binds_model_classes(self) -> None:
        with TemporaryDirectory() as tmp_dir:
            package_root = Path(tmp_dir)
            payload = [
                [
                    {
                        "type": "title",
                        "bbox": [0, 0, 10, 10],
                        "content": {
                            "title_content": [{"type": "text", "content": "Title"}],
                            "level": 1,
                        },
                    }
                ]
            ]
            payload_path = package_root / "sample_content_list_v2.json"
            payload_path.write_text(json.dumps(payload), encoding="utf-8")

            pages, source_path = runtime_facade.load_pages(package_root)

        self.assertEqual(source_path, payload_path)
        self.assertEqual(len(pages), 1)
        self.assertEqual(pages[0].blocks[0].text, "Title")


class LocalTranslationIntegrityTests(TestCase):
    """本地工具导出、严格响应和同步零回填边界。"""

    def make_case(self, *, technical=True):
        temporary = TemporaryDirectory(prefix="local_translation_integrity_")
        self.addCleanup(temporary.cleanup)
        root = Path(temporary.name)
        package = root / "package"
        package.mkdir()
        texts = ["Amplifier OPA228 uses 10 V.", "Supply 5 V."] if technical else ["Quiet amplifier", "Stable supply"]
        (package / "content_list_v2.json").write_text(json.dumps([
            {"type": "text", "text": text, "page_idx": index} for index, text in enumerate(texts)
        ]), encoding="utf-8")
        request_path = root / "request.json"
        response_path = root / "response.json"
        output = root / "out.docx"
        kwargs = dict(
            input_path=str(package), output_path=str(output), template_path=None, layout_config_path=None,
            subscript_config=None, no_template=True, use_mineru_header_footer=False,
            document_language="en", translation_mode="en2zh", translation_request_output=str(request_path),
            translation_max_pages=1, translation_max_chars=500,
        )
        with mock.patch.object(formal_mineru_conversion, "_build_openai_provider", side_effect=RuntimeError("离线导出")):
            with self.assertRaisesRegex(RuntimeError, "离线导出"):
                formal_mineru_conversion.run_conversion(**kwargs)
        payload = json.loads(request_path.read_text(encoding="utf-8"))
        kwargs["translation_response"] = str(response_path)
        return kwargs, payload, response_path, output, texts

    def test_v2_export_matches_actual_encoded_requests_and_page_plan(self):
        kwargs, payload, response, output, texts = self.make_case()
        self.assertEqual(payload["schema_version"], "document-ir-local-translation-request/v2")
        self.assertEqual(len(payload["batches"]), 2)
        self.assertNotIn("OPA228", json.dumps(payload))
        self.assertIn("MD2W_PT_v1", payload["items"][0]["text"])
        self.assertGreater(payload["batches"][0]["encoded_char_count"], payload["batches"][0]["char_count"])
        response.write_text(json.dumps({"batches": [
            {"batch_id": batch["batch_id"], "items": [{**item, "text": "译文：" + item["text"]} for item in batch["items"]]}
            for batch in payload["batches"]
        ]}), encoding="utf-8")
        translate = formal_mineru_conversion.LocalTranslationResponseProvider.translate
        with mock.patch.object(formal_mineru_conversion.LocalTranslationResponseProvider, "translate", autospec=True, side_effect=translate) as calls:
            result = formal_mineru_conversion.run_conversion(**kwargs)
        self.assertTrue(output.exists())
        actual = [[{"id": item.id, "text": item.text} for item in call.args[1].items] for call in calls.call_args_list]
        self.assertEqual(actual, [batch["items"] for batch in payload["batches"]])
        translated = json.loads(result.document_ir_path.read_text(encoding="utf-8"))
        self.assertEqual([block["text"] for block in translated["blocks"]], ["译文：" + text for text in texts])
        self.assertEqual(len(actual), 2)

    def test_invalid_local_envelope_rejected_before_any_apply_or_output(self):
        from app.domains.translation.document_ir_translator import DocumentIRTranslationService
        mutations = {
            "missing": lambda p: {"items": p["items"][:1]},
            "extra": lambda p: {"items": p["items"] + [{"id": "extra", "text": "extra"}]},
            "duplicate": lambda p: {"items": [p["items"][0], p["items"][0]]},
            "order": lambda p: {"items": list(reversed(p["items"]))},
            "extra_batch": lambda p: {"batches": [{"items": batch["items"]} for batch in p["batches"]] + [{"items": []}]},
            "missing_batch": lambda p: {"batches": [{"items": p["batches"][0]["items"]}]},
            "batch_order": lambda p: {"batches": [{"items": batch["items"]} for batch in reversed(p["batches"])]},
            "type": lambda p: {"items": [{**p["items"][0], "text": 42}, p["items"][1]]},
            "field": lambda p: {"items": [{**p["items"][0], "extra": "bad"}, p["items"][1]]},
            "wrapper": lambda p: {"response": {"items": p["items"]}},
        }
        for name, mutation in mutations.items():
            with self.subTest(name=name):
                kwargs, payload, response, output, _texts = self.make_case()
                response.write_text(json.dumps(mutation(payload)), encoding="utf-8")
                with mock.patch.object(DocumentIRTranslationService, "apply_replacements") as apply:
                    with self.assertRaises((ValueError, RuntimeError)):
                        formal_mineru_conversion.run_conversion(**kwargs)
                apply.assert_not_called()
                self.assertFalse(output.exists())
                self.assertFalse(output.with_suffix(".document-ir.json").exists())

    def test_old_technical_response_is_not_repaired_but_plain_text_response_remains_valid(self):
        from app.domains.translation.document_ir_translator import DocumentIRTranslationService, DocumentIRTranslationError
        kwargs, payload, response, output, texts = self.make_case()
        response.write_text(json.dumps({"items": [{**item, "text": text} for item, text in zip(payload["items"], texts)]}), encoding="utf-8")
        with mock.patch.object(DocumentIRTranslationService, "apply_replacements") as apply:
            with self.assertRaises(DocumentIRTranslationError) as raised:
                formal_mineru_conversion.run_conversion(**kwargs)
        self.assertEqual(raised.exception.code, "translation_integrity_placeholder_missing")
        apply.assert_not_called()
        self.assertFalse(output.exists())
        kwargs, payload, response, output, _texts = self.make_case(technical=False)
        response.write_text(json.dumps({"items": [{**item, "text": "安静且稳定"} for item in payload["items"]]}), encoding="utf-8")
        result = formal_mineru_conversion.run_conversion(**kwargs)
        self.assertTrue(output.exists())
        self.assertIn("安静且稳定", result.document_ir_path.read_text(encoding="utf-8"))

    def test_duplicate_json_keys_and_added_technical_content_have_zero_apply(self):
        from app.domains.translation.document_ir_translator import DocumentIRTranslationService
        for name in ("duplicate_key", "technical_injection"):
            with self.subTest(name=name):
                kwargs, payload, response, output, _texts = self.make_case()
                if name == "duplicate_key":
                    text = '{"items": [], "items": ' + json.dumps(payload["items"]) + '}'
                else:
                    payload["items"][1]["text"] += " extra 999 V"
                    text = json.dumps({"items": payload["items"]})
                response.write_text(text, encoding="utf-8")
                with mock.patch.object(DocumentIRTranslationService, "apply_replacements") as apply:
                    with self.assertRaises((ValueError, RuntimeError)):
                        formal_mineru_conversion.run_conversion(**kwargs)
                apply.assert_not_called()
                self.assertFalse(output.exists())

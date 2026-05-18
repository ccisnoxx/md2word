from __future__ import annotations

import re
from html import unescape
from typing import Any

from bs4 import BeautifulSoup, NavigableString, Tag
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


BODY_FONT = Pt(10.5)
SECTION_TITLE_FONT = Pt(12.5)
TOP_TITLE_FONT = Pt(16)

NUMBERED_HEADING_RE = re.compile(r"^(?P<num>\d+(?:\.\d+){0,5})\b")
CHAPTER_HEADING_RE = re.compile(
    r"^(?:第[一二三四五六七八九十百千零〇两\d]+(章|节|部分)\b|(?:chapter|section|part)\s+(?:[ivxlcdm]+|\d+)\b)",
    re.IGNORECASE,
)
HEADING_NUMBER_RE = re.compile(
    r"^(?:\d+[.、]|\d+\.\d+|附录[A-Z\d.\-]+|appendix\s+[A-Z\d.\-]+|(?:表|table|图|figure)\s*[A-Z\d.\-]+)",
    re.IGNORECASE,
)
TOC_TITLE_RE = re.compile(r"^(目录|table of contents|contents)\b", re.IGNORECASE)
TRAILING_PAGE_NOISE_RE = re.compile(r"(?:[.\s·。…]{2,}|\s+)\d+\s*$")
INLINE_TEX_RE = re.compile(r"\$(?P<dollar>[^$\n]+)\$|\\\((?P<paren>.*?)\\\)")
WRAPPER_RE = re.compile(r"\\(?:mathrm|mathsf|text)\s*\{([^{}]+)\}")
CONTROL_WORD_RE = re.compile(r"\\([A-Za-z]+)")
SIMPLE_EXPR_RE = re.compile(
    r"(?P<base>[A-Za-zΑ-Ωα-ωΔΨθρμπμΩ]+[A-Za-z0-9Α-Ωα-ωΔΨθρμπμΩ]*)"
    r"(?:_(?:\{(?P<sub_braced>[^{}]+)\}|(?P<sub_plain>[A-Za-z0-9()+/\-]+)))?"
    r"(?:\^(?:\{(?P<sup_braced>[^{}]+)\}|(?P<sup_plain>\([^)]+\)|[A-Za-z0-9()+/\-]+)))?"
)
SUP_ONLY_EXPR_RE = re.compile(
    r"(?P<base>[A-Za-zΑ-Ωα-ωΔΨθρμπμΩ]+[A-Za-z0-9Α-Ωα-ωΔΨθρμπμΩ]*)"
    r"\^(?P<sup>\([^)]+\)|[A-Za-z0-9()+/\-]+)"
)
INLINE_MARKUP_SPLIT_RE = re.compile(r"(<[^>]+>)")
SPACED_ALPHA_TOKEN_RE = re.compile(r"\b(?:[A-Za-z]\s+){1,}[A-Za-z]\b")
SPACED_DIGIT_TOKEN_RE = re.compile(r"\b(?:\d\s+){1,}\d\b")
ENGLISH_HEADING_TOKEN_RE = re.compile(r"[A-Za-z][A-Za-z0-9/&()\-]*")
ENGLISH_HEADING_STOPWORDS = {
    "a",
    "an",
    "and",
    "as",
    "at",
    "by",
    "for",
    "from",
    "in",
    "of",
    "on",
    "or",
    "the",
    "to",
    "with",
    "vs",
}
GREEK_AND_TEXT_COMMANDS = {
    "Delta": "Δ",
    "delta": "δ",
    "Theta": "Θ",
    "theta": "θ",
    "mu": "μ",
    "Omega": "Ω",
    "alpha": "α",
    "beta": "β",
    "gamma": "γ",
    "lambda": "λ",
    "rho": "ρ",
    "psi": "ψ",
    "phi": "φ",
    "pm": "±",
    "circ": "°",
    "vee": "V",
    "to": "to",
    "max": "max",
    "min": "min",
    "typ": "typ",
    "avg": "avg",
    "rms": "RMS",
    "pk": "pk",
    "pp": "pp",
}

_RUNTIME_MODULE: Any | None = None


def bind_runtime(module: Any) -> None:
    global _RUNTIME_MODULE
    _RUNTIME_MODULE = module


def _runtime_value(name: str, default: Any) -> Any:
    if _RUNTIME_MODULE is None:
        return default
    return getattr(_RUNTIME_MODULE, name, default)


def _clean_text(text: str) -> str:
    return " ".join(text.replace("\n", " ").split())


def normalize_toc_text(text: str) -> str:
    cleaned = _clean_text(text)
    cleaned = TRAILING_PAGE_NOISE_RE.sub("", cleaned)
    cleaned = cleaned.replace("..", " ").replace(" .", " ")
    cleaned = re.sub(r"\s{2,}", " ", cleaned)
    return cleaned.strip()


def fill_toc_entry_paragraph(paragraph: Any, text: str, render_config: Any) -> None:
    paragraph.text = ""
    normalize_toc_text_fn = _runtime_value("normalize_toc_text", normalize_toc_text)
    infer_heading_level_fn = _runtime_value("infer_heading_level", infer_heading_level)
    get_toc_style_name_fn = _runtime_value("get_toc_style_name", get_toc_style_name)
    apply_named_style_config_fn = _runtime_value("apply_named_style_config", apply_named_style_config)
    fill_rich_text_fn = _runtime_value("fill_rich_text", fill_rich_text)
    prepare_text_fn = _runtime_value("prepare_text", prepare_text)

    level = infer_heading_level_fn(normalize_toc_text_fn(text))
    style_name = get_toc_style_name_fn(level)
    apply_named_style_config_fn(paragraph, render_config, style_name)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(compute_toc_indent_pt(text, render_config))
    fill_rich_text_fn(paragraph, prepare_text_fn(text, render_config), render_config=render_config, style_name=style_name)


def fill_toc_title_paragraph(paragraph: Any, text: str, render_config: Any) -> None:
    paragraph.text = ""
    apply_named_style_config_fn = _runtime_value("apply_named_style_config", apply_named_style_config)
    fill_rich_text_fn = _runtime_value("fill_rich_text", fill_rich_text)
    prepare_text_fn = _runtime_value("prepare_text", prepare_text)

    apply_named_style_config_fn(paragraph, render_config, "TocTitle")
    fill_rich_text_fn(paragraph, prepare_text_fn(text, render_config), render_config=render_config, style_name="TocTitle", bold=True)


def get_toc_style_name(level: int) -> str:
    if level <= 1:
        return "TocLevel1"
    if level == 2:
        return "TocLevel2"
    return "TocLevel3"


def compute_toc_indent_pt(text: str, render_config: Any) -> float:
    level = _runtime_value("infer_heading_level", infer_heading_level)(_runtime_value("normalize_toc_text", normalize_toc_text)(text))
    style = _runtime_value("get_style_config", get_style_config)(render_config, "Normal")
    font_size = style.font_size if style and style.font_size else 10.5
    return max(level - 1, 0) * font_size * 2


def normalize_formula_text(text: str) -> str:
    if not text:
        return ""
    normalized = unescape(text)
    normalized = normalized.replace("\xa0", " ")
    normalized = re.sub(r"<eq>\s*(.*?)\s*</eq>", r"\1", normalized, flags=re.IGNORECASE | re.DOTALL)
    normalized = re.sub(r"\$+\s*(.*?)\s*\$+", r"\1", normalized, flags=re.DOTALL)
    normalized = normalize_tex_text(normalized)
    normalized = convert_inline_tex_to_html(normalized)
    normalized = re.sub(r"\s+,", ",", normalized)
    normalized = re.sub(r"([±+\-])\s+(\d)", r"\1\2", normalized)
    normalized = re.sub(r"(\d)\s*°\s*([CFK])\b", r"\1°\2", normalized)
    normalized = re.sub(r"([A-Za-zµμ])\s+Ω\b", r"\1Ω", normalized)
    normalized = re.sub(r"[ \t]+", " ", normalized)
    normalized = re.sub(r" ?\n ?", "\n", normalized)
    return normalized.strip()


def normalize_tex_text(text: str) -> str:
    normalized = text
    prev = None
    while prev != normalized:
        prev = normalized
        normalized = WRAPPER_RE.sub(lambda match: compact_math_token_runs(match.group(1) or ""), normalized)

    normalized = CONTROL_WORD_RE.sub(replace_control_word, normalized)
    normalized = normalized.replace(r"\_", "_")
    normalized = normalized.replace(r"\{", "{").replace(r"\}", "}")
    normalized = normalized.replace(r"\(", "(").replace(r"\)", ")")
    normalized = normalized.replace(r"\ ", " ")
    normalized = re.sub(r"\s*_\s*", "_", normalized)
    normalized = re.sub(r"\s*\^\s*", "^", normalized)
    normalized = re.sub(r"\^\s*\{\s*", "^{", normalized)
    normalized = re.sub(r"_\s*\{\s*", "_{", normalized)
    normalized = re.sub(r"\{\s*", "{", normalized)
    normalized = re.sub(r"\s*\}", "}", normalized)
    normalized = normalized.replace("^{°}", "°")
    normalized = compact_math_token_runs(normalized)
    return normalized


def convert_inline_tex_to_html(text: str) -> str:
    if not text:
        return ""

    def replacer(match: re.Match[str]) -> str:
        expr = match.group("dollar") if match.group("dollar") is not None else match.group("paren")
        return convert_simple_tex_expr(expr or match.group(0))

    converted = INLINE_TEX_RE.sub(replacer, text)
    converted = SIMPLE_EXPR_RE.sub(
        lambda match: convert_simple_tex_expr(match.group(0)),
        converted,
    )
    converted = SUP_ONLY_EXPR_RE.sub(
        lambda match: convert_simple_tex_expr(match.group(0)),
        converted,
    )
    return converted


def convert_simple_tex_expr(expr: str) -> str:
    cleaned = clean_tex_fragment(expr)
    match = SIMPLE_EXPR_RE.fullmatch(cleaned.strip())
    if not match:
        sup_only_match = SUP_ONLY_EXPR_RE.fullmatch(cleaned.strip())
        if not sup_only_match:
            return cleaned
        base = clean_tex_fragment(sup_only_match.group("base") or "")
        sup = clean_tex_fragment(sup_only_match.group("sup") or "")
        return f"{base}<sup>{sup}</sup>" if sup else base

    base = clean_tex_fragment(match.group("base") or "")
    sub = clean_tex_fragment(match.group("sub_braced") or match.group("sub_plain") or "")
    sup = clean_tex_fragment(match.group("sup_braced") or match.group("sup_plain") or "")

    if not sub and not sup:
        return base

    result = base
    if sub:
        result += f"<sub>{sub}</sub>"
    if sup:
        result += f"<sup>{sup}</sup>"
    return result


def clean_tex_fragment(text: str) -> str:
    cleaned = text.strip()
    prev = None
    while prev != cleaned:
        prev = cleaned
        cleaned = WRAPPER_RE.sub(lambda match: compact_math_token_runs(match.group(1) or ""), cleaned)

    cleaned = CONTROL_WORD_RE.sub(replace_control_word, cleaned)
    cleaned = cleaned.replace(r"\_", "_")
    cleaned = cleaned.replace(r"\{", "{").replace(r"\}", "}")
    cleaned = cleaned.replace(r"\(", "(").replace(r"\)", ")")
    cleaned = cleaned.replace(r"\ ", " ")
    cleaned = compact_math_token_runs(cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    cleaned = re.sub(r"([ΔδΘθμΩαβγλρψφ])\s+([A-Za-z])", r"\1\2", cleaned)
    return cleaned


def compact_math_token_runs(text: str) -> str:
    compacted = text
    prev = None
    while prev != compacted:
        prev = compacted
        compacted = SPACED_ALPHA_TOKEN_RE.sub(lambda match: match.group(0).replace(" ", ""), compacted)
        compacted = SPACED_DIGIT_TOKEN_RE.sub(lambda match: match.group(0).replace(" ", ""), compacted)
    return compacted


def replace_control_word(match: re.Match[str]) -> str:
    return GREEK_AND_TEXT_COMMANDS.get(match.group(1), match.group(0))


def prepare_text(text: str, render_config: Any | None) -> str:
    prepared = normalize_formula_text(text)
    if render_config and render_config.terminology:
        prepared = render_config.terminology.replace_text(prepared)
    return prepared


def apply_style_if_exists(paragraph: Any, style_name: str) -> None:
    try:
        paragraph.style = style_name
    except (KeyError, ValueError):
        return


def apply_named_style_config(paragraph: Any, render_config: Any, style_name: str) -> None:
    apply_style_if_exists(paragraph, style_name)
    config = render_config.paragraph_styles.get(style_name)
    if not config:
        return

    fmt = paragraph.paragraph_format
    if config.alignment:
        paragraph.alignment = parse_alignment(config.alignment)
    if config.space_before is not None:
        fmt.space_before = Pt(config.space_before)
    if config.space_after is not None:
        fmt.space_after = Pt(config.space_after)
    if config.line_spacing is not None:
        fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        fmt.line_spacing = config.line_spacing
    if config.first_line_indent_pt is not None:
        fmt.first_line_indent = Pt(config.first_line_indent_pt)
    if config.left_indent_pt is not None:
        fmt.left_indent = Pt(config.left_indent_pt)
    if config.right_indent_pt is not None:
        fmt.right_indent = Pt(config.right_indent_pt)


def parse_alignment(value: str) -> WD_ALIGN_PARAGRAPH:
    normalized = value.strip().lower()
    mapping = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    return mapping.get(normalized, WD_ALIGN_PARAGRAPH.LEFT)


def get_style_config(render_config: Any | None, style_name: str | None) -> Any | None:
    if not render_config or not style_name:
        return None
    return render_config.paragraph_styles.get(style_name)


def apply_run_defaults(run: Any, style_config: Any | None = None) -> None:
    if style_config and style_config.font_name:
        run.font.name = style_config.font_name
        r_fonts = run._element.rPr.rFonts if run._element.rPr is not None else None
        if r_fonts is not None:
            r_fonts.set(qn("w:eastAsia"), style_config.font_name)
            r_fonts.set(qn("w:ascii"), style_config.font_name)
            r_fonts.set(qn("w:hAnsi"), style_config.font_name)
    if style_config and style_config.font_size is not None:
        run.font.size = Pt(style_config.font_size)
    elif run.font.size is None:
        run.font.size = BODY_FONT
    if style_config and style_config.bold is not None:
        run.bold = style_config.bold if run.bold is None else run.bold
    if style_config and style_config.italic is not None:
        run.italic = style_config.italic if run.italic is None else run.italic
    if style_config and style_config.color:
        run.font.color.rgb = parse_rgb_color(style_config.color)


def parse_rgb_color(value: str) -> RGBColor:
    raw = value.strip().lstrip("#")
    if len(raw) != 6:
        return RGBColor(0, 0, 0)
    return RGBColor.from_string(raw.upper())


def fill_rich_text(
    paragraph: Any,
    text: str,
    *,
    render_config: Any | None = None,
    style_name: str | None = None,
    bold: bool = False,
    italic: bool = False,
) -> None:
    paragraph.text = ""
    _runtime_value("render_inline_html_into_paragraph", render_inline_html_into_paragraph)(
        paragraph,
        text,
        render_config=render_config,
        style_name=style_name,
        bold=bold,
        italic=italic,
    )


def render_inline_html_into_paragraph(
    paragraph: Any,
    text_or_html: str,
    *,
    render_config: Any | None = None,
    style_name: str | None = None,
    bold: bool = False,
    italic: bool = False,
    subscript: bool = False,
    superscript: bool = False,
) -> None:
    style_config = _runtime_value("get_style_config", get_style_config)(render_config, style_name)
    soup = _runtime_value("BeautifulSoup", BeautifulSoup)(f"<root>{text_or_html}</root>", "html.parser")
    root = soup.find("root")
    if root is None:
        run = paragraph.add_run(_runtime_value("normalize_formula_text", normalize_formula_text)(text_or_html))
        _runtime_value("apply_run_defaults", apply_run_defaults)(run, style_config)
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        if subscript:
            run.font.subscript = True
        if superscript:
            run.font.superscript = True
        return

    _runtime_value("render_html_children", render_html_children)(
        root,
        paragraph,
        style_config=style_config,
        bold=bold,
        italic=italic,
        subscript=subscript,
        superscript=superscript,
    )


def render_html_children(
    parent: Tag,
    paragraph: Any,
    *,
    style_config: Any | None,
    bold: bool,
    italic: bool,
    subscript: bool,
    superscript: bool,
) -> None:
    for child in parent.children:
        if isinstance(child, NavigableString):
            _runtime_value("render_normalized_text_fragment", render_normalized_text_fragment)(
                paragraph,
                str(child),
                style_config=style_config,
                bold=bold,
                italic=italic,
                subscript=subscript,
                superscript=superscript,
            )
            continue

        if not isinstance(child, Tag):
            continue

        if child.name == "br":
            paragraph.add_run("\n")
            continue

        render_html_children(
            child,
            paragraph,
            style_config=style_config,
            bold=bold or child.name in {"b", "strong", "th"},
            italic=italic or child.name in {"i", "em"},
            subscript=subscript or child.name == "sub",
            superscript=superscript or child.name == "sup",
        )


def render_normalized_text_fragment(
    paragraph: Any,
    text: str,
    *,
    style_config: Any | None,
    bold: bool,
    italic: bool,
    subscript: bool,
    superscript: bool,
) -> None:
    normalized = _runtime_value("normalize_formula_text", normalize_formula_text)(text)
    if not normalized:
        return

    if "<" not in normalized or ">" not in normalized:
        _runtime_value("append_text_run", append_text_run)(
            paragraph,
            normalized,
            style_config=style_config,
            bold=bold,
            italic=italic,
            subscript=subscript,
            superscript=superscript,
        )
        return

    if _runtime_value("render_inline_markup_fragment", render_inline_markup_fragment)(
        paragraph,
        normalized,
        style_config=style_config,
        bold=bold,
        italic=italic,
        subscript=subscript,
        superscript=superscript,
    ):
        return

    fragment = _runtime_value("BeautifulSoup", BeautifulSoup)(f"<root>{normalized}</root>", "html.parser").find("root")
    if fragment is None:
        _runtime_value("append_text_run", append_text_run)(
            paragraph,
            normalized,
            style_config=style_config,
            bold=bold,
            italic=italic,
            subscript=subscript,
            superscript=superscript,
        )
        return

    _runtime_value("render_markup_children", render_markup_children)(
        fragment,
        paragraph,
        style_config=style_config,
        bold=bold,
        italic=italic,
        subscript=subscript,
        superscript=superscript,
    )


def render_inline_markup_fragment(
    paragraph: Any,
    normalized: str,
    *,
    style_config: Any | None,
    bold: bool,
    italic: bool,
    subscript: bool,
    superscript: bool,
) -> bool:
    active_subscript = subscript
    active_superscript = superscript
    handled_markup = False

    for token in INLINE_MARKUP_SPLIT_RE.split(normalized):
        if not token:
            continue

        lowered = token.lower()
        if lowered == "<sub>":
            active_subscript = True
            handled_markup = True
            continue
        if lowered == "</sub>":
            active_subscript = subscript
            handled_markup = True
            continue
        if lowered == "<sup>":
            active_superscript = True
            handled_markup = True
            continue
        if lowered == "</sup>":
            active_superscript = superscript
            handled_markup = True
            continue
        if lowered in {"<br>", "<br/>", "<br />"}:
            paragraph.add_run("\n")
            handled_markup = True
            continue

        if token.startswith("<") and token.endswith(">"):
            return False

        _runtime_value("append_text_run", append_text_run)(
            paragraph,
            token,
            style_config=style_config,
            bold=bold,
            italic=italic,
            subscript=active_subscript,
            superscript=active_superscript,
        )

    return handled_markup


def render_markup_children(
    parent: Tag,
    paragraph: Any,
    *,
    style_config: Any | None,
    bold: bool,
    italic: bool,
    subscript: bool,
    superscript: bool,
) -> None:
    for child in parent.children:
        if isinstance(child, NavigableString):
            text = str(child)
            if not text:
                continue
            _runtime_value("append_text_run", append_text_run)(
                paragraph,
                text,
                style_config=style_config,
                bold=bold,
                italic=italic,
                subscript=subscript,
                superscript=superscript,
            )
            continue

        if not isinstance(child, Tag):
            continue

        if child.name == "br":
            paragraph.add_run("\n")
            continue

        render_markup_children(
            child,
            paragraph,
            style_config=style_config,
            bold=bold or child.name in {"b", "strong", "th"},
            italic=italic or child.name in {"i", "em"},
            subscript=subscript or child.name == "sub",
            superscript=superscript or child.name == "sup",
        )


def append_text_run(
    paragraph: Any,
    text: str,
    *,
    style_config: Any | None,
    bold: bool,
    italic: bool,
    subscript: bool,
    superscript: bool,
) -> None:
    run = paragraph.add_run(text)
    _runtime_value("apply_run_defaults", apply_run_defaults)(run, style_config)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True


def fill_title_paragraph(
    paragraph: Any,
    block: Any,
    page_width: float,
    render_config: Any,
    *,
    keep_with_next: bool = True,
) -> None:
    paragraph.text = ""
    paragraph.paragraph_format.keep_with_next = keep_with_next
    prepare_text_fn = _runtime_value("prepare_text", prepare_text)
    infer_heading_level_fn = _runtime_value("infer_heading_level", infer_heading_level)
    is_cover_title_fn = _runtime_value("is_cover_title", is_cover_title)
    apply_named_style_config_fn = _runtime_value("apply_named_style_config", apply_named_style_config)
    fill_rich_text_fn = _runtime_value("fill_rich_text", fill_rich_text)
    get_style_config_fn = _runtime_value("get_style_config", get_style_config)
    apply_run_defaults_fn = _runtime_value("apply_run_defaults", apply_run_defaults)

    prepared = prepare_text_fn(block.text, render_config)
    level = infer_heading_level_fn(prepared)

    if is_cover_title_fn(block, prepared, page_width):
        apply_named_style_config_fn(paragraph, render_config, "PartNumber")
        style_name = "PartNumber"
    else:
        style_name = f"Heading {level}"
        apply_named_style_config_fn(paragraph, render_config, style_name)

    fill_rich_text_fn(paragraph, prepared, render_config=render_config, style_name=style_name, bold=True)
    for run in paragraph.runs:
        apply_run_defaults_fn(run, get_style_config_fn(render_config, style_name))
        if is_cover_title_fn(block, prepared, page_width):
            if run.font.size is None:
                run.font.size = TOP_TITLE_FONT
        elif run.font.size is None:
            run.font.size = SECTION_TITLE_FONT


def is_cover_title(block: Any, text: str, page_width: float) -> bool:
    if block.page_idx not in {0, 1}:
        return False
    if block.y1 >= 220 or block.width < page_width * 0.22:
        return False
    if HEADING_NUMBER_RE.match(text):
        return False
    if TOC_TITLE_RE.match(text.strip()):
        return False
    return True


def fill_body_paragraph(paragraph: Any, text: str, render_config: Any) -> None:
    paragraph.text = ""
    apply_named_style_config_fn = _runtime_value("apply_named_style_config", apply_named_style_config)
    fill_rich_text_fn = _runtime_value("fill_rich_text", fill_rich_text)
    prepare_text_fn = _runtime_value("prepare_text", prepare_text)
    get_style_config_fn = _runtime_value("get_style_config", get_style_config)
    apply_run_defaults_fn = _runtime_value("apply_run_defaults", apply_run_defaults)

    apply_named_style_config_fn(paragraph, render_config, "Normal")
    fill_rich_text_fn(paragraph, prepare_text_fn(text, render_config), render_config=render_config, style_name="Normal")
    for run in paragraph.runs:
        apply_run_defaults_fn(run, get_style_config_fn(render_config, "Normal"))


def fill_list_paragraph(paragraph: Any, text: str, render_config: Any) -> None:
    paragraph.text = ""
    apply_named_style_config_fn = _runtime_value("apply_named_style_config", apply_named_style_config)
    fill_rich_text_fn = _runtime_value("fill_rich_text", fill_rich_text)
    prepare_text_fn = _runtime_value("prepare_text", prepare_text)
    get_style_config_fn = _runtime_value("get_style_config", get_style_config)
    apply_run_defaults_fn = _runtime_value("apply_run_defaults", apply_run_defaults)

    apply_named_style_config_fn(paragraph, render_config, "Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if paragraph.paragraph_format.left_indent is None:
        paragraph.paragraph_format.left_indent = Pt(10)
    if paragraph.paragraph_format.first_line_indent is None:
        paragraph.paragraph_format.first_line_indent = Pt(-10)
    fill_rich_text_fn(paragraph, f"• {prepare_text_fn(text, render_config)}", render_config=render_config, style_name="Normal")
    for run in paragraph.runs:
        apply_run_defaults_fn(run, get_style_config_fn(render_config, "Normal"))


def normalize_list_item(text: str) -> str:
    normalized = text.strip()
    while normalized.startswith(("•", "■", "", "-", "*")):
        normalized = normalized[1:].strip()
    return normalized


def infer_heading_level(text: str) -> int:
    text = _runtime_value("normalize_toc_text", normalize_toc_text)(text)
    match = NUMBERED_HEADING_RE.match(text)
    if match:
        numbering = match.group("num")
        return min(numbering.count(".") + 1, 3)
    if CHAPTER_HEADING_RE.match(text):
        return 1
    if re.match(r"^(?:附录|appendix\b)", text, re.IGNORECASE):
        return 1
    return 2


def looks_like_heading_text(text: str) -> bool:
    cleaned = _runtime_value("normalize_toc_text", normalize_toc_text)(text)
    if len(cleaned) > 80:
        return False
    if re.match(r"^\d+\.[A-Z]", cleaned):
        return False
    return bool(
        HEADING_NUMBER_RE.match(cleaned)
        or CHAPTER_HEADING_RE.match(cleaned)
        or looks_like_english_heading_text(cleaned)
    )


def looks_like_english_heading_text(text: str) -> bool:
    stripped = text.strip()
    if not stripped or len(stripped) > 60:
        return False
    if stripped[-1] in ".!?;:":
        return False

    tokens = ENGLISH_HEADING_TOKEN_RE.findall(stripped)
    if not 1 <= len(tokens) <= 8:
        return False

    significant_chars = sum(1 for char in stripped if char.isalpha() or char.isdigit())
    token_chars = sum(len(token) for token in tokens)
    if token_chars < significant_chars * 0.7:
        return False

    acceptable = 0
    for token in tokens:
        lowered = token.lower()
        if lowered in ENGLISH_HEADING_STOPWORDS or token.isupper() or token[0].isupper():
            acceptable += 1

    return acceptable >= max(len(tokens) - 1, 1)
